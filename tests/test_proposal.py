"""Unit tests for pxrd_review.lints and pxrd_review.proposal — synthetic text and a synthetic manuscript.

    python3 -m unittest tests.test_proposal -v
"""
import os, shutil, tempfile, unittest

from pxrd_review import lints as LI, proposal as PR
from tests.test_bv_check import RUTILE, _write


class Lints(unittest.TestCase):
    def test_spectroscopy(self):
        t = ('Raman spectroscopy was conducted on a Horiba XploRA PLUS with a 100× objective. The spectrum was obtained using a 248 nm laser. '
             'The broad band from 3600 to 2800 cm-1, with shoulders at 3483, 2922 and 2848 cm-1, is the O–H stretching of hydrogen-bonded water. '
             'A weak band at 2349 cm-1 is also assigned to water.')
        f = [m for m, _a in LI.spectro(t)]
        self.assertTrue(any('2848 and 2922' in m and 'C–H' in m for m in f), f)
        self.assertTrue(any('248 nm laser on a Horiba XploRA' in m and '532/638/785' in m for m in f), f)
        self.assertTrue(any('2349' in m and 'CO2' in m for m in f), f)
        self.assertEqual(LI.spectro('A LabRAM HR with a 532 nm laser. Bands at 3400 and 1630 cm-1 are the water modes; the 2920 cm-1 band is C–H of the epoxy.'), [])

    def test_notation(self):
        paras = ['Cleavage is expected on [110] and [010] based on the structure.', 'Twinning by rotation about [110] is common.',
                 'Table 2. Powder data for the mineral.', 'Table 2. Powder data for Mg-arsenate.', 'Figure 1. Crystals.']
        codes = {0: {'caption': 5, 'codes': {'2': ([[-1, 0, 0], [0, 1, 0], [0, 0, -1]], [1, 0, 1], '-x+1, y, -z+1')}},
                 1: {'caption': 6, 'codes': {'2': ([[-1, 0, 0], [0, 1, 0], [0, 0, -1]], [0, 0, 0], '-x, y, -z')}}}
        codes[2] = {'caption': 7, 'codes': codes[1]['codes']}                       # inherited: not a second definition
        f = [m for m, _a in LI.notation(paras + ['Prisms are elongated along (100) and striated parallel to (001).'], codes)]
        self.assertEqual(sum('zone symbol [110]' in m for m in f), 1, f)             # the cleavage, not the twin axis
        self.assertEqual(sum('given as the plane' in m for m in f), 2, f)              # a direction written as (hkl)
        self.assertTrue(any('Table 2 has 2 captions' in m for m in f), f)
        self.assertEqual(sum('symmetry code (2)' in m for m in f), 1, f)


    def test_crossrefs(self):
        paras = ['Table 1. Analytical data.', 'Table 2. Powder data.', 'Figure 1. Crystals.', 'Figure 2. The structure.', 'Figure 3. Raman spectrum.',
                 'The composition (Table 1) and the powder data (Tables 2 and 3) are given; the structure is shown in Figs. 2–3, and Fig. 5a shows twinning.',
                 'Compare Table 4 in Smith (2020).', 'Table 2 (continued).', 'References', 'Jones, A. (2019) Table 9 of wonders. Journal, 1, 1–2.']
        f = LI.crossrefs(paras)
        flags = [m for m, _a, s in f if s == 'flag']; notes = [m for m, _a, s in f if s == 'note']
        self.assertEqual(len(flags), 2, flags)
        self.assertTrue(any(m.startswith('cross-reference: Table 3 is cited') and 'table captions are 1–2' in m for m in flags), flags)
        self.assertTrue(any(m.startswith('cross-reference: Figure 5 is cited') and 'figure captions are 1–3' in m for m in flags), flags)
        self.assertEqual(notes, ["cross-reference: Figure 1 (‘Figure 1. Crystals.’) is never cited in the text"])
        # no captions of a kind: nothing is judged; a supplementary set travelling separately is left alone
        self.assertEqual(LI.crossrefs(['See Table 3 and Fig. 2.']), [])
        self.assertEqual([s for _m, _a, s in LI.crossrefs(['Table 1. Data.', 'Table 1 and Table S2 hold the data.'])], [])
        # a docx lint carries the flag separately from the notes
        self.assertEqual(LI._numbers('3a, b'), ['3']); self.assertEqual(LI._numbers('2–4'), ['2', '3', '4']); self.assertEqual(LI._numbers('S1 and S3'), ['S1', 'S3'])

    def test_evidence(self):
        t = 'The empirical formula is Ca2Fe3+2(SO4)3(OH)2·2H2O. Fe is trivalent by bond-valence sums. Analyses were made by EPMA.'
        f = [m for m, _a in LI.evidence(t)]
        self.assertEqual(len(f), 2, f)
        self.assertTrue(any('assigns a valence to Fe (Fe3+)' in m and 'rests on bond-valence' in m for m in f), f)
        self.assertTrue(any('carries H2O or OH' in m for m in f), f)
        self.assertEqual(LI.evidence(t + ' Mössbauer spectroscopy confirms Fe3+. The Raman band at 3450 cm-1 is the O–H stretch.'), [])
        self.assertEqual(LI.evidence('The formula is CaSO4. No water.'), [])


class Proposal(unittest.TestCase):
    def test_review_writes_report_and_copy(self):
        tmp = tempfile.mkdtemp(prefix='prop_')
        try:
            _write(tmp, 'testite.cif', RUTILE)
            from docx import Document
            doc = Document()
            doc.add_paragraph('Testite, ideal formula TiO2, is a new mineral (Smith et al., 2020). Cleavage on [110].')
            doc.add_paragraph('Density (calc.) = 4.250 g·cm–3 for the ideal formula.')
            doc.add_paragraph('The Ti2 site is octahedral.')
            doc.add_paragraph('References')
            doc.add_paragraph('Jones, A. (2019) Another paper. Journal, 1, 1–2.')
            path = os.path.join(tmp, 'testite manuscript.docx'); doc.save(path)
            docx, cif, cc = PR.find_files(tmp)
            self.assertEqual((os.path.basename(docx), os.path.basename(cif), cc), ('testite manuscript.docx', 'testite.cif', None))
            res = PR.review(tmp, quiet=True)
            text = open(res['report'], encoding='utf-8').read()
            self.assertIn('== references', text); self.assertIn('== .cif audit', text); self.assertIn('== lints', text)
            self.assertIn('site Ti2 is named in the text', text)
            self.assertIn('cleavage given as the zone symbol [110]', text)
            self.assertTrue(os.path.exists(res['copy']))
            self.assertGreaterEqual(res['comments'], 1)
            self.assertEqual(res['unplaced'], [])
            # the copy carries the comments; the source is untouched
            from docx import Document as D
            d2 = D(res['copy']); self.assertTrue(any('comments.xml' in str(p.partname) for p in d2.part.package.parts))
            self.assertEqual(len(Document(path).paragraphs), 5)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_a_pdf_manuscript_gets_a_report_and_no_copy(self):
        import pymupdf
        tmp = tempfile.mkdtemp(prefix='prop_pdf_')
        try:
            _write(tmp, 'testite.cif', RUTILE)
            doc = pymupdf.open(); page = doc.new_page(width=595, height=842)
            y = 72
            for line in ('Testite, ideal formula TiO2, is a new mineral (Smith et al., 2020).', 'Density (calc.) = 4.250 g/cm3 for the ideal formula.',
                         'The Ti2 site is octahedral. See Table 3.', 'Table 1. Data.', 'References', 'Jones, A. (2019) Another paper. Journal, 1, 1-2.'):
                page.insert_text((72, y), line, fontsize=10); y += 16
            path = os.path.join(tmp, 'testite proposal.pdf'); doc.save(path); doc.close()
            docx, cif, cc = PR.find_files(tmp)
            self.assertEqual((os.path.basename(docx), os.path.basename(cif)), ('testite proposal.pdf', 'testite.cif'))
            res = PR.review(tmp, quiet=True)
            self.assertIsNone(res['copy'])
            text = open(res['report'], encoding='utf-8').read()
            self.assertIn('report only — a .pdf manuscript takes no comments', text)
            self.assertIn('== .cif audit', text); self.assertIn('== lints', text)
            self.assertIn('Table 3 is cited', text)                                          # the lints run on a .pdf; the docx-only label check does not
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    unittest.main()
