"""annotate_review._fit_page: an output copy whose tables overrun the page gets a wider page — page
setup only; a page that already holds its tables is left alone."""
import unittest

from docx import Document
from docx.shared import Twips

from pxrd_review import annotate_review as A


def doc_with_table(width):
    d = Document()
    t = d.add_table(rows=1, cols=2)
    t.autofit = False
    tblW = t._tbl.tblPr.find(A._q('tblW'))
    if tblW is None:
        tblW = t._tbl.tblPr.makeelement(A._q('tblW'), {})
        t._tbl.tblPr.append(tblW)
    tblW.set(A._q('w'), str(width)); tblW.set(A._q('type'), 'dxa')
    t.cell(0, 0).text = 'PDFID :'
    return d


class FitPage(unittest.TestCase):
    def test_a_ten_inch_table_on_letter_portrait_turns_the_page(self):
        d = doc_with_table(14400)
        s = d.sections[-1]
        s.page_width, s.page_height = Twips(12240), Twips(15840)
        s.left_margin = s.right_margin = Twips(1440)
        self.assertTrue(A._fit_page(d))
        self.assertEqual((s.page_width.twips, s.page_height.twips), (15840, 12240))
        self.assertGreaterEqual(s.page_width.twips - s.left_margin.twips - s.right_margin.twips, 14400)
        self.assertEqual(d.tables[0].cell(0, 0).text, 'PDFID :')

    def test_a_page_that_holds_its_tables_is_untouched(self):
        d = doc_with_table(8000)
        s = d.sections[-1]
        before = (s.page_width, s.page_height, s.left_margin, s.right_margin)
        self.assertFalse(A._fit_page(d))
        self.assertEqual(before, (s.page_width, s.page_height, s.left_margin, s.right_margin))


if __name__ == '__main__':
    unittest.main()
