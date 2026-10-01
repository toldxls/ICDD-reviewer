"""Unit tests for pxrd_review.spectro — synthetic sentences in the forms papers use, and a synthetic uranyl structure.

    python3 -m unittest tests.test_spectro -v
"""
import os, shutil, tempfile, unittest

from pxrd_review import spectro as S
from tests.test_bv_check import _write

URANYL = """data_test
_chemical_name_mineral testite
_cell_length_a 10
_cell_length_b 10
_cell_length_c 10
_cell_angle_alpha 90
_cell_angle_beta 90
_cell_angle_gamma 90
_space_group_name_H-M_alt 'P 1'
loop_
_space_group_symop_operation_xyz
'x, y, z'
loop_
_atom_site_label
_atom_site_type_symbol
_atom_site_fract_x
_atom_site_fract_y
_atom_site_fract_z
U1 U 0.5 0.5 0.5
O1 O 0.5 0.5 0.677
O2 O 0.5 0.5 0.323
O3 O 0.735 0.5 0.5
O4 O 0.265 0.5 0.5
O5 O 0.5 0.735 0.5
O6 O 0.5 0.265 0.5
O7 O 0.015 0.5 0.5
H1 H 0.83 0.5 0.5
"""


def _kinds(text):
    return [(r['claim'].text, r['claim'].paired, r['status'], r['grade']) for r in S.check_claims(text)]


class Relations(unittest.TestCase):
    def test_bartlett_cooney(self):
        self.assertAlmostEqual(S.bc_nu1(860), 1.753, 3)
        self.assertAlmostEqual(S.bc_nu3(915), 1.774, 3)
        self.assertAlmostEqual(S.bc_nu1_inv(S.bc_nu1(831.0)), 831.0, 6)
        self.assertAlmostEqual(S.bc_nu3_inv(S.bc_nu3(927.0)), 927.0, 6)

    def test_libowitzky(self):
        self.assertAlmostEqual(S.lib_d(3013), 2.652, 3)
        self.assertAlmostEqual(S.lib_nu(S.lib_d(3401.0)), 3401.0, 6)
        self.assertIsNone(S.lib_d(3600))                          # above the curve's limit: a free O–H has no distance
        self.assertAlmostEqual(S.lib_nu_ho(S.lib_d_ho(3300.0)), 3300.0, 6)
        self.assertIsNone(S.lib_d_ho(3640))


class Text(unittest.TestCase):
    def test_wavenumber_forms(self):
        for form in ('845 cm–1', '845 cm−1', '845 cm-1', '845 cm1', '845cm1', '845 cm\x011', '845 cm 1', '845 cm.'):
            self.assertIn('845 cm-1', S.normalise('a band at %s' % form), form)
        self.assertNotIn('cm-1', S.normalise('a crystal 0.12 cm long'))

    def test_nu_forms(self):
        self.assertIn('ν1 (UO2)', S.normalise('the n1 (UO2)2+ stretch'))
        self.assertIn('ν3 (SO4', S.normalise('the υ3 (SO4)2– mode'))
        self.assertIn('ν3 (UO2)', S.normalise('the v3 (UO2)2+ mode'))
        self.assertEqual(S.normalise('in 3 samples and n 1 grains'), 'in 3 samples and n 1 grains')

    def test_hydrogen_bond_dots(self):
        for form in ('O⋅⋅⋅O', 'O···O', 'O. . .O', 'O…O', 'O· · ·O'):
            self.assertIn('O···O', S.normalise('the %s distance' % form), form)

    def test_sentences(self):
        s = S.sentences('The spectrum (Fig. 3) follows Smith et al. (2020). Bands at 3540 cm-1 are weak. (They are broad.) A third follows.')
        self.assertEqual(len(s), 4, s)

    def test_quote_is_the_papers_own_wording(self):
        raw, norm = S.split('The ν1 band at 845 cm–1 is strong. A second sentence.')
        self.assertIn('cm–1', raw[0]); self.assertIn('cm-1', norm[0])


class Bands(unittest.TestCase):
    def _b(self, t):
        return [(b.pos, b.family, b.mode) for b in S.bands(t)]

    def test_forward_assignment(self):
        self.assertEqual(self._b('Bands at 1210, 1188 and 1150 cm–1 are attributed to the triply degenerate ν3 (SO4)2– antisymmetric stretching vibrations '
                                 'and those at 1035 and 990 cm–1 to the ν1 (SO4)2– symmetric stretching vibrations.'),
                         [(1210.0, 'SO4', 3), (1188.0, 'SO4', 3), (1150.0, 'SO4', 3), (1035.0, 'SO4', 1), (990.0, 'SO4', 1)])

    def test_backward_assignment(self):
        self.assertEqual(self._b('The ν1 (UO2)2+ symmetric stretching vibration is present as a very strong band at 851 cm−1, with a weaker shoulder at 846 cm−1.'),
                         [(851.0, 'UO2', 1), (846.0, 'UO2', 1)])
        self.assertEqual(self._b('The n1 UO2 2+ symmetric stretching vibration is assigned to the very strong band at 812 cm1.'), [(812.0, 'UO2', 1)])

    def test_two_bands_two_assignments(self):
        self.assertEqual(self._b('(These are weaker than υ3, as usual in Raman.) A very weak band at 948 cm–1 and a very strong one at 853 cm–1 '
                                 'are assigned to υ3 (UO2)2+ antisymmetric and υ1 (UO2)2+ symmetric stretching vibrations.'),
                         [(948.0, 'UO2', 3), (853.0, 'UO2', 1)])

    def test_respectively(self):
        self.assertEqual(self._b('The bands at 528 and 392 cm–1 may be attributed to the ν4 and ν2 bending vibrations of AsO4 groups, respectively.'),
                         [(528.0, 'AsO4', 4), (392.0, 'AsO4', 2)])

    def test_unit_left_to_the_next_run(self):
        b = S.bands('The bands at 921 and 874 (IR) correspond to the split ν3 antisymmetric stretching vibrations of UO2, and a prominent Raman band '
                    'at 811 cm–1 is attributed to ν1 symmetric stretching vibrations of UO2, overlapping the symmetric stretching vibrations of SiO4 at ∼752 cm–1 (Raman).')
        self.assertEqual([(x.pos, x.technique, x.family, x.mode) for x in b],
                         [(921.0, 'ir', 'UO2', 3), (874.0, 'ir', 'UO2', 3), (811.0, 'raman', 'UO2', 1), (752.0, 'raman', 'SiO4', 1)])

    def test_oh_and_water(self):
        self.assertEqual(self._b('Raman bands at 3615, 3541 and 3490 cm–1 are attributed to υ O–H stretching vibrations of symmetrically non-equivalent H2O molecules. '
                                 'A very weak band at 1621 cm–1 is attributable to υ2 (δ) H–O–H bending vibrations.'),
                         [(3615.0, 'OH', 'stretch'), (3541.0, 'OH', 'stretch'), (3490.0, 'OH', 'stretch'), (1621.0, 'H2O', 2)])

    def test_bond_named_assignment_is_not_the_hydroxyl(self):
        b = self._b('Raman bands at 829 and 818 cm1 are due to n1 (UO2)2+ symmetric stretching vibrations and those at 690, 552 and 533 cm1 are assigned to δ U–OH bending vibrations.')
        self.assertEqual(b[:2], [(829.0, 'UO2', 1), (818.0, 'UO2', 1)])
        self.assertTrue(all(f != 'OH' for _p, f, _m in b[2:]), b)
        b = self._b('The spectrum shows bands of O–H stretching vibrations at 3580 cm–1, bands of C–O stretching vibrations at 760 and 1100 cm–1, and bands of Y–O vibrations at 180 and 270 cm–1.')
        self.assertEqual(b[0], (3580.0, 'OH', 'stretch'))
        self.assertTrue(all(f != 'OH' for _p, f, _m in b[1:]), b)

    def test_a_distance_is_not_an_assignment(self):
        b = self._b('An infrared band at 908 cm–1 (m-s) and a Raman band at 893 cm–1 (vw) are attributed to ν3 (UO2)2+ antisymmetric stretching vibrations, '
                    'corresponding to U–O (uranyl) bond lengths of 1.78 and 1.79 Å.')
        self.assertEqual(b, [(908.0, 'UO2', 3), (893.0, 'UO2', 3)])

    def test_ranges_and_regions_are_not_bands(self):
        self.assertEqual(self._b('Bands in the range 3000–3600 cm–1 correspond to O–H stretching vibrations and the ones at 1570 and 1610 cm–1 to bending vibrations of H2O molecules.'),
                         [(1570.0, 'H2O', 'bend'), (1610.0, 'H2O', 'bend')])
        self.assertEqual(self._b('The low-energy (<400 cm–1) and OH-stretching ranges are distinctive.'), [])
        self.assertEqual(S.band_ranges('A broad band from 3550 to 3100 cm–1 is present.')[0][:2], (3100.0, 3550.0))

    def test_semicolon_ends_a_run(self):
        b = self._b('The band at 3535 cm–1 is assigned to H2O stretching vibrations; 1385 and 1275 cm–1 to CO2 Fermi resonance; 1070 cm–1 (weak) to vibrations of the framework or HF librations.')
        self.assertEqual(b[0], (3535.0, 'H2O', 'stretch'))
        self.assertEqual([f for _p, f, _m in b[1:3]], [None, None])

    def test_shape_and_technique(self):
        b = S.bands('Strong Raman bands at 861 (vs) and 847 (s) cm–1 are attributed to the ν1 symmetric stretching vibration of the uranyl ion, as are weak infrared bands at 866 and 851 cm–1.')
        self.assertEqual([(x.pos, x.technique, x.shape) for x in b], [(861.0, 'raman', 'vs'), (847.0, 'raman', 's'), (866.0, 'ir', 'weak'), (851.0, 'ir', 'weak')])

    def test_another_works_values(self):
        b = S.bands('For the free SO4 2– ion the Raman frequencies ν1 ≈983 cm–1 and ν2 ≈452 cm–1 have been reported.')
        self.assertTrue(b and all(not x.own for x in b), b)

    def test_table_run_together_is_skipped(self):
        row = ' '.join('%d %.3f %d %d %d' % (k * 7 % 100, 1.0 + k * 0.137, k % 5, k % 3, k % 7) for k in range(1, 30))
        self.assertEqual(S.bands('I d h k l ' + row + ' 845 cm–1 ν1 (UO2)2+.'), [])


class BartlettCooney(unittest.TestCase):
    def test_band_beside_the_distance(self):
        t = ('The ν1 (UO2)2+ symmetric stretching vibration is present as a very strong band at 856 cm−1, with a shoulder at 852 cm−1. Bartlett and Cooney (1989) '
             'provided an empirical relationship to derive the U–O bond lengths from the band position, which gives 1.75 Å (856 cm−1) and 1.76 Å (852 cm−1), '
             'in excellent agreement with the average U–O bond length from the X-ray data: 1.770 Å.')
        self.assertEqual(_kinds(t), [('1.75', 'band', 'reproduced', None), ('1.76', 'band', 'reproduced', None)])

    def test_mode_beside_the_distance(self):
        t = ('Strong infrared bands at 915 cm–1 and 849 cm–1 are attributed to ν3 (UO2)2+ antisymmetric and ν1 (UO2)2+ symmetric stretching vibrations, respectively. '
             'The uranyl U–O bond lengths 1.773 (ν3) and 1.762 (ν1) Å, calculated using the empirical relation given by Bartlett and Cooney (1989), are comparable '
             'with bond lengths 1.792–1.793 Å obtained from the X-ray studies.')
        self.assertEqual(_kinds(t), [('1.773', 'mode', 'reproduced', None), ('1.762', 'mode', 'reproduced', None)])

    def test_other_forms(self):
        self.assertEqual(_kinds('According to Bartlett and Cooney (1989), approximate U–O bond lengths in the uranyl ion are 1.78 Å/904 cm–1, 1.78 Å/838 cm–1 and 1.79 Å/824 cm–1.'),
                         [('1.78', 'band', 'reproduced', None), ('1.78', 'band', 'reproduced', None), ('1.79', 'band', 'reproduced', None)])
        self.assertEqual(_kinds('Using the relation of Bartlett and Cooney (1989) we calculated the U–O bond lengths as follows (in Å): 1.77 (from 917 cm–1), 1.81 (870 cm–1) and 1.80 (808 cm–1).'),
                         [('1.77', 'band', 'reproduced', None), ('1.81', 'band', 'reproduced', None), ('1.80', 'band', 'reproduced', None)])
        self.assertEqual(_kinds('According to the empirical relationship of Bartlett and Cooney (1989), the ν1 (UO2)2+ band at 860 cm–1 corresponds to a U–O bond length of ∼1.75 Å, '
                                'in excellent agreement with U–O bond lengths from the X-ray data: 1.750(6) and 1.766(6) Å.'), [('1.75', 'single', 'reproduced', None)])

    def test_the_other_relation_gives_the_stated_value(self):
        t = ('Applying the empirical relationship of Bartlett & Cooney (1989), the intense v1 (UO2)2+ symmetric stretch at 824 cm1 corresponds to a U–O bond length '
             'of 1.79 Å, and 1.76 for v3 (858 cm1). In the structure the U–O bond lengths range between 1.75 and 1.76 Å.')
        r = S.check_claims(t)
        self.assertEqual([(x['claim'].text, x['status'], x['grade']) for x in r], [('1.76', 'sibling', 'flag'), ('1.79', 'reproduced', None)])
        self.assertIn(' vs 1.82 Å by the ν3 relation', r[0]['line'])
        self.assertIn('v3 (858 cm1)', r[0]['raw'])                 # the paper's wording, to be found again in the document

    def test_a_mode_read_from_the_context_never_makes_that_finding(self):
        # 808 is the paper's ν1; a reader that took it for ν3 would see 'the other relation': only a mode the sentence
        # itself puts beside the distance may say so
        t = ('The bands at 917 and 870 (IR) correspond to the split ν3 antisymmetric stretching vibrations of UO2, and a Raman band at 808 cm–1 is attributed to ν1. '
             'Using Bartlett and Cooney (1989) the U–O bond lengths are (in Å): 1.77 (from 917 cm–1), 1.81 (870 cm–1) and 1.80 (808 cm–1).')
        self.assertTrue(all(st == 'reproduced' for _t, _p, st, _g in _kinds(t)), _kinds(t))

    def test_unpaired_is_silent(self):
        self.assertEqual(_kinds('The inferred U–O bond length obtained from the empirical relation given by Bartlett & Cooney (1989) is ∼1.79 Å.'), [('1.79', '', 'unpaired', None)])
        # bands several sentences back (another spectrum, a page break between) are no ground for a remark
        t = ('A band at 924 cm–1 is assigned to ν3 (UO2)2+. ' + 'A filler sentence. ' * 3 +
             'Inferred approximate U–O bond lengths in uranyl are 1.73 and 1.76 Å, respectively (Bartlett and Cooney 1989).')
        self.assertTrue(all(g is None for _t, _p, _s, g in _kinds(t)), _kinds(t))

    def test_reference_list_entry_states_nothing(self):
        self.assertEqual(S.check_claims('Results. A band at 845 cm–1. ' * 30 + 'References Bartlett, J.R. and Cooney, R.P. (1989) On the determination of uranium–oxygen bond '
                                        'lengths in dioxouranium(VI) compounds by Raman spectroscopy. Journal of Molecular Structure, 193, 295–300 (1.76 Å).'), [])


class Libowitzky(unittest.TestCase):
    def test_lists_in_one_sentence(self):
        t = ('According to the equation of Libowitzky (1999), the O–H stretching bands with the absorption maxima at 3013, 3161, 3355 and 3401 cm–1 correspond to '
             'the O⋅⋅⋅O distances of 2.65, 2.69, 2.77 and 2.80 Å, respectively. These values are close to the D⋅⋅⋅A distances of 2.620 and 2.698 Å.')
        self.assertEqual([k[1:] for k in _kinds(t)], [('list', 'reproduced', None)] * 4)

    def test_one_pair_out(self):
        t = ('Based on the correlation reported by Libowitzky (1999), the D···A distances corresponding to the bands at 3225, 2451, and 3536 cm−1 are equal to '
             '2.71, 2.83, and 2.97 Å, respectively.')
        r = S.check_claims(t)
        self.assertEqual([(x['claim'].text, x['status'], x['grade']) for x in r], [('2.71', 'reproduced', None), ('2.83', 'outlier', 'flag'), ('2.97', 'reproduced', None)])
        self.assertIn('2451 cm⁻¹ vs 2.56 Å', r[1]['line'])
        self.assertIn('the other 2 pairs reproduce', r[1]['line'])

    def test_bands_in_the_sentence_before(self):
        t = ('A broad band with components at 3475, 3420, and 3262 cm−1 was assigned to the ν O–H stretching vibrations of the H2O groups. According to the empirical '
             'equation of Libowitzky (1999), O–H···O hydrogen bond lengths were calculated at approximately 2.87, 2.81, and 2.73 Å, which are comparable with the structure.')
        self.assertEqual([k[1:] for k in _kinds(t)], [('ctx', 'reproduced', None)] * 3)

    def test_another_order_without_respectively_is_reproduced(self):
        t = ('The O–H stretch region has maxima at 3338 and 3057 cm–1. According to Libowitzky (1999) these correspond to O⋅⋅⋅O distances of 2.66 and 2.76 Å.')
        self.assertTrue(all(st == 'reproduced' and g is None for _t, _p, st, g in _kinds(t)), _kinds(t))

    def test_band_beside_the_distance(self):
        t = 'Using the correlation curve of Libowitzky (1999), two O···O distances are expected, ~2.66 Å (3028 cm–1) and ~2.72 Å (3253 cm–1).'
        self.assertEqual(_kinds(t), [('2.66', 'band', 'reproduced', None), ('2.72', 'band', 'reproduced', None)])

    def test_a_range_is_counted_and_not_judged(self):
        t = 'Bands at 3620, 3546 and 3498 cm–1 are O–H stretches. The O–H⋅⋅⋅O hydrogen bond lengths, using the relation of Libowitzky (1999), vary in the range ∼3.2 to 2.9 Å.'
        self.assertEqual([k[1:] for k in _kinds(t)], [('range', 'range', None)] * 2)

    def test_h_acceptor_form(self):
        r = S.stated_distances('According to the correlation given by Libowitzky (1999), these vibrations correspond to H-bonds (H…Acceptor) in the range 2.1–1.7 Å.')
        self.assertEqual({c.relation for c in r}, {'lib_ho'})
        r = S.stated_distances('According to Libowitzky (1999), O–H···O hydrogen bond lengths vary from 2.82 to 2.63 Å.')
        self.assertEqual({c.relation for c in r}, {'lib_oo'})          # 'O–H···O' is the donor–acceptor distance, not H···O

    def test_a_class_boundary_is_not_a_distance(self):
        self.assertEqual(S.stated_distances('According to the Libowitzky (1999) classification, only these are strong bonds, with O···O < 2.7 Å.'), [])

    def test_a_proved_reading_makes_a_stated_pair_a_finding(self):
        t = ('Using Libowitzky (1999), the distances expected are ~2.66 Å (3028 cm–1) and ~2.72 Å (3253 cm–1). '
             'For the second sample the curve of Libowitzky (1999) gives 2.95 Å (3253 cm–1).')
        r = S.check_claims(t)
        self.assertEqual([(x['status'], x['grade']) for x in r], [('reproduced', None), ('reproduced', None), ('off', 'flag')])
        self.assertIn(' vs ', r[2]['line'])
        t1 = 'For the sample the curve of Libowitzky (1999) gives 2.95 Å (3253 cm–1).'      # the same pair alone: nothing vouches for the reading
        self.assertEqual([(x['status'], x['grade']) for x in S.check_claims(t1)], [('off', 'info')])


HEAD = 'Testite, a new mineral. The ideal formula is %s. '
RAMAN = ('The Raman spectrum shows a strong band at 832 cm–1 assigned to the ν1 (UO2)2+ stretch. Bands at 1010 and 1105 cm–1 are assigned to ν1 and ν3 of (SO4)2–, respectively. ')
OHS = 'Broad bands at 3250 and 3480 cm–1 are assigned to O–H stretching vibrations. '


class Absent(unittest.TestCase):
    """A band the sentence says the spectrum does NOT show is no band of the spectrum."""

    def test_forms(self):
        for t, lo, hi in (('There is no absorption band at ca. 1650 cm–1 in the spectrum.', 1650, 1650),
                          ('The H–O–H bend near 1600 cm–1 is absent.', 1600, 1600),
                          ('No bands were observed in the region of bending vibrations of water molecules (~1600 cm–1).', 1600, 1600),
                          ('Bands of H–O–H bending vibrations, seen in related minerals in the range 1700–1600 cm–1, were not detected.', 1600, 1700),
                          ('The absence of bands at 1430 cm–1 rules out carbonate.', 1430, 1430)):
            self.assertEqual([(a, b) for a, b, _s in S.absent_bands(t)], [(lo, hi)], t)
            self.assertEqual(S.bands(t), [], t)

    def test_what_is_there_stays(self):
        t = 'There are no hydroxyl bands, but stretches of molecular H2O are seen as a broad band at 3310 cm–1. A band at 1640 cm–1, which is not split, is assigned to H2O bending.'
        self.assertEqual(sorted(b.pos for b in S.bands(t)), [1640.0, 3310.0]); self.assertEqual(S.absent_bands(t), [])

    def test_a_region_is_no_band(self):
        for t in ('Bands below 320 cm–1 are assigned to lattice modes.', 'No significant Raman bands were observed above 1700 cm–1.', 'Bands at wavenumbers lower than 500 cm–1 are lattice modes.'):
            self.assertEqual(S.bands(t), [], t)
        self.assertEqual(S.absent_bands('No significant Raman bands were observed above 1700 cm–1.'), [])


class Formula(unittest.TestCase):
    """The formula's hydrogen species and light groups beside the bands stated: information, every line."""

    def _lines(self, formula, body, **kw):
        return S.species_lines(HEAD % formula + body, name='testite', **kw)

    def test_species_a_formula_writes(self):
        self.assertEqual(S.formula_species('Ca2(UO2)3(SO4)2(OH)6·4H2O'), {'H2O', 'OH'})
        self.assertEqual(S.formula_species('(NH4)2(H3O)[Al(CO3)2]'), {'NH4', 'H3O', 'CO3'})
        self.assertEqual(S.formula_species('Ca5(PO4)3F'), set())
        self.assertEqual(S.formula_species('Cu<sub>2</sub>(OH)<sub>3</sub>Cl'), {'OH'})

    def test_water_without_its_bending_band(self):
        ln = self._lines('Ca2(UO2)3(SO4)2(OH)6·4H2O', RAMAN + OHS)
        self.assertEqual(len(ln), 1); self.assertIn('O–H stretching bands are stated (3250, 3480 cm⁻¹) and no H–O–H bending band', ln[0]); self.assertIn('weak in a Raman spectrum', ln[0])
        self.assertNotIn('weak in a Raman', self._lines('Ca2(UO2)3(SO4)2(OH)6·4H2O', RAMAN.replace('Raman', 'infrared') + OHS)[0])
        for bend in ('A band at 1632 cm–1 is assigned to the H–O–H bending vibration. ', 'The band at 1615 is assigned to the ν2 (δ) bending vibrations of H2O. ',
                     'Bands in the range 1700–1600 cm–1 are due to water. '):
            self.assertEqual(self._lines('Ca2(UO2)3(SO4)2(OH)6·4H2O', RAMAN + OHS + bend), [], bend)
        self.assertEqual(self._lines('Ca(UO2)(SO4)F2', RAMAN + OHS.replace('O–H stretching vibrations', 'overtones')), [])

    def test_the_text_says_the_bend_is_not_there(self):
        ln = self._lines('Na(UO2)(SO4)(OH)·2H2O', RAMAN + OHS + 'No bands were observed in the region of bending vibrations of water molecules (~1600 cm–1). ')
        self.assertEqual(len(ln), 1); self.assertIn('the text states that the spectrum shows no band near 1600 cm⁻¹', ln[0]); self.assertIn('‘No bands were observed', ln[0])
        # a range the sentence does not tie to water is not about water; another mineral's spectrum is not this one's
        self.assertNotIn('shows no band', ' '.join(self._lines('Na(UO2)(SO4)(OH)·2H2O', RAMAN + OHS + 'Bands of carbonate (in the range 1350–1560 cm–1) are absent. ')))
        self.assertEqual([x for x in self._lines('Na(UO2)(SO4)(OH)·2H2O', RAMAN + OHS + 'Unlike that of otherite, the spectrum of testite has no peaks in the region 1500–1600 cm–1 (H–O–H bending). ')
                          if 'shows no band' in x], [])

    def test_a_spectrum_that_reaches_no_hydrogen_band(self):
        ln = self._lines('Ca2(UO2)3(SO4)2(OH)6·4H2O', RAMAN)
        self.assertEqual(len(ln), 1); self.assertIn('the formula carries H2O and OH', ln[0]); self.assertIn('the 3 bands stated (832–1105 cm⁻¹) reach neither', ln[0])
        self.assertEqual(self._lines('Ca(UO2)(SO4)F2', RAMAN), [])

    def test_water_stated_where_the_formula_has_hydroxyl(self):
        water = 'A band at 1634 cm–1 is assigned to the H–O–H bending vibration of H2O. '
        ln = self._lines('Cu(UO2)(SO4)(OH)2', RAMAN + OHS + water)
        self.assertEqual(len(ln), 1); self.assertIn('a band at 1634 cm⁻¹ is stated as molecular water; the formula carries OH and no H2O', ln[0])
        for guard in (water.replace('A band', 'A band of adsorbed water'), water.replace('is assigned', 'may be assigned'),
                      'In otherite a band at 1634 cm–1 is assigned to the H–O–H bending vibration of H2O. ', water + 'Otherite, Cu(UO2)(SO4)(OH)2·2H2O, is its hydrate. '):
            self.assertEqual(self._lines('Cu(UO2)(SO4)(OH)2', RAMAN + OHS + guard), [], guard)

    def test_hydrogen_bands_and_a_formula_without_hydrogen(self):
        ln = self._lines('Ca(UO2)(SO4)F2', RAMAN + OHS)
        self.assertEqual(len(ln), 1); self.assertIn('bands at 3250, 3480 cm⁻¹ are assigned to O–H or H2O; the formula carries no hydrogen (Ca(UO2)(SO4)F2)', ln[0])
        self.assertEqual(self._lines('Ca(UO2)(SO4)F2', RAMAN + OHS.replace('are assigned', 'of adsorbed moisture are assigned')), [])

    def test_a_light_group_with_no_band(self):
        ln = self._lines('(NH4)2(UO2)(SO4)2', RAMAN)
        self.assertEqual(len(ln), 1); self.assertIn('ammonium: the formula carries NH4', ln[0]); self.assertIn('1380–1490 cm⁻¹', ln[0])
        self.assertEqual(self._lines('(NH4)2(UO2)(SO4)2', RAMAN + 'A band at 1432 cm–1 is assigned to the ν4 bending of NH4. '), [])
        self.assertEqual(self._lines('(NH4)2(UO2)(SO4)2', RAMAN + 'A weak band is seen at 1440 cm–1. '), [])        # a band where it would stand, whatever it is called
        self.assertIn('carbonate: the formula carries CO3', self._lines('Na2Ca(CO3)2', 'Raman bands occur at 155, 260, 301 and 715 cm–1. ')[0])

    def test_a_group_the_mineral_has_no_element_for(self):
        body = RAMAN + 'Further bands are at 420 and 580 cm–1. '
        ln = self._lines('Ca(UO2)2(PO4)2', body)
        self.assertEqual(len(ln), 1); self.assertIn('bands at 1010, 1105 cm⁻¹ are assigned to SO4; neither the formula (Ca(UO2)2(PO4)2) nor the rest of the text carries S', ln[0])
        for guard in ('The analysis gave SO3 0.41 wt%. ', 'Minor sulfate substitutes for phosphate. '):
            self.assertEqual(self._lines('Ca(UO2)2(PO4)2', body + guard), [], guard)
        self.assertEqual(self._lines('Ca(UO2)2(PO4)2', body.replace('are assigned to ν1 and ν3', 'of an impurity are assigned to ν1 and ν3')), [])

    def test_the_section(self):
        lines = S.check_lines(HEAD % 'Ca2(UO2)3(SO4)2(OH)6·4H2O' + RAMAN + OHS)
        self.assertEqual(lines[0], 'spectroscopy: the stated bands beside the formula'); self.assertEqual(len(lines), 2)
        self.assertEqual(S.check_lines(HEAD % 'Ca2(UO2)3(SO4)2(OH)6·4H2O' + RAMAN + OHS + 'The H–O–H bending band is at 1630 cm–1. '), [])
        self.assertEqual(S.check_lines('No formula here. ' + RAMAN + OHS), [])


class Table(unittest.TestCase):
    """A band-assignment table of a manuscript .docx: band and assignment paired by row."""

    @classmethod
    def setUpClass(cls):
        import docx
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        cls.tmp = tempfile.mkdtemp()
        d = docx.Document()
        d.add_paragraph('Testite, a new mineral. The ideal formula is Na2Ca(UO2)(CO3)3·6H2O. The Raman band at 1086 cm–1 is assigned to the ν3 antisymmetric stretch of (CO3)2–, '
                        'and the band at 829 cm–1 to the ν1 (UO2)2+ stretch. A band at 3400 cm–1 is assigned to O–H stretching vibrations.')
        rows = [('Otherite', 'Testite', 'Calculated', 'Vibrational modes'), ('3700–3300', '3700–3100', '', 'O–H stretching'), ('1410', '1413 s', '1420', '§3 (CO3)2– asym. stretching'),
                ('1090', '1086', '1080', '§1 (CO3)2– sym. stretching'), ('888', '893', '', '§2 (CO3)2– bending or §3 (UO2)2+ asym. stretching'), ('825', '829 vs', '830', '§1 (UO2)2+ sym. stretching'),
                ('750', '752, 746', '', '4 (CO3)2– bending'), ('', '1150', '', '§4 (CO3)2– bending'), ('170', '189', '', 'lattice modes')]
        t = d.add_table(rows=len(rows), cols=4)
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                par = t.cell(i, j).paragraphs[0]
                for k, piece in enumerate(c.split('§')):
                    if k:                                           # the ν as Word sets it from the Symbol font: python-docx's .text drops it
                        sym = OxmlElement('w:sym'); sym.set(qn('w:font'), 'Symbol'); sym.set(qn('w:char'), 'F06E')
                        par.add_run()._r.append(sym)
                    par.add_run(piece)
        cls.path = os.path.join(cls.tmp, 'testite.docx'); d.save(cls.path)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_rows(self):
        tb = S.table_bands(self.path, 'testite')
        got = [(b.pos, b.family, b.mode, b.shape) for b in tb]
        self.assertEqual(got, [(1413.0, 'CO3', 3, 's'), (1086.0, 'CO3', 1, ''), (893.0, None, None, ''), (829.0, 'UO2', 1, 'vs'), (752.0, 'CO3', 4, ''), (746.0, 'CO3', 4, ''),
                               (1150.0, 'CO3', 4, ''), (189.0, 'lattice', 'lattice', '')])     # its own column; the ν Word keeps as a symbol; a ν lost altogether; two things in one row left open
        self.assertEqual(len(S.table_bands(self.path, '')), 14)                                   # no name to choose by: both minerals' columns, never the calculated one
        self.assertEqual(S.table_bands(os.path.join(self.tmp, 'nothing.pdf')), [])

    def test_beside_the_text_and_against_itself(self):
        from pxrd_review import paper_extract as PE
        text = PE.text_of(self.path)
        lines = S.check_lines(text, None, None, self.path)
        self.assertEqual(lines[0], 'spectroscopy: the stated bands beside the formula and the assignment table')
        self.assertTrue(any('table: the band at 1086 cm⁻¹ is ν1 CO3 in the table and ν3 CO3 in the text' in ln for ln in lines), lines)
        self.assertTrue(any('table: ν1 of CO3 at 1086 cm⁻¹ stands below ν4 at 1150 cm⁻¹' in ln for ln in lines), lines)
        self.assertFalse(any('829' in ln for ln in lines if ln.startswith('table:')))               # where table and text agree nothing is said
        self.assertTrue(any(ln.startswith('water:') for ln in lines))                              # the table's bands count: an O–H stretch stated, no bend


class Wording(unittest.TestCase):
    """Two regexes decide from a line's WORDS whether it is a finding (proposal._FLAG for the Word comments,
    review_gui._CALC_FLAG for the GUI): every finding must match both, every information line neither."""

    def _lines(self):
        flags = []; infos = []
        texts = [
            'Applying Bartlett & Cooney (1989), the ν1 (UO2)2+ stretch at 824 cm–1 corresponds to a U–O bond length of 1.79 Å, and 1.76 for ν3 (858 cm–1).',
            'Based on Libowitzky (1999), the D···A distances for the bands at 3225, 2451, and 3536 cm−1 are 2.71, 2.83, and 2.97 Å, respectively.',
            'Using Libowitzky (1999), the distances are ~2.66 Å (3028 cm–1) and ~2.72 Å (3253 cm–1). The curve of Libowitzky (1999) gives 2.95 Å (3253 cm–1).',
            'For the sample the curve of Libowitzky (1999) gives 2.95 Å (3253 cm–1).',
            'The O–H maxima are at 3338 and 3057 cm–1. According to Libowitzky (1999) the bands at 3338 and 3057 cm–1 give O⋅⋅⋅O distances of 2.66 and 2.76 Å, respectively.',
            'The ν1 (UO2)2+ band is at 845 cm–1. Using Bartlett and Cooney (1989) the U–O bond length is 1.70 Å (845 cm–1).',
            'According to Libowitzky (1999) the band at 3610 cm–1 corresponds to an O···O distance of 3.10 Å.',
        ]
        for t in texts:
            for r in S.check_claims(t):
                if r['line']:
                    (flags if r['grade'] == 'flag' else infos).append(r['line'])
        structure = {'uranyl': [('U1', [1.75, 1.76], 1.755), ('U2', [1.79, 1.81], 1.80)], 'source': 'the .cif',
                     'hbonds': [('O3', 'O7', 2.80, 'H'), ('O4', 'O8', 2.95, 'H')]}
        infos += S.structure_lines(structure, [822.0], [856.0], [2840.0, 3400.0, 3600.0])
        infos += S.structure_lines({'uranyl': [('U1', [1.75, 1.76], 1.755)], 'source': 'the bond table the paper prints'}, [850.0, 861.0], [], [])
        # the formula beside the bands, and the assignment table: information, every template
        water = 'A band at 1634 cm–1 is assigned to the H–O–H bending vibration of H2O. '
        for formula, body in (('Ca2(UO2)3(SO4)2(OH)6·4H2O', RAMAN + OHS), ('Ca2(UO2)3(SO4)2(OH)6·4H2O', RAMAN),
                              ('Na(UO2)(SO4)(OH)·2H2O', RAMAN + OHS + 'No bands were observed in the region of bending vibrations of water molecules (~1600 cm–1). '),
                              ('Cu(UO2)(SO4)(OH)2', RAMAN + OHS + water), ('Ca(UO2)(SO4)F2', RAMAN + OHS), ('(NH4)2(UO2)(SO4)2', RAMAN),
                              ('Ca(UO2)2(PO4)2', RAMAN + 'Further bands are at 420 and 580 cm–1. ')):
            got = S.species_lines(HEAD % formula + body, name='testite')
            assert got, (formula, body)
            infos += got
        infos += S.table_lines([S.Band(1086.0, None, '', 'CO3', 1, 'row', True), S.Band(1150.0, None, '', 'CO3', 4, 'row', True)],
                               [S.Band(1086.0, 'raman', '', 'CO3', 3, 'The band at 1086 cm-1 is assigned to ν3 of CO3.', True)])
        return flags, infos

    def test_flags_and_information(self):
        from pxrd_review import proposal as PR
        from pxrd_review.gui import review_gui as G
        flags, infos = self._lines()
        self.assertGreaterEqual(len(flags), 3); self.assertGreaterEqual(len(infos), 15)
        for ln in flags:
            self.assertTrue(PR._FLAG.search(ln), ln); self.assertEqual(G._calc_kind(ln), 'calc', ln)
        for ln in infos:
            self.assertFalse(PR._FLAG.search(ln), ln); self.assertEqual(G._calc_kind(ln), 'calcinfo', ln)


class Windows(unittest.TestCase):
    def test_shipped_windows_are_aggregates(self):
        w = S.band_windows()
        self.assertIn('UO2 ν1', w); self.assertIn('OH stretch', w)
        for name, entry in w.items():
            for tech, e in entry.items():
                self.assertIn(tech, ('any', 'raman', 'ir'))
                self.assertEqual(sorted(e), ['dropped', 'n', 'p25', 'p5', 'p50', 'p75', 'p95', 'papers'], name)   # numbers only: no paper, no mineral
                self.assertGreaterEqual(e['papers'], 8)
                self.assertTrue(e['p5'] <= e['p25'] <= e['p50'] <= e['p75'] <= e['p95'], name)
        u = w['UO2 ν1']['any']
        self.assertTrue(780 < u['p5'] < u['p95'] < 900, u)

    def test_suggest(self):
        self.assertEqual(S.suggest(840)[0][0], 'UO2 ν1')
        self.assertEqual([n for n, _w in S.suggest(1630, 'ir')], ['H2O bend'])
        self.assertEqual([n for n, _w in S.suggest(3400)], ['OH stretch'])
        self.assertEqual(S.suggest(2300), [])


class Structure(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.cif = _write(cls.tmp, 'uranyl.cif', URANYL)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_uranyl_and_hydrogen_bond(self):
        d = S.structure_distances(self.cif)
        self.assertEqual([(l, [round(x, 2) for x in ds]) for l, ds, _m in d['uranyl']], [('U1', [1.77, 1.77])])
        self.assertEqual([(D, A, round(x, 2)) for D, A, x, _h in d['hbonds']], [('O3', 'O7', 2.80)])

    def test_no_uranyl_without_a_linear_pair(self):
        bent = URANYL.replace('O2 O 0.5 0.5 0.323', 'O2 O 0.5 0.677 0.5').replace('O5 O 0.5 0.735 0.5', 'O5 O 0.5 0.5 0.265')
        self.assertEqual(S.structure_distances(_write(self.tmp, 'bent.cif', bent))['uranyl'], [])

    def test_section(self):
        t = ('The Raman spectrum was recorded with a 532 nm laser. The ν1 (UO2)2+ symmetric stretching vibration is present as a very strong band at 840 cm–1. '
             'Bartlett and Cooney (1989) provided an empirical relationship, which gives 1.77 Å (840 cm–1). Bands at 3350 and 3200 cm–1 are assigned to ν O–H stretching vibrations.')
        L = S.check_lines(t, cif=self.cif)
        self.assertTrue(L[0].startswith('spectroscopy: 1 distance credited to Bartlett & Cooney (1989) — 1 reproduced'), L[0])
        self.assertTrue(any(ln.startswith('uranyl: mean U–O 1.770 Å in the .cif; ν1 840 cm⁻¹ gives 1.77') for ln in L), L)
        self.assertTrue(any(ln.startswith('hydrogen bonds: 1 in the .cif, D···A 2.80 Å') for ln in L), L)
        self.assertEqual(S.check_lines('The crystals are yellow. No spectrum was recorded.', cif=self.cif), [])
        self.assertEqual(S.check_lines('A band at 1005 cm–1 is assigned to ν1 (SO4)2–.'), [])   # nothing credited, no structure: no section


if __name__ == '__main__':
    unittest.main()
