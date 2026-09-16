"""Manuscript mode of the review GUI — the /api/ms/* routes, through Flask's test client.

    python3 -m unittest tests.test_gui_ms -v

Builds a small manuscript + companion in a temp folder; no browser, no corpus."""
import json, os, re, shutil, tempfile, unittest, zipfile

from docx import Document


def _docx(path, paras):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    d.save(path)
    return path


class PdfManuscriptPages(unittest.TestCase):
    """'? look' on a calculation finding of a .pdf manuscript that names no table: the finding gets the
    page its own words are on (1-based) and the words to highlight; a line whose words are nowhere in
    the paper, and the readers line, are left without a page."""

    @classmethod
    def setUpClass(cls):
        import pymupdf
        cls.tmp = tempfile.mkdtemp(prefix='ms_pages_')
        cls.path = os.path.join(cls.tmp, 'paper.pdf')
        doc = pymupdf.open()
        doc.new_page().insert_text((72, 72), 'Testite, IMA 2024-001, is a new mineral from Test Hill.')
        doc.new_page().insert_text((72, 72), 'The Gladstone-Dale compatibility index is 0.011 (superior). a = 16.2184(3) A.')
        doc.save(cls.path); doc.close()

    @classmethod
    def tearDownClass(cls):
        from pxrd_review.gui import review_gui as G
        G.PW.shutdown()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_pages(self):
        from pxrd_review.gui import review_gui as G
        fs = [{'kind': 'calcinfo', 'label': 'Gladstone–Dale', 'msg': 'Gladstone–Dale: n 1.55, K_C 0.22', 'para': None, 'page': None, 'find': None},
              {'kind': 'calcinfo', 'label': 'cell', 'msg': 'cell: a=16.2184, b=16.2184, c=10.1357 (powder)', 'para': None, 'page': None, 'find': None},
              {'kind': 'calcinfo', 'label': 'cell', 'msg': "the .cif's cell (a=6.9035, b=7.5844) differs", 'para': None, 'page': None, 'find': None},
              {'kind': 'calcinfo', 'label': 'name', 'msg': 'name: testite is not on Mindat', 'para': None, 'page': None, 'find': None},
              {'kind': 'calcinfo', 'label': 'readers', 'msg': 'readers: table ✓', 'para': None, 'page': None, 'find': None},
              {'kind': 'calcinfo', 'label': 'composition', 'msg': 'composition: 2 constituents', 'para': None, 'page': 1, 'find': None}]
        G._ms_pdf_pages(self.path, fs)
        self.assertEqual((fs[0]['page'], fs[0]['find']), (2, 'compatibility|Gladstone'))
        self.assertEqual((fs[1]['page'], fs[1]['find']), (2, '16.2184|16.218'))
        self.assertIsNone(fs[2]['page'])                               # the .cif's a is not in the paper
        self.assertEqual((fs[3]['page'], fs[3]['find']), (1, 'IMA'))
        self.assertIsNone(fs[4]['page'])                               # the readers line names no place
        self.assertEqual(fs[5]['page'], 1)                             # a page already set is kept


class CalculationFindingAnchors(unittest.TestCase):
    """'? look' on a calculation finding of a .docx manuscript: the finding is anchored to the cell it
    names, in the table its section was read from, by the paragraph numbering the docx view uses."""

    @classmethod
    def setUpClass(cls):
        from docx import Document
        cls.tmp = tempfile.mkdtemp(prefix='ms_anchor_')
        cls.path = os.path.join(cls.tmp, 'paper.docx')
        d = Document()
        d.add_paragraph('Testite, a new mineral. The Gladstone-Dale compatibility index is 0.011 (superior).')
        d.add_paragraph('Unit cell a = 16.2184(3), c = 10.1357(2) Å, V = 2308.9 Å3.')
        d.add_paragraph('Table 1. Analytical data (in wt%) for testite.')
        t = d.add_table(rows=3, cols=2)
        for r, (a, b) in enumerate((('Const.', 'Mean'), ('SiO2', '40.1'), ('Na2O', '10.0'))):
            t.cell(r, 0).text = a; t.cell(r, 1).text = b
        d.add_paragraph('Table 3. Atom coordinates for testite.')
        t = d.add_table(rows=3, cols=2)
        for r, (a, b) in enumerate((('Atom', 'x'), ('Na2', '0.1234'), ('O8', '0.5678'))):
            t.cell(r, 0).text = a; t.cell(r, 1).text = b
        d.add_paragraph('Table 5. Bond valence analysis for testite.')
        t = d.add_table(rows=3, cols=3)
        for r, row in enumerate((('', 'Na1', 'Na2'), ('O1', '0.20', '0.59'), ('O8', '', '0.13'))):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
        d.save(cls.path)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def _idx(self, text, nth=0):
        from pxrd_review import refs_check as RC
        _doc, paras = RC.load_docx(self.path)
        return [p.idx for p in paras if p.text.strip() == text][nth]

    def test_anchors(self):
        from pxrd_review.gui import review_gui as G
        fs = [{'kind': 'calcinfo', 'label': 'composition', 'msg': 'composition: 2 constituents re-reduced', 'para': None, 'find': None},
              {'kind': 'calc', 'label': 'composition', 'msg': 'Si: paper 2.99, from the paper\'s own wt% 3.20', 'para': None, 'find': 'SiO2'},
              {'kind': 'calcinfo', 'label': 'bond valence', 'msg': 'bond-valence table 1: 5 cells compared, 0 disagree', 'para': None, 'find': None},
              {'kind': 'calc', 'label': 'bond valence', 'msg': 'table 1: O8–Na2 is blank but the .cif has that bond', 'para': None, 'find': 'O8|Na2'},
              {'kind': 'calc', 'label': 'bond valence', 'msg': 'table 1: Σ for Na1 1.02 vs 1.20', 'para': None, 'find': 'Na1'},
              {'kind': 'calcinfo', 'label': 'Gladstone–Dale', 'msg': 'Gladstone–Dale: n 1.55, K_C 0.22', 'para': None, 'find': None},
              {'kind': 'calcinfo', 'label': 'cell', 'msg': 'cell: a=16.2184, b=16.2184, c=10.1357 (powder) — V from the axes agrees', 'para': None, 'find': None},
              {'kind': 'calcinfo', 'label': 'readers', 'msg': 'readers: table ✓', 'para': None, 'find': None}]
        G._ms_docx_anchors(self.path, fs)
        self.assertEqual(fs[0]['para'], self._idx('Table 1. Analytical data (in wt%) for testite.'))     # a section head: the table's caption
        self.assertEqual(fs[1]['para'], self._idx('SiO2'))
        self.assertEqual(fs[2]['para'], self._idx('Table 5. Bond valence analysis for testite.'))
        self.assertEqual(fs[3]['para'], self._idx('O8', nth=1))       # the O8 ROW of the bond-valence table, not the coordinates table's O8
        self.assertEqual(fs[4]['para'], self._idx('Na1'))
        self.assertEqual(fs[5]['para'], 0)                             # the compatibility sentence
        self.assertEqual(fs[6]['para'], 1)                             # the paragraph that prints the a axis
        self.assertIsNone(fs[7]['para'])                               # the readers line names no place

    def test_a_clean_table_s_count_line_is_not_a_flag(self):
        from pxrd_review.gui import review_gui as G
        self.assertIsNone(G._CALC_FLAG.search('bond-valence table 1: 35 cells compared, 0 disagree (computed with Gagné & Hawthorne 2015)'))
        self.assertIsNotNone(G._CALC_FLAG.search('bond-valence table 1: 35 cells compared, 3 disagree (computed with Gagné & Hawthorne 2015)'))
        self.assertIsNotNone(G._CALC_FLAG.search('bond-valence table 1: 35 cells compared, 10 disagree'))
        self.assertEqual(G._calc_kind('bond-valence table 1: 35 cells compared, 0 disagree (computed with Gagné & Hawthorne 2015)'), 'calcinfo')
        # a blank cell: the CHECKER's wording says whether it is a difference (bv_check.BLANK_INFO); the GUI re-reads no number
        self.assertEqual(G._calc_kind('table 1: O8–Na2 is blank — the .cif has that bond at 0.03 vu, under the cutoff most tables print (not a difference)'), 'calcinfo')
        self.assertEqual(G._calc_kind('table 1: O8–Na2 is blank but the .cif has that bond (0.21 vu)'), 'calc')          # a bond a table prints: a missing cell
        self.assertEqual(G._calc_kind('table 1: Σ for As2/S2 5.48 vs 5.31 from the .cif (parameters: Gagné & Hawthorne 2015)'), 'calc')
        self.assertEqual(G._calc_kind('bond valence: the table vs the .cif — agrees best with X', head=True), 'calcinfo')
        self.assertFalse(hasattr(G, '_BLANK_CELL'))                                                                       # the severity rule lives in bv_check, not here

    def _build(self, name, build):
        from docx import Document
        d = Document(); build(d); p = os.path.join(self.tmp, name); d.save(p)
        return p

    @staticmethod
    def _table(d, rows):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
        return t

    @staticmethod
    def _paras(path):
        from pxrd_review import refs_check as RC
        _doc, paras = RC.load_docx(path)
        return {p.idx: p.text.strip() for p in paras}

    def test_superscript_footnote_mark_on_a_constituent_cell(self):
        """The reader (paper_extract._docx_cell_text) drops a one- or two-character superscript, so its
        term is 'SiO2' while load_docx reads 'SiO2a' / 'H2O[1]': the anchor must normalise the cell as
        the reader did, and a term missing from the analytical table must NOT fall through to a foreign
        table (the Gladstone–Dale constants list SiO2 too)."""
        from pxrd_review.gui import review_gui as G
        def build(d):
            d.add_paragraph('Table 1. Analytical data (wt%) for testite.')
            t = self._table(d, [['Const.', 'Mean'], ['SiO2', '40.1'], ['FeO', '10.0'], ['H2O', '5.0']])
            for cell, mark in ((t.cell(1, 0), 'a'), (t.cell(3, 0), '1')):
                r = cell.paragraphs[0].add_run(mark); r.font.superscript = True
            d.add_paragraph('Table 6. Gladstone-Dale constants used.')
            self._table(d, [['Oxide', 'k'], ['SiO2', '0.207'], ['FeO', '0.187'], ['H2O', '0.34'], ['MgO', '0.20']])
        p = self._build('sup.docx', build); m = self._paras(p)
        fs = [{'kind': 'calc', 'label': 'composition', 'msg': 'Si: paper 2.99, from the wt% 3.20', 'para': None, 'find': 'SiO2'},
              {'kind': 'calc', 'label': 'composition', 'msg': 'H: paper 2.0, from the wt% 2.2', 'para': None, 'find': 'H2O'},
              {'kind': 'calc', 'label': 'composition', 'msg': 'Mg: paper 0.1, from the wt% 0.0', 'para': None, 'find': 'MgO'}]
        G._ms_docx_anchors(p, fs)
        self.assertEqual(m[fs[0]['para']], 'SiO2a')                    # the analytical table's cell, not the Gladstone–Dale table's 'SiO2'
        self.assertEqual(m[fs[1]['para']], 'H2O[1]')
        self.assertEqual(m[fs[2]['para']], 'Table 1. Analytical data (wt%) for testite.')   # MgO is only in the foreign table: the caption, never that table

    def test_competing_captions_prefer_the_grid_whose_cells_are_the_labels(self):
        """'Selected bond distances (Å) and bond valences (vu)' comes first and names every site in its
        'Na1–O8' cells; the bond-valence grid after it has the sites as whole cells. 'O8–Na2 0.13 vs
        0.25' is the grid's O8 row, not the distance table's Na1–O8 (a different bond)."""
        from pxrd_review.gui import review_gui as G
        def build(d):
            d.add_paragraph('Table 4. Selected bond distances (Å) and bond valences (vu) for testite.')
            self._table(d, [['Bond', 'd', 'vu'], ['Na1–O1', '2.40', '0.20'], ['Na1–O8', '2.50', '0.15'], ['Na2–O8', '2.30', '0.13']])
            d.add_paragraph('Table 5. Bond-valence sums for testite.')
            self._table(d, [['', 'Na1', 'Na2', 'Σ'], ['O1', '0.20', '0.59', '1.9'], ['O8', '0.15', '0.13', '2.0'], ['Σ', '1.02', '1.10', '']])
        p = self._build('captions.docx', build); m = self._paras(p)
        fs = [{'kind': 'calcinfo', 'label': 'bond-valence table 1 vs the .cif', 'msg': 'bond-valence table 1: 4 cells compared, 1 disagree', 'para': None, 'find': None},
              {'kind': 'calc', 'label': 'bond-valence table 1 vs the .cif', 'msg': 'table 1: Σ for Na1 1.02 vs 1.20', 'para': None, 'find': 'Na1'},
              {'kind': 'calc', 'label': 'bond-valence table 1 vs the .cif', 'msg': 'table 1: O8–Na2 0.13 vs 0.25', 'para': None, 'find': 'O8|Na2'}]
        G._ms_docx_anchors(p, fs)
        self.assertEqual(m[fs[0]['para']], 'Table 5. Bond-valence sums for testite.')
        self.assertEqual(m[fs[1]['para']], 'Na1'); self.assertEqual(m[fs[2]['para']], 'O8')
        self.assertGreater(fs[2]['para'], fs[0]['para'])               # in the grid (after its caption), not in the distance table

    def test_table_n_picks_the_nth_grid_of_two_minerals(self):
        from pxrd_review.gui import review_gui as G
        def build(d):
            d.add_paragraph('Table 5. Bond-valence analysis for testite.')
            self._table(d, [['', 'Na1', 'Na2'], ['O1', '0.20', '0.59'], ['O8', '0.15', '0.13']])
            d.add_paragraph('Table 6. Bond-valence analysis for otherite.')
            self._table(d, [['', 'Na1', 'Na2'], ['O1', '0.21', '0.58'], ['O8', '0.16', '0.12']])
        p = self._build('two.docx', build); m = self._paras(p)
        fs = [{'kind': 'calc', 'label': 'bond valence', 'msg': 'table 2: O8–Na2 0.12 vs 0.25', 'para': None, 'find': 'O8|Na2'},
              {'kind': 'calc', 'label': 'bond valence', 'msg': 'table 1: O8–Na2 0.13 vs 0.25', 'para': None, 'find': 'O8|Na2'}]
        G._ms_docx_anchors(p, fs)
        self.assertEqual([m[f['para']] for f in fs], ['O8', 'O8'])
        self.assertGreater(fs[0]['para'], fs[1]['para'])               # the second grid's O8 row for 'table 2'

    def test_prose_anchors(self):
        """The species line is 'name: …' (paper_extract.verify); an integer axis ('a=16', %g) still anchors."""
        from pxrd_review.gui import review_gui as G
        p = _docx(os.path.join(self.tmp, 'prose.docx'), ['Testite is listed by Mindat as a valid species.', 'Unit cell a = 16, c = 10.1357 Å.'])
        fs = [{'kind': 'calcinfo', 'label': 'name', 'msg': 'name: Mindat lists testite [unverified]', 'para': None, 'find': None},
              {'kind': 'calcinfo', 'label': 'cell', 'msg': 'cell: a=16, b=16, c=10.1357 (powder) — nothing printed', 'para': None, 'find': None}]
        G._ms_docx_anchors(p, fs)
        self.assertEqual([f['para'] for f in fs], [0, 1])

    def test_anchors_through_a_content_control_around_a_row(self):
        """A w:sdt around a table row: load_docx numbers its paragraphs and the docx view must too, or
        every finding after the row lands early."""
        from docx.oxml.ns import qn
        from lxml import etree
        from pxrd_review.gui import review_gui as G
        def build(d):
            d.add_paragraph('Table 5. Bond valence analysis for testite.')
            t = self._table(d, [['', 'Na1', 'Na2'], ['O1', '0.20', '0.59'], ['O8', '0.15', '0.13']])
            tbl = t._tbl; tr = tbl.findall(qn('w:tr'))[1]
            sdt = etree.Element(qn('w:sdt')); sc = etree.SubElement(sdt, qn('w:sdtContent')); tbl.replace(tr, sdt); sc.append(tr)
            d.add_paragraph('Table 1. Analytical data (wt%) for testite.')
            self._table(d, [['Const.', 'Mean'], ['SiO2', '40.1']])
        p = self._build('sdt.docx', build); m = self._paras(p)
        fs = [{'kind': 'calc', 'label': 'bond valence', 'msg': 'table 1: O8–Na2 0.13 vs 0.25', 'para': None, 'find': 'O8|Na2'},
              {'kind': 'calc', 'label': 'composition', 'msg': 'Si: paper 2.99, from the wt% 3.20', 'para': None, 'find': 'SiO2'}]
        G._ms_docx_anchors(p, fs)
        self.assertEqual([m[f['para']] for f in fs], ['O8', 'SiO2'])
        html_ = G._docx_html(p)
        rendered = {int(a): re.sub('<[^>]+>', '', b).replace('&nbsp;', '').strip() for a, b in re.findall(r'<p data-p="(\d+)">(.*?)</p>', html_)}
        self.assertEqual(rendered[fs[0]['para']], 'O8'); self.assertEqual(rendered[fs[1]['para']], 'SiO2')


class DocxNumbering(unittest.TestCase):
    """_docx_html's <p data-p> numbering and refs_check.load_docx's paragraph indexes are the same
    sequence — the documented invariant every docx anchor rests on — on the shapes that once drifted
    (a content control around a row or a cell, a legacy VML text box) and the ones that did not."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix='ms_number_')

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    @staticmethod
    def _numbering(path):
        from pxrd_review import refs_check as RC
        from pxrd_review.gui import review_gui as G
        _d, paras = RC.load_docx(path)
        ld = [(p.idx, p.text.strip()[:25]) for p in paras if p.elem is not None]
        h = G._docx_html(path)
        hp = [(int(a), re.sub('<[^>]+>', '', b).replace('&nbsp;', '').strip()[:25]) for a, b in re.findall(r'<p data-p="(\d+)">(.*?)</p>', h)]
        return ld, hp

    def _shapes(self):
        from docx.oxml.ns import qn
        from lxml import etree
        V = 'urn:schemas-microsoft-com:vml'; MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
        def table(d, rows):
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    t.cell(r, c).text = v
            return t
        def plain(d):
            d.add_paragraph('cap'); table(d, [['A', ''], ['', 'B']]); d.add_paragraph('after')
        def nested(d):
            d.add_paragraph('cap'); t = table(d, [['']]); inner = t.cell(0, 0).add_table(rows=2, cols=1)
            inner.cell(0, 0).text = 'in0'; inner.cell(1, 0).text = 'in1'; d.add_paragraph('after')
        def merged(d):
            d.add_paragraph('cap'); t = table(d, [['r%dc%d' % (r, c) for c in range(3)] for r in range(3)])
            t.cell(0, 0).merge(t.cell(0, 2)); t.cell(1, 1).merge(t.cell(2, 1)); d.add_paragraph('after')
        def sdt_row(d):
            d.add_paragraph('cap'); t = table(d, [['h0', 'h1'], ['v0', 'v1']]); tbl = t._tbl; tr = tbl.findall(qn('w:tr'))[1]
            sdt = etree.Element(qn('w:sdt')); sc = etree.SubElement(sdt, qn('w:sdtContent')); tbl.replace(tr, sdt); sc.append(tr)
            d.add_paragraph('after')
        def sdt_cell(d):
            d.add_paragraph('cap'); t = table(d, [['c0', 'c1']]); tr = t._tbl.find(qn('w:tr')); tc = tr.findall(qn('w:tc'))[1]
            sdt = etree.Element(qn('w:sdt')); sc = etree.SubElement(sdt, qn('w:sdtContent')); tr.replace(tc, sdt); sc.append(tc)
            d.add_paragraph('after')
        def sdt_table(d):
            d.add_paragraph('cap'); t = table(d, [['c']]); body = d.element.body
            sdt = etree.Element(qn('w:sdt')); sc = etree.SubElement(sdt, qn('w:sdtContent')); body.replace(t._tbl, sdt); sc.append(t._tbl)
            d.add_paragraph('after')
        def box_p(host):
            txc = etree.SubElement(host, qn('w:txbxContent')); bp = etree.SubElement(txc, qn('w:p'))
            bt = etree.SubElement(etree.SubElement(bp, qn('w:r')), qn('w:t')); bt.text = 'IN THE BOX'
        def vml_box(d):
            r = d.add_paragraph('before box').add_run()
            pict = etree.SubElement(r._r, qn('w:pict')); shape = etree.SubElement(pict, '{%s}shape' % V)
            box_p(etree.SubElement(shape, '{%s}textbox' % V)); d.add_paragraph('after box'); table(d, [['cell']])
        def mc_box(d):
            r = d.add_paragraph('before box').add_run()
            ac = etree.SubElement(r._r, '{%s}AlternateContent' % MC)
            box_p(etree.SubElement(ac, '{%s}Choice' % MC)); box_p(etree.SubElement(ac, '{%s}Fallback' % MC))
            d.add_paragraph('after box'); table(d, [['cell']])
        return [('plain table', plain, 6), ('nested table', nested, 6), ('merged cells', merged, 12), ('sdt around a row', sdt_row, 6),
                ('sdt around a cell', sdt_cell, 4), ('sdt around a table', sdt_table, 3), ('legacy VML text box', vml_box, 3), ('mc:AlternateContent box', mc_box, 3)]

    def test_same_numbering_on_every_shape(self):
        from docx import Document
        for name, build, n in self._shapes():
            with self.subTest(shape=name):
                d = Document(); build(d); p = os.path.join(self.tmp, name.replace(' ', '_').replace(':', '') + '.docx'); d.save(p)
                ld, hp = self._numbering(p)
                self.assertEqual(ld, hp)
                self.assertEqual(len(ld), n)                   # and the count is what the shape holds: a text box's paragraph is numbered by neither side


class ManuscriptMode(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from pxrd_review.gui import review_gui as G
        cls.G = G
        G._ALLOWED_HOSTS = set()                      # pre-launch: the request gate lets the test client in
        cls.tmp = tempfile.mkdtemp(prefix='msgui_')
        _docx(os.path.join(cls.tmp, 'paper.docx'), [
            'Quartz is common (Smith 2019; Jones 2020). Kampf (1977) reported it.',
            'References',
            'Jones, A. (2020) A title. Journal 1, 1–2.',
            'Smith, J. (2019) A title. Journal 2, 3–4.',
            'Uncited, U. (2001) Never cited. Journal 3, 5–6.'])
        _docx(os.path.join(cls.tmp, 'table.docx'), ['Table 1. Data from Uncited (2001).'])
        with open(os.path.join(cls.tmp, 'x.cif'), 'w') as f:
            f.write('data_x\n')
        cls.c = G.app.test_client()
        assert G.ms_set_folder(cls.tmp)
        G.MS['athread'].join(30)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_state_and_doc(self):
        # order-independent: reset the companions another test may have set, analyse synchronously
        self.c.post('/api/ms/triage/paper', json={'findings': {}, 'companions': []})
        self.G.ms_analysis('paper'); self.G.ms_analysis('table')
        st = json.loads(self.c.get('/api/ms/state').data)
        self.assertEqual(sorted(f['key'] for f in st['files']), ['paper', 'table'])
        self.assertEqual(st['cifs'], ['x'])
        paper = next(f for f in st['files'] if f['key'] == 'paper')
        self.assertEqual({k: v for k, v in paper['summary'].items() if k in ('orphan', 'uncited', 'pair', 'form')}, {'orphan': 1, 'uncited': 1, 'pair': 0, 'form': 0})
        d = json.loads(self.c.get('/api/ms/doc/paper').data)
        kinds = sorted(f['kind'] for f in d['analysis']['findings'])
        self.assertEqual(kinds, ['orphan', 'uncited'])
        self.assertEqual(d['others'], ['table'])
        self.assertEqual(self.c.get('/api/ms/doc/nope').status_code, 404)

    def test_companion_triage_run_and_render(self):
        c = self.c
        # a companion file resolves the uncited entry
        r = json.loads(c.post('/api/ms/triage/paper', json={'findings': {}, 'companions': ['table']}).data)
        self.assertTrue(r['ok'] and r['reanalyse'])
        d = json.loads(c.get('/api/ms/doc/paper').data)
        self.assertEqual([f['kind'] for f in d['analysis']['findings']], ['orphan'])
        orphan = d['analysis']['findings'][0]
        self.assertEqual(orphan['fkey'], 'orphan:kampf:1977')
        # dismiss it with a note -> the run writes no comment for it
        c.post('/api/ms/triage/paper', json={'findings': {orphan['fkey']: {'verdict': 'dismiss', 'label': 'x'}},
                                             'companions': ['table'], 'reviewed': True})
        r = json.loads(c.post('/api/ms/run/paper').data)
        self.assertTrue(r['ok'])
        out = os.path.join(self.tmp, 'review_out', 'paper_refs.docx')
        self.assertTrue(os.path.exists(out))
        self.assertNotIn('word/comments.xml', zipfile.ZipFile(out).namelist())   # nothing left to write
        # un-dismiss with a note -> one comment carrying the note
        c.post('/api/ms/triage/paper', json={'findings': {orphan['fkey']: {'verdict': 'confirm', 'note': 'ask author', 'label': 'x'}},
                                             'companions': ['table']})
        self.assertTrue(json.loads(c.post('/api/ms/run/paper').data)['ok'])
        cx = zipfile.ZipFile(out).read('word/comments.xml').decode()
        self.assertEqual(len(re.findall(r'<w:comment ', cx)), 1)
        self.assertIn('reviewer: ask author', cx)
        # rendered copy: paragraph anchors, the tool's comment chip, the finding's paragraph index resolves
        h = json.loads(c.get('/api/ms/docx/paper.html?which=annotated').data)
        self.assertEqual(h['which'], 'annotated')
        self.assertIn('data-p="%d"' % orphan['para'], h['html'])
        self.assertIn('class="cmt"', h['html'])
        src = json.loads(c.get('/api/ms/docx/paper.html?which=source').data)
        self.assertNotIn('class="cmt"', src['html'])
        self.assertTrue(json.loads(c.get('/api/ms/report/paper').data)['text'].startswith('Reference check'))
        self.assertTrue(json.loads(c.post('/api/ms/export').data)['ok'])
        rep = open(os.path.join(self.tmp, 'review_out', 'ms_triage_report.txt'), encoding='utf-8').read()
        self.assertIn('paper.docx', rep)
        # the source docx is untouched
        self.assertNotIn('word/comments.xml', zipfile.ZipFile(os.path.join(self.tmp, 'paper.docx')).namelist())

    def test_default_companions_by_name(self):
        from pxrd_review.gui import review_gui as G
        old = dict(G.MS['triage'])
        try:
            G.MS['triage'].pop('paper', None)                       # nothing saved yet
            self.assertEqual(G._ms_companions('paper'), ['table'])  # 'table.docx' is pre-ticked by name
            self.assertEqual(G._ms_companions('table'), [])         # a table file has no companions of its own
            G.MS['triage']['paper'] = {'companions': []}            # the reviewer un-ticked it: respected
            self.assertEqual(G._ms_companions('paper'), [])
        finally:
            G.MS['triage'] = old

    def test_folder_validation(self):
        r = self.c.post('/api/ms/folder', json={'folder': '/nonexistent/x'})
        self.assertEqual(r.status_code, 400)
        empty = tempfile.mkdtemp(prefix='msgui_empty_')
        try:
            r = self.c.post('/api/ms/folder', json={'folder': empty})
            self.assertEqual(r.status_code, 400)
        finally:
            shutil.rmtree(empty, ignore_errors=True)


if __name__ == '__main__':
    unittest.main()
