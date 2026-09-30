"""Unit tests for the ideal-formula, Gladstone–Dale grid, basis-free ratio and density-statement readers.

    python3 -m unittest tests.test_ideal_gd -v
"""
import os, shutil, tempfile, unittest

from pxrd_review import paper_extract as PE, epma as EP

TEXT = ("The empirical formula, calculated on the basis of Mg = 1, is Mg1.00As0.91O7.75H9. "
        "The ideal formula is Mg(AsO3OH)·4H2O, which requires MgO 17.06, As2O5 48.63, H2O 34.41, total 100.00 wt%. "
        "Biaxial (–), α = 1.502(2), β = 1.512(2), γ = 1.518(2). "
        "Density (calc.) = 2.257 g·cm–3 for the ideal formula. "
        "The Gladstone-Dale compatibility, 1 – (Kp/Kc), is –0.043 (good) for the empirical formula and 0.031 (excellent) for the ideal formula.")


class Readers(unittest.TestCase):
    def test_density_statement_forms(self):
        for t, key, v in (('Density (calc.) = 2.257 g·cm–3 for the ideal formula.', 'D_calc', 2.257),
                          ('Density (for above formula)\t2.261 g cm–3', 'D_calc', 2.261),
                          ('Density (meas.) = 2.25(2) g/cm3 by flotation.', 'D_meas', 2.25)):
            self.assertEqual(PE.optics(t)[key], v, t)

    def test_every_stated_index(self):
        g = PE.gd_statement(TEXT)
        self.assertEqual(g['ci'], -0.043); self.assertEqual(g['category'], 'good')
        self.assertEqual([(x['ci'], x['category'], x['for']) for x in g['all']], [(-0.043, 'good', 'empirical'), (0.031, 'excellent', 'ideal')])

    def test_an_ideal_formula_with_a_comma_in_its_brackets_is_read_whole(self):
        f, counts = PE.ideal_formula('The ideal formula is Ba3(Mg,Fe)Si2O8, which requires BaO 60.')
        self.assertEqual(f, 'Ba3(Mg,Fe)Si2O8'); self.assertIn('Ba', counts)
        self.assertIsNone(PE.ideal_formula('The ideal formula is Ba3(Mg'))                     # cut short: not read

    def test_ideal_formula_and_its_wt(self):
        f, counts = PE.ideal_formula(TEXT)
        self.assertEqual(f, 'Mg(AsO3OH)·4H2O'); self.assertEqual(counts, {'Mg': 1.0, 'As': 1.0, 'O': 8.0, 'H': 9.0})
        wt = PE.ideal_wt(counts)
        self.assertAlmostEqual(wt['MgO'], 17.06, places=1); self.assertAlmostEqual(wt['As2O5'], 48.63, places=1); self.assertAlmostEqual(wt['H2O'], 34.31, places=1)

    def test_ideal_wt_check_finds_the_typo_and_the_total(self):
        L = PE.ideal_wt_check(TEXT, {'epma': {}})
        self.assertTrue(L and L[0].startswith('ideal formula: Mg(AsO3OH)·4H2O gives'), L)
        self.assertTrue(any('H2O 34.41 in the text vs 34.31' in x for x in L), L)
        self.assertTrue(any('add to 100.10 vs the total printed, 100.00' in x for x in L), L)
        self.assertFalse(any('MgO' in x and ' vs ' in x for x in L), L)                       # 17.06 is right
        # the table's Ideal column too
        ex = {'epma': {'head_cells': [('Constituent', 70.0), ('Mean', 120.0), ('Ideal', 400.0)],
                       'rows_all': [{'constituent': 'H2O', 'all': [0.0, 34.61], 'xs': [120.0, 400.0]}]}}
        L = PE.ideal_wt_check(TEXT, ex)
        self.assertTrue(any("H2O 34.61 in the table's Ideal column vs 34.31" in x for x in L), L)
        self.assertEqual(PE.ideal_wt_check('No ideal formula here. Density 3.2.', {'epma': {}}), [])

    def test_basis_free_ratios(self):
        ex = {'epma': {'rows': [{'constituent': 'MgO', 'mean': 21.82, 'sd': 0.65}, {'constituent': 'As2O5', 'mean': 56.74, 'sd': 1.73}]}}
        L = PE.basis_free_ratios(ex, {'Mg': 1.0, 'As': 1.0, 'O': 8.0, 'H': 9.0})
        self.assertEqual(len(L), 2, L); self.assertIn('Mg:As = 1.097', L[1]); self.assertIn('information', L[1])      # 2σ on the s.d. alone
        L = PE.basis_free_ratios(ex, {'Mg': 1.0, 'As': 1.0, 'O': 8.0, 'H': 9.0}, n_points=11)
        self.assertIn('whatever the basis', L[1]); self.assertIn('n = 11', L[1])                                        # 7σ on the s.d. of the means
        self.assertEqual(PE.basis_free_ratios(ex, {'Mg': 1.1, 'As': 1.0, 'O': 8.0}), [])                              # within scatter: nothing

    def test_gd_grid(self):
        ex = {'optics': PE.optics(TEXT), 'epma': {'rows': [{'constituent': 'MgO', 'mean': 17.66}, {'constituent': 'As2O5', 'mean': 45.92}, {'constituent': 'H2O', 'mean': 35.99}], 'total': 99.57}, '_text': TEXT}
        L = PE.gd_grid(ex, {'counts': {'Mg': 1.0, 'As': 0.91, 'O': 7.75, 'H': 9.0}}, PE.gd_statement(TEXT))
        self.assertTrue(L and L[0].startswith('Gladstone–Dale, the statements one by one'), L)
        self.assertEqual(len(L), 3, L)
        self.assertTrue(any('stated +0.031 (excellent) for the ideal formula' in x for x in L), L)
        self.assertTrue(any("not of the ideal formula it is stated for" in x or 'reproduced by no combination' in x or 'slack' in x for x in L), L)


class BasisSensitivity(unittest.TestCase):
    def test_alternatives(self):
        tmp = tempfile.mkdtemp(prefix='ep_')
        try:
            csv = os.path.join(tmp, 'p.csv')
            with open(csv, 'w') as f:
                f.write('MgO,CoO,As2O5\n21.82,0.54,56.74\n22.10,0.50,57.20\n')
            ds, red, table, text = EP.prepare(csv, 'Mg+Co=1', adds=['H2O=structure:4.5'])
            alts = EP.basis_sensitivity(ds, red, adds=[EP._parse_add('H2O=structure:4.5')])
            labels = [a[0] for a in alts]
            self.assertEqual(labels[0], 'Mg+Co = 1.0 apfu')
            self.assertIn('As = 1 apfu', labels); self.assertIn('2 cations apfu', labels)
            self.assertTrue(alts[0][1].startswith('Mg0.99Co0.01As0.9'), alts[0])
            as1 = dict((a[0], a[1]) for a in alts)['As = 1 apfu']
            self.assertTrue(as1.startswith('Mg1.1'), as1)                                    # the Mg:As ratio is the same on every basis
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class ChargeBalance(unittest.TestCase):
    EX = {'name': None, 'epma': {'rows': [{'constituent': 'MgO', 'mean': 17.7}, {'constituent': 'As2O5', 'mean': 45.9}, {'constituent': 'H2O', 'mean': 36.0}]}}

    def test_balanced_formula_is_silent(self):
        t = 'The empirical formula, calculated on the basis of Mg = 1, is Mg1.00As1.00O8.00H9.00.'
        self.assertEqual(PE.charge_balance_check(t, self.EX), [])

    def test_an_o_count_one_short_is_a_finding(self):
        t = 'The empirical formula, calculated on the basis of Mg = 1, is Mg1.00As1.00O7.00H9.00.'
        L = PE.charge_balance_check(t, self.EX)
        self.assertEqual(len(L), 2, L)
        self.assertIn('O7.00 printed vs O8.00 from the charges', L[1]); self.assertIn('As+5 from the analysis', L[1])

    def test_a_quarter_o_is_information(self):
        t = 'The empirical formula, calculated on the basis of Mg = 1, is Mg1.00As1.00O7.70H9.00.'
        L = PE.charge_balance_check(t, self.EX)
        self.assertEqual(len(L), 1, L); self.assertIn('information', L[0]); self.assertNotIn(' vs ', L[0])

    def test_a_variable_valence_with_no_source_stops_the_sum(self):
        t = 'The empirical formula is Fe1.00As1.00O5.00.'
        L = PE.charge_balance_check(t, {'name': None, 'epma': {'rows': [{'constituent': 'As2O5', 'mean': 50.0}]}})
        self.assertEqual(L, ['charge balance: not summed — the valence of Fe is stated nowhere the tool reads'])
        # the formula's own superscript decides
        t = 'The empirical formula is Fe3+1.00As1.00O4.00.'
        self.assertEqual(PE.charge_balance_check(t, {'name': None, 'epma': {'rows': [{'constituent': 'As2O5', 'mean': 50.0}]}}), [])


class ParameterSetFit(unittest.TestCase):
    URANYL = """data_u
_cell_length_a 10
_cell_length_b 10
_cell_length_c 10
_cell_angle_alpha 90
_cell_angle_beta 90
_cell_angle_gamma 90
_space_group_name_H-M_alt 'P 1'
loop_
_space_group_symop_operation_xyz
'x, y, z'
loop_
_atom_site_label
_atom_site_type_symbol
_atom_site_fract_x
_atom_site_fract_y
_atom_site_fract_z
U1 U6+ 0 0 0
O1 O 0.18 0 0
O2 O 0 0.235 0
"""

    def test_the_set_behind_printed_valences_is_named_and_a_cited_other_set_is_a_finding(self):
        import math, tempfile, shutil, os
        from pxrd_review import bv_check as B
        tmp = tempfile.mkdtemp(prefix='setfit_')
        try:
            path = os.path.join(tmp, 'u.cif'); open(path, 'w').write(self.URANYL)
            st = B.Structure(path); P = B.Params()
            rows = {rid: (r0, b) for r0, b, rid, det in P.table[('U', 6, 'O', -2)]}
            Rs = (1.78, 1.80, 2.30, 2.45, 2.50)
            burns = [('U1', 'O1', R, round(math.exp((rows['r'][0] - R) / rows['r'][1]), 2)) for R in Rs]
            L = PE.set_fit_lines(burns, st, cited=None)
            self.assertEqual(len(L), 2, L); self.assertIn('Burns et al. (1997)', L[1]); self.assertNotIn(' vs ', L[1])
            L = PE.set_fit_lines(burns, st, cited='gh')                      # the text cites Gagné & Hawthorne; the numbers are Burns'
            self.assertIn('follow Burns et al. (1997)', L[1]); self.assertIn(' vs the Gagné and Hawthorne (2015) the text cites', L[1])
            gh = [('U1', 'O1', R, round(math.exp((rows['bs'][0] - R) / rows['bs'][1]), 2)) for R in Rs]
            L = PE.set_fit_lines(gh, st, cited='gh'); self.assertIn('Gagné and Hawthorne (2015)', L[1]); self.assertNotIn(' vs ', L[1])
            L = PE.set_fit_lines([('U1', 'O1', R, 0.9) for R in (1.8, 2.4)], st, None); self.assertIn('no set the tool carries reproduces', L[1])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class EpmaTableLint(unittest.TestCase):
    def test_range_sd_n_and_beam(self):
        text = ('Chemical analyses (mean of 6 analyses) were made with a 10 μm beam diameter. Crystals are up to 8 μm across. '
                'The composition is given in Table 1.')
        ex = {'epma': {'caption': 'Table 1. Chemical data (wt%) for testite (mean of 6 analyses).',
                       'rows': [{'constituent': 'MgO', 'mean': 21.5, 'sd': 0.25, 'range': (22.1, 23.9)},        # the mean outside its range
                                {'constituent': 'As2O5', 'mean': 50.0, 'sd': 0.20, 'range': (47.0, 53.0)},      # s.d. 0.20 vs at least 3/√5 = 1.34
                                {'constituent': 'CoO', 'mean': 2.0, 'sd': 1.4, 'range': (0.5, 3.6)},           # within the bound
                                {'constituent': 'H2O', 'mean': 30.0, 'sd': None, 'range': None}]}}
        L = PE.epma_table_lint(ex, text)
        self.assertTrue(L and L[0].startswith('analytical table: 4 constituents, n = 6'), L)
        self.assertTrue(any(x.startswith('MgO: mean 21.5 vs its range 22.1–23.9') for x in L), L)
        self.assertTrue(any(x.startswith('As2O5: s.d. 0.2 vs at least 1.34') for x in L), L)
        self.assertFalse(any('CoO' in x for x in L), L)                                                    # scatter alone is not judged: the s.d. column is not read reliably enough
        self.assertTrue(any('beam (10 µm) is as wide as the largest grains named (8 µm)' in x for x in L), L)
        # within Samuelson's bound, no n printed: nothing
        ex2 = {'epma': {'rows': [{'constituent': 'MgO', 'mean': 22.9, 'sd': 0.7, 'range': (22.1, 23.9)}]}}
        self.assertEqual(PE.epma_table_lint(ex2, 'No count here.'), [])
        self.assertEqual(PE.analyses_count('The mean of 12 analyses is given.'), 12); self.assertEqual(PE.analyses_count('nothing'), None)


class Dominance(unittest.TestCase):
    def test_a_leader_inside_the_scatter_is_a_finding(self):
        text = 'The empirical formula, based on 4 O apfu, is (Mn0.524Ca0.476)Σ1.00(Mg0.90Fe2+0.10)Σ1.00Si1.00O4.'
        ex = {'name': None, 'epma': {'rows': [{'constituent': 'MnO', 'mean': 18.1, 'sd': 1.2, 'range': (16.2, 19.8)},
                                              {'constituent': 'CaO', 'mean': 13.0, 'sd': 1.0, 'range': (11.7, 14.5)},
                                              {'constituent': 'MgO', 'mean': 17.7, 'sd': 0.3, 'range': (17.2, 18.1)},
                                              {'constituent': 'FeO', 'mean': 3.5, 'sd': 0.2, 'range': (3.2, 3.8)},
                                              {'constituent': 'SiO2', 'mean': 29.3, 'sd': 0.3, 'range': None}]}}
        L = PE.dominance_check(text, ex)
        self.assertTrue(L and L[0].startswith('dominance within the scatter'), L)
        self.assertTrue(any(x.startswith('(Mn0.524Ca0.476): Mn leads Ca by 0.048 apfu vs the analytical scatter') and 'not established' in x for x in L), L)
        self.assertTrue(any('the ranges cross' in x and 'Mn' in x for x in L), L)
        self.assertFalse(any('Mg' in x for x in L), L)                                          # 0.90 vs 0.10: not in doubt
        # a clear leader: nothing
        self.assertEqual(PE.dominance_check('The empirical formula is (Mn0.80Ca0.20)Σ1.00Si1.00O3.', ex), [])


class TwoV(unittest.TestCase):
    def test_two_stated_values_and_the_indices(self):
        text = ('Optically biaxial (+), α = 1.600(2), β = 1.610(2), γ = 1.630(2), 2V(meas.) = 70.5°. '
                'Table 7. Optical data. 2V 63° 2V(calc) 71.6°.')
        L = PE.optics_2v_lines(text)
        self.assertTrue(L and L[0].startswith('optics: 2V stated 70.5° (meas), 63°, 71.6° (calc); from the indices 71.3°'), L)
        self.assertTrue(any(x.startswith('2V is given as 63° vs 70.5° in different places') for x in L), L)
        self.assertEqual(PE.optics_2v_lines('Optically biaxial (+), α = 1.600, β = 1.610, γ = 1.630, 2V(meas.) = 72°.'), [])
        L = PE.optics_2v_lines('Biaxial (−): α = 1.617(3) β = 1.632(3) \uf067 = 1.637(3) 2Vx (calc.) = 70.5\uf0b0 Dispersion: none.')   # a Symbol-font γ and °
        self.assertTrue(any(x.startswith('2V(calc) 70.5° stated vs 59.5° from the indices') for x in L), L)
        self.assertTrue(any('from the indices' in x and 'information' in x for x in PE.optics_2v_lines('α = 1.600, β = 1.610, γ = 1.630, 2V = 40°.')))


SS_CIF = """data_ss
_cell_length_a 8
_cell_length_b 8
_cell_length_c 8
_cell_angle_alpha 90
_cell_angle_beta 90
_cell_angle_gamma 90
_cell_volume 512
_cell_formula_units_z 1
_space_group_name_H-M_alt 'P 1'
loop_
_space_group_symop_operation_xyz
'x, y, z'
loop_
_atom_site_label
_atom_site_type_symbol
_atom_site_fract_x
_atom_site_fract_y
_atom_site_fract_z
_atom_site_occupancy
_atom_site_U_iso_or_equiv
M1 Mn 0 0 0 0.82 0.010
M1B Fe 0 0 0 0.18 0.010
M3 Zn 0.5 0.5 0.5 1.0 0.010
O1 O 0.25 0 0 1.0 0.015
"""


def _ss_pdf(path, rows, caption='Table 7. Refined site-scattering values (epfu) and assigned site populations for testite.',
            header=(('Site', 40), ('Refined', 90), ('Assigned site population', 160), ('Calculated', 330), ('<M–O>', 400))):
    """A site-scattering table as a page: rows = [[(text, x), …] per line]; a row's wrapped line has no label."""
    import pymupdf
    doc = pymupdf.open(); page = doc.new_page(width=595, height=842)
    y = 60
    page.insert_text((40, y), caption, fontsize=9); y += 16
    for t, x in header:
        page.insert_text((x, y), t, fontsize=9)
    y += 14
    for cells in rows:
        for t, x in cells:
            page.insert_text((x, y), t, fontsize=9)
        y += 14
    doc.save(path); doc.close()


class SiteScattering(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.d, ignore_errors=True)

    def test_population_forms(self):
        e = lambda s, b=None: round(PE._population_electrons(s, 58.59, b)[0], 2)
        self.assertEqual(e('Mn1.64Fe0.36'), 50.36)
        self.assertEqual(e('Mn2+ 0.75Ca0.14Sr0.11'), 25.73)                    # the valence split from its coefficient by a superscript run
        self.assertEqual(e('1.68 Mg 0.35 Fe2', True), 29.26)                    # the coefficient before its symbol, the + of Fe2+ lost
        self.assertEqual(e('0.71 Na + 0.28□+ 0.01 Ca', True), 8.01)             # a vacancy counts no electrons
        self.assertEqual(e('(Sr1.07Ca0.26Na0.04Pb0.02)Σ1.39Ln3+ 0.61'), 83.68)  # a lanthanide group at the note's electrons
        self.assertEqual(e('0.36 Mg 0.79 Fe2+ 0.72 Fe3++ 0.13 Ti 0.01 Zn', True), 46.74)
        self.assertEqual(e('2.00Mn', False), 50.0)                              # a number glued before its symbol in a symbol-first table
        self.assertEqual(e('(H2O)0.89Na0.11'), 10.11)
        self.assertEqual(PE._population_electrons('Ln0.835Ca0.165')[3], ['Ln'])  # no note: unknown

    def test_table_against_itself_and_the_cif(self):
        pdf = os.path.join(self.d, 't.pdf')
        _ss_pdf(pdf, [[('M1', 40), ('50.1', 90), ('Mn1.64Fe0.36', 160), ('50.4', 330), ('2.150', 400)],
                      [('M2', 40), ('34.2', 90), ('Mg1.20Mn0.60Fe0.20', 160), ('33.5', 330), ('2.080', 400)],
                      [('M3', 40), ('18.0', 90), ('Zn0.43Mg0.41Cu0.16', 160), ('22.46', 330), ('2.073', 400)],
                      [('T', 40), ('14.0', 90), ('Si', 160), ('14.0', 330), ('1.620', 400)]])
        tabs = PE.site_scattering_tables(pdf)
        self.assertEqual(len(tabs), 1)
        self.assertEqual([r['label'] for r in tabs[0]['rows']], ['M1', 'M2', 'M3', 'T'])
        self.assertEqual(tabs[0]['rows'][0]['groups'], ['Mn1.64Fe0.36'])
        L = PE.site_scattering_check(pdf, {})
        self.assertTrue(L[0].startswith('site scattering: 3 of 4 assigned populations give the calculated scattering printed'), L)
        self.assertTrue(any(x.startswith('M2: calculated site scattering 33.5 printed vs 34.60 from the population Mg1.20Mn0.60Fe0.20') for x in L), L)   # 34.6 is right; 33.5 is not
        self.assertTrue(any(x.startswith('M3: refined site scattering 18.0 against 22.46 calculated from the assigned population') and x.endswith('(-20 %) (information)') for x in L), L)
        cif = os.path.join(self.d, 's.cif')
        with open(cif, 'w') as f:
            f.write(SS_CIF)
        L = PE.site_scattering_check(pdf, {}, cif)
        self.assertTrue(any(x.startswith('M3: refined site scattering 18.0 (18.0 e per atom) vs 30.0 e per atom in the .cif (Zn 1.000)') for x in L), L)   # 67 % off: a finding
        self.assertFalse(any(x.startswith('M1:') for x in L), L)                # 25.05 e per atom printed, 25.18 in the .cif

    def test_wrapped_population_coefficient_first(self):
        pdf = os.path.join(self.d, 'w.pdf')
        _ss_pdf(pdf, [[('X', 40), ('15.69(9)', 90), ('0.61 Ca + 0.35 Na + 0.04 □', 160), ('16.10', 330)],
                      [('Y', 40), ('47.26(24)', 90), ('1.50 Mg + 0.47 Fe2+ + 0.71 Al +', 160), ('47.04', 330)],
                      [('0.14 Fe3+ + 0.18 Ti', 160)],
                      [('Z', 40), ('79.40(24)', 90), ('4.54 Al + 0.18 Fe3+ + 1.27 Mg', 160), ('79.19', 330)]],
                caption='Table 7. Refined site-scattering values and optimised site-populations for testite.',
                header=(('Site', 40), ('Refined', 90), ('Optimised site-population', 160), ('Calculated', 330)))
        tabs = PE.site_scattering_tables(pdf)
        self.assertTrue(tabs[0]['before'])
        self.assertEqual(tabs[0]['rows'][1]['groups'], ['1.50 Mg 0.47 Fe2+ 0.71 Al 0.14 Fe3+ 0.18 Ti'])
        L = PE.site_scattering_check(pdf, {})
        self.assertTrue(L[0].startswith('site scattering: 3 of 3 assigned populations'), L)
        self.assertEqual(len(L), 1, L)

    def test_manuscript_docx_table(self):
        from docx import Document
        path = os.path.join(self.d, 'm.docx')
        doc = Document()
        doc.add_paragraph('Table 5. Refined site-scattering values (epfu) and assigned site populations for testite.')
        t = doc.add_table(rows=1, cols=4)
        for c, h in zip(t.rows[0].cells, ('Site', 'Refined site scattering', 'Assigned site population', 'Calculated site scattering')):
            c.text = h
        for row in (('M1', '50.1', 'Mn1.64Fe0.36', '50.4'), ('M2', '34.2', 'Mg1.20Mn0.60Fe0.20', '33.5'), ('T', '14.0', 'Si', '14.0')):
            for c, v in zip(t.add_row().cells, row):
                c.text = v
        doc.save(path)
        L = PE.site_scattering_check(path, {})
        self.assertTrue(L and L[0].startswith('site scattering: 2 of 3'), L)
        self.assertTrue(any(x.startswith('M2: calculated site scattering 33.5 printed vs 34.60') for x in L), L)

    def test_electrons_per_formula_unit_in_prose(self):
        text = ('The empirical formula, calculated on the basis of 7 O apfu, is Ca1.00Mg2.01Si2.99O7(OH)0.98. '
                'The number of electrons per formula unit derived from EMPA and SREF (86.0 and 88.5 epfu, respectively) agree.')
        L = PE.epfu_lines(text, {'name': 'testite'})
        self.assertTrue(any(x.startswith('86 electrons per formula unit stated: the empirical formula gives 86.0 (cations)') for x in L), L)
        self.assertFalse(any(x.startswith('88.5') for x in L), L)                # a count no formula gives: left alone
        self.assertEqual(PE.epfu_lines('No count here. The empirical formula is Ca1.00Mg2.01Si2.99O7(OH)0.98.', {}), [])


if __name__ == '__main__':
    unittest.main()
