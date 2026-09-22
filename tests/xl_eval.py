"""A small evaluator for the formulas pxrd_review writes into its workbooks — enough to prove in a
test that a sheet's live formulas give the numbers the Python gave. Not a spreadsheet engine: cell
and range references (other sheets too), arithmetic, and SUM / SUMIF / SUMPRODUCT / AVERAGE /
STDEV / MIN / MAX / ABS / IF / EXP / TEXT and & (joining text), whole-column ranges (bonds!A:A) and quoted sheet names."""
import re, math, statistics
import openpyxl
from openpyxl.utils import range_boundaries

_REF = re.compile(r"(?:(?:'([^']+)'|(\w+))!)?(?:(\$?[A-Z]{1,2}\$?\d+)(?::(\$?[A-Z]{1,2}\$?\d+))?|([A-Z]{1,2}):([A-Z]{1,2})(?![A-Za-z(]))")

def _flat(args):
    out = []
    for a in args:
        out += _flat(a) if isinstance(a, list) else [a]
    return out

def _nums(args):
    return [v for v in _flat(args) if isinstance(v, (int, float)) and not isinstance(v, bool)]

def _sumif(rng, crit, vals):
    return sum(v for k, v in zip(rng, vals) if k == crit and isinstance(v, (int, float)))

_FUNCS = {'SUM': lambda *a: sum(_nums(a)), 'AVERAGE': lambda *a: statistics.mean(_nums(a)), 'STDEV': lambda *a: statistics.stdev(_nums(a)),
          'MIN': lambda *a: min(_nums(a) or [0]), 'MAX': lambda *a: max(_nums(a) or [0]),   # of no numbers: 0, as Excel
          'ABS': abs, 'EXP': math.exp, 'ROUND': lambda x, n=0: math.copysign(math.floor(abs(x) * 10 ** n + 0.5) / 10 ** n, x),   # half away from zero, as Excel
          'MEDIAN': lambda *a: statistics.median(_nums(a)) if _nums(a) else 0, 'COUNT': lambda *a: len(_nums(a)),
          'LN': math.log, 'ISNUMBER': lambda v: isinstance(v, (int, float)) and not isinstance(v, bool), 'COUNTIFS': lambda *a: _countifs(a), 'SUMIFS': lambda vals, *a: _countifs(a, vals),
          'INDEX': lambda rng, i: rng[int(i) - 1] if i else '#N/A', 'MATCH': lambda v, rng, _t=0: (rng.index(v) + 1) if v in rng else 0,   # IF is eager here, lazy in Excel: an unused branch must not raise
          'LEFT': lambda t, n_: str(t)[:int(n_)], 'TRUE': True, 'FALSE': False,
          '_AND': lambda *a: all(a), '_OR': lambda *a: any(a), 'TEXT': lambda v, f: _text(v, f), 'IF': lambda c, a, b: a if c else b, 'SUMIF': _sumif,
          'SUMPRODUCT': lambda x, y: sum(p * q for p, q in zip(x, y) if isinstance(p, (int, float)) and isinstance(q, (int, float)))}

def _countifs(pairs, vals=None):
    """COUNTIFS(range, criterion, ...) — or SUMIFS when `vals` is given. Criteria: a value, or "<>" (not blank)."""
    n = len(pairs[0]); keep = [True] * n
    def hit(v, crit):
        if crit == '<>':
            return v not in (None, '')
        if isinstance(crit, str) and isinstance(v, str):
            # Excel matches text case-insensitively, '*' any run and '?' one character ("does not*")
            return re.fullmatch(re.escape(crit).replace(r'\*', '.*').replace(r'\?', '.'), v, re.I | re.S) is not None
        return v == crit
    for rng, crit in zip(pairs[0::2], pairs[1::2]):
        for i, v in enumerate(rng):
            keep[i] = keep[i] and hit(v, crit)
    if vals is None:
        return sum(keep)
    return sum(v for v, k in zip(vals, keep) if k and isinstance(v, (int, float)))

def _text(v, f):
    f0 = f.split(';')[0]
    d = f0.split('.')[1] if '.' in f0 else ''
    out = '%.*f' % (len(d), v)
    if '#' in d:
        out = out.rstrip('0').rstrip('.')
    return ('+' if f0.startswith('+') and v >= 0 else '') + out

class Book:
    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path); self.memo = {}
    def value(self, sheet, ref):
        key = (sheet, ref.replace('$', ''))
        if key not in self.memo:
            self.memo[key] = None                                      # a circular reference reads as blank rather than recursing
            v = self.wb[sheet][key[1]].value
            self.memo[key] = self._eval(sheet, v[1:]) if isinstance(v, str) and v.startswith('=') else v
        return self.memo[key]
    def fills(self, sheet):
        """{cell coordinate: [fill colours whose conditional-format rule is TRUE there]} — every FormulaRule of the sheet, its
        formula shifted row by row as Excel shifts a relative reference from the range's top-left cell. The colours a reviewer
        sees are otherwise untested (issue #11)."""
        from openpyxl.utils import range_boundaries
        ws = self.wb[sheet]; out = {}
        for rng in ws.conditional_formatting:
            for rule in rng.rules:
                if rule.type != 'expression' or not rule.formula:
                    continue
                colour = rule.dxf.fill.fgColor.rgb if rule.dxf is not None and rule.dxf.fill is not None and rule.dxf.fill.fgColor is not None else '?'
                for cr in str(rng.sqref).split():
                    c1, r1, c2, r2 = range_boundaries(cr)
                    for r in range(r1, r2 + 1):
                        for c in range(c1, c2 + 1):
                            f = re.sub(r'(?<![\$A-Za-z])([A-Z]{1,2})(\$?)(\d+)(?![\d(])', lambda m: '%s%s%d' % (m.group(1), m.group(2), int(m.group(3)) if m.group(2) else int(m.group(3)) + r - r1), rule.formula[0])
                            f = re.sub(r'\$([A-Z]{1,2})(\d+)(?![\d(])', lambda m: '$%s%d' % (m.group(1), int(m.group(2)) + r - r1), f)   # $F4: the column fixed, the row relative
                            try:
                                if self._eval(sheet, f):
                                    out.setdefault(ws.cell(r, c).coordinate, []).append(colour[-6:])
                            except Exception:
                                out.setdefault(ws.cell(r, c).coordinate, []).append('ERR')
        return out

    def _eval(self, sheet, expr):
        env = dict(_FUNCS); n = [0]
        def sub(m):
            sh = m.group(1) or m.group(2) or sheet; name = '_r%d' % n[0]; n[0] += 1
            if m.group(5):                                             # a whole column: down to the sheet's last row
                ws = self.wb[sh]
                env[name] = [self.value(sh, '%s%d' % (m.group(5), r)) for r in range(1, ws.max_row + 1)]
            elif m.group(4):
                c1, r1, c2, r2 = range_boundaries((m.group(3) + ':' + m.group(4)).replace('$', ''))
                env[name] = [self.value(sh, self.wb[sh].cell(r, c).coordinate) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
            else:
                v = self.value(sh, m.group(3)); env[name] = 0 if v is None else v
            return name
        parts = re.split(r'("[^"]*")', expr)                           # references are not looked for inside strings
        code = ''.join(p if p.startswith('"') else _REF.sub(sub, p).replace('<>', '!=').replace('^', '**').replace('&', '+').replace('AND(', '_AND(').replace('OR(', '_OR(') for p in parts)
        code = re.sub(r'(?<![<>!=])=(?!=)', '==', code)
        return self._lazy(code, env)

    def _lazy(self, code, env):
        """Excel's IF evaluates only the branch it takes (IF(D6=0,"no K_C",1-B10/D6) must not divide): an expression
        that IS an IF call is taken apart at its top-level commas and only the chosen branch is evaluated."""
        code = code.strip()
        if code.startswith('IF(') and code.endswith(')'):
            depth = 0; quote = False; args = []; start = 3; whole = True
            for i, ch in enumerate(code):
                if ch == '"':
                    quote = not quote
                elif quote:
                    continue
                elif ch == '(':
                    depth += 1
                elif ch == ')':
                    depth -= 1
                    if depth == 0 and i != len(code) - 1:
                        whole = False; break                           # 'IF(a,b,c)+1': not one call
                elif ch == ',' and depth == 1:
                    args.append(code[start:i]); start = i + 1
            if whole and len(args) == 2:
                args.append(code[start:-1])
                return self._lazy(args[1] if self._lazy(args[0], env) else args[2], env)
        return eval(code, {'__builtins__': {}}, env)
