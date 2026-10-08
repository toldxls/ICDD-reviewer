"""Unit tests for pxrd_review.spectrum and spectrum_formats — synthetic spectra with known answers, files written here.

    python3 -m unittest tests.test_spectrum -v

Every spectrum is made from known bands plus seeded noise (random.Random), so what a fit, a baseline or the cosmic-ray
filter should return is known exactly. No instrument file is in the repo: the binary readers are tested on files this
module writes to the published layouts, which proves the reader against the layout and not against an instrument.
"""
import contextlib, io, math, os, random, shutil, struct, tempfile, unittest

from pxrd_review import spectrum as SM, spectrum_formats as F

TRUTH = [{'shape': 'pvoigt', 'pos': 831.0, 'height': 1000.0, 'fwhm': 12.0, 'eta': 0.6},
         {'shape': 'pvoigt', 'pos': 846.0, 'height': 600.0, 'fwhm': 10.0, 'eta': 0.4},
         {'shape': 'pvoigt', 'pos': 915.0, 'height': 150.0, 'fwhm': 14.0, 'eta': 0.5}]


OH = [{'shape': 'pvoigt', 'pos': 3050.0, 'height': 120.0, 'fwhm': 110.0, 'eta': 0.2}, {'shape': 'pvoigt', 'pos': 3210.0, 'height': 260.0, 'fwhm': 120.0, 'eta': 0.2},
      {'shape': 'pvoigt', 'pos': 3350.0, 'height': 340.0, 'fwhm': 100.0, 'eta': 0.2}, {'shape': 'pvoigt', 'pos': 3470.0, 'height': 280.0, 'fwhm': 80.0, 'eta': 0.2},
      {'shape': 'pvoigt', 'pos': 3565.0, 'height': 150.0, 'fwhm': 40.0, 'eta': 0.2}]


def synth(seed, comps, lo=700.0, hi=1000.0, dx=0.5, noise=4.0, base=lambda v: 0.0, spikes=0):
    rng = random.Random(seed)
    x = [lo + dx * i for i in range(int((hi - lo) / dx) + 1)]
    y = [base(v) + sum(SM.profile(c['shape'], v, c['pos'], c['height'], c['fwhm'], c.get('eta', 0.5)) for c in comps) + rng.gauss(0, noise) for v in x]
    at = []
    for _ in range(spikes):
        i = rng.randrange(10, len(x) - 10); w = rng.choice((1, 1, 2)); h = rng.uniform(150, 2000)
        for k in range(w):
            y[i + k] += h * (1.0 if k == 0 else 0.6)
        at.append(i)
    return x, y, at


def start(comps):
    return [dict(c, pos=c['pos'] + 2, height=c['height'] * 0.8, fwhm=c['fwhm'] * 1.3, eta=0.5) for c in comps]


class Formats(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.x = [400 + 0.5 * i for i in range(600)]
        cls.y = [100 + 900 * math.exp(-((v - 540) / 6.0) ** 2) + 3 * math.sin(v) for v in cls.x]
        cls.s = F._series('test', cls.x, cls.y, y_unit='ABSORBANCE')

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def _file(self, name, text, mode='w'):
        p = os.path.join(self.tmp, name)
        with open(p, mode, **({} if 'b' in mode else {'encoding': 'utf-8'})) as f:
            f.write(text)
        return p

    def _same(self, s, tol=1e-3):
        self.assertEqual(len(s.x), len(self.x))
        self.assertLess(max(abs(a - b) for a, b in zip(s.x, self.x)), 1e-3)
        self.assertLess(max(abs(a - b) for a, b in zip(s.y, self.y)), tol)

    def test_text_forms(self):
        rows = list(zip(self.x, self.y))
        self._same(F.load(self._file('a.txt', ''.join('%.2f\t%.4f\n' % r for r in rows)))[0])
        self._same(F.load(self._file('b.csv', 'Wavenumber,Intensity\n' + ''.join('%.2f,%.4f\n' % r for r in reversed(rows))))[0])      # descending x
        s = F.load(self._file('c.txt', '#Acq. time (s)=\t10\n#Laser=\t532\n' + ''.join(('%.2f\t%.4f\n' % r).replace('.', ',') for r in rows)))[0]
        self._same(s); self.assertEqual(s.meta['format'], 'labspec'); self.assertIn('a comma is read as the decimal mark', s.meta['notes'])
        s = F.load(self._file('d.txt', '##NAMES=Testite\n##RRUFFID=R000000\n' + ''.join('%.4f, %.4f\n' % r for r in rows) + '##END=\n'))[0]
        self._same(s); self.assertEqual((s.meta['format'], s.meta['names']), ('rruff', 'Testite'))
        two = F.load(self._file('e.csv', 'cm-1;Sample 1;Sample 2\n' + ''.join('%.2f;%.4f;%.4f\n' % (a, b, 2 * b) for a, b in rows)))
        self.assertEqual([t.name for t in two], ['e.csv [Sample 1]', 'e.csv [Sample 2]'])

    def test_text_refusals(self):
        with self.assertRaises(ValueError):
            F.load(self._file('words.txt', 'no numbers here\nat all\n'))
        with self.assertRaises(ValueError):
            F.load(self._file('empty.txt', ''))
        with self.assertRaises(ValueError) as cm:
            F.load(self._file('x.l6s', b'\x00\x01binary', 'wb'))
        self.assertIn('export the spectrum as text', str(cm.exception))

    def test_jcamp_round_trip_and_its_own_checks(self):
        for form in ('DIF', 'AFFN'):
            s = F.load(self._file(form + '.jdx', F.write_jcamp(self.s, form)))[0]
            self._same(s); self.assertEqual(s.meta['notes'], [], form)
            self.assertEqual(s.meta['y_unit'], 'ABSORBANCE')
        bad = F.write_jcamp(self.s, 'DIF').replace('##FIRSTX=400.0', '##FIRSTX=410.0')
        self.assertTrue(any('departs from the even spacing' in n for n in F.read_jcamp(bad)[0].meta['notes']))
        short = F.write_jcamp(self.s, 'AFFN').replace('##NPOINTS=600', '##NPOINTS=590')
        self.assertTrue(any('NPOINTS says 590' in n for n in F.read_jcamp(short)[0].meta['notes']))

    def test_jcamp_compressed_line_of_the_specification(self):
        # the DIFDUP example every description of the format prints
        vals, last_dif = F._asdf_line('1000 @VKT%TLkj%J%KLJ%njKjL%kL%jJULJ%kLK1%lLMNPNPRLJ0QTOJ1P')
        self.assertEqual(vals[1:], [0, 0, 0, 0, 2, 4, 4, 4, 7, 5, 4, 4, 5, 5, 7, 10, 11, 11, 6, 5, 7, 6, 9, 9, 7, 10, 10, 9, 10, 11, 12, 15, 16, 16, 14, 17, 38, 38,
                                    35, 38, 42, 47, 54, 59, 66, 75, 78, 88, 96, 104, 110, 121, 128])
        self.assertTrue(last_dif)
        self.assertEqual(F._asdf_line('100 1 2-3 4.5E1')[0], [100.0, 1.0, 2.0, -3.0, 45.0])       # AFFN, a sign as the separator (PAC), an exponent
        with self.assertRaises(ValueError):
            F._asdf_line('100 1 2 ? 3')

    def test_spc_round_trip(self):
        for as_float in (True, False):
            s = F.load(self._file('f%d.spc' % as_float, F.write_spc(self.s, as_float), 'wb'))[0]
            self._same(s, 1e-3)
            self.assertIn('has not been checked against a file from an instrument', s.meta['unverified'])
        old = bytearray(F.write_spc(self.s)); old[1] = 0x4d
        with self.assertRaises(ValueError):
            F.load(self._file('old.spc', bytes(old), 'wb'))

    def test_opus_and_spa_to_their_layouts(self):
        n = len(self.x)
        # OPUS: header, a directory of (type, length in words, offset), a parameter block and a float32 data block
        def par(name, kind, blob):
            return name.encode() + b'\x00' + struct.pack('<HH', kind, len(blob) // 2) + blob
        params = (par('NPT', 0, struct.pack('<i', n)) + par('FXV', 1, struct.pack('<d', self.x[0])) + par('LXV', 1, struct.pack('<d', self.x[-1])) +
                  par('CSF', 1, struct.pack('<d', 2.0)) + par('END', 0, b''))
        data = struct.pack('<%df' % n, *[v / 2.0 for v in self.y])
        off_p = 24 + 3 * 12; off_d = off_p + len(params)
        head = b'\x0a\x0a\xfe\xfe' + struct.pack('<d', 920622.0) + struct.pack('<iii', 24, 3, 2)
        directory = struct.pack('<iii', 15 | (16 << 8), len(data) // 4, off_d) + struct.pack('<iii', 31 | (16 << 8), len(params) // 4 + 1, off_p) + b'\x00' * 12
        s = F.load(self._file('s.0', head + directory + params + data, 'wb'))[0]
        self._same(s, 1e-2); self.assertEqual(s.meta['format'], 'opus'); self.assertIn('absorbance', s.name)
        # SPA: the title, a directory from byte 304 (count at 294), the header entry (type 2) and the data entry (type 3)
        raw = bytearray(1200 + 4 * n)
        raw[0:18] = b'Spectral Data File'; raw[30:37] = b'testite'
        struct.pack_into('<H', raw, 294, 2)
        raw[304] = 2; struct.pack_into('<I', raw, 306, 600)
        raw[320] = 3; struct.pack_into('<II', raw, 322, 1200, 4 * n)
        struct.pack_into('<I', raw, 604, n); struct.pack_into('<ff', raw, 616, self.x[-1], self.x[0])        # an FTIR spectrum is stored high to low
        struct.pack_into('<%df' % n, raw, 1200, *reversed(self.y))
        s = F.load(self._file('s.spa', bytes(raw), 'wb'))[0]
        self._same(s, 1e-2); self.assertEqual((s.meta['format'], s.name), ('spa', 'testite'))
        with self.assertRaises(ValueError):
            F.load(self._file('bad.spa', b'Spectral Data File' + b'\x00' * 700, 'wb'))

    def test_workbook_blocks(self):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Raman'
        for i, (a, b) in enumerate(zip(self.x, self.y)):
            ws.append([a, b, b + 500, None, self.x[-1 - i], 'note' if i == 3 else None, self.y[-1 - i]])
        ir = wb.create_sheet('IR'); ir.append(['wavenumber', 'A'])
        for a, b in zip(reversed(self.x), reversed(self.y)):
            ir.append([a, b])
        p = os.path.join(self.tmp, 'w.xlsx'); wb.save(p)
        got = F.load(p)
        self.assertEqual(len(got), 3, [s.name for s in got])                # A:B, its offset copy A:C, and the sheet IR — E has no y beside it
        self._same(got[0]); self.assertTrue(got[1].meta['offset_copy']); self.assertIn('shifted by 500', got[1].meta['notes'][0])
        self.assertEqual(got[2].name, 'w.xlsx | IR [A]'); self._same(got[2])
        self.assertEqual(len(F.load(p, 'ir')), 1)
        with self.assertRaises(ValueError):
            F.load(p, 'nosuch')

    def test_text_is_read_not_assumed(self):
        rows = list(zip(self.x, self.y))
        tsv = lambda rr, f='%.2f\t%.3f': '\n'.join(f % r for r in rr) + '\n'
        # the axis is the column that runs one way — by its heading, else the first that is no running index
        s = F.read_text('idx,wn,int\n' + '\n'.join('%d,%.2f,%.3f' % (i, a, b) for i, (a, b) in enumerate(rows)))
        self.assertEqual(len(s), 1); self._same(s[0]); self.assertIn('the axis is column 2 (wn)', s[0].meta['notes'][0]); self.assertIn('column 1 left out', s[0].meta['notes'][1])
        s = F.read_text('Intensity\tWavenumber\n' + '\n'.join('%.3f\t%.2f' % (b, a) for a, b in rows))
        self.assertEqual(len(s), 1); self._same(s[0])
        # spectra one after another in a file are several spectra, never one sorted together
        s = F.read_text(tsv(rows) + '\nsecond\n' + tsv([(a, 2 * b) for a, b in rows]))
        self.assertEqual([t.name for t in s], ['spectrum [spectrum 1 of 2]', 'spectrum [spectrum 2 of 2]']); self._same(s[0]); self.assertAlmostEqual(s[1].y[280], 2 * self.y[280], 2)
        self.assertFalse(any('sorted' in n for t in s for n in t.meta['notes']))
        # a stray pair of numbers before or after the spectrum is no point of it
        s = F.read_text('10 2\n' + tsv(rows) + '1 2\n')
        self._same(s[0]); self.assertIn('2 lines of two numbers standing apart', s[0].meta['notes'][0])
        # thousands separators, quoted cells, a typographic minus: read — and a line that cannot be read is SAID
        big = [(900 + 0.5 * i, 10.0 + i) for i in range(400)]
        s = F.read_text('\n'.join('{:,.2f}\t{:.1f}'.format(a, b) for a, b in big))
        self.assertEqual((len(s[0].x), s[0].x[0], s[0].x[-1]), (400, 900.0, 1099.5))
        s = F.read_text('\n'.join('"{:,.2f}","{:.1f}"'.format(a, b) for a, b in big))
        self.assertEqual((len(s[0].x), s[0].x[-1]), (400, 1099.5))
        s = F.read_text('\n'.join('%.1f;%s' % (a, ('{:.3f}'.format(b)).replace('.', ',')) for a, b in [(1000 + i, 1234.5 + i) for i in range(50)]).replace('1234,500', '1.234,500'))
        self.assertAlmostEqual(s[0].y[0], 1234.5, 6)
        s = F.read_text(tsv([(a, b - 500) for a, b in rows]).replace('-', '\u2212'))
        self.assertEqual(len(s[0].x), 600); self.assertLess(min(s[0].y), 0)
        s = F.read_text('\n'.join('{:,d}\t{:.1f}'.format(int(900 + i), 10.0 + i) for i in range(300)))      # 999, then 1,000 with no decimals: the axis runs on
        self.assertEqual((len(s), len(s[0].x), s[0].x[-1]), (1, 300, 1199.0)); self.assertIn('thousands separator', s[0].meta['notes'][0])
        s = F.read_text('\n'.join('%.1f\t%s' % (100 + i, 'n/a' if i == 7 else '%.1f' % i) for i in range(60)))
        self.assertEqual(len(s[0].x), 59); self.assertIn('1 line that begins with a number could not be read', s[0].meta['notes'][0])
        s = F.read_text('\n'.join(('%.1f,%.2f' % (a, b)).replace('.', ',') for a, b in rows))             # the comma as decimal mark AND separator
        self.assertEqual(len(s), 1); self.assertAlmostEqual(s[0].x[3], self.x[3], 6); self.assertIn('each pair of columns was read as one number', s[0].meta['notes'][-1])

    def test_text_units_and_encodings(self):
        nm = 'Wavelength (nm)\tCounts\n' + '\n'.join('%.2f\t%.1f' % (540 + 0.05 * i, 100 + i % 7) for i in range(400))
        s = F.read_text(nm)[0]
        self.assertEqual(s.meta['x_unit'], 'nm')
        prep = SM.prepare_series([s], laser=532.0)
        self.assertAlmostEqual(prep['x'][0], 1e7 / 532 - 1e7 / 540, 3); self.assertTrue(any('converted from nm' in n for n in prep['notes']))
        self.assertTrue(any('give the laser wavelength' in n for n in SM.prepare_series([s])['notes']))
        self.assertEqual(F.read_text(nm.replace('Wavelength (nm)', 'Wavelength (µm)'))[0].meta['x_unit'], 'um')
        body = 'Raman shift (cm\u00b9)\tIntensity \u00b5V\n' + '\n'.join('%.2f\t%.3f' % r for r in zip(self.x, self.y)) + '\n'
        for pad in ('', ' '):                                             # an even and an odd count of bytes: Windows-1252 is not UTF-16
            self._same(F.load(self._file('cp%d.txt' % len(pad), (body + pad).encode('cp1252'), 'wb'))[0])
        self._same(F.load(self._file('u16.txt', body.encode('utf-16-le'), 'wb'))[0])       # UTF-16 with no byte-order mark
        self._same(F.load(self._file('u16b.txt', body.encode('utf-16'), 'wb'))[0])

    def test_jcamp_e_is_a_digit_in_compressed_data(self):
        self.assertEqual(F._asdf_line('400E23J5'), ([400.0, 523.0, 538.0], True))
        self.assertEqual(F._asdf_line('400e23j5'), ([400.0, -523.0, -538.0], True))
        self.assertEqual(F._asdf_line('400 1.5E+3 J5', True)[0], [400.0, 1500.0, 1515.0])
        five = F._series('five', self.x, [0.05 + v / 1000.0 for v in self.y])                 # ordinates that begin with the digit 5
        for form, per in (('DIF', 8), ('DIF', 7), ('AFFN', 8)):
            back = F.read_jcamp(F.write_jcamp(five, form, per))[0]
            self.assertEqual((len(back.x), back.meta['notes']), (600, []))
            self.assertLess(max(abs(a - b) for a, b in zip(back.y, five.y)), 2e-6)
        head = '##TITLE=t\n##JCAMP-DX=4.24\n##XUNITS=1/CM\n##FIRSTX=400\n##LASTX=409\n##NPOINTS=10\n'
        got = F.read_jcamp(head + '##XYPOINTS=(XY..XY)\n' + '\n'.join('%d, %d; %d, %d' % (400 + i, i, 401 + i, i + 1) for i in range(0, 10, 2)) + '\n##END=\n')[0]
        self.assertEqual((got.x, got.y), ([float(v) for v in range(400, 410)], [float(v) for v in range(10)]))
        # DIF lines written WITHOUT the Y-check: the x printed on a line says its first ordinate is a new point
        got = F.read_jcamp(head + '##XYDATA=(X++(Y..Y))\n400@JJJJ\n405EJJJJ\n##END=\n')[0]
        self.assertEqual((got.y, got.meta['notes']), ([float(v) for v in range(10)], []))
        with self.assertRaises(ValueError):
            F.read_jcamp('##TITLE=t\n##XYDATA=(X++(Y..Y))\n;09\tk1\n02% 8936k02j4\n##END=\n')

    def test_workbook_cells_that_are_no_number(self):
        import openpyxl
        def book(fill):
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'S'
            fill(ws)
            p = os.path.join(self.tmp, 'c.xlsx'); wb.save(p)
            return F.load(p)
        def holes(ws):
            ws.append(['wavenumber', 'y'])
            for i, (a, b) in enumerate(zip(self.x, self.y)):
                ws.append([None if i == 250 else a, '#N/A' if i in (40, 41) else b])
        got = book(holes)
        self.assertEqual((len(got), len(got[0].x), got[0].x[-1]), (1, 597, self.x[-1])); self.assertIn('3 rows with a cell that is no number left out', got[0].meta['notes'])
        self.assertFalse(got[0].meta['offset_copy'])
        def text_and_index(ws):
            for i, (a, b) in enumerate(zip(self.x, self.y)):
                ws.append([i + 1, '%.2f' % a, b])
        got = book(text_and_index)
        self.assertEqual(len(got), 1); self._same(got[0])                  # the index column is no axis; numbers typed as text are numbers
        def y_first(ws):
            for a, b in zip(self.x, self.y):
                ws.append([b, -a])
        got = book(y_first)
        self.assertEqual((len(got), got[0].x[0]), (1, -self.x[-1]))        # intensity left of the axis; an anti-Stokes axis
        def stacked(ws):
            for a, b in zip(self.x, self.y):
                ws.append([a, b])
            ws.append([]); ws.append(['second', None])
            for a, b in zip(self.x, self.y):
                ws.append([a + 1000, 2 * b])
        got = book(stacked)
        self.assertEqual([(len(t.x), t.x[0]) for t in got], [(600, 400.0), (600, 1400.0)])
        with self.assertRaises(ValueError) as e:
            F.load(self._file('fake.xlsx', 'not a workbook'))
        self.assertIn('could not be opened as a workbook', str(e.exception))
        with self.assertRaises(ValueError) as e:
            F.load(self._file('cut.spc', F.write_spc(self.s)[:700], 'wb'))
        self.assertIn('cut short or damaged', str(e.exception))

    def test_units_and_kind(self):
        nm = F._series('r', [532.0 + 0.05 * i for i in range(400)], [1.0] * 400, x_unit='nm')
        s = F.raman_shift(nm, 532.0)
        self.assertAlmostEqual(s.x[0], 0.0, 6); self.assertAlmostEqual(s.x[-1], 1e7 / 532 - 1e7 / 551.95, 3); self.assertEqual(s.meta['x_unit'], 'cm-1')
        t = F._series('t', self.x, [98.0 - 60 * math.exp(-((v - 540) / 8.0) ** 2) for v in self.x])
        self.assertEqual(F.guess_kind(t), 'ir-transmittance')
        self.assertEqual(F.guess_kind(self.s), 'ir-absorbance')              # its stated unit
        a = SM.to_absorbance(t.y)
        self.assertAlmostEqual(max(a), -math.log10(0.38), 3)


class Cleaning(unittest.TestCase):
    def test_cosmic_rays_go_and_bands_stay(self):
        hit = total = false = 0
        for seed in range(12):
            x, y, at = synth(200 + seed, TRUTH, spikes=4)
            y2, found = SM.despike(x, y)
            total += len(at); hit += sum(any(a - 1 <= i <= b + 1 for a, b in found) for i in at)
            false += sum(1 for a, b in found if not any(a - 1 <= i <= b + 1 for i in at))
            self.assertLess(max(y2[i] - TRUTH[0]['height'] for i in range(len(y2))), 60)       # nothing left standing above the strongest band
        self.assertEqual((hit, false), (total, 0))
        # a real band three points wide is not a cosmic ray: it rises and falls over more points than a spike may
        x, y, _ = synth(1, [{'shape': 'lorentz', 'pos': 850.0, 'height': 5000.0, 'fwhm': 3.0}])
        self.assertEqual(SM.despike(x, y)[1], [])

    def test_a_sharp_band_is_no_cosmic_ray(self):
        # every sharp band of a low-dispersion grating is two to four points wide: it has flanks, a cosmic ray has none
        rng = random.Random(4)
        for shape, fw in (('gauss', 2.0), ('gauss', 2.5), ('gauss', 3.0), ('lorentz', 3.0), ('lorentz', 4.0)):
            for trial in range(12):
                c = 464.0 + rng.random()
                x = [300.0 + i for i in range(400)]
                y = [200 + SM.profile(shape, v, c, 1000.0, fw, 0.5) + rng.gauss(0, 5) for v in x]
                out, found = SM.despike(x, y)
                self.assertEqual(found, [], (shape, fw, c)); self.assertEqual(out, y)
        x = [300.0 + i for i in range(400)]; y = [200 + rng.gauss(0, 5) for _ in x]; y[100] += 900; y[250] += 700; y[251] += 400
        self.assertEqual(SM.despike(x, y)[1], [(100, 100), (250, 251)])                       # one point, and two
        y2 = [200 + SM.profile('gauss', v, 464.3, 1000.0, 2.5, 0.5) + rng.gauss(0, 5) for v in x]
        self.assertEqual(len(SM.despike(x, y2, resolution=8.0)[1]), 1)                         # … unless the resolution says nothing can be that narrow
        sparse = [float(rng.random() < 0.2) for _ in x]                                        # counts so sparse that most differences are zero
        self.assertEqual(SM.despike(x, sparse)[1], [])

    def test_snip_on_a_finely_sampled_spectrum(self):
        x = [0.1 * i for i in range(20000)]
        y = [50 + 0.01 * v + 300 * math.exp(-((v - 1000) / 40.0) ** 2) for v in x]
        b = SM.baseline(x, y, 'snip', width=300.0)                                             # 3000 points either side: on every fifth point, read back
        self.assertLess(max(abs(q - (50 + 0.01 * v)) for v, q in zip(x, b)), 1.5)

    def test_noise_is_the_noise(self):
        for sigma in (2.0, 10.0):
            x, y, _ = synth(3, TRUTH, noise=sigma, base=lambda v: 200 + 0.5 * v)
            self.assertAlmostEqual(SM.noise(y) / sigma, 1.0, delta=0.15)

    def test_baselines(self):
        true_area = sum(SM.area(c['shape'], c['height'], c['fwhm'], c['eta']) for c in TRUTH)
        line = lambda v: 300 + 0.8 * (v - 700)
        x, y, _ = synth(5, TRUTH, base=line)
        for kind, kw, tol in (('als', dict(lam=1e5, p=0.005), 0.12), ('snip', dict(width=100), 0.10), ('rubber', {}, 0.10)):
            b = SM.baseline(x, y, kind, **kw)
            got = sum((y[i] - b[i]) * 0.5 for i in range(len(x)))
            self.assertLess(abs(got - true_area) / true_area, tol, kind)
            self.assertLess(max(abs(b[i] - line(x[i])) for i in range(len(x)) if not 790 < x[i] < 950), 45, kind)   # on the line where there is no band
        b = SM.baseline(x, y, 'anchors', anchors=[705, 760, 960, 995])
        self.assertLess(max(abs(b[i] - line(x[i])) for i in range(len(x))), 12)
        self.assertEqual(SM.baseline(x, y, 'none'), [0.0] * len(x))
        with self.assertRaises(ValueError):
            SM.baseline(x, y, 'spline')

    def test_savgol(self):
        y = [0.5 * i * i - 3 * i + 2 for i in range(40)]                     # a parabola is reproduced exactly, and its second derivative is 1
        self.assertLess(max(abs(a - b) for a, b in zip(SM.savgol(y, 9, 2), y)), 1e-8)
        self.assertLess(max(abs(v - 1.0) for v in SM.savgol(y, 9, 2, deriv=2)), 1e-8)

    def test_find_peaks(self):
        x, y, _ = synth(4, TRUTH)
        got = sorted(p['pos'] for p in SM.find_peaks(x, y))
        self.assertEqual(len(got), 3)
        for a, b in zip(got, (831.0, 846.0, 915.0)):
            self.assertAlmostEqual(a, b, delta=3.0)
        self.assertEqual(SM.find_peaks(*synth(4, [])[:2]), [])               # noise alone holds no band


class Fitting(unittest.TestCase):
    def test_positions_within_their_uncertainties(self):
        within = total = 0; chi = []
        for seed in range(25):
            x, y, _ = synth(seed, TRUTH)
            f = SM.fit(x, y, start(TRUTH))
            self.assertTrue(f.converged); chi.append(f.chi2)
            for c, t in zip(f.comps, TRUTH):
                total += 1; within += abs(c['pos'] - t['pos']) <= 3 * c['esd']['pos']
                self.assertAlmostEqual(c['fwhm'], t['fwhm'], delta=1.0)
        self.assertGreaterEqual(within, total - 1)
        self.assertAlmostEqual(sum(chi) / len(chi), 1.0, delta=0.12)          # the model leaves the noise and nothing else
        self.assertEqual(f.warnings, [])

    def test_shapes_and_areas(self):
        for shape, unit in (('gauss', SM.G_AREA), ('lorentz', SM.L_AREA)):
            c = [{'shape': shape, 'pos': 850.0, 'height': 500.0, 'fwhm': 10.0}]
            x, y, _ = synth(2, c, lo=600, hi=1100)
            f = SM.fit(x, y, start(c), local_baseline='none')
            self.assertAlmostEqual(f.comps[0]['area'], 500.0 * 10.0 * unit, delta=0.03 * 500 * 10 * unit)
            self.assertAlmostEqual(f.comps[0]['pos'], 850.0, delta=0.2)
        # a Gaussian fitted as a pseudo-Voigt: η goes to its bound and the fit still converges
        g = [{'shape': 'gauss', 'pos': 840.0, 'height': 800.0, 'fwhm': 12.0}]
        x, y, _ = synth(7, g)
        f = SM.fit(x, y, [dict(g[0], shape='pvoigt', eta=0.5)])
        self.assertTrue(f.converged); self.assertLess(f.comps[0]['eta'], 0.05)

    def test_local_baseline_and_fixed_parameters(self):
        x, y, _ = synth(8, TRUTH[:1], base=lambda v: 50 + 0.2 * (v - 850))
        f = SM.fit(x, y, start(TRUTH[:1]), local_baseline='linear')
        self.assertAlmostEqual(f.base[0], 50.0, delta=1.5); self.assertAlmostEqual(f.base[1], 0.2, delta=0.02)
        f = SM.fit(x, y, [dict(TRUTH[0], fix=['pos', 'eta'], height=700.0)], local_baseline='linear')
        self.assertEqual((f.comps[0]['pos'], f.comps[0]['eta']), (831.0, 0.6)); self.assertIsNone(f.comps[0]['esd']['pos'])
        pair = [dict(TRUTH[0], fwhm=11.0, eta=0.5), dict(TRUTH[1], fwhm=11.0, eta=0.5)]
        x, y, _ = synth(9, pair)
        f = SM.fit(x, y, start(pair), share_width=True)
        self.assertEqual(f.comps[0]['fwhm'], f.comps[1]['fwhm']); self.assertAlmostEqual(f.comps[0]['fwhm'], 11.0, delta=0.3)

    def test_refusals(self):
        x, y, _ = synth(1, TRUTH)
        with self.assertRaises(ValueError):
            SM.fit(x, y, [])
        with self.assertRaises(ValueError):
            SM.fit(x, y, start(TRUTH), lo=830, hi=832)
        with self.assertRaises(ValueError):
            SM.fit(x, y, [dict(TRUTH[0], shape='voigt')])
        with self.assertRaises(ValueError):
            SM.fit(x, y, [dict(TRUTH[0], fix=['pos', 'height', 'fwhm', 'eta'])], local_baseline='none')

    def test_what_a_fit_should_not_be_trusted_for(self):
        close = [dict(TRUTH[0], pos=840.0, fwhm=12.0, eta=0.5, height=800.0), dict(TRUTH[0], pos=843.0, fwhm=12.0, eta=0.5, height=500.0)]
        x, y, _ = synth(9, close)
        w = ' | '.join(SM.fit(x, y, [dict(c) for c in close]).warnings)
        self.assertIn('under half their width', w); self.assertIn('correlated at', w)
        w = ' | '.join(SM.fit(x, y, [dict(c) for c in close], resolution=4.0).warnings)
        self.assertIn('closer than the resolution', w)
        x, y, _ = synth(10, TRUTH[:2])                                       # two bands, one component: the residual says so
        self.assertTrue(any('long stretches of one sign' in m for m in SM.fit(x, y, start(TRUTH[:1])).warnings))
        weak = [dict(TRUTH[0], height=8.0)]
        x, y, _ = synth(11, weak)
        self.assertTrue(any('noise deviations high' in m for m in SM.fit(x, y, [dict(weak[0])]).warnings))

    def test_the_number_of_bands_the_data_support(self):
        cases = ((TRUTH, 3), (TRUTH[:1], 1),
                 ([dict(TRUTH[0], pos=840.0, eta=0.5), dict(TRUTH[0], pos=852.0, height=400.0, eta=0.5)], 2),      # one width apart
                 ([dict(TRUTH[0], pos=840.0, height=800.0, eta=0.5), dict(TRUTH[0], pos=847.5, height=200.0, eta=0.5)], 2))   # a shoulder, 0.6 of a width away
        for comps, n in cases:
            for seed in range(4):
                f, trials = SM.auto_fit(*synth(100 + seed, comps)[:2], max_n=6)
                self.assertEqual(len(f.comps), n, (n, trials))
        f, _t = SM.auto_fit(*synth(100, TRUTH)[:2], n=2)
        self.assertEqual(len(f.comps), 2)
        with self.assertRaises(ValueError):
            SM.auto_fit([1.0, 2.0, 3.0], [1.0, 2.0, 1.0])

    def test_an_envelope_of_broad_bands_comes_apart(self):
        # five overlapping O–H components 40–120 cm⁻¹ wide: the second derivative at a 6 cm⁻¹ smoothing sees none of them
        # (one band 435 wide was the fit), at 25–100 cm⁻¹ all five
        for seed, noise in ((0, 4.0), (1, 12.0)):
            x, y, _ = synth(seed, OH, lo=2700, hi=3800, dx=1.0, noise=noise)
            self.assertGreaterEqual(len(SM.find_bands(x, y)), 4)                # the seeds; the residual supplies a fifth where one is missed
            f, _trials = SM.auto_fit(x, y, local_baseline='linear')
            self.assertEqual(len(f.comps), 5, [(c['pos'], c['fwhm']) for c in f.comps])
            for c, t in zip(f.comps, OH):
                self.assertAlmostEqual(c['pos'], t['pos'], delta=6.0); self.assertAlmostEqual(c['fwhm'], t['fwhm'], delta=10.0)
            self.assertAlmostEqual(f.chi2, 1.0, delta=0.2)

    def test_a_ripple_is_no_component(self):
        # a significant ripple a hundredth of the band it rides on (atmospheric water on an infrared O–H band)
        x, y, _ = synth(2, OH[1:2], lo=2800, hi=3600, dx=1.0, noise=0.2)
        y = [v + 3.0 * math.sin(a / 2.5) for a, v in zip(x, y)]
        got = SM.find_bands(x, y)
        self.assertEqual(len(got), 1, got); self.assertAlmostEqual(got[0]['pos'], 3210.0, delta=40.0)

    def test_width_limit_and_linear_heights(self):
        x, y, _ = synth(3, OH, lo=2700, hi=3800, dx=1.0, noise=4.0)
        f, _t = SM.auto_fit(x, y, n=2, local_baseline='linear', max_fwhm=150.0)
        self.assertTrue(all(c['fwhm'] <= 150.0 + 1e-6 for c in f.comps))
        self.assertTrue(any('rests on the width limit' in w for w in f.warnings), f.warnings)
        start = [dict(c, height=1.0) for c in OH]                              # positions and widths right, heights unknown
        got = SM.linear_heights(x, y, start, 'linear')
        for c, t in zip(got, OH):
            self.assertAlmostEqual(c['height'], t['height'], delta=0.05 * t['height'] + 6)
        own = SM.fit(x, y, [dict(c, wmax=60.0) if k == 2 else dict(c) for k, c in enumerate(OH)], local_baseline='linear')
        self.assertLessEqual(own.comps[2]['fwhm'], 60.0 + 1e-6); self.assertEqual(own.comps[2]['at_limit'], 'max')

    def test_what_the_residual_still_holds(self):
        x, y, _ = synth(4, OH, lo=2700, hi=3800, dx=1.0, noise=4.0)
        f = SM.fit(x, y, [dict(c) for c in OH if c['pos'] != 3470.0], local_baseline='linear')      # one band left out
        r = SM.residual_bands(x, y, f)
        self.assertTrue(any(abs(q['pos'] - 3470.0) < 30 for q in r[:3]), r[:3])
        grown, _t = SM.auto_fit(x, y, local_baseline='linear')                                       # … and the automatic fit puts it back
        self.assertEqual([round(c['pos'], -1) for c in grown.comps], [3050.0, 3210.0, 3350.0, 3470.0, 3560.0])
        full = SM.fit(x, y, [dict(c) for c in OH], local_baseline='linear')
        self.assertEqual(SM.residual_bands(x, y, full), [])                                           # nothing left: nothing proposed
        d2 = SM.second_derivative(x, y, 50.0)
        self.assertEqual(len(d2), len(x)); self.assertGreater(d2[x.index(3350.0)], 0)                 # bands point up

    def test_a_band_driven_to_zero_is_taken_out(self):
        x, y, _ = synth(6, TRUTH)
        f = SM.fit(x, y, start(TRUTH) + [dict(TRUTH[0], pos=975.0, height=40.0, fwhm=6.0, fix=['pos', 'fwhm'])])   # a band where there is nothing
        g, dropped = SM.without_dead(x, y, f)
        if f.comps[-1]['height'] == 0.0:
            self.assertEqual((len(g.comps), [round(p) for p in dropped]), (3, [975])); self.assertIn('driven to zero height', g.warnings[0])
        else:
            self.assertEqual((len(g.comps), dropped), (4, []))
        same, none = SM.without_dead(x, y, SM.fit(x, y, start(TRUTH)))
        self.assertEqual((len(same.comps), none), (3, []))

    def test_evaluating_without_fitting(self):
        x, y, _ = synth(7, TRUTH)
        off = [dict(c, pos=c['pos'] + 1.0) for c in TRUTH]
        f = SM.fit(x, y, off, max_iter=0)
        self.assertEqual([c['pos'] for c in f.comps], [c['pos'] for c in off]); self.assertEqual(f.iterations, 0)
        self.assertGreater(f.chi2, SM.fit(x, y, off).chi2)
        self.assertTrue(all(c['esd']['pos'] for c in f.comps))                                         # uncertainties at the point given

    def test_evaluating_changes_nothing(self):
        # the model as it stood: the local baseline it ended on, widths outside a limit, unequal η under 'one η', and a
        # model with every parameter held
        x, y, _ = synth(8, TRUTH, base=lambda v: 300 + 0.4 * (v - 850))
        f = SM.fit(x, y, start(TRUTH), local_baseline='linear')
        bare = [SM._bare(c) for c in f.comps]
        self.assertGreater(SM.fit(x, y, bare, local_baseline='linear', max_iter=0).chi2, 100 * f.chi2)          # the baseline left at zero: another model
        for b0 in ((f.base[0], f.base[1], f.x0), (f.base[0] + 30 * f.base[1], f.base[1], f.x0 + 30)):           # … about any x
            e = SM.fit(x, y, bare, local_baseline='linear', max_iter=0, base0=b0)
            self.assertAlmostEqual(e.chi2, f.chi2, 6); self.assertAlmostEqual(e.base[0], f.base[0], 6)
        e = SM.fit(x, y, bare, local_baseline='linear', max_iter=0, base0=f.base, max_fwhm=5.0, share_eta=True, share_width=True)
        self.assertEqual([(c['fwhm'], c['eta']) for c in e.comps], [(c['fwhm'], c['eta']) for c in f.comps])
        held = [dict(c, fix=['pos', 'height', 'fwhm', 'eta']) for c in bare]
        e = SM.fit(x, y, held, local_baseline='none', max_iter=0)
        self.assertEqual((e.k, e.comps[0]['esd']['pos']), (0, None))
        with self.assertRaises(ValueError):
            SM.fit(x, y, held, local_baseline='none')

    def test_what_is_held_stays_under_one_width_and_one_shape(self):
        x, y, _ = synth(3, TRUTH)
        c = start(TRUTH); c[1] = dict(c[1], fwhm=10.0, eta=0.9, fix=['fwhm', 'eta'])
        f = SM.fit(x, y, c, share_eta=True)
        self.assertEqual([q['eta'] for q in f.comps], [0.9] * 3); self.assertEqual(f.comps[1]['fwhm'], 10.0); self.assertNotIn('eta', f.names)
        f = SM.fit(x, y, c, share_width=True)
        self.assertEqual([q['fwhm'] for q in f.comps], [10.0] * 3); self.assertNotIn('fwhm', f.names)
        f = SM.fit(x, y, start(TRUTH), share_width=True)
        self.assertEqual(len({round(q['fwhm'], 9) for q in f.comps}), 1); self.assertIn('fwhm', f.names)

    def test_bands_on_a_background_that_was_not_taken_off(self):
        x, y, _ = synth(3, TRUTH)
        for bg in (6000.0, 20000.0, 1e6):
            f, _t = SM.auto_fit(x, [v + bg for v in y], lo=780, hi=960)
            self.assertEqual([round(c['pos']) for c in f.comps], [831, 846, 915], bg)
            self.assertLess(f.chi2, 1.3)

    def test_an_edge_band_stays_when_the_fit_needs_it(self):
        # beside a strong band just outside the region the band on the edge IS the model of its foot: without it the real
        # bands are bent to make up for it (the 915 band came out one point wide). It stays, and is said to be what it is.
        x, y, _ = synth(4, TRUTH + [{'shape': 'pvoigt', 'pos': 1012.0, 'height': 3000.0, 'fwhm': 30.0, 'eta': 0.8}], lo=700, hi=1040)
        f, _t = SM.auto_fit(x, y, lo=780, hi=1000, local_baseline='const')
        real = [c for c in f.comps if c['pos'] < 950]
        self.assertEqual([round(c['pos']) for c in real], [831, 846, 915]); self.assertAlmostEqual(real[2]['fwhm'], 14.0, delta=1.5)
        self.assertTrue(any(c['pos'] > 990 for c in f.comps)); self.assertTrue(any('stands on the edge of the region' in w for w in f.warnings))

    def test_twins_envelopes_and_rulers(self):
        self.assertTrue(SM.twins({'pos': 3300.0, 'fwhm': 80.0}, {'pos': 3320.0, 'fwhm': 100.0}))
        self.assertFalse(SM.twins({'pos': 3300.0, 'fwhm': 30.0}, {'pos': 3310.0, 'fwhm': 280.0}))             # a narrow band on a broad one
        self.assertFalse(SM.twins({'pos': 3300.0, 'fwhm': 80.0}, {'pos': 3360.0, 'fwhm': 90.0}))
        x, y, _ = synth(0, OH, lo=2700, hi=3800, dx=1.0)
        one, _t = SM.auto_fit(x, y, n=1, local_baseline='linear')
        hint = SM.envelope_hint(x, y, one)
        self.assertGreaterEqual(len(hint['inside']), 2); self.assertLess(hint['cap'], 0.8 * one.comps[0]['fwhm'])
        five, _t = SM.auto_fit(x, y, local_baseline='linear')
        self.assertIsNone(SM.envelope_hint(x, y, five))
        from pxrd_review import spectro as SP
        oo, nu1, nu3 = SM.distance_rulers()
        self.assertEqual((oo['ticks'][0][0], oo['ticks'][-1][0]), SM.RULER_OO); self.assertGreater(oo['ticks'][0][1], 2700)
        for d, nu in oo['ticks'][::10]:
            self.assertAlmostEqual(SP.lib_d(nu), d, 3)
        for r, back in ((nu1, SP.bc_nu1), (nu3, SP.bc_nu3)):
            for d, nu in r['ticks'][::5]:
                self.assertAlmostEqual(back(nu), d, 4)

    def test_many_warnings_of_a_kind_are_counted(self):
        x, y, _ = synth(8, OH, lo=2700, hi=3800, dx=1.0)
        crowd = [dict(OH[2], pos=3300.0 + 12.0 * k, height=60.0, fwhm=100.0) for k in range(8)]         # eight bands 12 apart, 100 wide
        w = SM.fit(x, y, crowd, max_iter=0).warnings
        self.assertEqual(sum('under half their width' in m for m in w), 3)
        self.assertTrue(any(m == '4 more pairs of bands are closer than half their width' for m in w), w)

    def test_seeds_from_a_structure(self):
        structure = {'uranyl': [('U1', [1.76, 1.77], 1.765)], 'hbonds': [('O3', 'O7', 2.80, 'H'), ('O4', 'O8', 2.30, 'H')]}
        ticks = SM.expected_bands(structure)
        from pxrd_review import spectro as SP
        self.assertEqual([p for p, _l in ticks], [SP.bc_nu1_inv(1.765), SP.bc_nu3_inv(1.765), SP.lib_nu(2.80)])    # ν1, ν3, one O–H (2.30 Å is under the curve's range)
        self.assertEqual([round(p) for p, _l in ticks], [847, 928, 3403])
        self.assertIn('U1', ticks[0][1])

    def test_beside_the_manuscript(self):
        comps = [{'pos': 831.2}, {'pos': 846.4}, {'pos': 915.3}]
        c = SM.compare_stated(comps, [832.0, 846.0, 880.0])
        self.assertEqual([(s, round(p, 1)) for s, p, _d in c['matched']], [(832.0, 831.2), (846.0, 846.4)])
        self.assertEqual((c['not_fitted'], c['not_stated']), ([880.0], [915.3]))
        i = SM.inferred([{'pos': 840.0}, {'pos': 3400.0}, {'pos': 3605.0}, {'pos': 2300.0}])
        self.assertTrue(i[0]['distance'].startswith('U–O 1.77')); self.assertEqual(i[1]['distance'], 'O···O 2.80 Å')
        self.assertIn('free O–H', i[2]['distance']); self.assertEqual((i[3]['suggest'], i[3]['distance']), ([], ''))


class EndToEnd(unittest.TestCase):
    def test_prepare_fit_export(self):
        tmp = tempfile.mkdtemp()
        try:
            x, y, _ = synth(21, TRUTH, lo=300, hi=1300, base=lambda v: 400 + 0.3 * v, spikes=3)
            p = os.path.join(tmp, 'testite raman.txt')
            with open(p, 'w') as f:
                f.write(''.join('%.2f\t%.3f\n' % r for r in zip(x, y)))
            prep = SM.prepare(p, lo=780, hi=960, kind='raman', despike_on=True, baseline_kind='als')
            self.assertEqual(len(prep['x']), 361); self.assertGreaterEqual(len(SM.prepare(p, despike_on=True)['spikes']), 3)
            f, _trials = SM.auto_fit(prep['x'], prep['y'], local_baseline='linear')
            self.assertEqual([round(c['pos']) for c in f.comps], [831, 846, 915])
            out = SM.export(prep, [f], os.path.join(tmp, 'review_out'), 'testite_raman', {'baseline': 'als'}, word=True, xlsx=True, ticks=[(848.0, 'ν1 U1')])
            self.assertEqual(sorted(out), ['docx', 'json', 'svg', 'txt', 'xlsx'])
            self.assertTrue(all(os.path.getsize(v) > 200 for v in out.values()))
            text = open(out['txt'], encoding='utf-8').read()
            self.assertIn('UO2 ν1', text); self.assertIn('U–O 1.7', text)
            self.assertTrue(open(out['svg'], encoding='utf-8').read().startswith('<svg'))
            with self.assertRaises(ValueError):
                SM.prepare(p, lo=5000, hi=6000)
            with self.assertRaises(ValueError):
                SM.prepare(p, series=3)
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                self.assertEqual(SM.main([p, '--range', '780', '960', '--kind', 'raman', '--baseline', 'als', '--fit', 'auto', '--out', os.path.join(tmp, 'o2')]), 0)
            self.assertIn('models tried: 3 (BIC', buf.getvalue())
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_a_smoothed_spectrum_is_said_to_be_one(self):
        tmp = tempfile.mkdtemp()
        try:
            x, y, _ = synth(1, TRUTH, noise=0.0)
            p = os.path.join(tmp, 's.txt')
            with open(p, 'w') as f:
                f.write(''.join('%.2f\t%.6f\n' % r for r in zip(x, y)))
            self.assertTrue(any('smoothed or interpolated' in n for n in SM.prepare(p)['notes']))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    unittest.main()


class BackgroundHump(unittest.TestCase):
    def test_a_band_over_half_the_region_is_named_a_background(self):
        import math, random
        from pxrd_review import spectrum as S
        random.seed(3)
        xs = [200.0 + k for k in range(1001)]
        ys = [50 * math.exp(-((v - 700) / 400) ** 2) + 20 * math.exp(-((v - 500) / 6) ** 2) + random.gauss(0, 0.2) for v in xs]
        f = S.fit(xs, ys, [{'shape': 'gauss', 'pos': 700, 'height': 50, 'fwhm': 600, 'eta': 0.5}, {'shape': 'gauss', 'pos': 500, 'height': 20, 'fwhm': 10, 'eta': 0.5}])
        wide = [w for w in f.warnings if 'background hump' in w]
        self.assertEqual(len(wide), 1, f.warnings); self.assertIn('over half the region', wide[0]); self.assertIn('set a baseline', wide[0])
        g = S.fit(xs, ys, [{'shape': 'gauss', 'pos': 700, 'height': 50, 'fwhm': 600, 'eta': 0.5, 'fix': ['fwhm']}, {'shape': 'gauss', 'pos': 500, 'height': 20, 'fwhm': 10, 'eta': 0.5}], max_fwhm=600)
        self.assertTrue(any('background hump' in w for w in g.warnings))
