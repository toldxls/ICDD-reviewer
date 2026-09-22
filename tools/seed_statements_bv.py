"""Diagnosis accuracy of the bond-valence check sheet: the table the tool itself would print, spoiled in ONE known way.

    python3 tools/seed_statements_bv.py "<folders with .cif files>" OUT_DIR [TAG] [--limit N] [--jobs N]

For every .cif that computes, the anion × cation grid is written from the tool's own cells (so, unspoiled,
every cell agrees by construction), then spoiled, and the workbook's check sheet is read:

  none     as computed                       right = no cell and no Σ differs
  cell     one cell + 0.20 vu                right = THAT cell differs and no other; its column reads 'a distance, a ×n … or a mistyped valence'
  column   one cation's column × 0.85        right = its column reads 'the whole column differs' (columns of >= 3 cells)
  sum      one row's Σ mis-added by + 0.30   right = that Σ differs as ARITHMETIC, and no cell differs

No paper is read; nothing leaves the machine; workbooks go to OUT_DIR and are removed.
"""
import os, sys, glob, json, random, hashlib, argparse
import multiprocessing as mp
from concurrent.futures import ProcessPoolExecutor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT); sys.path.insert(0, os.path.join(ROOT, 'tests'))
from pxrd_review import bv_check as B

_CTX = mp.get_context('spawn')

def grid(st, result, cells, spoil=None, scale=None, sum_off=None):
    cats = [r[0].label for r in result if r[0].element != 'H']
    occ = {r[0].label: min(getattr(r[0], 'occ_total', 1.0) or 1.0, 1.0) for r in result}
    rows = [['Atom'] + cats + ['Σ']]
    for an in st.anions:
        if not any((an.label, c) in cells for c in cats):
            continue
        row = [an.label]; tot = 0.0
        for c in cats:
            segs = cells.get((an.label, c))
            if not segs:
                row.append(''); continue
            k = (scale or {}).get(c, 1.0)
            vals = [(s_ * k + (spoil or {}).get((an.label, c), 0.0), nd, na) for s_, nd, na in segs]
            row.append(', '.join('%.2f%s' % (v, B._mark(nd, na)) for v, nd, na in vals))
            tot += sum(round(v, 2) * (na if isinstance(na, int) else 1) * occ.get(c, 1.0) for v, nd, na in vals)   # a row Σ as the tool forms it: weighted by the cation's occupancy
        rows.append(row + ['%.2f' % (tot + ((sum_off or {}).get(an.label, 0.0)))])
    return rows

def read(path):
    import openpyxl
    from xl_eval import Book
    b = Book(path)
    if 'check' not in b.wb.sheetnames:
        return None
    wc = b.wb['check']; out = {'cells': {}, 'sums': {}, 'cols': {}}
    block = 'cells'
    for r in range(4, wc.max_row + 1):
        a = wc.cell(r, 1).value
        if a == 'table' and wc.cell(r, 2).value == 'Σ of':
            block = 'sums'; continue
        if a == 'cation column':
            block = 'cols'; continue
        if a is None or (isinstance(a, str) and a.startswith('WHERE')):
            continue
        if block == 'cells' and isinstance(a, int):
            out['cells'][(wc.cell(r, 2).value, wc.cell(r, 3).value)] = wc.cell(r, 9).value
        elif block == 'sums' and isinstance(a, int):
            out['sums'][wc.cell(r, 3).value] = (wc.cell(r, 8).value, str(wc.cell(r, 9).value or ''))
        elif block == 'cols' and isinstance(a, str) and wc.cell(r, 5).value:
            try:
                out['cols'][a] = b.value('check', 'E%d' % r)
            except Exception as e:
                out['cols'][a] = 'ERR %s' % e
    return out

def _run_one(job):
    cif, base, out_dir = job
    try:
        st = B.Structure(cif); P = B.Params(prefer='gh')
        result, anion_sum, cells, hbonds = B.compute(st, P)
    except Exception:
        return None
    cats = [r[0] for r in result if r[0].element != 'H']
    full = {c.label for c in cats if (getattr(c, 'occ_total', 1.0) or 1.0) >= 0.98 and '/' not in c.label
            and len([sp for sp in (getattr(c, 'species', None) or []) if (getattr(sp, 'occ', 0) or 0) > 0.02]) <= 1}   # one species, fully occupied: no convention excuses a scaled column
    single = [(a, c) for (a, c), segs in cells.items() if len(segs) == 1 and c in full and segs[0][0] >= 0.10]
    if len(single) < 3:
        return None
    rng = random.Random(int(hashlib.sha1(base.encode()).hexdigest()[:8], 16)); runs = []
    def sheet(tag, **kw):
        path = os.path.join(out_dir, '%s__%s.xlsx' % (base[:40], tag))
        B.write_xlsx(st, P, result, anion_sum, cells, hbonds, path, tables=[grid(st, result, cells, **kw)], params_label='Gagné & Hawthorne 2015')
        try:
            return read(path)
        finally:
            os.remove(path)
    try:
        s0 = sheet('none')
        if s0 is None:
            return None
        bad0 = [k for k, v in s0['cells'].items() if v == 'differs'] + [k for k, v in s0['sums'].items() if v[0] == 'differs']
        runs.append({'kind': 'none', 'class': 'RIGHT' if not bad0 else 'FALSE-ALARM', 'said': [str(x) for x in bad0[:4]]})
        if bad0:
            return {'base': base, 'runs': runs}                          # its own table does not pass: nothing seeded on top is readable
        a, c = rng.choice(single)
        s1 = sheet('cell', spoil={(a, c): 0.20})
        d1 = [k for k, v in s1['cells'].items() if v == 'differs']
        col = s1['cols'].get(c, '')
        runs.append({'kind': 'cell', 'target': '%s–%s' % (a, c), 'said': [str(x) for x in d1[:4]] + [col[:60]],
                     'class': ('RIGHT' if 'mistyped valence' in col or not col else 'PARTIAL') if d1 == [(a, c)] else 'MISSED' if (a, c) not in d1 else 'EXTRA'})
        ncol = {}
        for (a_, c_), v in s0['cells'].items():
            if v == 'agrees':
                ncol[c_] = ncol.get(c_, 0) + 1
        wide = [c_ for c_, n_ in ncol.items() if n_ >= 3 and c_ in full]
        if wide:
            c2 = rng.choice(wide); s2 = sheet('col', scale={c2: 0.85}); col = s2['cols'].get(c2, '')
            other = [k for k, v in s2['cells'].items() if v == 'differs' and k[1] != c2]
            runs.append({'kind': 'column', 'target': c2, 'said': [col[:70]] + [str(x) for x in other[:3]],
                         'class': 'RIGHT' if 'whole column differs' in col and not other else 'PARTIAL' if 'whole column' in col else 'WRONG-KIND' if col.startswith('PROBLEM') else 'MISSED'})
        rows_with_sum = [an for an in s0['sums'] if any(k[0] == an for k in s0['cells'])]
        if rows_with_sum:
            a3 = rng.choice(rows_with_sum); s3 = sheet('sum', sum_off={a3: 0.30})
            v = s3['sums'].get(a3, ('', '')); d3 = [k for k, x in s3['cells'].items() if x == 'differs']
            runs.append({'kind': 'sum', 'target': a3, 'said': [v[0], v[1][:60]] + [str(x) for x in d3[:3]],
                         'class': 'RIGHT' if v[0] == 'differs' and 'ARITHMETIC' in v[1] and not d3 else 'PARTIAL' if v[0] == 'differs' else 'MISSED'})
    except Exception as e:
        return {'base': base, 'fail': str(e)[:140]}
    return {'base': base, 'runs': runs}

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('dirs'); ap.add_argument('out_dir'); ap.add_argument('tag', nargs='?', default=''); ap.add_argument('--limit', type=int); ap.add_argument('--jobs', type=int, default=0)
    a = ap.parse_args(argv)
    seen = {};
    for d in a.dirs.split(','):
        for p in sorted(glob.glob(os.path.join(d, '**', '*.cif'), recursive=True)):
            seen.setdefault(os.path.basename(p), p)
    wb = os.path.join(a.out_dir, 'seedbv' + a.tag); os.makedirs(wb, exist_ok=True)
    jobs = [(p, b, wb) for b, p in sorted(seen.items())][:a.limit]
    with ProcessPoolExecutor(max_workers=a.jobs or os.cpu_count() or 4, mp_context=_CTX) as pool:
        recs = [r for r in pool.map(_run_one, jobs, chunksize=2) if r]
    fails = [r for r in recs if 'fail' in r]; recs = [r for r in recs if 'runs' in r]
    CL = ('RIGHT', 'PARTIAL', 'MISSED', 'EXTRA', 'WRONG-KIND', 'FALSE-ALARM')
    L = ['DIAGNOSIS ACCURACY of the bond-valence check sheet%s — %d structures' % ((' ' + a.tag) if a.tag else '', len(recs)), '',
         '  %-8s %5s  ' % ('seeded', 'n') + ' '.join('%11s' % c for c in CL)]
    for kind in ('none', 'cell', 'column', 'sum'):
        xs = [x for r in recs for x in r['runs'] if x['kind'] == kind]
        if xs:
            L.append('  %-8s %5d  ' % (kind, len(xs)) + ' '.join('%10.0f%%' % (100.0 * sum(1 for x in xs if x['class'] == c) / len(xs)) for c in CL))
    L += ['', 'WORKLIST']
    for r in recs:
        for x in r['runs']:
            if x['class'] != 'RIGHT':
                L.append('  %-11s %-7s %-26s %-12s %s' % (x['class'], x['kind'], r['base'][:26], x.get('target', ''), ' | '.join(x['said'])[:150]))
    L += ['', 'failed: %d' % len(fails)] + ['  %s %s' % (r['base'], r['fail']) for r in fails[:8]]
    json.dump(recs, open(os.path.join(a.out_dir, 'seed_statements_bv%s.json' % a.tag), 'w'), default=str)
    open(os.path.join(a.out_dir, 'seed_statements_bv%s.txt' % a.tag), 'w', encoding='utf-8').write('\n'.join(L) + '\n')
    print('\n'.join(L[:60]))

if __name__ == '__main__':
    main()
