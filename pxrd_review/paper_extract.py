#!/usr/bin/env python3
"""
paper_extract — read a mineral description .pdf for what the Tables mode needs, so the tabs can be
filled from the paper and its own calculations re-done its own way:

  • the analytical table: mean wt% per constituent, with range, s.d. and standard where printed;
  • the basis the authors normalised the formula on ("on the basis of 21 O apfu", "Si + Al = 4"),
    and how they treated what the probe cannot give (H2O by difference / by stoichiometry / for
    charge balance, Fe3+/Fe2+ by charge balance, CO2 from the structure …) — as the sentences;
  • the optical data: the mean refractive index (n, or the mean of α β γ / ω ε) and the measured
    and calculated densities;
  • the bond-valence parameter set the paper cites (Gagné & Hawthorne / Brese & O'Keeffe / Brown &
    Altermatt, Burns for U6+, Ferraris & Ivaldi for hydrogen bonds);
  • the powder table: observed (Iobs, dobs) and calculated (dcalc, Icalc, hkl) lines.

    python3 -m pxrd_review.paper_extract paper.pdf [--out DIR]     # prints what it found, writes the data files

Everything is best effort and every value carries the line or sentence it was read from, so the
reviewer can see why. Nothing here decides — the EPMA / GD / PXRD tools take these as inputs.
"""
import os, re, sys, json, math, argparse
from collections import OrderedDict

from pxrd_review import epma as EP

# ----------------------------------------------------------------------------- pdf text and lines

def page_lines(page, x_lo=None, x_hi=None):
    """The page's words grouped into lines (by baseline, ±3 pt), each line's words by x."""
    words = page.get_text('words')
    words = [w for w in words if (x_lo is None or w[0] >= x_lo) and (x_hi is None or w[2] <= x_hi)]
    words.sort(key=lambda w: ((w[1] + w[3]) / 2, w[0]))
    # pass 1: words on one baseline (centres within 1.2 pt) — a label and its numbers stay together
    # even when the other page column's line sits 3 pt off
    tight = []
    for w in words:
        yc = (w[1] + w[3]) / 2
        if tight and abs(tight[-1]['y'] - yc) <= 1.2:
            tight[-1]['w'].append(w)
        else:
            tight.append({'y': yc, 'top': w[1], 'bot': w[3], 'w': [w]})
    # pass 2: a subscript ('P2O5') deepens a word's box, a superscript ('Fe3+') raises it: adjacent
    # groups whose top, centre or bottom agree within 3 pt are one line
    lines = []
    for g in tight:
        if lines and (abs(lines[-1]['y'] - g['y']) <= 3.0 or abs(lines[-1]['top'] - g['top']) <= 3.0 or abs(lines[-1]['bot'] - g['bot']) <= 3.0):
            lines[-1]['w'] += g['w']
        else:
            lines.append(g)
    for ln in lines:
        ln['w'].sort(key=lambda w: w[0])
        # a subscript digit set as its own word ('B2O' 'a' + '3' below the baseline) is glued into
        # its parent; a superscript digit (a numbered footnote) is dropped — neither is a value
        out = []
        for w in ln['w']:
            if re.fullmatch(r'\d+\.\d+[a-d*†‡§]{1,2}', w[4]):
                w = (w[0], w[1], w[2], w[3], re.sub(r'[a-d*†‡§]+$', '', w[4]))     # '11.41c', '3.19*': a footnote mark on the value
            if out and re.fullmatch(r'\d{1,2}', w[4]) and w[0] - out[-1][2] < 3.5 and (w[3] - w[1]) < 0.85 * (out[-1][3] - out[-1][1]) \
                    and re.match(r'^[A-Z(]', out[-1][4]):
                pc = (out[-1][1] + out[-1][3]) / 2
                if w[1] >= pc - 1:                                  # below the centre: a subscript
                    m = re.match(r'^(.*?)([a-d*†‡§]*)$', out[-1][4])
                    out[-1] = (out[-1][0], out[-1][1], w[2], max(out[-1][3], w[3]), m.group(1) + w[4] + m.group(2))
                    continue
                if w[3] <= pc + 1:                                  # above the centre: a footnote number
                    continue
            out.append(w)
        ln['w'] = out
    return lines

_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
_MC_FALLBACK = '{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback'

def _para_text(p):
    """A paragraph's text with tracked changes accepted — w:ins kept, w:del / w:moveFrom dropped —
    the way refs_check reads a manuscript. python-docx's Paragraph.text joins only the direct-child
    runs, so a sentence inserted under Track Changes (the normal state of a manuscript under
    revision) would vanish from the formula / basis / optics readers. The legacy twin of a text box
    (mc:Fallback) is skipped so its text is not read twice."""
    out = []
    for el in p.iter():
        tag = el.tag
        if tag == _W + 't':
            if any(a.tag in (_W + 'del', _W + 'moveFrom', _MC_FALLBACK) for a in el.iterancestors()):
                continue
            out.append(el.text or '')
        elif tag == _W + 'tab':
            out.append(' ')
        elif tag in (_W + 'br', _W + 'cr'):
            out.append('\n')
    return ''.join(out)

def _docx_body(path):
    """A manuscript .docx in body order: [('p', text) | ('t', python-docx Table)]."""
    from docx import Document
    from docx.table import Table
    doc = Document(path); out = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            out.append(('p', _para_text(child)))
        elif tag == 'tbl':
            out.append(('t', Table(child, doc)))
    return out

def _docx_cell_text(tc):
    """A Word table cell as the reader wants it: subscripts inline ('SiO2'), a superscript of one
    or two characters (a footnote mark, a charge) dropped, the runs joined."""
    out = []
    for para in tc.iter(_W + 'p'):
        for r in para.iter(_W + 'r'):
            t = ''.join(x.text or '' for x in r.iter(_W + 't'))
            rpr = r.find(_W + 'rPr'); va = rpr.find(_W + 'vertAlign') if rpr is not None else None
            if va is not None and va.get(_W + 'val') == 'superscript' and len(t.strip()) <= 2:
                continue
            out.append(t)
        out.append(' ')
    return re.sub(r'\s+', ' ', ''.join(out)).strip()

def _grid_lines(caption, rows, y0=70.0):
    """A table as lines of words, the shape page_lines gives a pdf page: the caption first, then one
    line per row with each cell's words at its column's x. rows: [[(column, text), …], …].
    Column positions as a typeset table would have them: each column as wide as its widest cell, so
    'h k l' sit close together and a range column is wide. The gutter is 8 pt: with a wider one a
    signed two-digit index ('-10', three characters) pushes the 'h' and 'k' header words past the
    30 pt within which _powder_columns reads 'h k l' as one column."""
    width = {}
    for cells in rows:
        for col, t in cells:
            width[col] = max(width.get(col, 0), 6.0 * len(t) + 8.0)
    col_x = {}; x = 40.0
    for col in range(max(width) + 1 if width else 0):
        col_x[col] = x; x += width.get(col, 20.0)
    def line_of(cells, y):                                              # cells: [(column, text)]
        ws = []
        for col, text in cells:
            x = col_x.get(col, 40.0 + 60.0 * col)
            for tok in text.split():
                ws.append((x, y - 8.0, x + 6.0 * len(tok), y + 2.0, tok)); x += 6.0 * (len(tok) + 1)
        return {'y': y, 'top': y - 8.0, 'bot': y + 2.0, 'w': ws}
    lines = []; y = y0
    if caption:
        lines.append(line_of([(0, caption)], y)); y += 14.0
    for cells in rows:
        lines.append(line_of(cells, y)); y += 14.0
    return lines

def _empty_line(y):
    return {'y': y, 'top': y - 8.0, 'bot': y + 2.0, 'w': []}

def _docx_pages(path):
    """A manuscript .docx as 'pages', one per table: the caption paragraph before it as the first
    line, then one line per row with each cell's words at its column's x — the shape page_lines
    gives a pdf page, so the table readers run on a manuscript unchanged."""
    pages = []; prev = ''
    for kind, item in _docx_body(path):
        if kind == 'p':
            if item.strip():
                prev = item.strip()
            continue
        rows = []
        for row in item.rows:
            seen = []; cells = []
            for col, c in enumerate(row.cells):
                if any(c._tc is x for x in seen):
                    continue                                            # a merged cell repeats across its span
                seen.append(c._tc); t = _docx_cell_text(c._tc)
                if t:
                    cells.append((col, t))
            if cells:
                rows.append(cells)
        lines = _grid_lines(prev if re.match(r'^(Table|TABLE|Tab\.|Таблица)\s*\d', prev) else '', rows)
        prev = ''
        if len(lines) >= 2:
            pages.append(lines)
    return pages

_pages_reader = None
_pages_mode = 'fallback'

def set_pages_reader(fn, mode='fallback'):
    """Install a second pdf page reader (fn(path) -> pages as page_lines gives them); None removes
    it. mode 'fallback' (the default): the table readers try fitz's page first and the other
    reader's page only where fitz read no table; 'replace': every page comes from the other reader.
    A .docx is never routed through it. A hook only — the tool ships no second reader: a local
    layout model (docling) was measured on the corpus in 0.5.6 and read the tables worse than the
    pdf text on both the powder and the coordinates tables, and was removed."""
    global _pages_reader, _pages_mode
    _pages_reader = fn; _pages_mode = mode

def _pages(path):
    """The document's pages as lines of words: a pdf's pages, or a manuscript .docx's tables."""
    if path.lower().endswith('.docx'):
        return _docx_pages(path)
    if _pages_reader is not None and _pages_mode == 'replace':
        return _pages_reader(path)
    import fitz
    with fitz.open(path) as doc:
        return [page_lines(page) for page in doc]

def text_of(pdf):
    if pdf.lower().endswith('.docx'):                                   # a manuscript: its paragraphs, then its tables' text
        parts = []; cells = []
        for kind, item in _docx_body(pdf):
            if kind == 'p':
                parts.append(item)
            else:
                for row in item.rows:
                    cells.append(' '.join(_docx_cell_text(c._tc) for c in row.cells))
        t = ' '.join(parts + cells).replace('þ', '+')                   # prose first: the sentence readers prefer it to a table's footnote
    else:
        import fitz
        doc = fitz.open(pdf)
        t = ' '.join(page.get_text() for page in doc).replace('þ', '+')   # a journal font prints '+' as 'þ'
    t = re.sub(r'-\n(?=[a-z])', '', t)                     # de-hyphenate line breaks
    t = re.sub(r'(?<=[A-Za-z\)])\s*¼\s*(?=\d)', ' = ', t)   # a journal font that prints '=' as '¼' ("O ¼ 32")
    return re.sub(r'\s+', ' ', t)

# ----------------------------------------------------------------------------- the analytical table

_CONST = re.compile(r'^(\(NH4\)2O|H2O[+\-]?|H2O\+?|CO2|SO3|SO2|HS|[A-Z][a-z]?\d*O\d*|[A-Z][a-z]|[A-Z])[*†‡§¹²³]*$')
_TOTAL = re.compile(r'^(Total:?|Sum:?|Σ|Сумма|Итого|[-–−]?O\s*[=≡]\s*(F|Cl|S|F,Cl|Cl,F)(,Cl)?)$', re.I)
_NUM = re.compile(r'^[-–−]?\d+\.\d+$|^\d+$')
_NUM_ESD = re.compile(r'^(\d+(?:\.\d+)?)\((\d+(?:\.\d+)?)\)$')
_RANGE1 = re.compile(r'^(\d+\.\d+)\s*[-–—]\s*(\d+\.\d+)$')
_NA = re.compile(r'^(n\.?d\.?|b\.?d\.?l?\.?|[-–—]|bdl|nd|n/a)$', re.I)

_CYR = str.maketrans('ОСНРКВАЕМТаеорсух', 'OCHPKBAEMTaeopcyx')       # Cyrillic lookalikes in a Russian-typeset 'Na2О'

def _constituent_ok(tok):
    t = re.sub(r'\([^)]*\)$', '', tok.translate(_CYR))               # 'Fe2O3(tot)', 'H2O(calc)': a qualifier
    t = re.sub(r'[*†‡§¹²³]+$', '', t)
    if re.fullmatch(r'(?:S|Se|Te|F|Cl|Br|I)2?[\-–−]', t):
        t = re.sub(r'\d?[\-–−]$', '', t)                           # 'S2–', 'Cl–': the anion
    if _TOTAL.match(t):
        return t, 'total'
    m = _CONST.match(t)
    if m:
        try:
            EP.parse_constituent(m.group(1).replace('(NH4)2O', 'N2H8O').rstrip('+-'))
        except ValueError:
            m = None                                             # 'Sb' is antimony; 'Xa' is nothing
    if not m and re.fullmatch(r'.+[a-d]', t):
        m = _CONST.match(t[:-1])                                 # a footnote letter: 'SiO2a', 'Fe2O3b'
    if not m and re.fullmatch(r'.+\d\)', t):
        m = _CONST.match(t[:-2])                                 # a footnote number: 'V2O31)', 'MnO1)'
    if not m:
        return None, None
    c = m.group(1)
    m2 = re.fullmatch(r'([A-Z][a-z]?)2O', c)
    if m2 and m2.group(1) not in ('Na', 'K', 'Li', 'Rb', 'Cs', 'Tl', 'Ag', 'Cu', 'H', 'N'):
        usual = EP._USUAL_OXIDE.get(m2.group(1), '')        # the O count was a lost subscript: Fe2O -> Fe2O3, P2O -> P2O5
        c = usual if usual.startswith(m2.group(1) + '2O') else m2.group(1) + '2O3'
    m3 = re.fullmatch(r'(Na|K|Li|Rb|Cs|Tl|Ag)O', c)
    if m3:
        c = m3.group(1) + '2O'                              # the 2 was a lost subscript
    try:
        EP.parse_constituent('N2H8O' if c == '(NH4)2O' else c.rstrip('+-'))
    except ValueError:
        return None, None
    return c, 'constituent'

def _numbers(tokens):
    """Fold 'a – b' triples into ranges; return [(kind, value)] with kind num | range | text."""
    out = []; i = 0
    while i < len(tokens):
        t = tokens[i].replace('−', '-').replace('–', '-').replace('—', '-')
        m = _RANGE1.match(t)
        if m:
            out.append(('range', (float(m.group(1)), float(m.group(2))))); i += 1; continue
        if _NUM.match(t) and i + 2 < len(tokens) and tokens[i + 1] in ('-', '–', '—', '−') and _NUM.match(tokens[i + 2].replace('−', '-')):
            out.append(('range', (float(t), float(tokens[i + 2].replace('−', '-'))))); i += 3; continue
        if _NUM.match(t):
            out.append(('num', float(t))); i += 1; continue
        if re.fullmatch(r'\(\d+\.\d+\)', t):
            out.append(('num', float(t[1:-1]))); i += 1; continue      # '(13.16)': a derived value (total S as SO3) — still the cell's value
        m = _NUM_ESD.match(t)
        if m:                                               # '38.23(58)': the value, then its s.d. in the last digits
            v, e = m.group(1), m.group(2)
            sd = float(e) if '.' in e else (int(e) * 10 ** (-(len(v) - v.index('.') - 1)) if '.' in v else float(e))   # '29(3)': an esd on an integer
            out.append(('num', float(v))); out.append(('esd', round(sd, 4))); i += 1; continue
        if _NA.match(t):
            out.append(('na', None)); i += 1; continue
        out.append(('text', tokens[i])); i += 1
    return out

_STOP = {'with', 'the', 'and', 'of', 'in', 'is', 'are', 'to', 'a', 'an', 'for', 'by', 'on', 'from', 'that', 'this', 'as', 'at', 'was', 'were', 'which', 'it'}

def _row_at(ws, x_col):
    """One table row in a line of words: the first constituent token (at x_col when known) that
    numbers follow. Two-column pages merge body text into the line, so the cells stop at the last
    number or at a short name within 130 pt of it. -> (constituent, kind, values, x0) or None."""
    merged = []; k = 0                                   # 'S' 'b' printed as two words -> 'Sb'
    while k < len(ws):
        w = ws[k]
        if k + 1 < len(ws) and re.fullmatch(r'[A-Z]', w[4]) and re.fullmatch(r'[a-z]', ws[k + 1][4]) and ws[k + 1][0] - w[2] < 3 \
                and _constituent_ok(w[4] + ws[k + 1][4])[0]:
            merged.append((w[0], w[1], ws[k + 1][2], w[3], w[4] + ws[k + 1][4])); k += 2
        else:
            merged.append(w); k += 1
    ws = merged
    toks = [w[4] for w in ws]
    for k, w in enumerate(ws):
        c, kind = _constituent_ok(toks[k])
        if c is None and toks[k] in ('O', '-O', '–O', '−O') and k + 2 < len(toks) and toks[k + 1] in ('=', '≡'):
            c, kind = 'O=' + toks[k + 2], 'total'; rest = ws[k + 3:]
        elif c is not None:
            rest = ws[k + 1:]
        else:
            continue
        if x_col is not None and abs(w[0] - x_col) > 30:
            continue
        vals = _numbers([r[4] for r in rest])
        if not any(kk in ('num', 'range') for kk, _ in vals):
            continue
        # cells end at the last number plus a short name (the standard): body text from the other
        # page column — a comma, a stop word, more than three words — is not a cell
        last_num_x = max(r[2] for r, (kk, _) in zip(rest, vals) if kk in ('num', 'range', 'na'))
        keep = []; words = 0
        for r, (kk, v) in zip(rest, vals):
            if kk in ('num', 'range', 'na'):
                keep.append((kk, v)); words = 0; continue
            if kk != 'text':
                keep.append((kk, v)); continue
            if not keep and re.fullmatch(r'[a-d*†‡§¹²³]{1,2}', v):
                continue                                    # a footnote mark printed as its own word: 'TiO2 a 15.36'
            if r[0] > last_num_x + 130 or words >= 3 or v.lower() in _STOP or v.endswith((',', '.', ';', ':')):
                break
            keep.append((kk, v)); words += 1
        # a second table block on the same line (side-by-side tables) starts at the next constituent
        cut = len(keep)
        for idx, (kk, v) in enumerate(keep[1:], 1):
            if kk == 'text' and _constituent_ok(v)[0] and idx + 1 < len(keep) and keep[idx + 1][0] in ('num', 'range'):
                cut = idx; break
        if not any(kk in ('num', 'range', 'na') for kk, _ in keep[:cut]):
            continue                                            # 'A total of 16 scans …': prose, not the Total row — try the next token
        return c, kind, keep[:cut], w[0]
    return None

_MEANROW = re.compile(r'^(mean|average|aver\.?|avg\.?|среднее)[:.]?$', re.I)
_STATROW = re.compile(r'^(range|s\.?d\.?|σ|min\.?|max\.?|esd|standard|stdev|st\.?dev|n|apfu|wt\.?%)', re.I)

def _numlike(t):
    """'12.3', '29(3)', '(13.16)', '−0.5': a value cell."""
    t = t.replace('−', '-')
    return bool(_NUM.match(t) or _NUM_ESD.match(t) or re.fullmatch(r'\(\d+\.\d+\)', t))

def _num_x(w):
    t = w[4].replace('−', '-')
    if not _numlike(t):
        return None
    return float(re.match(r'-?\d+\.?\d*', t.lstrip('(')).group(0)), (w[0] + w[2]) / 2

def _transposed(lines, pno, name=''):
    """The other layout: constituents across a header line ('Constituent Nb2O5 MgO FeOa MnO TiO2
    Total'), the analyses down the rows, a Mean/Average row (else the rows are averaged). Cells
    are matched to the constituent columns by x. -> a candidate like epma_table's, or None."""
    for i, ln in enumerate(lines):
        ws = ln['w']; cols = []
        for w in ws:
            c, kind = _constituent_ok(w[4])
            if kind == 'constituent' and c != 'O':
                cols.append((c, (w[0] + w[2]) / 2))
        if len(cols) < 3 or len(cols) < 0.5 * len(ws):
            continue
        if any(re.fullmatch(r'x|y|z|U ?eq|U ?iso|s\.o\.f\.?|occ\.?|Wyck\w*|Site|Atom', w[4]) for w in ws):
            continue                                                            # a structure table
        oxide_els = {m_.group(0) for m_ in (re.match(r'[A-Z][a-z]?', c) for c, _ in cols if re.search(r'O\d*$', c) and c != 'O') if m_}
        if sum(1 for c, _ in cols if re.search(r'[A-Za-z]\d*O\d*$', c) and not c.startswith('H2O')) >= 3:
            cols = [(c, x) for c, x in cols if not (re.fullmatch(r'[A-Z][a-z]?', c) and c in oxide_els and c not in ('F', 'Cl', 'Br', 'I', 'S', 'Se', 'Te'))]   # 'CaO MgO MnO … Ca Mg Mn': the apfu columns beside the oxides
        n_twin = len(cols)
        seen_ = set(); cols = [t for t in cols if not (t[0] in seen_ or seen_.add(t[0]))]            # 'H2O*' and the apfu 'H2O': the first column of a constituent
        if len(cols) < 3 or len(cols) < 0.5 * n_twin:
            continue                                                            # mostly repeats ('Na Ca K Na OH− F− Cl− … Na Ca K Na'): a site table, not a header
        total_x = next((((w[0] + w[2]) / 2) for w in ws if _TOTAL.match(w[4])), None)
        rows = []; mean_row = None; mean_rows = []; j = i + 1; gap = 0
        while j < len(lines) and j < i + 45:
            wsj = lines[j]['w']; toks = [w[4] for w in wsj]
            nums = [nx for nx in (_num_x(w) for w in wsj) if nx]
            if len(nums) >= max(3, len(cols) - 1):
                vals = {}
                for v, x in nums:
                    c, cx = min(cols, key=lambda t: abs(t[1] - x))
                    if abs(cx - x) < 25 and c not in vals:
                        vals[c] = v
                tot = next((v for v, x in nums if total_x is not None and abs(x - total_x) < 25), None)
                if len(vals) >= 3:
                    label = ' '.join(t for t in toks if not _num_x((0, 0, 0, 0, t)))[:60]
                    if toks and _MEANROW.match(toks[0]):
                        mean_rows.append((vals, tot, label))
                        if mean_row is None:
                            mean_row = (vals, tot)
                    elif not (toks and _STATROW.match(toks[0])):
                        rows.append((vals, tot, label))
                gap = 0
            elif rows or mean_row:
                gap += 1
                if gap > 2:
                    break
            j += 1
        if mean_row is None and not rows:
            continue
        if mean_row:
            vals, tot = mean_row
        else:
            vals = {c: round(sum(r[0][c] for r in rows if c in r[0]) / max(1, sum(1 for r in rows if c in r[0])), 3) for c, _ in cols if any(c in r[0] for r in rows)}
            tots = [r[1] for r in rows if r[1] is not None]; tot = round(sum(tots) / len(tots), 2) if tots else None
        labelled_rows = [(r_[2], {('N2H8O' if c == '(NH4)2O' else c.rstrip('+-')): v for c, v in r_[0].items()}) for r_ in (mean_rows if mean_rows else rows)]
        out = [{'constituent': 'N2H8O' if c == '(NH4)2O' else c.rstrip('+-'), 'mean': v, 'range': None, 'sd': None, 'standard': None} for c, v in vals.items()]
        tot_guess = sum(vals.values())
        alts = [{('N2H8O' if c == '(NH4)2O' else c.rstrip('+-')): v for c, v in r_[0].items()} for r_ in (mean_rows[1:] if len(mean_rows) > 1 else rows)]
        cap = _caption(lines, i)
        return {'rows': out, 'total': tot, 'header': ' '.join(w[4] for w in ws)[:200], 'page': pno + 1, 'n': len(out),
                'score': len(out) + 4 + (5 if 88 <= tot_guess <= 104 else 0) + _caption_score(cap, name), 'transposed': True, 'alts': alts, 'caption': cap[:160],
                'labelled_rows': labelled_rows, 'head_cells': [(w[4], (w[0] + w[2]) / 2) for w in ws]}
    return None

def _drop_totals(rows):
    """A constituent listed twice with another oxide of the same element between or after
    (FeO total, then FeO + Fe2O3 from Mössbauer or charge balance): the first is the total."""
    els = {}
    for r in rows:
        if _parses(r['constituent']):
            els.setdefault(EP.parse_constituent(r['constituent']).element, []).append(r['constituent'])
    out = list(rows)
    for el, cs in els.items():
        if len(cs) >= 3 and len(set(cs)) >= 2:
            first = next(r for r in out if r['constituent'] == cs[0])
            later = [r for r in out if r is not first and _parses(r['constituent']) and EP.parse_constituent(r['constituent']).element == el]
            x0 = (first.get('xs') or [None])[0]
            if x0 is None or any(any(abs(x - x0) <= 30 for x in (r.get('xs') or [])) for r in later):
                out.remove(first)                                  # the split has a value in the total's own column: the total goes
    return out

_PROSE_CONST = r'(\(NH4\)2O|H2O[+\-]?|CO2|SO3|[A-Z][a-z]?\d*O\d*|[A-Z][a-z]?)'
_PROSE_PAREN = r'(?:\s*\((?:\d+|\d+\.\d+\s*[–\-]\s*\d+\.\d+(?:[,;]\s*\d+\.\d+)?)\))?'          # '(8)' an esd; '(8.03–9.01, 0.25)' the range and s.d.
_PROSE_CORE = _PROSE_CONST + r'\s*=?\s*(\d+\.\d+)' + _PROSE_PAREN + r'\s*(?:wt\.?\s*%)?\s*'
_PROSE_SEP = r'(?:[,;]|\band\b|(?=\s+(?:\(NH4\)2O|H2O|CO2|SO3|[A-Z][a-z]?\d*O\d*|[A-Z][a-z]?)\s*\d+\.\d+))\s*(?:and\s+)?'
_PROSE_ITEM = re.compile(_PROSE_CORE)
_PROSE_RUN = re.compile('(?:' + _PROSE_CORE + _PROSE_SEP + '){3,}' + _PROSE_CORE)
_PROSE_CORE_R = r'(\d+\.\d+)' + _PROSE_PAREN + r'\s*(?:wt\.?\s*%\s*)?' + _PROSE_CONST + r'(?![a-z])\s*'   # '0.03 Na2O, 3.70 K2O': the value first
_PROSE_SEP_R = r'(?:[,;]|\band\b)\s*(?:and\s+)?'
_PROSE_ITEM_R = re.compile(_PROSE_CORE_R)
_PROSE_RUN_R = re.compile('(?:' + _PROSE_CORE_R + _PROSE_SEP_R + '){3,}' + _PROSE_CORE_R)

def prose_table(text):
    """A composition given in the running text ('MnO 14.78, Ce2O3 34.19, P2O5 29.57, and H2O 21.46,
    total 100.00'): the longest run of at least four constituent–value pairs, as a table candidate."""
    best = None
    runs = [(m, [(c, v) for c, v in _PROSE_ITEM.findall(m.group(0))]) for m in _PROSE_RUN.finditer(text)]
    runs += [(m, [(c, v) for v, c in _PROSE_ITEM_R.findall(m.group(0))]) for m in _PROSE_RUN_R.finditer(text)]
    for m, pairs in runs:
        seg = m.group(0); rows = []
        if re.search(r'requires|ideal|theoretical|calculated for|end[- ]member', text[max(0, m.start() - 80):m.start()], re.I):
            continue                                                # the ideal formula's composition, not the analysis
        for c, v in pairs:
            c2, kind = _constituent_ok(c)
            if kind == 'constituent' and c2 not in [r['constituent'] for r in rows]:
                rows.append({'constituent': 'N2H8O' if c2 == '(NH4)2O' else c2.rstrip('+-'), 'mean': float(v), 'range': None, 'sd': None, 'standard': None})
        if rows and max(r['mean'] for r in rows) < 5.0:
            continue                                                # 'K0.89 Na0.05 Y0.02': apfu, a formula — not a composition
        if len(rows) >= 4 and (best is None or len(rows) > best['n']):
            tot = None
            mt = re.search(r'(?:total|sum)\s*[=:]?\s*(\d+\.\d+)', text[m.end():m.end() + 40], re.I)
            if mt:
                tot = float(mt.group(1))
            best = {'rows': rows, 'total': tot, 'header': 'composition given in the text: ' + seg[:120], 'page': None, 'n': len(rows),
                    'score': len(rows) + (5 if 88 <= sum(r['mean'] for r in rows) <= 104 else 0), 'prose': True, 'end': m.end()}
    return best

def _prose_before(text, ctx):
    """The composition the paper lists in the sentences right before its formula sentence ('… SiO2 29.35,
    F 2.41, H2O 0.26, total 101.18. The empirical formula is …'): the formula's context holds only its
    last 200 characters, and a list of twenty oxides is longer than that, so the run is looked for in
    the 1400 characters before the formula and must end within 350 of it. -> a table candidate or None."""
    if not ctx:
        return None
    tn = _norm_text(text); tail = ctx[-150:]; p = tn.find(tail)
    if p < 0:
        return prose_table(ctx)
    end = p + len(tail); window = tn[max(0, end - 1400):end]
    pt = prose_table(window)
    if pt and pt.get('end', 0) >= len(window) - 350:
        return pt
    return prose_table(ctx)

_CAP_YES = re.compile(r'chemical|composition|analy|EPMA|EMPA|electron|microprobe|wt\.?\s*%|WDS|LA-ICP', re.I)
_CAP_NO = re.compile(r'coordinat|displacement|bond|powder|diffraction|crystal data|refinement|Raman|infrared|\bIR\b|unit[- ]cell|reflection|optical|physical propert|Mössbauer|site popul|occupanc', re.I)

def _caption(lines, i):
    """The 'Table N. …' line within eight lines above a block (or the block's own first lines)."""
    for k in range(i - 1, max(-1, i - 9), -1):
        toks = [w[4] for w in lines[k]['w']]
        if toks and re.match(r'^(Table|TABLE|Tab\.|Таблица)$', toks[0]) and len(toks) > 2:
            return ' '.join(toks)
    return ''

def _caption_score(cap, name=''):
    """A human finds the analytical table by its caption: +6 for 'chemical composition / EPMA /
    analytical results', −8 for coordinates, bonds, powder data, crystal data; +3 when it names the
    headline mineral, −4 when it names another mineral and not the headline (supporting phases)."""
    if not cap:
        return 0
    m = re.match(r'^(?:Table|TABLE|Tab\.|Таблица)\s*(\d+)', cap)
    n = int(m.group(1)) if m else None
    prior = 3 if n in (1, 2) else (-2 if n is not None and n >= 4 else 0)        # the analytical table is Table 1 or 2 in almost every paper (owner)
    who = 0
    if name:
        stem = re.sub(r'-\(.*\)$', '', name).lower()[:max(5, len(name) - 3)]
        low = cap.lower()
        others = [w for w in re.findall(r'[^\W\d_]{5,}ite', low) if stem not in w and w not in _NOT_MINERALS]
        who = 3 if stem in low else (-4 if others else 0)
        if re.search(r'associated|accompanying|coexisting|co-existing|host|matrix|other (?:minerals|phases)', low):
            who = -5                                                  # 'phases associated with zoisite-(Pb)': the supporting phases, not the headline
    return prior + who + (6 if _CAP_YES.search(cap) and not _structural_caption(cap) else 0) - (8 if _structural_caption(cap) else 0)

_CAP_STRONG = re.compile(r'chemical|composition|EPMA|EMPA|microprobe|wt\.?\s*%|WDS|LA-ICP|electron', re.I)

def _structural_caption(cap):
    """'Bond-valence analysis', 'Structure refinement data': structural unless a chemical word says otherwise ('Electron microprobe analyses')."""
    return bool(cap) and bool(_CAP_NO.search(cap)) and not _CAP_STRONG.search(cap)

def epma_table(pdf, name=''):
    """The analytical table with the most constituent rows: {'rows': [{'constituent', 'mean',
    'range', 'sd', 'standard'}], 'total', 'header', 'page'}; None when no table is found.
    Both layouts are read: constituents down the rows (usual) or across a header line."""
    best = None; cands = []
    for pno, lines in enumerate(_pages(pdf)):
        tr = _transposed(lines, pno, name)
        if tr:
            cands.append(tr)
        if tr and (best is None or tr['score'] > best['score']):
            best = tr
        i = 0
        while i < len(lines):
            block = []; j = i; x_col = None; gap = 0; last_row = i; merged_above = False
            while j < len(lines):
                ws = lines[j]['w']; toks = [w[4] for w in ws]
                if not toks:
                    break
                row = _row_at(ws, x_col)
                if row:
                    c, kind, vals, x0 = row
                    if x_col is None:
                        x_col = x0
                    ws_row = list(ws)
                    if j >= 1 and kind == 'constituent' and vals:
                        prev = lines[j - 1]['w']; ptoks = [w[4] for w in prev]
                        first_x = next((w[0] for w in ws if _NUM.match(w[4].replace('−', '-')) or _NUM_ESD.match(w[4])), None)
                        if 1 <= len(ptoks) <= 2 and all(_NUM.match(t_.replace('−', '-')) for t_ in ptoks) and 0 < lines[j]['y'] - lines[j - 1]['y'] <= 8 \
                                and first_x is not None and prev[0][0] < first_x - 15 and prev[0][0] > x0 + 20 and not _row_at(prev, x_col):
                            vals = [(kk, v) for kk, v in _numbers(ptoks)] + vals; ws_row = list(prev) + ws_row; merged_above = True   # the mean above the label; the range and the other columns beside it
                    if j + 1 < len(lines) and kind == 'constituent':
                        nxt = lines[j + 1]['w']; ntoks = [w[4] for w in nxt]
                        n_numlike = sum(1 for t_ in ntoks if _NUM.match(t_.replace('−', '-')) or _RANGE1.match(t_.replace('−', '-')) or _NA.match(t_) or re.fullmatch(r'\(\d+\.\d+[-–]\d+\.\d+\)', t_))
                        fx = next((w[0] for w in ws if _numlike(w[4])), None)
                        under_row = ntoks[0].startswith('(') or bool(_RANGE1.match(ntoks[0].replace('−', '-'))) or (fx is not None and nxt[0][0] >= fx - 10)
                        if n_numlike >= 3 and n_numlike >= 0.8 * len(ntoks) and under_row and _row_at(nxt, x_col) is None and not any(_constituent_ok(t_)[0] for t_ in ntoks):
                            more = _numbers([t_.strip('()') for t_ in ntoks])           # a two-line cell: the mean above, '(range)' and the other columns below
                            vals = vals + [(kk, v) for kk, v in more if kk in ('num', 'range', 'na')]
                            ws_row = ws_row + list(nxt); j += 1
                    block.append((c, kind, vals, ws_row)); j += 1; gap = 0; last_row = j
                elif block and gap < 3:
                    gap += 1; j += 1                         # the other page column's lines, a wrapped name
                else:
                    break
            j = last_row
            block = [b for b in block if not (b[1] == 'constituent' and re.fullmatch(r'[A-Z][a-z]?', b[0]) and next((v for kk, v in b[2] if kk == 'num'), 0) > 110)]
            n_const = sum(1 for c, k, v, _ in block if k == 'constituent')                     # 'Ba 201': ppm, not wt%
            if n_const >= 3:
                head = []; head_ws = []
                for k in (i - 1, i - 2):
                    if k >= 0:
                        head = [w[4] for w in lines[k]['w']] + head; head_ws = list(lines[k]['w']) + head_ws
                head_cells = []
                above = [i - 3, i - 2, i - 1]
                for k in (i - 4, i - 5, i - 6):                                  # a names row further up ('Burnettite Melilite Paqueite …' over 'Constituent wt% n = 5 SD'): header-shaped, below the caption
                    if k < 0 or lines[k]['y'] < 60:
                        break
                    toks_k = [w[4] for w in lines[k]['w']]
                    if re.match(r'^(Table|TABLE|Tab\.|Таблица)$', toks_k[0] if toks_k else '') or len(toks_k) > 14 or re.search(r'[.;:]$', toks_k[-1] if toks_k else '.'):
                        break
                    if not re.search(r'[A-Za-z]{4,}ite\b|\bn\s*=|sample|analys|holotype|type', ' '.join(toks_k), re.I):
                        break
                    above.insert(0, k)
                for k in above:                                                  # a caption ('Table 1. Composition of X') is not a column header, even mid-line
                    if k < 0:
                        continue
                    ws_k = lines[k]['w']
                    if lines[k]['y'] < 60:
                        continue                                                 # the running head of the page, not a column header
                    cut = next((q for q, w in enumerate(ws_k) if re.match(r'^(Table|TABLE|Tab\.|Таблица)$', w[4]) and q + 1 < len(ws_k) and re.match(r'^\d+\.?$', ws_k[q + 1][4])), None)
                    head_cells += [(w[4], (w[0] + w[2]) / 2) for w in (ws_k[:cut] if cut is not None else ws_k)]
                # which of several candidate blocks is THE analytical table: the header reads like one
                # (wt%, Mean, Range, S.D., Standard) and the means add to about 100
                first_nums = [next((v for k, v in vals if k == 'num'), None) for c, kind, vals, _ in block if kind == 'constituent']
                tot_guess = sum(v for v in first_nums if v is not None)
                score = n_const + (4 if re.search(r'wt\.?\s*%|mean|range|s\.?d\.?|standard|average', ' '.join(head), re.I) else 0) + (5 if 88 <= tot_guess <= 104 else 0)
                n_range = sum(1 for c, kind, vals, _ in block if kind == 'constituent' and any(kk == 'range' for kk, _ in vals))
                score += 3 if n_range >= n_const / 2 else 0
                score -= 6 if any(v is not None and v > 110 for v in first_nums) else 0        # ppm (a trace-element table)
                cap = _caption(lines, i); score += _caption_score(cap, name)
                if re.search(r'µg\s*/?\s*g|μg\s*/?\s*g|ppm|trace[- ]element', cap, re.I) and not re.search(r'wt\.?\s*%', cap, re.I):
                    i = j; continue                                              # 'Trace element composition (µg g−1)': not the analytical table
            if n_const >= 3 and len(re.findall(r'(?<![A-Za-z])(x|y|z|U ?eq|U ?iso|s\.o\.f\.?|occ\.?|Wyck\w*|Site|Atom|Q|Ueq|Uiso)(?![A-Za-z])', ' '.join(head))) >= 2:
                i = j; continue                                                  # the atom-coordinates table
            syms = [c for c, k, v, _ in block if k == 'constituent']
            firsts = [next((str(v) for kk, v in vals if kk == 'num'), '') for c, kk_, vals, _ in block if kk_ == 'constituent']
            top = max((syms.count(x) for x in set(syms)), default=0)
            if n_const >= 3 and ((top >= 3 and top >= len(syms) / 2) or sum(1 for v in firsts if re.search(r'\.\d{4,}', v)) > len(firsts) / 2):
                i = j; continue                                                  # Si Si Si … or 0.8641-style values: a structure table
            if n_const >= 3:
                rows = []; total = None
                has_std = bool(re.search(r'standard|std|prob', ' '.join(head), re.I))   # names count only under such a column
                n_oxide = sum(1 for c, k, v, _ in block if k == 'constituent' and re.search(r'[A-Za-z]\d*O\d*$', c) and c != 'O' and not c.startswith('H2O'))
                oxide_els = {m_.group(0) for m_ in (re.match(r'[A-Z][a-z]?', c) for c, k, v, _ in block if k == 'constituent' and re.search(r'O\d*$', c)) if m_}
                # the column the means sit in: under a 'Mean' / 'Average' header token when there is one
                mean_like = [(0 if re.match(r'^(mean|aver(?:age)?\.?|avg\.?)$', w[4], re.I) else 1, (w[0] + w[2]) / 2) for w in head_ws
                             if re.match(r'^(mean|aver(?:age)?\.?|avg\.?|wt\.?%?)$', w[4], re.I)] if head_ws else []
                mean_x = min(mean_like)[1] if mean_like else None                   # 'Mean' outranks a 'wt.%' over another column; leftmost among equals
                head_low = [h.lower() for h in head]; first_is_mean = False
                if mean_x is not None and any(h in ('constituent', 'constituents', 'oxide', 'element', 'component') for h in head_low):
                    k0 = max(i_ for i_, h in enumerate(head_low) if h in ('constituent', 'constituents', 'oxide', 'element', 'component'))
                    if k0 + 1 < len(head_low) and re.match(r'^(mean|aver(?:age)?\.?|avg\.?|wt\.?%?)$', head_low[k0 + 1]):
                        mean_x = None; first_is_mean = True                    # Mean is the first column: the first number is it
                ints = [int(t) for t in head if re.fullmatch(r'\d{1,2}', t)]
                point_cols = len(ints) >= 4 and mean_x is None and not first_is_mean and not merged_above and ints == sorted(ints) and len(set(ints)) == len(ints)   # points 1…n as columns
                keep_cols = None
                if point_cols:                                                   # only the analyses of the same phase as column 1: the major constituent within 15 %
                    major = max((b for b in block if b[1] == 'constituent'), key=lambda b: next((v for kk, v in b[2] if kk == 'num'), 0), default=None)
                    if major:
                        nums_m = [v for kk, v in major[2] if kk == 'num']
                        if nums_m:
                            keep_cols = [q for q, v in enumerate(nums_m) if abs(v - nums_m[0]) <= 0.15 * nums_m[0]]
                past_total = False; first_val = {}; rows_nd = []
                for c, kind, vals, ws_row in block:
                    if c == 'O':
                        continue                                                 # an 'O = F' remnant, never a constituent
                    v0 = next((v for k_, v in vals if k_ == 'num'), None)
                    if kind == 'constituent' and c in first_val and v0 is not None and v0 < first_val[c] / 2:
                        continue                                                 # the same constituent again, much smaller: the apfu block (Ag 72.29 … Ag 8.01)
                    if kind == 'constituent' and past_total and c in first_val and re.fullmatch(r'[A-Z][a-z]?', c):
                        continue                                                 # a bare element again below the Total: the apfu block ('Na 0.22' under 'Na n.d. n.d. 0.18'); an oxide there is the FeO / Fe2O3 split
                    if kind == 'constituent' and (v0 is not None or (vals and vals[0][0] == 'na')):
                        first_val.setdefault(c, v0 if v0 is not None else 0.0)
                    if past_total and re.fullmatch(r'[A-Z][a-z]?', c) and c in oxide_els:
                        continue                                                 # below the Total: the apfu block (S 1.99 under SO3 54.76)
                    if n_oxide >= 3 and re.fullmatch(r'[A-Z][a-z]?', c) and c in oxide_els and c not in ('F', 'Cl', 'Br', 'I', 'S', 'Se', 'Te'):
                        continue                                                 # an apfu row (Si 5.936) beside the oxide row of the same element
                    nums_all = [v if k == 'num' else 0.0 for k, v in vals if k in ('num', 'na')]   # positional: 'n.d.' = 0 in its column, the columns stay aligned
                    nums = [v for k, v in vals if k == 'num']                   # the measured numbers: the mean, the points' average, the s.d.
                    nd_first = bool(vals) and vals[0][0] == 'na'
                    if nd_first and (kind != 'constituent' or not nums):
                        continue                                                 # not detected anywhere: nothing to use
                    rng = next((v for k, v in vals if k == 'range'), None)
                    texts = [v for k, v in vals if k == 'text' and not re.fullmatch(r'[a-d*†‡§]+', v)] if has_std else []
                    if kind == 'total':
                        if (c.lower().startswith(('total', 'sum')) or c in ('Σ', 'Сумма', 'Итого')) and nums:
                            total = nums[0]; past_total = True
                        continue
                    if not nums:
                        continue
                    mean = 0.0 if nd_first else nums[0]                          # n.d. in the first column: no mean of its own, the row serves the named columns only
                    if point_cols and len(nums) >= 3:
                        use_ = [nums[q] for q in (keep_cols or range(len(nums))) if q < len(nums)] or nums
                        mean = round(sum(use_) / len(use_), 3)                # no Mean column: the points' average (of the same phase)
                    if mean_x is not None:                                      # the value under the Mean column
                        cand = [(abs((w[0] + w[2]) / 2 - mean_x), w[4]) for w in ws_row if _numlike(w[4])]
                        if cand and min(cand)[0] < 22:
                            mean = float(re.match(r'-?\d+\.?\d*', min(cand)[1].replace('−', '-').lstrip('(')).group(0))
                    sd = None
                    # s.d. is the number after the range (Mean | Range | S.D.), else the second number when
                    # the header says so
                    esd = next((v for k, v in vals if k == 'esd'), None)
                    if esd is not None:
                        sd = esd
                    elif rng is not None:
                        after = [v for k, v in vals[[k for k, _ in vals].index('range') + 1:] if k == 'num']
                        sd = after[0] if after else None
                    elif len(nums) >= 2 and any(re.match(r'(S\.?D\.?|σ|e\.?s\.?d)', h, re.I) for h in head):
                        sd = nums[1]
                    xs = [(w[0] + w[2]) / 2 for w in ws_row if _numlike(w[4]) or _NA.match(w[4])]
                    if c == 'HS':                                               # hydrosulfide, wt% of HS: the sulfur (the H is informational)
                        k_hs = 32.06 / 33.068; c = 'S'; mean = round(mean * k_hs, 3); nums_all = [round(v * k_hs, 3) for v in nums_all]
                        rng = (round(rng[0] * k_hs, 3), round(rng[1] * k_hs, 3)) if rng else None
                    (rows_nd if nd_first else rows).append({'constituent': 'N2H8O' if c == '(NH4)2O' else c.rstrip('+-'), 'mean': mean, 'range': rng, 'sd': sd,
                                 'standard': ' '.join(texts) if texts else None, 'all': nums_all, 'xs': xs})
                cand = {'rows': _drop_totals(rows), 'total': total, 'header': ' '.join(head)[:200], 'page': pno + 1, 'n': n_const, 'score': score, 'caption': cap[:160], 'head_cells': head_cells, 'rows_all': rows + rows_nd, 'label_x': x_col}
                cands.append(cand)
                if best is None or score > best['score']:
                    best = cand
                i = j; continue
            i += 1
    pt = prose_table(text_of(pdf))
    if pt:
        cands.append(pt)
        if best is None or (pt['score'] > best['score'] and (best.get('total') is None or best['score'] <= 0)):
            best = pt                                                            # the composition is in the text
    if best is not None and best['score'] <= 0 and _structural_caption(best.get('caption') or ''):
        return None                                                              # only a bond-valence / coordinates table was read: no analytical table
    if pdf.lower().endswith('.docx'):
        for c_ in cands:
            c_['page'] = None                                                    # a manuscript's tables have no page
    if best is not None:
        best['candidates'] = sorted((c for c in cands if c is not best), key=lambda c: -c['score'])[:6]
    return best

# ----------------------------------------------------------------------------- the paper's method

_BASIS = [
    (r'(?:basis of|based on|normali[sz]ed (?:to|on(?: the basis of)?)) (\d+(?:\.\d+)?) (?:O|oxygen|oxygens)(?: atoms)?(?: per formula unit| apfu| pfu)?', 'O'),
    (r'(?:basis of|based on|normali[sz]ed (?:to|on(?: the basis of)?)) (\d+(?:\.\d+)?) (?:anions?|\(O ?\+ ?(?:F|OH|Cl)[^)]*\)|O ?\+ ?(?:F|OH|Cl)|total anions?)', 'O'),
    (r'(?:basis of|based on|normali[sz]ed (?:to|on(?: the basis of)?)) (\d+(?:\.\d+)?) (?:total )?cations', 'cations'),
    (r'(?:basis of|based on|normali[sz]ed (?:to|on(?: the basis of)?)) (\d+(?:\.\d+)?) ((?:[A-Z][a-z]? ?\+ ?)*[A-Z][a-z]?)(?: atoms| apfu| pfu| atom)?\b', 'element'),
    (r'(?:basis of|based on|normali[sz]ed (?:to|on(?: the basis of)?)) ((?:[A-Z][a-z]? ?\+ ?)*[A-Z][a-z]?) ?= ?(\d+(?:\.\d+)?)', 'element2'),
    (r'(?:basis of|based on|normali[sz]ed (?:to|on(?: the basis of)?)) \(?(?:O ?\+ ?(?:F|OH|Cl|S)(?: ?\+ ?(?:F|OH|Cl|S))*|anions?|total anions?)\)? ?= ?(\d+(?:\.\d+)?)', 'O'),   # 'basis of O + F = 2 apfu'
]
_SKIP_EL = {'O', 'OXYGEN', 'ANIONS', 'CATIONS', 'H2O', 'H', 'OH', 'APFU', 'PFU'}

def basis_statement(text):
    """The first basis the paper states near its empirical formula: (basis tuple for epma, the
    sentence) — e.g. (('O', 21.0), '… on the basis of 21 O apfu …'); (None, '') when none."""
    hits = []
    for pat, kind in _BASIS:
        for m in re.finditer(pat, text, re.I):
            lo = max(0, m.start() - 160); pre = text[lo:m.start()]
            cut = max(pre.rfind('. '), pre.rfind('; '))
            sent = text[(lo + cut + 2) if cut >= 0 else lo: m.end() + 80].strip()
            sent = sent[:sent.find('. ', len(sent) - 80) + 1] if sent.find('. ', len(sent) - 80) > 0 else sent
            if kind == 'element':
                spec, n = m.group(2).replace(' ', ''), float(m.group(1))
                if spec.upper() in _SKIP_EL:
                    continue
                b = ('element', spec, n)
            elif kind == 'element2':
                spec = m.group(1).replace(' ', '')
                anion_only = set(re.findall(r'[A-Z][a-z]?', spec)) <= {'O', 'F', 'Cl', 'S'} and 'O' in spec      # 'O + F = 2': the anion basis
                b = ('O', float(m.group(2))) if (spec.upper() == 'O' or anion_only) else ('element', spec, float(m.group(2)))
            else:
                b = (kind, float(m.group(1)))
            near = 1 if re.search(r'empirical formula|formula', text[max(0, m.start() - 200): m.end() + 200], re.I) else 0
            hits.append((-near, m.start(), b, sent))
    if not hits:
        return None, ''
    hits.sort(key=lambda h: (h[0], h[1]))
    return hits[0][2], hits[0][3]

def method_statements(text):
    """Sentences on how the unmeasured constituents were handled. -> {'sentences': [...],
    'h2o': 'difference' | 'stoichiometry' | 'structure' | 'charge' | 'measured' | None,
    'charge': 'Fe' | 'H2O' | None, 'calculated': [constituents]}"""
    sents = re.split(r'(?<=[.;])\s+', text)
    out = {'sentences': [], 'h2o': None, 'charge': None, 'calculated': []}
    for s in sents:
        if len(s) > 400 or len(re.findall(r'\d+\.\d+', s)) > 6:
            continue
        if re.search(r'\b(calculated|by difference|stoichiometr|charge[- ]balance|charge balance|by analogy|assuming)\b', s, re.I) and \
                re.search(r'H2O|CO2|B2O3|Li2O|BeO|Fe2O3|FeO|Fe3\+|Fe2\+|Mn3\+|OH|water|hydrogen|carbon|boron|lithium', s, re.I):
            out['sentences'].append(s.strip())
            if re.search(r'H2O|water|hydrogen', s, re.I):
                if re.search(r'by difference', s, re.I): out['h2o'] = 'difference'
                elif re.search(r'charge', s, re.I): out['h2o'] = out['h2o'] or 'charge'; out['charge'] = out['charge'] or 'H2O'
                elif re.search(r'stoichiometr|ideal formula', s, re.I): out['h2o'] = out['h2o'] or 'stoichiometry'
                elif re.search(r'structure|structural|refinement|crystal[- ]structure', s, re.I): out['h2o'] = out['h2o'] or 'structure'
            if re.search(r'Fe2O3|FeO|Fe3\+|Fe2\+', s, re.I) and re.search(r'charge', s, re.I):
                out['charge'] = 'Fe'
            for c in ('H2O', 'CO2', 'B2O3', 'Li2O', 'BeO', 'Fe2O3'):
                if re.search(r'\b' + c + r'\b', s) and re.search(r'calculated|by difference|stoichiometr|charge', s, re.I) and c not in out['calculated']:
                    out['calculated'].append(c)
        if re.search(r'\bH2O\b[^.]{0,60}(TGA|thermogravimetr|Penfield|CHN|measured directly)', s, re.I):
            out['h2o'] = 'measured'; out['sentences'].append(s.strip())
    return out

# ----------------------------------------------------------------------------- optics, density

_F = r'(\d\.\d{2,4})'

def optics(text):
    """{'n': mean refractive index or None, 'n_from': how, 'D_meas', 'D_calc', 'sentences'}"""
    out = {'n': None, 'n_from': '', 'D_meas': None, 'D_calc': None, 'sentences': []}
    t = text.replace('−', '-')
    m = re.search(r'\b(?:n\s*=|n\s*\(?\s*(?:mean|average|calc)\s*\)?\s*=|nmean\s*=|n\s*≈)\s*' + _F, t)
    abg = re.findall(r'\b(?:n?α|nα|alpha)\s*=\s*' + _F + r'.{0,60}?\b(?:n?β|nβ|beta)\s*=\s*' + _F + r'.{0,60}?\b(?:n?γ|nγ|gamma)\s*=\s*' + _F, t)
    we = re.findall(r'\b(?:n?ω|nω|omega)\s*=\s*' + _F + r'.{0,60}?\b(?:n?ε|nε|epsilon)\s*=\s*' + _F, t)
    if abg:
        a, b, g = (float(x) for x in abg[0]); out['n'] = round((a + b + g) / 3, 4); out['n_from'] = 'mean of α %s, β %s, γ %s' % abg[0]
    elif we:
        w, e = (float(x) for x in we[0]); out['n'] = round((2 * w + e) / 3, 4); out['n_from'] = 'mean of ω %s (×2), ε %s' % we[0]
    elif m:
        out['n'] = float(m.group(1)); out['n_from'] = 'n = %s' % m.group(1)
    for key, pat in (('D_meas', r'(?:D\s*meas\.?|Dmeas|measured density|density[^.]{0,50}?(?:measured|floatation|flotation|pycnomet)[^.]{0,40}?)\s*(?:=|is|of|was|:)?\s*' + _F),
                     ('D_calc', r'(?:D\s*calc\.?|Dcalc|calculated density|density[^.]{0,30}?calculated[^.]{0,40}?)\s*(?:=|is|of|was|:)?\s*' + _F)):
        mm = re.search(pat, t, re.I)
        if mm:
            out[key] = float(mm.group(1)); out['sentences'].append(t[max(0, mm.start() - 40): mm.end() + 30].strip())
    if out['n'] and out['n_from']:
        out['sentences'].append(out['n_from'])
    return out

# ----------------------------------------------------------------------------- bond-valence parameters

def bv_statement(text):
    """{'params': 'gh'|'bo'|'ba'|None, 'u6': 'burns'|'params'|None, 'hb': 'oo'|None, 'sentences'}"""
    out = {'params': None, 'u6': None, 'hb': None, 'sentences': []}
    for s in re.split(r'(?<=[.;])\s+', text):
        if len(s) > 400 or not re.search(r'bond[- ]valence|valence units|hydrogen[- ]bond', s, re.I):
            continue
        hit = False
        if re.search(r'Gagn[eé]', s): out['params'] = out['params'] or 'gh'; hit = True
        if re.search(r"Brese", s): out['params'] = out['params'] or 'bo'; hit = True
        if re.search(r'Brown (?:and|&) Altermatt', s): out['params'] = out['params'] or 'ba'; hit = True
        if re.search(r'Burns', s) and re.search(r'U\s*6\s*\+|U6\+|uran', s, re.I): out['u6'] = 'burns'; hit = True
        elif re.search(r'U\s*6\s*\+|U6\+|U6þ', s) and re.search(r'Gagn[eé]', s): out['u6'] = out['u6'] or 'params'; hit = True
        if re.search(r'Ferraris', s): out['hb'] = 'oo'; hit = True
        if hit:
            out['sentences'].append(s.strip())
    if out['params'] == 'gh' and out['u6'] is None and re.search(r'Gagn[eé][^.]{0,120}U6?\+', text):
        out['u6'] = 'params'
    return out

# ----------------------------------------------------------------------------- the powder table

_HKL = re.compile(r'^[-−]?\d{1,2}$')
_PHEAD = re.compile(r'^(I|d|2θ|2theta)_?\(?(obs|calc|meas|c|o)\)?\.?$|^(hkl|h|k|l)$', re.I)   # the legacy single-token form (kept for the survey scripts)
# the header vocabulary of a powder table, as the corpus writes it: a quantity (I, I/I0, I/Imax, Irel,
# d, dhkl, 2θ) with or without a nature suffix (obs / meas / exp / o | calc / cal / c) that may be
# attached ('dcalc', 'd(obs)', 'I/I0meas'), or stand alone as the next word ('dhkl calc',
# 'I/Imax (calc)'); footnote marks after a label ('dcalc*'); units as their own words ('(Å)', '[Å]',
# '(%)'); the indices as one token ('hkl') or as the letters h k l (h k i l for a hexagonal table),
# however far apart the columns are set
_PH_QTY = re.compile(r'^(?:100[⋅·×*x]?)?(I(?:/I(?:0|o|max))?|d(?:hkl)?|2θ|2theta|2th|2q|2h)(?:[_/\-]?(obs|calc|cal|clac|cacl|meas|meass|exp|est|rel|c|o)[a-f]?)?$', re.I)   # '2q', '2h': 2θ in a font that lost its Greek; 'Icalca': a footnote letter; 'Iest': estimated by eye; 'Dclac', 'Imeass': as typeset
_PH_SUFFIX = re.compile(r'^(obs|calc|cal|clac|cacl|meas|meass|exp|est)[a-f]?$', re.I)
_PROSE = {'and', 'the', 'for', 'with', 'are', 'was', 'were', 'from', 'that', 'this', 'not', 'only', 'which', 'has', 'have', 'been', 'also'}   # a running sentence, not a table
_PH_WORDS = {'obs', 'calc', 'cal', 'meas', 'exp', 'rel', 'hkl', 'theta', 'int', 'intensity', 'irel', 'dhkl', 'imax', 'index', 'indices', 'bold', 'sample', 'synthetic', 'natural', 'ideal'}   # words a header line may carry beside its labels
_PH_NATURE = {'obs': 'obs', 'meas': 'obs', 'meass': 'obs', 'exp': 'obs', 'est': 'obs', 'o': 'obs', 'calc': 'calc', 'cal': 'calc', 'clac': 'calc', 'cacl': 'calc', 'c': 'calc'}

def _header_labels(ws):
    """The labelled columns of one line -> ([(x centre, ('d' | 'I', 'obs' | 'calc' | None) | 'hkl',
    (x first, x last))], the other words of three or more letters among them). h k l (or h k i l) as
    separate words are one column, its span from the 'h' word to the 'l' word; a 2θ column is
    skipped; a quantity's nature is None when the header does not say ('d', 'I', 'Irel', 'dhkl')."""
    cols = []; run = []; used = set()                               # run: the index letters seen so far; used: the words read as labels
    def flush():
        letters = ''.join(l for _, _, l in run)
        if letters in ('hkl', 'hkil'):
            xs = [x for x, _, _ in run]
            cols.append([sum(xs) / len(xs), 'hkl', (run[0][1][0], run[-1][1][2])])   # the span: the 'h' word's left edge to the 'l' word's right edge
        del run[:]
    for wi, w in enumerate(ws):
        t2 = re.sub(r'[()\[\]]', '', w[4])                            # '(meas.)' -> 'meas'; 'dcalc**' -> 'dcalc'
        t2 = re.sub(r'[*†‡§#]+$', '', t2).rstrip('.')
        xc = (w[0] + w[2]) / 2
        if t2 in ('h', 'k', 'l', 'i', 'H', 'K', 'L'):
            used.add(wi)
            letter = t2.lower()
            if run and ('hkl'.startswith(''.join(l for _, _, l in run) + letter) or 'hkil'.startswith(''.join(l for _, _, l in run) + letter)):
                run.append((xc, w, letter)); continue
            flush()
            if letter == 'h':
                run.append((xc, w, letter))
            continue
        flush()
        if t2.lower() == 'hkl':
            used.add(wi); cols.append([xc, 'hkl', (w[0], w[2])]); continue   # one word: the three digits sit under its box
        m = _PH_QTY.match(t2)
        if m and t2 == 'D':
            m = None                                                   # a bare 'D' ('D(2θ)'): a difference or a density column, never the d spacing — 'Dobs' is
        if m and m.group(1)[0] == '2':
            used.add(wi); continue                                     # a 2θ column is never read (no λ to turn it into d) — and '2θ' in the prose beside a table must not open one
        if m:
            used.add(wi); q = 'I' if m.group(1)[0] in 'iI' else 'd'    # '100⋅I/Imax' is an intensity
            cols.append([xc, (q, _PH_NATURE.get(re.sub(r'[a-f]$', '', (m.group(2) or '').lower()) if m.group(2) and m.group(2).lower() not in _PH_NATURE else (m.group(2) or '').lower())), (w[0], w[2])]); continue
        m = _PH_SUFFIX.match(t2)
        if m and cols and isinstance(cols[-1][1], tuple) and cols[-1][1][1] is None:
            used.add(wi); cols[-1][1] = (cols[-1][1][0], _PH_NATURE[m.group(1).lower()])   # 'dhkl calc', 'I/Imax (calc)': the nature as the next word
            cols[-1][2] = (cols[-1][2][0], w[2]); cols[-1][0] = (cols[-1][2][0] + w[2]) / 2   # … and the column is the pair of words, its values under their middle
    flush()
    if not cols:
        return [], []
    lo = min(c[2][0] for c in cols) - 15; hi = max(c[2][1] for c in cols) + 15
    foreign = [w[4].rstrip('.,;:').lower() for wi, w in enumerate(ws)   # a sentence that mentions dobs, Iobs and hkl is no header: real ones hold labels, units and marks
               if wi not in used and lo <= (w[0] + w[2]) / 2 <= hi and re.fullmatch(r'[A-Za-z]{3,}[.,;:]?', w[4]) and w[4].rstrip('.,;:').lower() not in _PH_WORDS]
    return [(xc, lab, span) for xc, lab, span in cols], foreign

def _powder_columns(ws, caption=''):
    """The header line of a powder table -> [(x centre, label, (x first, x last))] with labels Iobs
    dobs dcalc Icalc hkl; [] when the line is no such header (no d, no hkl, or a sentence). A
    quantity with no nature is calculated when the caption says the table is, observed otherwise."""
    bare = 'calc' if re.search(r'\bcalc', caption, re.I) and not re.search(r'\b(obs|meas|exp)', caption, re.I) else 'obs'
    cols, foreign = _header_labels(ws)
    out = [(xc, lab if lab == 'hkl' else lab[0] + (lab[1] or bare), span) for xc, lab, span in cols]
    labs = [c[1] for c in out]
    if not (('dobs' in labs or 'dcalc' in labs) and 'hkl' in labs):
        return []
    if len(foreign) >= 2 or any(f in _PROSE for f in foreign):
        return []
    return out

def _powder_rows_by_columns(lines, start, cols):
    """Rows under a header: tokens go to the nearest column; several blocks per line are handled
    by walking the columns in order — a block ends at each hkl column, or, when the header puts the
    indices first ('h k l dcalc Icalc | h k l dcalc Icalc'), begins at each."""
    obs, calc = [], []
    order = [c for c in cols]
    leading = order[0][1] == 'hkl'
    x_lo = min(c[2][0] for c in order) - 15; x_hi = max(c[2][1] for c in order) + 15   # the table's own width: on a two-column page the other column's prose runs 20–30 pt beside it
    # an index column's territory: from halfway to the column on its left to halfway to the one on
    # its right (the three digits under one 'hkl' word spread wider than the word), for index-like
    # tokens only — before the nearest centre decides, which would hand the 'h' digits of a wide
    # h k l block to the column on its left (the merged centre sits nearer 'l')
    terr = {}
    for k, c in enumerate(order):
        if c[1] == 'hkl':
            terr[k] = (min(c[2][0] - 8, max(c[0] - 40, (order[k - 1][0] + c[0]) / 2)) if k else min(c[2][0] - 8, c[0] - 40),
                       max(c[2][1] + 8, min(c[0] + 40, (c[0] + order[k + 1][0]) / 2)) if k + 1 < len(order) else max(c[2][1] + 8, c[0] + 40))
    def in_table(w):                                             # within the header's width — or an index-like token within an index column's territory (the digits under one 'hkl' word spread past it; the prose of the next page column must not)
        xc = (w[0] + w[2]) / 2
        if x_lo <= xc <= x_hi:
            return True
        return bool(re.fullmatch(r'-?\d{1,3}|-?\d{1,2}(?:\.-?\d{1,2}){2}', w[4].replace('−', '-'))) and any(lo <= xc <= hi for lo, hi in terr.values())
    n = 0
    first = True
    for ln in lines[start:]:
        ws = [w for w in ln['w'] if in_table(w)]
        if not ws:
            continue
        toks = [w[4].replace('−', '-') for w in ws]
        if not any(re.fullmatch(r'\d+\.\d+', t) for t in toks):
            if n > 3:
                break                                            # the table ended
            continue
        if first:                                                # a sentence that happens to say 'dobs', 'Iobs' and 'hkl' is no header: the first row under a real one is numbers, not words
            first = False
            words = [t.rstrip('.,;:').lower() for t in toks if re.fullmatch(r'[A-Za-z]{3,}[.,;:]?', t)]
            if len(words) >= 2 or any(t in _PROSE for t in words):
                return obs, calc
        cells = {}
        for w in ws:
            xc = (w[0] + w[2]) / 2
            t = w[4].replace('−', '-')
            k = next((i for i, (lo, hi) in terr.items() if lo <= xc <= hi), None) if re.fullmatch(r'-?\d{1,3}|-?\d{1,2}(?:\.-?\d{1,2}){2}', t) else None
            if k is None:
                k = min(range(len(order)), key=lambda i: abs(order[i][0] - xc))
                if abs(order[k][0] - xc) > 40:
                    continue
            cells.setdefault(k, []).append(w[4].replace('−', '-'))
        block = {}
        for k in range(len(order)):
            lab = order[k][1]
            if lab == 'hkl' and leading and block:
                _powder_emit(block, obs, calc); n += 1; block = {}
            if k in cells:
                block[lab] = cells[k]
            if lab == 'hkl' and not leading:
                if block:
                    _powder_emit(block, obs, calc); n += 1
                block = {}
        if leading and block:
            _powder_emit(block, obs, calc); n += 1
    return obs, calc

def _num1(v):
    try:
        return float(v[0]) if v else None
    except ValueError:
        return None

def _powder_emit(block, obs, calc):
    raw = [p for t in block.get('hkl', []) for p in _split_glued(t)]
    hk = [t for t in raw if _HKL.match(t)]
    hkl = tuple(int(t) for t in ((hk[0], hk[1], hk[3]) if len(hk) == 4 else hk[-3:])) if len(hk) >= 3 else None   # h k i l: i is redundant
    if hkl is None and len(raw) == 1 and re.fullmatch(r'-?\d{3}', raw[0]):
        hkl = tuple(int(c) for c in raw[0].lstrip('-'))                # '001' — as the only token: three of them are three intensities gone astray
    if hkl is None and len(raw) == 1 and re.fullmatch(r'-?\d{1,2}(?:\.-?\d{1,2}){2}', raw[0]):
        hkl = tuple(int(c) for c in raw[0].split('.'))                 # '2.1.10': two-digit indices written with dots
    if hkl is None and 2 <= len(raw) <= 3 and all(t.isdigit() for t in raw) and len(''.join(raw)) == 3:
        hkl = tuple(int(c) for c in ''.join(raw))                      # '01 1': one index broken off the others by the font
    for dk, ik in (('dobs', 'Iobs'), ('dcalc', 'Icalc')):
        dv, iv = (block.get(dk) or [''])[0], (block.get(ik) or [''])[0]
        if re.fullmatch(r'\d+', dv) and re.fullmatch(r'\d+\.\d{3,}', iv):
            block[dk], block[ik] = block[ik], block[dk]                # the header names the columns in the other order than the numbers stand
    d_o, i_o, d_c, i_c = (_num1(block.get(k)) for k in ('dobs', 'Iobs', 'dcalc', 'Icalc'))
    if d_o is not None:
        obs.append((d_o, i_o))
    if d_c is not None and hkl is not None:
        calc.append((d_c, i_c, hkl))

# ----------------------------------------------------------------------------- the powder table by its contents
# The header's spelling is the fragile part of reading a powder table — every journal writes the
# labels its own way, some put 'h k l' on one line and the d and I labels on another, some have no
# header at all. The columns can be typed from what stands in them: a column of floats between
# 0.5 and 40 that falls down the rows is d; a column bounded by 100 is an intensity; three
# neighbouring columns of small integers are h k l (four with h + k + i = 0 are h k i l); floats that
# rise down the rows are 2θ. The header words, when found, only say which d and I are observed and
# which calculated; the caption and the number of decimals stand in when they are missing.

_PT_TOKEN = re.compile(r'^([-−–¯]?)(\d+)(?:\.(\d+))?(?:\((\d+)\))?[a-z*†‡§]?[,;]?$')
_PT_DOTTED = re.compile(r'^[-−]?\d{1,2}(?:\.[-−]?\d{1,2}){2}$')
_PT_IWORDS = ('vvs', 'vs', 's', 'ms', 'm', 'mw', 'w', 'vw', 'vvw', 'b', 'br', 'sh', 'd', 'tr')
_PT_GLUED = re.compile(r'^[-−]?\d{1,2}(?:[-−]\d{1,2})+$')

def _split_glued(t):
    """'0-1' -> ['0', '-1'], '-1-3' -> ['-1', '-3'], '2-41' -> ['2', '-4', '1']: a negative index the
    extraction glued to the one before it. Two-digit parts are split when fewer than three come out."""
    t = t.replace('−', '-')
    if not _PT_GLUED.match(t):
        return [t]
    parts = re.findall(r'-?\d+', t)
    if len(parts) < 3 and any(len(p.lstrip('-')) > 1 for p in parts):
        out = []
        for p in parts:
            sign, digits = ('-', p[1:]) if p.startswith('-') else ('', p)
            out += [sign + digits[0]] + list(digits[1:])
        parts = out
    return parts

def _split_words(ws):
    """The words of a line with glued indices split, each part given its share of the word's box."""
    out = []
    for w in ws:
        parts = _split_glued(w[4]) if _PT_GLUED.match(w[4].replace('−', '-')) else [w[4]]
        if len(parts) == 1:
            out.append(w); continue
        step = (w[2] - w[0]) / len(parts)
        for k, p in enumerate(parts):
            out.append((w[0] + k * step, w[1], w[0] + (k + 1) * step, w[3], p))
    return out

def _pt_token(t):
    """A table token -> ('num', value, decimals, signed, has_esd) | ('dot', 'h.k.l') | ('word', t) | None."""
    t = t.replace('̄', '-').replace('¯', '-')
    if _PT_DOTTED.match(t.replace('−', '-')):
        return ('dot', t.replace('−', '-'))
    m = _PT_TOKEN.match(t)
    if m:
        v = float(m.group(2) + ('.' + m.group(3) if m.group(3) else ''))
        return ('num', -v if m.group(1) else v, len(m.group(3) or ''), bool(m.group(1)), bool(m.group(4)))
    if t.lower() in _PT_IWORDS or re.fullmatch(r'n\.?d\.?|[–—-]', t):
        return ('word', t)
    return None

def _pt_regions(lines):
    """Runs of lines that read like table rows (≥3 numeric tokens; gaps of up to two other lines),
    at least four rows long -> [(first, last)]. A header line inside a run splits it."""
    tab = []
    for ln in lines:
        kinds = [_pt_token(w[4]) for w in ln['w']]
        n = sum(1 for k in kinds if k and k[0] in ('num', 'dot'))
        hdr = len(_header_labels(ln['w'])[0]) >= 2 and n < 3
        tab.append('h' if hdr else 't' if n >= 3 else '-')
    out = []; i = 0
    while i < len(tab):
        if tab[i] != 't':
            i += 1; continue
        j = i; last = i
        while j < len(tab) and tab[j] != 'h' and (tab[j] == 't' or j - last <= 2):
            if tab[j] == 't':
                last = j
            j += 1
        if sum(1 for k in range(i, last + 1) if tab[k] == 't') >= 4:
            out.append((i, last))
        i = last + 1
    return out

def _pt_columns(lines, first, last):
    """The x-clusters of the numeric tokens of a region: [{'x', 'lo', 'hi', 'cells': {line: [tokens]}}],
    left to right, those present on at least 40 % of the region's rows."""
    pts = []
    rows = 0
    for li in range(first, last + 1):
        ws = _split_words(lines[li]['w'])
        if sum(1 for w in ws if (_pt_token(w[4]) or (None,))[0] in ('num', 'dot')) < 2:
            continue
        rows += 1
        for w in ws:
            k = _pt_token(w[4])
            if k and k[0] != 'word' or (k and k[0] == 'word' and k[1].lower() in _PT_IWORDS):
                pts.append(((w[0] + w[2]) / 2, w[0], w[2], li, w[4]))
    pts.sort()
    cols = []
    for x, x0, x1, li, t in pts:
        if cols and x - cols[-1]['xs'][-1] <= 9.0:
            c = cols[-1]
        else:
            c = {'xs': [], 'lo': x0, 'hi': x1, 'cells': {}}; cols.append(c)
        c['xs'].append(x); c['lo'] = min(c['lo'], x0); c['hi'] = max(c['hi'], x1)
        c['cells'].setdefault(li, []).append(t)
    out = []
    for c in cols:
        if rows and len(c['cells']) >= max(3, 0.15 * rows):           # observed lines are sparser than calculated ones; a stray number in the prose is rarer still
            c['x'] = sum(c['xs']) / len(c['xs']); del c['xs']; out.append(c)
    return out, rows

def _pt_type(col):
    """What a column holds: 'd' | 'theta' | 'I' | 'idx' | 'idx3' | 'dot' | 'hkl3' (three small ints per
    cell) | 'rownum' | 'other'."""
    cells = list(col['cells'].values())
    n = len(cells)
    per = [len(c) for c in cells]
    toks = [_pt_token(t) for c in cells for t in c]
    nums = [k for k in toks if k and k[0] == 'num']
    if n and sum(1 for p in per if 2 <= p <= 4) >= 0.8 * n and sum(1 for p in per if p == 3) >= 0.6 * n \
            and sum(1 for k in toks if k and k[0] == 'num' and k[2] == 0 and abs(k[1]) <= 30) >= 0.9 * len(toks):
        return 'hkl3'                                                  # the three indices set so close they cluster as one column (a row or two glued or split)
    raw = [t.replace('−', '-') for c in cells for t in c]
    three = [t for t in raw if re.fullmatch(r'-?\d{3}', t)]
    dotted = sum(1 for t in raw if _PT_DOTTED.match(t))
    if n and len(three) + dotted >= 0.8 * len(raw) and len(raw) == n and (dotted or any(t.lstrip('-').startswith('0') for t in three) or sum(1 for t in three if abs(int(t)) > 100) >= 0.5 * len(three)):
        return 'idx3'                                                  # '001', '110', '2.1.10': indices as one word — an intensity column is not all three-digit, nor above 100
    if not nums or len(nums) < 0.7 * len(toks):
        return 'other'
    vals = [k[1] for k in nums]
    ints = sum(1 for k in nums if k[2] == 0); dec2 = sum(1 for k in nums if k[2] >= 2)
    signed = any(k[3] for k in nums)
    seq = [vals[i] for i in range(len(vals))]
    dec = sum(1 for a, b in zip(seq, seq[1:]) if a > b); inc = sum(1 for a, b in zip(seq, seq[1:]) if a < b)
    pairs = max(1, len(seq) - 1)
    if ints == len(nums) and len(seq) >= 4 and all(b - a == 1 for a, b in zip(seq, seq[1:])) and seq[0] in (1, 0):
        return 'rownum'
    if dec2 >= 0.8 * len(nums) and sum(1 for v in vals if 0.5 <= v <= 40) >= 0.9 * len(vals) and (dec >= 0.5 * pairs or len(seq) <= 3):
        return 'd'
    if dec2 >= 0.5 * len(nums) and sum(1 for v in vals if 2 <= v <= 170) >= 0.9 * len(vals) and inc >= 0.6 * pairs and not signed:
        return 'theta'
    if ints == len(nums) and sum(1 for v in vals if abs(v) <= 30) >= 0.95 * len(vals):
        return 'idx'
    if sum(1 for k in nums if k[2] <= 2) >= 0.9 * len(nums) and sum(1 for v in vals if 0 <= v <= 100) >= 0.9 * len(vals) and not signed and dec < 0.8 * pairs:
        return 'I'
    if sum(1 for k in nums if k[2] <= 1) >= 0.9 * len(nums) and sum(1 for v in vals if 0 <= v <= 1000) >= 0.95 * len(vals) and not signed and dec < 0.8 * pairs:
        return 'I'
    return 'other'

def _pt_labels(lines, first):
    """The header words above a region: the labelled columns of the up to three lines above it
    (a caption or another table's rows stop the search) -> [(x, label, span)], any line."""
    out = []
    for k in range(first - 1, max(-1, first - 4), -1):
        ws = lines[k]['w']
        if not ws:
            continue
        toks = [w[4] for w in ws]
        if re.match(r'^(Table|TABLE|Tab\.|Таблица)$', toks[0]) or sum(1 for w in ws if (_pt_token(w[4]) or (None,))[0] in ('num', 'dot')) >= 3:
            break
        labs, _ = _header_labels(ws)
        out += labs
    return out

def _pt_blocks(cols, labels, bare):
    """Typed columns -> blocks, each {'d': [(col, nature)], 'I': [(col, nature)], 'hkl': [cols] | None}."""
    typed = []
    for c in cols:
        c['type'] = _pt_type(c)
    # h k l groups: runs of neighbouring idx columns, three (or four with h + k + i = 0) at a time
    i = 0; groups = []
    while i < len(cols):
        if cols[i]['type'] == 'idx':
            j = i
            while j + 1 < len(cols) and cols[j + 1]['type'] == 'idx' and cols[j + 1]['x'] - cols[j]['x'] <= 75:
                j += 1
            run = []
            for c in cols[i:j + 1]:
                if any(lab != 'hkl' and lab[0] == 'I' and abs(lx - c['x']) <= 25 for lx, lab, _ in labels):
                    c['type'] = 'I'                                    # a small-valued intensity column beside the indices is not one of them when its header says so
                else:
                    run.append(c)
            while len(run) >= 3:
                take = run[:4] if len(run) >= 4 else run[:3]
                if len(take) == 4:
                    rows = [li for li in take[0]['cells'] if all(li in c['cells'] for c in take)]
                    def iv(c, li):
                        k = _pt_token(c['cells'][li][0]); return k[1] if k and k[0] == 'num' else None
                    hkil = sum(1 for li in rows if all(iv(c, li) is not None for c in take[:3]) and abs(iv(take[0], li) + iv(take[1], li)) == abs(iv(take[2], li)))   # i = −(h + k); the overbar is lost in extraction, so by magnitude
                    if not (rows and hkil >= 0.7 * len(rows)):
                        k0 = min(range(len(run) - 2), key=lambda k: run[k + 2]['x'] - run[k]['x'])   # not h k i l: the three set closest together are the indices
                        take = run[k0:k0 + 3]
                        for c in run[:k0]:
                            c['type'] = 'I'
                        run = run[k0:]
                for c in take:
                    c['type'] = 'hkl'
                groups.append(take); run = run[len(take):]
            for c in run:                                              # a leftover small-int column beside the indices: an intensity when it looks like one
                c['type'] = 'I' if not any(_pt_token(t)[3] for cc in [c] for t in [x for v in cc['cells'].values() for x in v] if _pt_token(t) and _pt_token(t)[0] == 'num') else 'other'
            i = j + 1
        else:
            i += 1
    for c in cols:
        if c['type'] in ('idx3', 'dot', 'hkl3'):
            groups.append([c]); c['type'] = 'hkl'
    # labels: the nearest compatible header word within 25 pt of the column (a group: of its middle)
    def label_for(x, kinds):
        best = None
        for lx, lab, span in labels:
            kind = lab if lab == 'hkl' else lab[0]
            if kind in kinds:
                dist = 0 if span[0] - 4 <= x <= span[1] + 4 else abs(lx - x)
                if dist <= 25 and (best is None or dist < best[0]):
                    best = (dist, lab)
        return best[1] if best else None
    blocks = []; cur = None
    def new():
        return {'d': [], 'I': [], 'hkl': None, 'labelled': False, 'closed': False}
    for c in cols:
        t = c['type']
        if t not in ('d', 'I', 'hkl'):
            continue
        if t == 'hkl':
            grp = next(g for g in groups if c in g)
            if grp[0] is not c:
                continue                                               # the group's other members
            if cur is None or cur['hkl'] is not None:
                cur = new(); blocks.append(cur)
            cur['hkl'] = grp
            if cur['d'] or cur['I']:
                cur['closed'] = True                                   # the indices end a block ('Iobs dobs dcalc hkl | Iobs …'); leading ones open it
            continue
        if cur is None or cur['closed'] or len(cur[t]) >= 2:
            cur = new(); blocks.append(cur)
        lab = label_for(c['x'], t)
        cur[t].append((c, lab[1] if lab else None))
        if t == 'd' and lab:
            cur['labelled'] = True
    # natures: the label, else two d columns by their decimals (the calculated one has more), else the caption
    for b in blocks:
        ds = b['d']
        if len(ds) == 2 and any(n is None for _, n in ds):
            (c1, n1), (c2, n2) = ds
            if n1 is None and n2 is None:
                dc = lambda c: sorted(k[2] for k in (_pt_token(t) for v in c['cells'].values() for t in v) if k and k[0] == 'num')
                m1 = dc(c1); m2 = dc(c2); m1 = m1[len(m1) // 2] if m1 else 0; m2 = m2[len(m2) // 2] if m2 else 0
                n1, n2 = ('obs', 'calc') if m1 <= m2 else ('calc', 'obs')
            elif n1 is None:
                n1 = 'calc' if n2 == 'obs' else 'obs'
            else:
                n2 = 'calc' if n1 == 'obs' else 'obs'
            b['d'] = [(c1, n1), (c2, n2)]
        elif len(ds) == 1 and ds[0][1] is None:
            b['d'] = [(ds[0][0], bare)]
        for k, (c, n) in enumerate(b['I']):
            if n is None:
                n = b['d'][k][1] if k < len(b['d']) else (b['d'][0][1] if b['d'] else bare)
                b['I'][k] = (c, n)
    return [b for b in blocks if b['d'] and (b['hkl'] or (b['I'] and b['labelled']))]

def _pt_num(c, li):
    k = _pt_token(c['cells'][li][0]) if li in c['cells'] else None
    return k[1] if k and k[0] == 'num' else None

def _pt_hkl(grp, li):
    if len(grp) == 1:
        raw = grp[0]['cells'].get(li, [])
        if len(raw) == 1:
            t = raw[0].replace('−', '-').replace('̄', '-').replace('¯', '-')
            if re.fullmatch(r'-?\d{3}', t):
                return tuple(int(ch) for ch in t.lstrip('-')) if not t.startswith('-') else (-int(t[1]), int(t[2]), int(t[3]))
            if _PT_DOTTED.match(t):
                return tuple(int(x) for x in t.split('.'))
        ks = [_pt_token(t) for t in raw]
        if len(ks) in (3, 4) and all(k and k[0] == 'num' and k[2] == 0 for k in ks):
            v = [int(k[1]) for k in ks]
            return (v[0], v[1], v[3]) if len(v) == 4 else tuple(v)
        return None
    vs = []
    for c in grp:
        k = _pt_token(c['cells'][li][0]) if li in c['cells'] and len(c['cells'][li]) == 1 else None
        if not (k and k[0] == 'num' and k[2] == 0):
            return None
        vs.append(int(k[1]))
    return (vs[0], vs[1], vs[3]) if len(vs) == 4 else tuple(vs)

def _pt_read(lines, caption_of):
    """Observed (d, I) and calculated (d, I, hkl) lines from one page's table regions; caption_of(i)
    is the caption over line i (a 'Table N. Cont.' resolved to Table N's)."""
    obs, calc = [], []
    for first, last in _pt_regions(lines):
        cols, rows = _pt_columns(lines, first, last)
        if not cols:
            continue
        cap = caption_of(first)
        bare_here = 'calc' if re.search(r'\bcalc', cap, re.I) and not re.search(r'\b(obs|meas|exp)', cap, re.I) else 'obs'
        blocks = _pt_blocks(cols, _pt_labels(lines, first), bare_here)
        obs_done = False                                               # a second sample's columns (no indices of their own, after a block that had them) are not the holotype's: its observed lines stay out
        for b in blocks:
            if b['hkl'] is None and obs_done:
                continue
            # a bond-length or a chemical table has a word on every row: those are not powder rows
            xs = [c['x'] for c, _ in b['d'] + b['I']] + [c['x'] for c in (b['hkl'] or [])]
            lo, hi = min(xs) - 20, max(xs) + 20
            words = rows_in = 0
            for li in range(first, last + 1):
                ws = [w for w in lines[li]['w'] if lo <= (w[0] + w[2]) / 2 <= hi]
                if ws:
                    rows_in += 1
                    words += 1 if any(re.fullmatch(r'[A-Za-z][A-Za-z–\-]{2,}.*', w[4]) and w[4].lower() not in _PT_IWORDS for w in ws) else 0
            if rows_in and words > 0.2 * rows_in:
                continue
            d_obs = next((c for c, n in b['d'] if n == 'obs'), None); d_calc = next((c for c, n in b['d'] if n == 'calc'), None)
            i_obs = next((c for c, n in b['I'] if n == 'obs'), None); i_calc = next((c for c, n in b['I'] if n == 'calc'), None)
            for li in range(first, last + 1):
                if d_obs is not None and li in d_obs['cells']:
                    v = _pt_num(d_obs, li)
                    if v is not None:
                        obs.append((v, _pt_num(i_obs, li) if i_obs is not None else None)); obs_done = True
                if d_calc is not None and b['hkl'] and li in d_calc['cells']:
                    v = _pt_num(d_calc, li); h = _pt_hkl(b['hkl'], li)
                    if v is not None and h is not None:
                        calc.append((v, _pt_num(i_calc, li) if i_calc is not None else None, h))
    return obs, calc

def pxrd_table(pdf, with_pages=False):
    """Observed (d, I) and calculated (d, I, hkl) powder lines from the paper's table — by a header
    line the reader knows (Iobs dobs dcalc Icalc h k l, in any order, one or more blocks side by
    side) where the page has one, else by what the columns hold. With with_pages, also the 1-based
    numbers of the pages the lines came from."""
    pages = _pages(pdf)
    obs, calc, hit = [], [], []
    caps = {}                                                          # 'Table 4. Cont.' on a later page: the caption is the first page's
    for lines in pages:
        for ln in lines:
            toks = [w[4] for w in ln['w']]
            if len(toks) > 2 and re.match(r'^(Table|TABLE|Tab\.|Таблица)$', toks[0]) and re.match(r'^\d+\.?$', toks[1]) and not re.match(r'^\(?[Cc]ont', toks[2]):
                caps.setdefault(toks[1].rstrip('.'), ' '.join(toks))
    def full_caption(lines, i):
        cap = _caption(lines, i)
        m = re.match(r'^(?:Table|TABLE|Tab\.|Таблица)\s*(\d+)\.?\s*\(?[Cc]ont', cap)
        return caps.get(m.group(1), cap) if m else cap
    def read_page(lines):
        o_all, c_all = [], []
        i = 0
        while i < len(lines):                                          # by a header line, where the page has one the reader knows
            cols = _powder_columns(lines[i]['w'], full_caption(lines, i))
            if cols:
                o, c = _powder_rows_by_columns(lines, i + 1, cols)
                if o or c:
                    o_all += o; c_all += c
                    i += 1 + max(len(o), len(c)); continue
            i += 1
        if not (o_all or c_all):                                       # else by what the columns hold (a two-line header, an unlabelled table)
            o_all, c_all = _pt_read(lines, lambda i, lines=lines: full_caption(lines, i))
        return o_all, c_all
    other = None                                                       # the second reader's pages, only when a page needs them
    for pno, lines in enumerate(pages):
        o, c = read_page(lines)
        if not (o or c) and _pages_reader is not None and _pages_mode == 'fallback' and not pdf.lower().endswith('.docx'):
            if other is None:
                try:
                    other = _pages_reader(pdf)
                except Exception:
                    other = []
            if pno < len(other) and other[pno] is not lines:
                o, c = read_page(other[pno])
        obs += o; calc += c
        if o or c:
            hit.append(pno + 1)
    seen = set(); obs2 = []
    for d, I in obs:
        if (round(d, 3), I) in seen:
            continue
        seen.add((round(d, 3), I)); obs2.append((d, I))
    return (obs2, calc, hit) if with_pages else (obs2, calc)

# ----------------------------------------------------------------------------- the powder table vs the cell
# Every calculated line of a powder table is a statement about the cell: d follows from h k l and
# the cell parameters exactly. Recomputing it is the one rigorous check of the table — a line that
# does not follow is mis-indexed or mistyped, a table that does not follow at all was computed
# from another cell (or another setting) than the one given.
_CELL_TOL = 0.005

def _cell_floats(cell):
    """{'a','b','c','α','β','γ'} as the .cif / the text state them (esds allowed) -> floats, the
    angles defaulting to 90; None when an axis is missing."""
    try:
        v = {k: float(re.sub(r'\(.*', '', str(cell[k]).replace(' ', ''))) for k in ('a', 'b', 'c')}
        for k in ('α', 'β', 'γ'):
            v[k] = float(re.sub(r'\(.*', '', str(cell[k]).replace(' ', ''))) if cell.get(k) else 90.0
        return v if all(v[k] > 0 for k in v) else None
    except (KeyError, ValueError, TypeError):
        return None

def _paper_cells(text):
    """The cells the text states, the powder one first -> [(context, cell floats)]."""
    from pxrd_review import cell_lambda_check as CL
    out = []; seen = set()
    for cc in CL.find_cells(text):
        variants = [{'a': cc.a, 'b': cc.b, 'c': cc.c, 'α': cc.al, 'β': cc.be, 'γ': cc.ga}]
        if cc.b is None and cc.c is not None and cc.ga is None:       # 'a = 4.59, c = 2.96': tetragonal (γ 90) or hexagonal / trigonal (γ 120) — the table decides
            variants = [{'a': cc.a, 'b': cc.a, 'c': cc.c, 'α': cc.al, 'β': cc.be, 'γ': '90'},
                        {'a': cc.a, 'b': cc.a, 'c': cc.c, 'α': cc.al, 'β': cc.be, 'γ': '120'}]
        for v in variants:
            cell = _cell_floats(v)
            if not cell:
                continue
            key = tuple(round(cell[k], 3) for k in ('a', 'b', 'c', 'α', 'β', 'γ'))
            if key in seen:
                continue
            seen.add(key); out.append((cc.context or 'stated', cell))
    return sorted(out, key=lambda x: 0 if x[0] == 'powder' else 1)

def _cell_str(cell):
    ang = ', '.join('%s=%g' % (k, cell[k]) for k in ('α', 'β', 'γ') if abs(cell[k] - 90) > 1e-6)
    return 'a=%g, b=%g, c=%g%s' % (cell['a'], cell['b'], cell['c'], (', ' + ang) if ang else '')

def _d_variants(cell, hkl):
    """d for (hkl) and for the sign changes that give a different d — an overbar lost in the
    extraction turns −1 into 1, which is the extraction's fault, not the paper's."""
    from pxrd_review.extra_checks import _dstar2
    h, k, l = hkl; out = []
    for sh in ((1, 1, 1), (-1, 1, 1), (1, -1, 1), (1, 1, -1)):
        s2 = _dstar2(cell['a'], cell['b'], cell['c'], cell['α'], cell['β'], cell['γ'], h * sh[0], k * sh[1], l * sh[2])
        if s2 > 0:
            out.append((1 / math.sqrt(s2), sh != (1, 1, 1)))
    return out

def _cell_score(cell, calc):
    rows = []
    for d, I, hkl in calc:
        if hkl == (0, 0, 0) or not d or d <= 0:
            continue
        ds = _d_variants(cell, hkl)
        if not ds:
            continue
        dc, flipped = min(ds, key=lambda x: abs(x[0] - d))
        rows.append((d, I, hkl, dc, (dc - d) / d, flipped))
    return len(rows), sum(1 for r in rows if abs(r[4]) <= _CELL_LOOSE), rows

_CELL_LOOSE = 0.015                                                    # a table computed with a cell slightly different from the stated one sits within this
_CELL_WILD = 0.15                                                      # beyond this the reader paired a d with the wrong indices — its fault, not the paper's

_CELL_SRC = {'.cif': 'the .cif cell', 'powder': "the paper's powder cell", 'single': "the paper's single-crystal cell"}

def _reflection_ds(cell, dmin):
    """Every d-spacing the cell allows down to dmin (no extinction rule applied: a line the table
    leaves unindexed is judged by whether ANY reflection of the cell sits there), sorted."""
    from pxrd_review import extra_checks as X
    a, b, c = cell['a'], cell['b'], cell['c']; al, be, ga = cell['α'], cell['β'], cell['γ']
    dmin = max(dmin, 0.8)
    hm, km, lm = (int(1.2 * x / dmin) + 1 for x in (a, b, c))
    while (2 * hm + 1) * (2 * km + 1) * (2 * lm + 1) > 300000:                    # a huge cell: the enumeration is capped, the finest lines go unchecked
        dmin *= 1.25; hm, km, lm = (int(1.2 * x / dmin) + 1 for x in (a, b, c))
    ds = []
    for h in range(-hm, hm + 1):
        for k in range(-km, km + 1):
            for l in range(0, lm + 1):                                            # Friedel: (hkl) and (-h-k-l) share a d
                if l == 0 and (k < 0 or (k == 0 and h <= 0)):
                    continue
                s2 = X._dstar2(a, b, c, al, be, ga, h, k, l)
                if s2 > 0 and s2 <= 1 / (dmin * dmin):
                    ds.append(1 / math.sqrt(s2))
    ds.sort()
    return ds

def _on_cell(ds, d, tol):
    """Whether some reflection of the cell (ds sorted) lies within tol of d."""
    import bisect
    i = bisect.bisect_left(ds, d * (1 - tol))
    return i < len(ds) and ds[i] <= d * (1 + tol)

def cell_check(path, cif=None, text=None, table=None):
    """The paper's powder table against a unit cell: every calculated line's d recomputed from its
    h k l and the cell — the .cif's when given and it reproduces the table, else the cell the paper
    states (the powder one first) — and every observed line looked for among the calculated ones.
    A paper is blamed only for isolated lines off by more than five times the table's own median
    deviation (never under 0.5 %, up to 15 %: a mistyped or mis-indexed line); a
    table whose lines sit 0.5–1.5 % off throughout was computed with a slightly different cell, and
    one that follows no cell the paper gives is reported as such, both unverified; a line off by
    more than 15 % is the reader's pairing, counted and left out.
    -> {'status': 'checked' | 'none' (no table) | 'unindexed' (no h k l) | 'nocell', 'head', 'lines',
    'source', 'cell', 'n', 'agree' (within 0.5 %), 'loose' (within 1.5 %), 'bad': [(d, I, hkl,
    d_cell, dev, sign_flipped)] (the lines blamed), 'wild': n, 'unmatched_obs', 'pages', 'tried'}."""
    obs, calc, pages = table if table is not None else pxrd_table(path, with_pages=True)   # extract() has read it already
    where = ' (p%d)' % pages[0] if pages and not path.lower().endswith('.docx') else ''   # a manuscript's "pages" are its tables
    if not obs and not calc:
        return {'status': 'none', 'lines': [], 'pages': pages}
    if not calc:
        return {'status': 'unindexed', 'pages': pages, 'lines': [],
                'head': 'powder table%s: %d observed lines, none indexed — nothing to check against a cell' % (where, len(obs))}
    cells = []
    if cif:
        from pxrd_review import extra_checks as X
        cc = _cell_floats((X.parse_cif(cif) or {}).get('cell') or {})
        if cc:
            cells.append(('.cif', cc))
    cells += _paper_cells(text_of(path) if text is None else text)[:6]
    if not cells:
        return {'status': 'nocell', 'pages': pages, 'n': len(calc), 'lines': [],
                'head': 'powder table%s: %d indexed lines but no cell to check them against (no .cif, none stated in the text)' % (where, len(calc))}
    scored = [(lab, cell) + _cell_score(cell, calc) for lab, cell in cells]
    scored = [x for x in scored if x[2]]
    if not scored:
        return {'status': 'nocell', 'pages': pages, 'n': len(calc), 'lines': [],
                'head': 'powder table%s: %d indexed lines but no cell to check them against' % (where, len(calc))}
    good = [x for x in scored if x[3] >= 0.8 * x[2]]
    ref = next((x for x in good if x[0] == '.cif'), None) or (good[0] if good else max(scored, key=lambda x: x[3] / x[2]))
    lab, cell, n, loose, rows = ref
    src = _CELL_SRC.get(lab, 'the cell the paper states')
    tight = [r for r in rows if abs(r[4]) <= _CELL_TOL]
    wild = [r for r in rows if abs(r[4]) > _CELL_WILD]
    devs = sorted(abs(r[4]) for r in rows if abs(r[4]) <= _CELL_LOOSE)
    tol = max(_CELL_TOL, 5 * devs[len(devs) // 2]) if devs else _CELL_TOL   # an outlier is judged against the table's own scatter: five times its median deviation, never under 0.5 %
    systematic = n - len(tight) > max(5, 0.1 * n)                     # most lines beyond 0.5 %: the cell, not the lines
    off = [r for r in rows if (max(tol, 0.03) if systematic else tol) < abs(r[4]) <= _CELL_WILD]
    lines = []
    cif_row = next((x for x in scored if x[0] == '.cif'), None)
    if cif_row and lab != '.cif':
        lines.append("the .cif's cell (%s) reproduces only %d of %d lines — another setting or determination than the powder cell; the table was checked against the paper's" % (_cell_str(cif_row[1]), cif_row[3], cif_row[2]))
    bad = []
    if loose < 0.8 * n:
        head = 'powder table%s: %d indexed lines vs %s (%s) — only %d follow it within 1.5 %%' % (where, n, src, _cell_str(cell), loose)
        lines.append('the table follows no cell the paper gives (best of %d tried) — the cell, the indexing, or the reading of the table is off [unverified]' % len(scored))
    else:
        head = 'powder table%s: %d indexed lines vs %s (%s) — %d agree within 0.5 %%, %d within 1.5 %%' % (where, n, src, _cell_str(cell), len(tight), loose)
        if systematic:
            shifted = [r for r in rows if _CELL_TOL < abs(r[4]) <= _CELL_LOOSE]
            lines.append('%d of %d lines sit 0.5–%.1f %% off — computed with a cell slightly different from the one given, or the cell is quoted to too few figures [unverified]'
                         % (len(shifted), n, max(abs(r[4]) for r in shifted) * 100 if shifted else _CELL_LOOSE * 100))
        bad = sorted(off, key=lambda r: -abs(r[4]))
        for d, I, hkl, dc, dev, flipped in bad:
            lines.append('%g (%d %d %d) does not follow the cell: it gives %.4f (%+.1f %%)%s' % (d, hkl[0], hkl[1], hkl[2], dc, dev * 100, ' — a weak line' if I is not None and I <= 5 else ''))
    if wild:
        lines.append('%d line%s the reader could not pair with the cell (off by more than 15 %%: a row it read across two blocks) — left out [unverified]'
                     % (len(wild), 's' if len(wild) > 1 else ''))     # the reader's fault, not the paper's: never a flag
    unmatched = []; off_cell = []
    if obs and loose >= 0.8 * n:
        cds = [dc for _, _, _, dc, _, _ in rows] + [d for d, _, _ in calc]
        tol_obs = _CELL_LOOSE if systematic else tol                     # an observed line pairs within the table's own scatter
        unmatched = [d for d, I in obs if not any(abs(d - dc) / d <= tol_obs for dc in cds)]
        if unmatched:
            # a table lists observed lines it never indexes (weak ones, or all of them in a separate
            # table): they are judged against every reflection the cell allows — the reading of the
            # column is vouched for when they sit on the cell, and a line on no reflection is worth a look
            try:
                allowed = _reflection_ds(cell, 0.98 * min(d for d, _ in obs))
                off_cell = [d for d in unmatched if not _on_cell(allowed, d, tol_obs)]
            except Exception:
                off_cell = []
            on_cell = len(unmatched) - len(off_cell)
            if on_cell:
                lines.append('%d observed line%s without a calculated partner in the table lie%s on reflections of the cell — left unindexed by the table' % (on_cell, 's' if on_cell > 1 else '', '' if on_cell > 1 else 's'))
            for d in off_cell[:6]:
                lines.append('observed %g: no calculated line within %.1f %% and no reflection of the cell there — a typo, or another phase [unverified]' % (d, 100 * tol_obs))
            if len(off_cell) > 6:
                lines.append('… and %d more observed lines on no reflection of the cell [unverified]' % (len(off_cell) - 6))
    return {'status': 'checked', 'source': lab, 'cell': cell, 'n': n, 'agree': len(tight), 'loose': loose, 'bad': bad, 'wild': len(wild), 'unmatched_obs': unmatched, 'off_cell': off_cell,
            'head': head, 'lines': lines, 'pages': pages, 'tried': [(x[0], x[3], x[2]) for x in scored]}

# ----------------------------------------------------------------------------- the paper's empirical formula

def empirical_formula(text):
    return _formula(text)[:4]

def _formula(text):
    fs = _formulas(text)
    return fs[0] if fs else ('', {}, [], {}, '', '')

def _attributed_to(ctx):
    """The mineral a formula sentence attributes its formula to ('the empirical formula of hancockite
    is …'), lowercased, or ''."""
    low = (ctx or '').lower()
    ms = list(re.finditer(r'formulae?\s*(?:\([^)]*\)\s*)?(?:of|for)\s+(?:the\s+)?(?:holotype\s+|type\s+)?([a-zà-ÿč]{4,}ite(?:-\([a-z]+\))?)', low)) or \
         list(re.finditer(r'([a-zà-ÿč]{4,}ite(?:-\([a-z]+\))?)[^.]{0,40}(?:has|have|gave|gives|yield\w*|with|is)\b[^.]{0,30}formula', low))
    m = ms[-1] if ms else None                                         # the attribution nearest the formula
    return m.group(1) if m and m.group(1).split('-')[0] not in _NOT_MINERALS else ''

def _formulas(text, name=''):
    """Every formula sentence of the paper that parses, in text order, distinct by its counts:
    [(formula text, counts, issues, charges, kind, context), …]. A paper's abstract and body may
    print different formulas — the check uses the one the table reproduces and reports the other.
    A formula the sentence attributes to another mineral (a supporting phase) is left out."""
    stem = re.sub(r'-\(.*\)$', '', name or '').lower()[:8]
    out = []; cited = []
    for cand in _formula_iter(text):
        who = _attributed_to(cand[5])
        if stem and who and stem not in who:
            continue                                                  # 'the empirical formula of hancockite is …'
        if any(_same_counts(cand[1], o[1]) for o in out + cited):
            continue
        if re.search(r'formula[e]?\s+of\s+[\w-]+\s*\([^()]*(?:1[89]|20)\d\d[a-z]?\)[^.()]{0,80}?\s(?:is|was|=|:)', cand[5][-200:], re.I):
            cited.append(cand); continue                              # 'the formula of koragoite (Voloshin et al. 1992) is …': the original description's, not this paper's
        out.append(cand)
        if len(out) >= 4:
            break
    out += cited[:max(0, 4 - len(out))]
    texts = [o[0].replace(' ', '') for o in out]
    return [o for o, t_ in zip(out, texts) if not any(u != t_ and u.startswith(t_) for u in texts)]   # 'Pb4.95Ca0.02Sr0.02' cut by a page break

def _same_counts(a, b):
    keys = set(a) | set(b)
    return all(abs(a.get(k, 0.0) - b.get(k, 0.0)) <= 0.011 + 0.01 * max(a.get(k, 0.0), b.get(k, 0.0)) for k in keys if k not in ('H', 'O'))

def _norm_text(text):
    """The paper's text as the formula finder reads it: control and zero-width codes out, 'þ' = '+', a
    font's ':' for '.' and ð Þ for brackets, a charge superscript set mid-number rejoined."""
    t = re.sub(r'[\x00-\x08\x0b-\x1f\u200b\u200c\u200d\u2060\ufeff]', '', text.replace('þ', '+').replace('\xad', ''))   # soft hyphens, zero-width and glyph control codes
    t = re.sub(r'(\d):(\d)', r'\1.\2', t).replace('ð', '(').replace('Þ', ')')        # a font that prints '.' as ':' and brackets as ð Þ
    t = re.sub(r'([A-Z][a-z]?\d*)\. (\d\+) (\d+)(?=[A-Za-z□\)\]])', r'\1.\3', t)     # 'Fe0. 2+ 20o0.47': the charge superscript set between the digits of 0.20
    return t

def _formula_iter(text):
    """The empirical formula the paper states ("The empirical formula (based on 28 O apfu) is …"):
    (formula text, {element: apfu}, issues). Journal notation is turned into the ICDD form the
    epma parser reads: 'Fe3+1.52' / 'Fe3þ 1.52' charges, Σ sums, '·8H2O', '□' vacancies.
    The formula is the first formula-shaped run (a bracket, or an element with a decimal count)
    within 300 characters after 'empirical … formula' that parses — the wording between varies
    (is / = / : / being / as follows / can be written as / with the parenthetical basis)."""
    t = _norm_text(text)
    # the empirical formula first; a structural / crystal-chemical one (site populations from the
    # refinement) only when the paper gives no other — its coefficients need not match the analysis
    triggers = [(m, 'empirical') for m in re.finditer(r'empirical (?:mineral |chemical |crystal[- ]chemical |holotype |structural )?formula[e]?|эмпирическ\w+ (?:кристаллохимическ\w+ |химическ\w+ )?формул\w*', t, re.I)]
    triggers += [(m, 'structural') for m in re.finditer(r'(?<!empirical )(?:crystal[- ]chemical|structural) (?:chemical )?formula[e]?', t, re.I)]
    triggers += [(m, 'stated') for m in re.finditer(r'(?<!empirical )(?<!crystal )(?<!ideal )(?<!simplified )(?<!general )chemical formula(?: of [\w-]+)?(?:\s*(?:is|=|:))?', t, re.I)]   # 'the chemical formula of ciriottiite is …', a crystal-data table's 'Chemical formula …'
    triggers += [(m, 'total') for m in re.finditer(r'(?:total|bulk|raw|uncorrected|measured|overall)\s+composition(?: of the (?:sample|material|analys[ei]s))?[^.]{0,30}?(?:results? in|is|gives|corresponds to|of)', t, re.I)]   # 'the total composition of the sample results in K0.10Na2.95…'
    triggers += [(m, 'stated') for m in re.finditer(r'(?:gives|give|giving|yield(?:ed|s|ing)?|resulted in|resulting in|leading to|lead(?:s)? to|corresponds? to|with)(?:\s+(?:the|an?))?(?:\s+following)?(?:\s+(?:average|mean|general|resulting))?(?:\s+(?:chemical|crystal[- ]chemical|unit|mineral))?\s+formula(?:e)?(?:\s*(?:of|is|:))?|average (?:chemical )?formula(?:\s*(?:of|is|:))?', t, re.I)]   # 'gives the formula …', 'with an average formula of …'
    for m, kind in triggers:
        if kind == 'stated' and re.search(r'site[- ]scattering|structure refinement|site populations?|refined (?:site|occupanc)', t[max(0, m.start() - 260):m.start()], re.I):
            kind = 'structural'                                       # 'EMPA and site scattering values … yielded the formula': the refinement's, not the analysis alone
        window = t[m.end():m.end() + 300]; run = t[m.end():m.end() + 1200]
        for c in re.finditer(r'(?<![A-Za-z0-9])(?:[A-Z][a-z]?\d*)?[\(\[{□]|(?<![A-Za-z])[A-Z][a-z]?(?:\d?\+)?\s?\d+\.\d+', window):
            if re.search(r'(?:based|derived|obtained|calculated) (?:up)?on the (?:crystal[- ])?structure (?:refinement|determination)|from the (?:structure |crystal-structure )?refinement|site[- ]scattering', window[:c.start()], re.I):
                kind = 'structural'                                   # 'the empirical formula … based on the structure refinement is …': the refinement's numbers
            f = run[c.start():]                                     # the formula starts within 300 chars; it may run much longer
            f = re.split(r'\s(?=[a-z])|\s\(?(?=[A-Z][a-z]{2,}\b)|, |; |\. |\s=\s', f)[0]   # the formula ends where prose ('Crystal B', '(Sample 1)'), the next clause or '= (structural form)' starts
            if not re.search(r'\d\.\d', f[:30]) or len(re.findall(r'[A-Z][a-z]?\d*\.\d+', f)) < 1:
                continue                                                # '(C + S) = 2 apfu', '(atoms per formula unit)'
            decs = re.findall(r'\d+\.(\d+)', f)
            if kind != 'empirical' and sum(1 for d in decs if d.rstrip('0') not in ('', '5')) < 2 and sum(1 for d in decs if len(d) >= 2) < 2:
                continue                                                # 'Ca4(Al0.5Si0.5)2Si4O16': an ideal formula, not the analysis (Au3.00Tl1.01Te2.00 is one)
            if kind == 'empirical' and not decs:
                continue                                                # no decimal at all: the ideal one after all
            norm = _journal_to_icdd(f)
            counts, ox, issues = EP.parse_icdd_formula(norm, has_sulfur='S' in counts_guess(norm))
            issues = [x for x in issues if not re.match(r'unrecognised symbol [A-Z]\b', x)]  # a bare site label, not a fault
            if counts and sum(1 for k in counts if k not in ('H', 'O')) >= 1 and not any('garbled' in x for x in issues):
                ox_all = {}
                for el, chg in re.findall(r'([A-Z][a-z]?)(?:\d+(?:\.\d+)?)? \+(\d)', norm):
                    ox_all.setdefault(el, set()).add(int(chg))
                yield f.strip(), counts, issues, ox_all, kind, t[max(0, m.start() - 200):m.end() + c.start() + 40]
                break

_MULTI = {'Fe': (2, 3), 'Mn': (2, 3, 4), 'V': (3, 4, 5), 'Cr': (3, 6), 'Cu': (1, 2), 'Ti': (3, 4), 'Ce': (3, 4), 'Eu': (2, 3),
          'U': (4, 6), 'As': (3, 5), 'Sb': (3, 5), 'S': (4, 6), 'Se': (4, 6), 'Te': (4, 6), 'Co': (2, 3), 'Ni': (2, 3), 'Sn': (2, 4),
          'Pb': (2, 4), 'Tl': (1, 3), 'Bi': (3, 5), 'Nb': (4, 5), 'W': (4, 6), 'Mo': (4, 6)}

def ideal_oxidation(html):
    """Mindat's ideal formula ('KFe<sup>3+</sup><sub>3</sub>(S<sup>6+</sup>O<sub>4</sub>)…') -> {element: {charges}}."""
    out = {}
    for el, chg, sign in re.findall(r'([A-Z][a-z]?)(?:<sub>[^<]*</sub>)?<sup>(\d?)([+\-−])</sup>', html or ''):
        out.setdefault(el, set()).add(int(chg or 1) * (1 if sign == '+' else -1))
    return out

def species_record(name):
    """The species in Mindat's offline cache: {'formula', 'ox': {el: {charges}}, 'elements'} or None.
    Exact name only — a variety or another Levinson suffix has other elements."""
    if not name:
        return None
    try:
        from pxrd_review import extra_checks as XC
        rec = XC.mindat_struct(name, exact=True)
    except Exception:
        return None
    if not rec or not rec.get('formula'):
        return None
    return {'formula': rec['formula'], 'ox': ideal_oxidation(rec['formula']), 'elements': set(rec.get('elements') or [])}

def oxide_for(el, v):
    """The oxide of `el` at charge v: Fe 3 -> Fe2O3, Mn 4 -> MnO2, Cu 1 -> Cu2O."""
    if v <= 0:
        return None
    return '%sO%s' % (el, '' if v == 2 else v // 2) if v % 2 == 0 else '%s2O%s' % (el, '' if v == 1 else v)

def oxide_alternatives(wt, ox_paper, species):
    """Table oxides whose valence disagrees with the paper's formula (or, failing that, the
    species' ideal formula): [(old constituent, new constituent, evidence)]. The wt% is carried
    over per cation (FeO 10.00 -> Fe2O3 11.11). An element the formula carries in two states is
    left alone — the split was the authors' own charge balance."""
    out = []
    for c in wt:
        try:
            k = EP.parse_constituent(c)
        except ValueError:
            continue
        el = k.element
        if k.n_o == 0 or k.kind in ('water', 'element-anion', 'element') or el not in _MULTI:
            continue
        vt = k.charge
        ve = None; why = ''
        pv = ox_paper.get(el)
        if pv and len(pv) == 1:
            ve, why = next(iter(pv)), 'the formula writes %s%d+' % (el, next(iter(pv)))
        elif pv:
            continue                                            # Fe2+ and Fe3+ both: the authors' own split
        elif species and el in species['ox'] and len(species['ox'][el]) == 1:
            ve = next(iter(species['ox'][el])); why = "Mindat's ideal formula has %s%d+" % (el, ve)
        if ve is None or ve == vt or ve not in _MULTI[el] or ve <= 0:
            continue
        new = oxide_for(el, ve)
        if new and new != c and new not in wt:
            out.append((c, new, why))
    return out

def _convert(wt, old, new):
    a, b = EP.parse_constituent(old), EP.parse_constituent(new)
    w2 = dict(wt); v = w2.pop(old)
    w2[new] = round(v * (b.mw / b.n_cat) / (a.mw / a.n_cat), 3)
    return w2

_LN = ('La', 'Ce', 'Pr', 'Nd', 'Sm', 'Eu', 'Gd', 'Tb', 'Dy', 'Ho', 'Er', 'Tm', 'Yb', 'Lu')

def counts_guess(norm):
    return set(re.findall(r'[A-Z][a-z]?', norm))

def _journal_to_icdd(f):
    """'(NH4)1.895Na0.065' -> '( N H4 )1.895 Na0.065'; 'Fe3+1.52' -> 'Fe1.52 +3'; 'Σ2.000' kept; '·8H2O' -> '!8 H2 O'."""
    f = f.replace('{', '(').replace('}', ')')                                            # '{[(NH4)2.13K0.87]Σ3.00(H2O)}{…}': braces are brackets
    f = re.sub(r'(?<![A-Za-z])(S|Se|Te|Fe|Mn|Cu|Ti|V|Cr|As|Sb|U|Ce|Eu|Co|Ni|Sn|Pb|Tl|Bi|Nb|W|Mo)([1-6]) (\d+\.\d+)(?![\d.])', r'\1 \3', f)   # 'S2 2.60': a charge whose sign the text layer lost (S2− 2.60)
    f = re.split(r'\s*\(?\bZ\s*=', f)[0]                                               # '(Z = 2)' and what follows
    f = re.sub(r'^\(\s*[A-Za-z\-+, ]+\)\s*', '', f)                                      # '(SREF) Cu2Fe0.84…', '(O + F) (Na2.74…': a tag or the basis, not a group
    f = re.sub(r'(?<![A-Za-z(])((?:P|S|As|Si|V|Se|Cr|Mo|W|B|C|N)O([2-4]))(\d+\.\d{2})(?![\d.])', r'(\1)\3', f)   # 'PO43.02' = (PO4)3.02, the brackets flattened
    f = re.sub(r'(\d\.\d{2})(?=\d+\.\d+\s*H2O)', r'\1·', f)                             # '(OH)3.524.13H2O': the hydrate dot lost
    f = re.sub(r'([)\]])(\d)(\d{1,2})(?=\s*H2O)', lambda m_: m_.group(0) if int(m_.group(2) + m_.group(3)) <= 12 else m_.group(1) + m_.group(2) + '·' + m_.group(3), f)   # '(PO4)22H2O' = (PO4)2·2H2O
    f = re.sub(r'(\d)\(\d+\)', r'\1', f)                                                # esds after counts: Cu9.71(8)
    def _dot(m_):                                                                        # '·8H2O' / ')·4.13H2O' / 'x.8H2O': the hydrate dot; '4.13' after it keeps its decimal point
        before = re.sub(r'\d+$', '', f[:m_.start()]).rstrip()
        return m_.group(0) if m_.group(0) == '.' and before.endswith(('!', '·', '•', '⋅', '∙')) else ' ! '
    f = re.sub(r'[·•⋅∙]|(?<=[)\]A-Za-z0-9])\.(?=\s*\d*\.?\d*\s*H2O(?![\d.]))', _dot, f).replace('□', '?').replace('∑', 'Σ').replace('−', '-').replace('–', '-')   # 'La0.01H2O0.21': a count, not a hydrate dot
    f = re.sub(r'<\s*(?=\d)', '', f)                                                      # 'Pt4+ <0.01': below detection, the bound stands
    if len(re.findall(r'\d+\.\d+\s*[-–]\s*\d+\.\d+', f)) >= 2:
        f = re.sub(r'(\d+\.\d+)\s*[-–]\s*(\d+\.\d+)', lambda m_: '%.3f' % ((float(m_.group(1)) + float(m_.group(2))) / 2), f)   # 'Ge0.91-0.97': ranges, their middles
    else:
        f = re.sub(r'(\d+\.\d+)\s*[-–]\s*(\d+\.\d+)', r'\1 ?\2', f)                     # 'Th0.01–0.51': a vacancy glyph printed as a dash
    f = re.sub(r'\(([A-Z][a-z]?),([A-Z][a-z]?)\)(?:\d\+)?\s*(\d+\.\d+)', r'\1\3', f)     # '(Th,U)4+ 0.54': the first element carries the count
    f = re.sub(r'·\s*n\s*H2O', '', f)                                                    # '·nH2O': an unstated hydrate
    f = re.sub(r'([)\]])\s*R(?=\s*\d)', r'\1 Σ', f)                                     # a Σ printed as R
    f = re.sub(r'([)\]])\s*6\s*=\s*(?=\d)', r'\1 Σ', f)                                  # ')6=0.96': a Σ printed as 6 (font), with its '='
    f = re.sub(r'Σ\s*=\s*', 'Σ', f)                                                       # 'Σ=48.49'
    f = re.sub(r'(?<![A-Za-z])[AMXTRZ]\d?(?:\+[AMXTRZ]\d?)?\s*(?=[\(\[]|Σ|6\d)', ' ', f)   # site labels A1[…], M2+M3(…), M3 Σ3.95(…), ')61.00Y('
    f = re.sub(r'(?<![A-Za-z])[YVW](?=[\(\[])', ' ', f)                                  # Y / V / W site labels right before a bracket (the elements keep their counts)
    f = re.sub(r'(?<![A-Za-z])[XYZTVWAMR](?=[A-Z][a-z]?\d*\.\d)', ' ', f)                # 'ZAl6.00[': a label glued to the element it holds
    f = re.sub(r'(?<![A-Za-z])A(?=\d*\.\d)', '?', f)                                      # 'A1.91': a vacancy printed as A (font)
    f = re.sub(r'(?<=[\d\s(])[ho](?=\d*\.\d)', '?', f)                                  # a vacancy printed as 'h' or 'o' (font)
    labelled = {}                                                                        # 'Mg1(Mg1.42…)Mg2(Mg1.71…)': sites named after their element
    for el, d in re.findall(r'(?<![A-Za-z])([A-Z][a-z]?)(\d)(?=[\(\[])', f):
        labelled.setdefault(el, set()).add(d)
    def _site(m_):
        el, body = m_.group(1), m_.group(4)
        named = '1' in labelled.get(el, ()) or len(labelled.get(el, ())) >= 2                # numbered from 1, or a run: a label; 'Si4(' is a count
        return ' ' + m_.group(3) + body if named and re.search(r'(?<![A-Za-z])' + el + r'\d*\.\d', body) else m_.group(0)
    f = re.sub(r'(?<![A-Za-z])([A-Z][a-z]?)(\d)([\(\[])([^)\]]*)', _site, f)
    f = re.sub(r'\s*Σ\s*\d+(?:\.\d+)?', lambda m_: m_.group(0) if f[:m_.start()].rstrip().endswith((')', ']')) else ' ', f)   # a Σ without a bracket: the sum of the run before it
    f = re.sub(r'([A-Z][a-z]?)\s*\+\s*(\d+\.\d+)', r'\1\2 +1', f)                      # 'Cu+ 0.99': a charge without a digit
    if re.search(r'[)\]]\s*6\d{1,2}\.\d', f):                                            # ']60.90', ')616.15', ')62 O54': a Σ printed as '6' (font)
        f = re.sub(r'([)\]])\s*6(\d{1,2}(?:\.\d+)?)(?![\d.])', r'\1 Σ\2', f)
    f = re.sub(r'(?<![A-Za-z])(REE|Ln|TR)\*?(?![a-z])', 'Ce', f)                          # a grouped lanthanide stands in as Ce ('Ln0.10', 'REE*0.01')
    f = re.sub(r'\(\s*[Σ6]?\s*Me\s*=\s*[\d.]+\s*\)', ' ', f)                             # '(ΣMe = 3.97)' after the formula
    f = re.sub(r'(?<![A-Za-z])Me(?![a-z])', ' ', f)                                       # 'Me' is a metal placeholder, never an element
    f = re.sub(r'([A-Z][a-z]?)\+(\d)\s*(\d+\.\d+)', r'\1\3 +\2', f)                     # 'Fe+3 0.23': the sign before the digit
    f = re.sub(r'([A-Z][a-z]?)(\d+\.\d+)\s+(\d)\+', r'\1\2 +\3', f)                     # 'Fe0.25 3+': the charge after the count
    f = re.sub(r'(?<![(A-Za-z])((?:P|S|As|Si|V|C|B|Se|Cr|Mo|W)O\d)\)', r'(\1)', f)          # 'PO4)1.91': the opening bracket lost
    f = re.sub(r'([A-Z][a-z]?)\s*(\d(?:\.\d)?)\s*\+\s*(\d+\.\d+|\d+)', r'\1\3 +\2', f)   # Fe3+1.52, 'Mn 2+ 1.89' -> Mn1.89 +2; Fe2.5+ 0.25
    f = re.sub(r'([A-Z][a-z]?)(\d)\+', r'\1 +\2', f)                                  # Fe3+ (no count)
    # ')P2.01' after (Mn1.20Mg0.49Fe0.27Zn0.05), or after a nested group: a Σ printed as P — only where
    # the value is the group's own cation (+ vacancy) sum, so a real phosphorus count is never touched
    for m_ in reversed(list(re.finditer(r'\)P(\d\.\d{2})', f))):
        depth = 0; start = None
        for q in range(m_.start(), -1, -1):
            if f[q] == ')': depth += 1
            elif f[q] == '(':
                depth -= 1
                if depth == 0:
                    start = q; break
        if start is None:
            continue
        body = f[start + 1:m_.start()]
        if 'P' in re.findall(r'(?<![A-Za-z])[A-Z][a-z]?', body) or len(re.findall(r'\d\.\d+', body)) < 2:
            continue
        try:
            cnt = EP.parse_icdd_formula(body)[0]
        except Exception:
            continue
        parts = sum(v for e_, v in cnt.items() if e_ not in ('O', 'H')) + sum(float(x) for x in re.findall(r'\?(\d+\.\d+)', body))
        val = float(m_.group(1))
        if abs(parts - val) <= max(0.03, 0.03 * val):
            f = f[:m_.start() + 1] + ' Σ' + m_.group(1) + f[m_.end():]
    f = re.sub(r'(?<=[\d)])\s+(\d\.\d+)(?=\s*[A-Z)\]])', r' ?\1', f)                        # '(Mg1.24Ca0.69 0.06Mn0.01)': a vacancy glyph lost before its count

    f = re.sub(r'\(NH4\)', '( N H4 )', f)
    f = re.sub(r'(?<=!)\s*(\d+(?:\.\d+)?)?\s*H2O', lambda m: ' %s H2 O' % (m.group(1) or ''), f)
    f = re.sub(r'\(OH\)', '( O H )', f)
    f = re.sub(r'(?<![(\s!])H2O(\d+\.\d+)', r'( H2 O )\1', f)                             # 'H2O0.21' inside a group: a water group with its count
    f = re.sub(r'H2O', ' H2 O ', f)
    f = re.sub(r'([A-Z][a-z]?)(\d+\.\d+|\d+)', r'\1\2 ', f)                          # a space after every count
    f = re.sub(r'([)\]])\s*Σ\s*', r'\1 Σ', f)
    return re.sub(r'\s+', ' ', f).strip(' .')

# ----------------------------------------------------------------------------- the checks (manuscript tool)

def check_composition(ex, text):
    """Re-do the paper's reduction from its own table, basis and method, against its own empirical
    formula. When the paper prints several formulas (abstract and body), the one the table
    reproduces is the reference and every other one that disagrees with it is reported."""
    fs = _formulas(text, ex.get('name') or '')
    if not fs:
        return _check_formula(ex, text, ('', {}, [], {}, '', ''))
    results = [(fc, _check_formula(ex, text, fc)) for fc in fs[:3]]
    chosen = next(((fc, r_) for fc, r_ in results if r_ and r_['ok'] and r_['verified']), None) or results[0]
    fc, res = chosen
    if res is None:
        return res
    corrected = re.compile(r'impurit|admixture|admixed|subtract|removed|corrected for|contribution', re.I)
    emp = next((f_ for f_ in fs if f_[4] == 'empirical'), None)
    if fc[4] == 'total' and emp is not None and corrected.search(emp[5]):
        res['lines'].append('  the uncorrected composition the paper gives (%s) reproduces from the table; the paper then subtracts admixed phases before writing its formula %s — that corrected formula cannot be re-derived from the table alone' % (fc[0][:60], emp[0][:50]))
        res['ok'] = False; res['verified'] = False
        return res
    if fc[4] == 'empirical' and not res['ok'] and corrected.search(fc[5]):
        res['lines'].append('  not verifiable: the paper corrects the analysis for impurities or admixed phases before writing its formula — the residuals above are not evidence')
        res['verified'] = False
        return res
    stem = re.sub(r'-\(.*\)$', '', ex.get('name') or '').lower()[:8]
    if any('chosen by fit' in l or 'not the first (mean) column' in l or 'reproduces the formula; the one first read' in l or 'analyses under it averaged' in l for l in res['lines']):
        return res                                                    # the reference column or table was found by fit or fallback: no formula comparison
    for other, r_o in results:
        if other is fc or not other[1]:
            continue
        if re.search(r'formula[es]\b', other[5][-160:], re.I) or re.search(r'formula[es]\b', fc[5][-160:], re.I):
            continue                                                  # 'the empirical formulas X and Y': two minerals or samples in one sentence
        if re.search(r'rang(?:ing|es|e) (?:from|between)|varies from|between', other[5][-120:], re.I):
            continue                                                  # 'the empirical formula ranging from X to Y': endpoints, not the mean
        a, b = fc[1], other[1]
        shared = [k for k in a if k in b and k not in ('H', 'O')]
        n_dec = sum(1 for d in re.findall(r'\d+\.(\d+)', other[0]) if len(d) >= 2)
        if len(shared) < 3 or n_dec < 2 or other[4] != 'empirical':
            continue                                                  # an ideal or simplified formula, or one from a weaker trigger — not another empirical one (integers from a normalisation are fine)
        major_a = {k for k, v in a.items() if v >= 0.05 and k not in ('H', 'O')}; major_b = {k for k, v in b.items() if v >= 0.05 and k not in ('H', 'O')}
        if not (major_a | major_b) <= (set(a) & set(b)) or other[4] == 'structural':
            continue                                                  # another phase (a major element the other lacks), or the refinement's formula
        ctx_o = other[5].lower()
        who = _attributed_to(ctx_o)
        if who and stem and stem not in who:
            continue                                                  # the sentence attributes this formula to another mineral
        tags_a = set(sample_hints(fc[5][-200:])['codes']); tags_b = set(sample_hints(other[5][-200:])['codes'])
        orig = re.compile(r'original (?:\w+ )?(?:sample|material|specimen|description)', re.I)
        if (tags_a and tags_b and not tags_a & tags_b) or (bool(orig.search(fc[5][-200:])) != bool(orig.search(other[5][-200:]))):
            continue                                                  # 'beraunite (FR)' against 'beraunite (NM)', or the original sample's formula: two samples, not one paper disagreeing with itself
        diffs = ['%s %.3g vs %.3g' % (k, b[k], a[k]) for k in shared if abs(a[k] - b[k]) > max(0.02, 0.1 * max(a[k], b[k]))]
        missing = [k for k in b if k not in a and k not in ('H', 'O') and b[k] >= 0.02] + [k for k in a if k not in b and k not in ('H', 'O') and a[k] >= 0.02]
        if diffs or missing:
            tn = _norm_text(text); p_o = tn.find(other[0][:30]); p_f = tn.find(fc[0][:30])   # the formulas were cut from the normalised text
            if p_o < 0:
                continue                                                    # cannot place the other formula: no abstract / body call
            in_abstract = p_o < min(4000, p_f if p_f >= 0 else 4000)
            if in_abstract:                                             # the abstract's formula against the body's: a fault
                res['lines'].append('  the formula in the abstract does not agree with the one the table reproduces: %s%s' % (
                    '; '.join(diffs[:6]), ('; only one of them carries ' + ', '.join(missing[:4])) if missing else ''))
                res['ok'] = False
            else:                                                       # a second formula in the body: often an alternative normalisation — a note
                res['lines'].append('  another formula sentence in the body differs from the one the table reproduces (an alternative normalisation?): %s%s' % (
                    '; '.join(d.replace(' vs ', ' / ') for d in diffs[:6]), ('; only one of them carries ' + ', '.join(missing[:4])) if missing else ''))
    return res

def _check_formula(ex, text, fcand):
    """Re-do the paper's reduction from its own table, basis and method, against its own empirical
    formula. -> {'ok', 'lines', 'formula', 'basis', 'result'}; None when the paper gives no table
    or no formula."""
    e = ex.get('epma')
    if not e or not e['rows']:
        return None
    ftxt, counts, f_issues, ox_paper, f_kind, f_ctx = fcand
    wt = {}
    for r in e['rows']:
        wt.setdefault(r['constituent'], r['mean'])                # a repeated constituent: the first (wt%) row
    species = species_record(ex.get('name'))
    excluded = [e_ for e_ in excluded_elements(text) if e_ not in counts]
    if excluded:
        wt = {c: v for c, v in wt.items() if not (_parses(c) and EP.parse_constituent(c).element in excluded)}   # 'calculated without Al'
        e = _without_elements(e, excluded)                        # … and out of every other column and candidate table the check may fall back to
    if not counts:
        return _derived_composition(ex, wt, species)
    if 'S' in counts and not any(_parses(c) and EP.parse_constituent(c).element == 'S' for c in wt) and re.search(r'\sS\d\.\d\d', ftxt):
        counts = {k: v for k, v in counts.items() if k != 'S'}      # 'Fe2+ 0.14Ti0.18 S3.00': a Σ printed as S, the brackets lost; the table has no sulfur
    grouped_ln = bool(re.search(r'(?<![A-Za-z])(REE|Ln|TR)\*?\d*\.\d', ftxt)) and not any(el in counts for el in _LN if el != 'Ce')
    def _fold_ln(wt_):
        """A formula that writes the lanthanides as one REE: the table's Ln2O3 rows as one Ce2O3-equivalent."""
        if not grouped_ln:
            return wt_
        ln_rows = [c for c in wt_ if _parses(c) and EP.parse_constituent(c).element in _LN]
        if not ln_rows:
            return wt_
        wt_ = dict(wt_); ce = EP.parse_constituent('Ce2O3')
        tot = sum(wt_[c] * (ce.mw / EP.parse_constituent(c).mw) * (EP.parse_constituent(c).n_cat / 2) for c in ln_rows)
        for c in ln_rows:
            wt_.pop(c)
        wt_['Ce2O3'] = round(tot, 3)
        return wt_
    wt = _fold_ln(wt)
    bases = [ex['basis']] if ex.get('basis') else []
    groups = {}
    for grp, tot in re.findall(r'[\(\[]([^()\[\]]*)[\)\]]\s*Σ\s*(\d+(?:\.\d+)?)', _journal_to_icdd(ftxt)):
        els = tuple(e_ for e_ in dict.fromkeys(re.findall(r'[A-Z][a-z]?', grp)) if e_ in EP.ATOMIC_WEIGHTS and e_ not in ('O', 'H'))
        if len(els) >= 2 and float(tot) > 0.5:
            groups[els] = groups.get(els, 0.0) + float(tot)       # two sites of the same elements: one sum
    for els, tot in groups.items():
        inside = sum(sum(EP.parse_icdd_formula(grp)[0].get(e_, 0.0) for e_ in els) for grp, t_ in re.findall(r'[\(\[]([^()\[\]]*)[\)\]]\s*Σ\s*(\d+(?:\.\d+)?)', _journal_to_icdd(ftxt))
                     if tuple(e_ for e_ in dict.fromkeys(re.findall(r'[A-Z][a-z]?', grp)) if e_ in EP.ATOMIC_WEIGHTS and e_ not in ('O', 'H')) == els)
        if abs(inside - sum(counts.get(e_, 0.0) for e_ in els)) > 0.05:
            continue                                             # an element of the group occurs elsewhere too (Si4 outside (S1.61Si0.32)): not a basis
        if not set(els) <= set(EP.parse_constituent(c).element for c in wt if _parses(c)):
            continue                                             # a group with an unanalysed element (B, Be, Li) cannot be the basis
        try:
            b_ = EP._parse_basis('%s=%.3f' % ('+'.join(els), tot))
            if b_ not in bases:
                bases.append(b_)                                 # 'normalised to Σ(Pb + REE) = 2'
        except Exception:
            pass
    has_s = any(c in ('S', 'SO3', 'SO2') for c in wt)
    # a flag needs a deviation beyond what rounding, atomic weights and oxide conventions give: 5 % / 0.03 apfu
    stated = ex.get('basis')
    def _rep(wt_, bs):
        return EP.replicate_formula(wt_, counts, bs, ex.get('name') or 'paper', tol_abs=0.03, tol_rel=0.05) if bs else None
    def _best(wt_):
        # The basis the paper states wins whenever it reproduces the formula. A Σ-group sum read off
        # the formula ('(Sc0.90Mn0.08Fe0.01Pb0.01)Σ1.00' -> Sc+Mn+Fe+Pb=1) normalises cations to their
        # own printed sum, so it reproduces them by construction and outscores a true anion basis in
        # the third decimal — on the corpus that discarded a correct 'based on 6 O apfu' 58 times in
        # 245 papers and reported the paper's own basis as failing. Group sums and derived candidates
        # are a fallback for a stated basis that genuinely fails, never a competitor to one that works.
        r_ = _rep(wt_, [stated]) if stated else None
        if r_ is not None and r_['score'] <= 0.03 and not r_.get('factor'):
            return r_
        r_ = _rep(wt_, bases) or r_
        if r_ is None or r_['score'] > 0.03 or r_.get('factor'):
            alt = _rep(wt_, [b for b in EP.basis_candidates(counts) if b not in bases])
            if alt is not None and (r_ is None or alt['score'] < r_['score'] - 0.01):
                r_ = alt
        return r_
    converted = []; column_note = None
    need = set(counts) - {'H', 'O'} - set(EP.CALCULATED_ELEMENTS)

    def _resolve(wt0, e0):
        """One starting column through the whole chain: the best basis, another table of the paper
        that carries every element of the formula, the other columns, the oxide forms."""
        notes = []; cn = None; wt_ = _fold_ln(dict(wt0)); e_ = e0
        r_ = _best(wt_)
        have = set(EP.parse_constituent(c).element for c in wt_ if _parses(c))
        if (r_ is None or r_['score'] > 0.03 or not need <= have) and e_.get('candidates'):
            for cand in e_['candidates']:
                wt_c = {}
                for row in cand['rows']:
                    wt_c.setdefault(row['constituent'], row['mean'])
                have_c = set(EP.parse_constituent(c).element for c in wt_c if _parses(c))
                if len(wt_c) < 3 or not need <= have_c or wt_c == wt_:
                    continue
                r_c = _best(wt_c)
                if r_c is not None and (r_ is None or r_c['score'] < r_['score'] - 0.01):
                    r_ = r_c; wt_ = wt_c; e_ = dict(e_, rows=cand['rows'], total=cand['total'], header=cand['header'], page=cand['page'])
                    notes.append('the %s reproduces the formula; the one first read did not carry every element of it' % (
                        'composition given in the text' if (cand.get('header') or '').startswith('composition given in the text')
                        else ('table on page %d' % cand['page']) if cand.get('page') else 'other table of the manuscript'))
                    break
        if r_ is None or r_['score'] > 0.03:
            for k_, alt in enumerate(table_alternatives(e_, wt_)):
                r_alt = _best(alt)
                if r_alt is not None and (r_ is None or r_alt['score'] < r_['score'] - 0.01):
                    r_ = r_alt; wt_ = alt
                    cn = 'the wt%% that reproduce the formula are column %d of the table, not the first (mean) column — a normalised column, or another sample or mineral in the same table' % (k_ + 2)
        # the oxide the table reports vs the charge the formula (or the species) gives: FeO with an Fe3+
        # formula reduces as Fe2O3 — adopted only when it reproduces the formula better
        for old, new, why_ in oxide_alternatives(wt_, ox_paper, species):
            wt2 = _convert(wt_, old, new); r2 = _best(wt2)
            if r2 is not None and (r_ is None or r2['score'] < r_['score'] - 0.005):
                r_ = r2; wt_ = wt2; notes.append('%s in the table taken as %s for the reduction (%s)' % (old, new, why_))
        return r_, wt_, e_, notes, cn

    def _clean(r_):
        return r_ is not None and not r_['diffs']

    hints = sample_hints(f_ctx, text); hints['domains'] = legend_columns(text, ex.get('name')) + headline_domains(text, ex.get('name'))
    cands_named = headline_columns(e, ex.get('name'), hints)
    pt_ctx = _prose_before(text, f_ctx) if f_ctx else None
    if pt_ctx and len(pt_ctx['rows']) >= 4:
        cands_named.insert(0, ({r_['constituent']: r_['mean'] for r_ in pt_ctx['rows']}, 'the wt% the formula sentence itself lists'))   # '… Ni 17.09, Fe 9.76, … which corresponds to (Rh4.50…)': the paper's own word on which analysis, before any header
    r, wt_used, e_used, notes_used, column_note = _resolve(wt, e)
    first_named = None
    if _clean(r):
        cands_named = []                                          # the mean column reproduces the formula: it is the source (owner) — no other column is tried
    for named, why in cands_named:
        if named == wt:
            continue
        r_n, wt_n, e_n, notes_n, cn_n = _resolve(named, e)
        if r_n is None:
            continue
        if first_named is None:
            first_named = (r_n, wt_n, e_n, ['the wt%% column used is the one headed by %s' % why] + notes_n, cn_n, why)
        if _clean(r_n):
            first_named = (r_n, wt_n, e_n, ['the wt%% column used is the one headed by %s' % why] + notes_n, cn_n, why); break
    if first_named is not None:
        r_n, wt_n, e_n, notes_n, cn_n, why = first_named
        if _clean(r_n) or not _clean(r):
            r, wt_used, e_used, notes_used, column_note = r_n, wt_n, e_n, notes_n, cn_n
        elif _clean(r):
            notes_used = ['the column headed by %s does not reproduce the formula; the mean column does and is used' % why] + notes_used
    wt = wt_used; e = e_used; converted += notes_used
    if r is None:
        return {'ok': False, 'verified': False, 'lines': ['composition: not verifiable — the wt% table read (%d constituents) and the formula read could not be reconciled on any basis; check the table and the formula sentence by eye' % len(wt)],
                'formula': ftxt, 'basis': None, 'result': None, 'doubts': ['not reconcilable on any basis'], 'wt': wt, 'counts': counts}
    # confidence: only a clean read with a specific deviation is a finding; anything doubtful is a note
    doubts = []; notes = []
    an = [x for x in f_issues if x.startswith('anion group sum')]
    if an:
        notes.append('the anion group of the formula does not add to its Σ (%s) — the cations were checked regardless' % an[0][17:]); f_issues = [x for x in f_issues if x not in an]
    if f_issues == ['unbalanced bracket in the formula'] and r['score'] <= 0.02:
        notes.append('one bracket of the formula is lost in the pdf text; the coefficients read agree'); f_issues = []
    elif len(f_issues) == 1 and f_issues[0].startswith('formula group sums do not add up') and f_issues[0].count('Σ') == 1 and r['score'] <= 0.02 and not r['diffs']:
        notes.append('one Σ of the formula does not add up (%s) — a misprint; every coefficient agrees' % f_issues[0][33:]); f_issues = []
    elif f_issues:
        doubts.append('the formula did not parse cleanly (%s)' % '; '.join(f_issues))
    if len(wt) < 3:
        doubts.append('only %d constituents were read' % len(wt))
    tot = e.get('total') if (e.get('total') or 0) >= 50 else sum(wt.values())      # a Total of 8.00 is the apfu block's
    if not (85.0 <= tot <= 112.0):                       # analyses do run 95–105; far outside that the table was misread
        doubts.append('the wt%% read add to %.1f' % tot)
    if r['score'] > 0.06 and len([d for d in r['diffs'] if d[2] is not None and d[1] >= 0.1]) >= 2:
        doubts.append('the cations deviate %.0f%% overall — a basis or table-reading problem rather than one slip' % (100 * r['score']))
    def _factor(ratio):
        return any(abs(ratio - f_) <= 0.08 * f_ for f_ in (0.1, 0.2, 0.25, 0.333, 0.5, 2.0, 3.0, 4.0, 5.0, 10.0)) or not 0.4 <= ratio <= 2.5
    factor_like = [d for d in r['diffs'] if d[0] != 'H' and d[2] is not None and ((d[1] >= 0.1 and not 0.6 <= d[2] / d[1] <= 1.6 and _factor(d[2] / d[1])) or (d[1] < 0.1 and d[2] > 5 * max(d[1], 0.02)))]
    if factor_like:
        doubts.append('%s is off by a factor (%.2f vs %.2f) — a multiplier or notation problem in the read, not a coefficient slip' % (factor_like[0][0], factor_like[0][1], factor_like[0][2]))
    if column_note and r['diffs']:
        doubts.append('the wt%% column was chosen by fit, not by its header — a residual disagreement there is not evidence')
    if any('analyses under it averaged' in c_ for c_ in converted) and r['diffs']:
        doubts.append('the analyses under the name were averaged by the tool — the paper\'s own mean may treat the iron split or the H2O differently; a residual there is not evidence')
    if f_kind == 'structural' and r['diffs']:
        doubts.append('the formula read is the structural (site-population) one, the paper stating no empirical formula — its coefficients come from the refinement, not the analysis')
    if any('reproduces the formula' in c_ for c_ in converted) and r['diffs']:
        doubts.append('the table was chosen over the first read because it carries every element of the formula — a residual disagreement there is not evidence')
    hdr = e.get('header') or ''
    n_means = max(len(re.findall(r'\b(mean|average|aver\.?|avg\.?)\b', hdr, re.I)), len(re.findall(r'\bsample\b|#\d|\bn\s*=\s*\d', hdr, re.I)))
    if n_means >= 2 and r['score'] > 0.02 and r['diffs']:
        doubts.append('the table holds %d samples and the first fits the formula only to %.0f%% — the formula may belong to another' % (n_means, 100 * r['score']))
    # the doubt this raises is "we may have grabbed a column of numbers that is not the analysis".
    # A transposed table settles it by construction — its header IS the constituent row — and so
    # does a caption that names the table as the chemical one; neither carries the words below.
    hdr_unknown = not re.search(r'wt\.?\s*%|mean|average|range|s\.?d\.?|standard|oxide|constituent|element|component|composition|analys', hdr, re.I)
    if r['diffs'] and hdr_unknown and not e.get('prose') and not e.get('transposed') and not _CAP_YES.search(e.get('caption') or ''):
        doubts.append('the table\'s header was not recognised — the values used may be one analysis rather than the mean')
    # H comes from hydrate and hydroxyl notation the text extraction mangles, and is usually the
    # authors' own calculation: informational, never a flag on its own
    h_only = [d for d in r['diffs'] if d[0] == 'H']
    r['diffs'] = [d for d in r['diffs'] if d[0] != 'H']
    # a trace (< 0.1 apfu) is a finding only when it is off by half or more: 0.08 vs 0.05 is rounding
    r['diffs'] = [d for d in r['diffs'] if not (d[1] < 0.1 and d[2] is not None and abs(d[2] - d[1]) < 0.5 * d[1])]
    missing = [el for el, v, t, note in r['diffs'] if t is None]
    if missing:
        doubts.append('the table read has no %s although the formula carries it — the table was probably read incompletely' % ', '.join(missing))
    extra = sorted(set(EP.parse_constituent(c).element for c, v in wt.items() if _parses(c) and v >= 1.0 and EP.parse_constituent(c).kind != 'water') - set(counts) - {'O', 'H'})
    if extra:
        doubts.append('the table has %s at 1 wt%% or more that the formula read does not carry — a simplified or another formula sentence was read' % ', '.join(extra))
    # A basis the paper does not state is one the tool chose, and a cation or element sum that is
    # not a round number was not chosen — it was read off the formula's own printed Σ (Ni+Co+Fe =
    # 18.03, Pd+Cu = 2.02). Normalising a formula to the sum of its own cations reproduces those
    # cations by construction, so a deviation against it is not evidence of anything. No paper
    # normalises to 18.03; a stated basis, or an inferred round one, still can be a finding.
    if r['diffs'] and not ex.get('basis') and r['basis'] and r['basis'][0] in ('cations', 'element'):
        n_ = r['basis'][-1]
        if abs(n_ - round(n_)) > 0.02:
            doubts.append('the paper states no basis and the one that reproduces its formula, %s, is the formula\'s own '
                          'printed sum — a deviation against it is circular' % EP._basis_label(r['basis']))
    verified = not doubts
    lines = []
    b = r['basis']
    lines.append('composition: %d constituents re-reduced on %s%s → rms deviation of the cations %.1f%%' % (
        len(wt), EP._basis_label(b), '' if ex.get('basis') == b else (' (the paper states %s; that does not reproduce its formula)' % EP._basis_label(ex['basis']) if ex.get('basis') else ' (basis inferred: the paper does not state one)'), 100 * r['score']))
    if r.get('factor'):
        lines.append('the published coefficients are the replicated ones ÷ %.3f — a different basis than the one read' % r['factor'])
    if excluded:
        notes.append('%s left out of the reduction, as the paper states its formula was calculated without it' % ', '.join(excluded))
    for c_ in converted + notes:
        lines.append('  ' + c_)
    if column_note:
        lines.append('  ' + column_note)
    lines += species_lines(species, set(EP.parse_constituent(c).element for c in wt if _parses(c)), set(counts), ox_paper)
    for el, v, t, note in r['diffs']:
        lines.append('  %s: paper %.3f, from the paper\'s own wt%% %s (%s)%s' % (el, v, ('%.3f' % t) if t is not None else '—', note, '' if verified else ' [unverified]'))
    if r['unanalysed']:
        lines.append('  calculated by the authors, not analysed: %s' % ', '.join(r['unanalysed']))
    for el, v, t, note in h_only:
        lines.append('  H (informational): paper %.3f, from the paper\'s own wt%% %s — hydrate/hydroxyl notation is read approximately' % (v, ('%.3f' % t) if t is not None else '—'))
    for d in doubts:
        lines.append('  not verifiable: ' + d)
    ok = not r['diffs'] and not f_issues
    if ok:
        lines.append('  every coefficient of the published formula follows from the published wt%')
    basis_flag = False
    if ok and verified and ex.get('basis') and b and not _same_basis(b, ex['basis']) and not r.get('factor') \
            and _basis_flaggable(ex['basis'], b) and _basis_is_the_formulas(ex.get('basis_sentence') or '', f_ctx):
        # the paper says 'on the basis of 12 O' and its coefficients follow from 13 anions, or from 8 cations:
        # a reviewer confirms which is the slip (owner: flag it when they differ). Only with a clean reduction
        # on the other basis, and the basis sentence being the formula's own
        basis_flag = True
        lines.append('  the paper states its formula is calculated on %s, but every coefficient follows from %s — the stated basis does not reproduce the formula; one of the two is a slip'
                     % (EP._basis_label(ex['basis']), EP._basis_label(b)))
    return {'ok': ok, 'verified': verified, 'lines': lines, 'formula': ftxt, 'basis': b, 'result': r, 'doubts': doubts, 'wt': wt, 'counts': counts, 'basis_flag': basis_flag}

def _basis_flaggable(stated, found):
    """Whether the basis that reproduces the formula is one a paper states — a whole-number anion or
    cation count, a single element ('P = 3'), or the stated group with another number ('S+Se+Te = 36'
    against a stated 18) — and not the formula's own printed group sum ('Na+Ca+K+Zn = 0.97', which
    reproduces its cations by construction; on the corpus those were 30 of 43 first flags)."""
    whole = lambda n: abs(float(n) - round(float(n))) <= 0.02
    if found[0] == 'element':
        els = set(re.findall(r'[A-Z][a-z]?', found[1])); n = found[2]
        if len(els) == 1 and whole(n):
            return True
        return stated[0] == 'element' and set(re.findall(r'[A-Z][a-z]?', stated[1])) == els and whole(n) and whole(stated[2])
    return whole(found[-1])

def _basis_is_the_formulas(basis_sentence, f_ctx):
    """Whether the basis sentence read is the formula sentence's own (or names the formula) rather
    than a basis quoted elsewhere — for another phase, a cited formula, or a site-population step."""
    if not basis_sentence or not f_ctx:
        return False
    if re.search(r'\b(?:empirical|chemical)?\s*formula', basis_sentence, re.I) is None:
        return False
    key = re.sub(r'\s+', ' ', basis_sentence[:60]).strip()
    return key[:40] in re.sub(r'\s+', ' ', f_ctx) or re.sub(r'\s+', ' ', f_ctx[-120:]).strip()[:40] in re.sub(r'\s+', ' ', basis_sentence)

def _without_elements(e, excluded):
    """The table candidate without the constituents of the given elements ('calculated without Al'),
    in its rows, its positional rows and its candidate tables alike."""
    def keep(r):
        return not (_parses(r['constituent']) and EP.parse_constituent(r['constituent']).element in excluded)
    out = dict(e); out['rows'] = [r for r in e['rows'] if keep(r)]
    if e.get('rows_all'):
        out['rows_all'] = [r for r in e['rows_all'] if keep(r)]
    if e.get('alts'):
        out['alts'] = [{c: v for c, v in a.items() if keep({'constituent': c})} for a in e['alts']]
    if e.get('candidates'):
        out['candidates'] = [_without_elements(c_, excluded) for c_ in e['candidates']]
    return out

def table_alternatives(e, wt):
    """The other numeric columns of a row-layout table, or the other Mean / analysis rows of a
    transposed one, as wt% dicts — a table of several minerals or samples: [wt dict, …]."""
    if e.get('alts'):
        return [a for a in e['alts'] if len(a) >= 3 and a != wt]
    out = []; base_sum = sum(v for v in wt.values() if v is not None) or 100.0
    rows_pos = _drop_totals(e['rows_all']) if e.get('rows_all') else e['rows']      # the n.d.-first rows too: their other columns hold values
    for k in range(1, 8):
        alt = {r['constituent']: r['all'][k] for r in rows_pos if len(r.get('all') or []) > k and r['all'][k] > 0}
        if len(alt) >= 3 and alt != wt and alt not in out and sum(alt.values()) >= 0.5 * base_sum:
            out.append(alt)                                    # not an s.d. or apfu column
    seen = {}
    for r in e['rows']:                                        # a constituent listed twice ('Fe 2.0' from prose, then the row): the other value
        if r['constituent'] in seen and r['mean'] != seen[r['constituent']]:
            alt = dict(wt); alt[r['constituent']] = r['mean']
            if alt not in out:
                out.append(alt)
        seen.setdefault(r['constituent'], r['mean'])
    return out

_NOT_SAMPLES = {'holotype', 'the', 'this', 'table', 'museum', 'natural', 'history', 'catalogue', 'university', 'mineral', 'specimen', 'deposited',
                'collection', 'figure', 'photo', 'field', 'department', 'published', 'keywords', 'raman', 'nomenclature', 'classification',
                'international', 'abstract', 'introduction', 'results', 'discussion', 'references', 'acknowledgements', 'journal', 'american',
                'mineralogist', 'canadian', 'european', 'occurrence', 'appearance', 'physical', 'optical', 'properties', 'chemical', 'crystal',
                'structure', 'powder', 'diffraction', 'sample', 'samples', 'material', 'cotype', 'commission', 'association', 'minerals'}

def holotype_words(text):
    """Capitalised words within 120 characters of 'holotype' / 'type material' anywhere in the paper —
    a locality or sample name that may head the table's column."""
    out = []
    for m in re.finditer(r'holotype|type (?:specimen|material|sample)', text or '', re.I):
        seg = text[max(0, m.start() - 120):m.end() + 120]
        at = m.start() - max(0, m.start() - 120); pos = 0
        for part in re.split(r'(?i)cotype|paratype|neotype', seg):        # 'the holotype in Madagascar, and a cotype from Myanmar': the cotype's words are not the holotype's
            if pos <= at < pos + len(part) + 7:
                seg = part; break
            pos += len(part) + 7
        out += [w for w in re.findall(r'\b([A-Z][a-zà-ÿ]{4,}|[A-Z]{2,6}[-–]?\d[\w\-–/]*)\b', seg) if w.lower() not in _NOT_SAMPLES]
    return list(dict.fromkeys(out))

def headline_domains(text, name):
    """Domain / crystal / grain / sample letters the paper assigns to the headline mineral: each
    'domain(s) B and C' / 'domains A–C' phrase belongs to the nearest mineral name in its sentence,
    so 'ferriandrosite-(Ce) (domains A–C) and associated vielleaureite-(Ce) (domain D)' gives A, B, C.
    Letters named in more sentences come first ('Domains B and C correspond to … ferriandrosite')."""
    stem = re.sub(r'-\(.*\)$', '', name or '').lower()
    if len(stem) < 5 or not text:
        return []
    stem = stem[:max(5, len(stem) - 2)]
    t = re.sub(r'\b(i\.e|e\.g|et al|Fig|Figs|Tab|cf|ca|vs|no)\.', r'\1', text)
    word = r'(?:domain|crystal|grain|sample|zone|area|spot|specimen|analys[ie]s|point)s?'
    letters = r'((?:[IVX]{1,4}(?![A-Z])|[A-Z]\d{0,2})(?:\s*(?:,|and|&|–|-|to)\s*(?:[IVX]{1,4}(?![A-Z])|[A-Z]\d{0,2}))*)'
    score = {}
    for sent in re.split(r'(?<=[.;])\s+', t):
        low = sent.lower()
        if stem not in low:
            continue
        names = [(m.start(), m.group(0)) for m in re.finditer(r'\b[a-zà-ÿč]{4,}ite(?:-\([a-z]+\))?', low) if m.group(0).split('-')[0] not in _NOT_MINERALS]
        if not names:
            continue
        for m in re.finditer(word + r'\s+' + letters + r'(?![a-z])', sent, re.I):
            run = m.group(1)
            if re.match(r'^(?:domain|crystal|grain|sample|zone|area|spot|specimen|analys|point)', run, re.I):
                continue
            nearest = min(names, key=lambda nm: abs(nm[0] - m.start()))[1]
            if stem not in nearest:
                continue                                              # 'domain D' belongs to vielleaureite
            codes = re.findall(r'[IVX]{1,4}(?![A-Z])|[A-Z]\d{0,2}', run)
            rng = re.match(r'^([A-Z])\s*(?:–|-|to)\s*([A-Z])$', run.strip())
            if rng and ord(rng.group(2)) > ord(rng.group(1)):
                codes = [chr(c_) for c_ in range(ord(rng.group(1)), ord(rng.group(2)) + 1)]   # 'A–C' = A, B, C
            for c_ in codes:
                score[c_] = score.get(c_, 0) + 1
    return sorted(score, key=lambda c_: (-score[c_], c_))[:6]

def excluded_elements(text):
    """Elements the paper says it left out of the reduction: 'the empirical formula was calculated
    without Al', 'excluding Si and Al', 'Al was excluded' -> ['Al']."""
    out = []
    for m in re.finditer(r'(?:formula|reduction|calculation|normali[sz]ation)[^.]{0,80}?(?:calculated|computed|normali[sz]ed|derived)?[^.]{0,40}?\b(?:without|excluding|omitting|ignoring|disregarding)\s+((?:[A-Z][a-z]?\d?O?\d?(?:\s*,\s*|\s+and\s+|\s*/\s*)?){1,4})', text):
        out += re.findall(r'[A-Z][a-z]?', m.group(1))
    for m in re.finditer(r'\b([A-Z][a-z]?)(?:2O\d|O\d?)?\s+(?:was|were|is|are)\s+(?:therefore\s+|thus\s+)?(?:excluded|omitted|not included|left out|disregarded|ignored)', text):
        out.append(m.group(1))
    return [e_ for e_ in dict.fromkeys(out) if e_ in EP.ATOMIC_WEIGHTS and e_ not in ('O', 'H')]

def legend_columns(text, name):
    """A table legend '1, 2, 8 – fluorpyromorphite (1 – holotype, mean of 8 spot analyses; 2 – F-richest
    spot; 8 – cotype)' -> the column numbers of the headline mineral, the holotype's first: ['1', '2', '8']."""
    stem = re.sub(r'-\(.*\)$', '', name or '').lower()
    if len(stem) < 5 or not text:
        return []
    stem = stem[:max(5, len(stem) - 2)]
    out = []
    low = text.lower()
    def _under_figure(pos):
        prev = re.findall(r'\b(table|tab\.|таблица|fig\.?|figure|рис\.?)\s*\d+[a-z]?[.:]', low[max(0, pos - 700):pos])
        return bool(prev) and prev[-1].startswith(('fig', 'рис'))              # the nearest caption start is 'Fig. 2.': a figure legend ('Figure 4B)' mid-sentence is not one)
    for m in re.finditer(r'(?<![\d.])(\d{1,2}(?:\s*(?:,|and|–|-)\s*\d{1,2})*)\s*[–-]\s*([a-zà-ÿč]{4,}ite(?:-\([a-z]+\))?)', low):
        if stem in m.group(2) and not _under_figure(m.start()):
            nums = re.findall(r'\d{1,2}', m.group(1))
            rng = re.match(r'^(\d{1,2})\s*[–-]\s*(\d{1,2})$', m.group(1).strip())
            if rng and int(rng.group(2)) > int(rng.group(1)):
                nums = [str(k_) for k_ in range(int(rng.group(1)), int(rng.group(2)) + 1)]
            out += [n_ for n_ in nums if n_ not in out]
    for m in re.finditer(r'([a-zà-ÿč]{4,}ite(?:-\([a-z]+\))?)\s*\((\d{1,2})\s*[–—-]\s*[a-z]', low):      # 'zoharite (3—aggregate, Figure 3C)'
        if stem in m.group(1) and m.group(2) not in out and not _under_figure(m.start()):
            out.append(m.group(2))
    holo = re.findall(r'\(?(\d{1,2})\s*[–-]\s*holotype', low)
    return [h_ for h_ in holo if h_ in out or not out] + [n_ for n_ in out if n_ not in holo]

def sample_hints(ctx, text=''):
    """What the formula sentence says about which analyses it rests on: sample / specimen codes,
    'holotype', 'mean of 12 analyses' -> {'codes': [...], 'n': 12 or None, 'holotype': bool}."""
    codes = re.findall(r'(?:sample|specimen|crystal|grain|fragment|analysis|analyses)\s*(?:no\.?|number|#)?\s*([A-Za-z]{0,6}[-–]?\d[\w\-–/.]*)', ctx, re.I)
    codes += re.findall(r'\(([A-Z]{1,6}[-–]?\d[\w\-–/]*)\)', ctx)
    codes += re.findall(r'(?<=[a-z]) \(([A-Z]{2,4})\)', re.split(r'\.\s+(?=[A-Z])', ctx)[-1])   # 'beraunite (NM)': a sample tag, from the formula's own sentence only
    codes += re.findall(r'(?:domain|crystal|grain|sample|zone|area|spot|type)\s+([A-Z]|[IVX]{1,4}|\d{1,2})\b', ctx, re.I)   # 'domain A', 'crystal II'
    codes += re.findall(r'(?<![\w-])([A-Z]{1,4}[-–][A-Z]*\d\w*|[A-Z]{2,6}\d{3,}[\w-]*)(?![\w-])', ctx)                  # 'A-WP1', 'NRM19331765'
    codes += [w for w in re.findall(r'\b([A-Z][a-zà-ÿ]{4,})\b', ctx) if w.lower() not in ('the', 'this', 'table', 'based', 'chemical', 'empirical', 'formula', 'analyses', 'electron', 'microprobe')]   # a locality or sample name
    n = None
    m = re.search(r'(?:mean|average)\s+of\s+(\d+)|(\d+)\s+(?:spot\s+)?analys[ei]s|n\s*=\s*(\d+)', ctx, re.I)
    if m:
        n = int(next(g for g in m.groups() if g))
    return {'codes': [c_.strip('.,;') for c_ in codes if len(c_) >= 1], 'n': n, 'holotype': bool(re.search(r'holotype|type (?:specimen|material)', ctx, re.I)),
            'holotype_words': holotype_words(text), 'domains': []}

def headline_column(e, name, hints):
    cols = headline_columns(e, name, hints)
    return cols[0] if cols else (None, None)

def headline_columns(e, name, hints):
    """The wt% columns the paper may have computed the formula from, by evidence rather than by fit,
    best evidence first: the column headed by a domain the text assigns to the headline mineral, by
    its name or Levinson suffix, by a sample code the formula sentence cites, by 'holotype', or by
    the number of analyses it states. -> [(wt dict, what named it), …]."""
    cells = [(t.strip('()[],;:'), x) for t, x in (e.get('head_cells') or [])]
    found = []
    if not cells:
        return found
    stem = re.sub(r'-\(.*\)$', '', name or '').lower()
    keys = []
    if len(stem) >= 5:
        keys.append((stem[:max(5, len(stem) - 2)], 'the headline mineral, %s' % name))
    suf = re.search(r'-\((\w+)\)$', name or '')
    if suf:
        keys.append((suf.group(1).lower(), 'the Levinson suffix of the headline mineral, (%s)' % suf.group(1).capitalize()))   # '(Nd)1 (Y)2 (Ce)3' columns
    pos = 0
    for c_ in hints.get('domains', []):
        entry = (c_.lower(), ('column %s, which the table legend assigns to %s' % (c_, name)) if c_.isdigit() else ('the domain the text assigns to %s, %s' % (name, c_)))
        if c_.isdigit():
            keys.insert(pos, entry); pos += 1                                  # legend numbers are explicit: before the name-headed span, in the legend's order (holotype first)
        else:
            keys.append(entry)
    for c_ in hints.get('codes', []):
        keys.append((c_.lower(), 'the sample the formula cites, %s' % c_))
    if hints.get('holotype'):
        keys.append(('holotype', 'the holotype'))
    for w in hints.get('holotype_words', [])[:12]:
        keys.append((w.lower(), 'the holotype (%s)' % w))
    if hints.get('n'):
        keys.append(('n = %d' % hints['n'], '%d analyses, as the formula sentence says' % hints['n']))
        keys.append(('n=%d' % hints['n'], '%d analyses' % hints['n'])); keys.append(('(%d)' % hints['n'], '%d analyses' % hints['n']))
    joined = ' '.join(t for t, x in cells).lower()
    n_digit_cells = sum(1 for t, x in cells if re.fullmatch(r'\d{1,2}', t))
    for key, why in keys:
        if key not in joined:
            continue
        if key.isdigit() and n_digit_cells < 2:
            continue                                                           # a legend number names a column only of a numbered header
        # the x of the cell (or run of cells) carrying the key; a short code must be a whole cell
        x_hit = None
        for k_ in range(len(cells)):
            run = ' '.join(t for t, x in cells[k_:k_ + 4]).lower()
            if (len(key) <= 3 and (re.sub(r'[^a-z]', '', cells[k_][0].lower()) == key or (key.isdigit() and cells[k_][0].strip('()[].,') == key))) or (len(key) > 3 and (key in cells[k_][0].lower() or run.startswith(key))):
                x_hit = cells[k_][1]; break
        if x_hit is None:
            continue
        # the value column of the group the key heads: the last Mean / Ave / wt% token at or before the key's span
        means = sorted(x for t, x in cells if re.match(r'^(mean|aver(?:age)?\.?|avg\.?|av\.?|wt\.?%?|\(wt%\))$', t, re.I) and x > (e.get('label_x') or 0) + 30)
        under = [x for x in means if x <= x_hit + 25]
        x_col = under[-1] if under and x_hit - under[-1] <= 110 else x_hit
        span = None
        if key.isdigit():
            under = [x_hit]; x_col = x_hit                                     # '3' from the legend: that one column
        if not under:                                                      # no Mean column: the analyses under the name's span are averaged
            labelled = [x for t, x in cells if not re.fullmatch(r'[\d.,()–\-]+|n|=|mean|average|range|s\.?d\.?|σ|s\.u\.|wt\.?%?|e\.?s\.?d\.?', t, re.I)]
            names_x = sorted(x for x in labelled if x > x_hit + 40); prev_x = [x for x in labelled if x < x_hit - 40]
            span = ((max(prev_x) + 25) if prev_x else x_hit - 70, (names_x[0] - 25) if names_x else 1e9)   # between the neighbouring labelled columns
            inside = [x for x in means if span[0] <= x < span[1]]
            if inside:
                x_col = inside[0]; span = None                             # 'Madagascar: 5 6 13 Mean (n = 3)': the group's own Mean cell, not the spots averaged
            elif any(re.fullmatch(r'range|s\.?d\.?|σ|e\.?s\.?d\.?|\(n|n', t, re.I) and x_hit < x < span[1] for t, x in cells):
                x_col = x_hit; span = None                                 # 'CF9a1 | Range (n = 16) | SD': one value column with its statistics beside it
        if e.get('transposed'):
            for label, vals in e.get('labelled_rows') or []:
                if key in (label or '').lower() and len(vals) >= 3 and not any(vals == w_ for w_, _ in found):
                    found.append((vals, why))
            continue
        wt = {}
        all_rows = e.get('rows_all') or e['rows']
        oxide_els = {m_.group(0) for m_ in (re.match(r'[A-Z][a-z]?', r_['constituent']) for r_ in all_rows if re.search(r'O\d*$', r_['constituent'])) if m_}
        for row in all_rows:                                       # a later row of the same constituent (the split after a total) overrides
            if not row.get('xs'):
                continue
            if re.fullmatch(r'[A-Z][a-z]?', row['constituent']) and row['constituent'] in oxide_els and row['constituent'] not in ('F', 'Cl', 'Br', 'I', 'S', 'Se', 'Te'):
                continue                                                   # 'Gd 0.14 0.01 0.18' under 'Gd2O3': the apfu block
            if span:
                vals_ = [row['all'][q] for q in range(min(len(row['xs']), len(row.get('all') or []))) if span[0] <= row['xs'][q] < span[1]]
                if len(vals_) >= 2:
                    wt[row['constituent']] = round(sum(vals_) / len(vals_), 3); continue
            k_best = min(range(len(row['xs'])), key=lambda q: abs(row['xs'][q] - x_col))
            if abs(row['xs'][k_best] - x_col) <= 30 and k_best < len(row.get('all') or []) and row['all'][k_best] > 0:
                wt[row['constituent']] = row['all'][k_best]                # a 0 is an 'n.d.' cell: not a value
        if len(wt) >= 3 and sum(wt.values()) >= 50 and not any(wt == w_ for w_, _ in found):
            found.append((wt, why + (' (the analyses under it averaged)' if span else '')))
    return found

def _parses(c):
    try:
        EP.parse_constituent(c); return True
    except ValueError:
        return False

def species_lines(species, table_els, formula_els, ox_paper):
    """What the species' ideal formula (Mindat) says about the table and the formula: an essential
    element neither carries; a charge the formula states that the species never has."""
    if not species:
        return []
    L = []
    essential = species['elements'] - {'H', 'O'} - set(EP.CALCULATED_ELEMENTS)
    gone = sorted(essential - table_els - formula_els)
    if gone:
        L.append("  the ideal formula (Mindat: %s) carries %s, absent from both the wt%% table and the formula" % (_plain_html(species['formula']), ', '.join(gone)))
    for el, vs in ox_paper.items():
        if el in species['ox'] and vs and not (vs & species['ox'][el]):
            L.append("  the formula writes %s; Mindat's ideal formula has %s" % ('/'.join('%s%d+' % (el, x) for x in sorted(vs)), '/'.join('%s%d+' % (el, x) for x in sorted(species['ox'][el]))))
    return L

def _plain_html(h):
    return re.sub(r'<[^>]+>', '', h or '')

def _species_mass(rec):
    """The molar mass of Mindat's ideal formula, read from its HTML ('Ca<sub>2</sub>(UO<sub>2</sub>)<sub>3</sub>
    (CO<sub>3</sub>)<sub>5</sub>·8H<sub>2</sub>O'): subscripts expanded, charges dropped, groups and the hydrate
    read by the formula parser. A token scan of the raw HTML dropped every subscript (abelsonite came out at
    86 instead of 519) and that number stood in as 'an ideal formula' for the density oracle. None when it
    does not parse."""
    h = (rec or {}).get('formula') or ''
    if not h:
        return None
    plain = re.sub(r'<sub>\s*(\d+(?:\.\d+)?)\s*</sub>', r'\1', h)
    plain = re.sub(r'<sup>[^<]*</sup>', '', plain)
    plain = _plain_html(plain).replace('&middot;', '·').replace('&nbsp;', ' ')
    try:
        norm = _journal_to_icdd(plain)
        counts, _ox, issues = EP.parse_icdd_formula(norm, has_sulfur='S' in counts_guess(norm))
    except Exception:
        return None
    if not counts or any('garbled' in x for x in issues):
        return None
    return _formula_mass_counts(counts)

def _derived_composition(ex, wt, species):
    """No empirical formula could be read: reduce the table on the paper's basis (the species'
    ideal formula deciding an ambiguous oxide) and print the formula that gives, for the reviewer
    to hold against the paper's own by eye."""
    if len(wt) < 3 or not ex.get('basis'):
        return None
    conv = []
    for old, new, why in oxide_alternatives(wt, {}, species):
        wt = _convert(wt, old, new); conv.append('%s as %s (%s)' % (old, new, why))
    cons, vals = [], []
    for c, v in wt.items():
        if _parses(c):
            cons.append(EP.parse_constituent(c)); vals.append(v)
    try:
        red = EP.reduce(EP.Dataset(cons, [vals], ['mean'], {}, ex.get('name') or 'paper', None), ex['basis'])
    except Exception:
        return None
    lines = ['composition: no empirical formula sentence could be read from the paper; from its wt%% table on %s the tool derives %s [unverified]' % (EP._basis_label(ex['basis']), red.formula())]
    if conv:
        lines.append('  ' + '; '.join(conv))
    lines += species_lines(species, set(k.element for k in cons), set(), {})
    lines.append('  not verifiable: compare the derived formula with the paper\'s by eye — the formula sentence takes a form the tool does not read')
    return {'ok': False, 'verified': False, 'lines': lines, 'formula': '', 'basis': ex['basis'], 'result': None, 'doubts': ['no formula read'], 'derived': red.formula(), 'wt': wt, 'counts': {}}

# ----------------------------------------------------------------------------- what was read, and what vouches for it
# Every value the readers take from a paper is a guess about the text until something outside
# the text agrees with it. Each field carries a record: the value, where it was read, which reader
# read it, and what vouched for it — the composition re-derived from the table, the powder table
# against the cell, the bond-valence table against the .cif, Gladstone–Dale for the optics, the
# cell against its own volume and density, Mindat for the species. A field an oracle adjudicated
# is 'agrees' or 'disagrees'; one it looked at with doubts is 'unverified'; one nothing could check
# is 'nooracle'; one not read is 'none'. A misread value that fails its oracle is then a doubt to
# report, never a verdict.
FIELDS = ('name', 'formula', 'basis', 'method', 'epma', 'optics.n', 'optics.D_meas', 'optics.D_calc', 'cell', 'bv.params', 'pxrd.obs', 'pxrd.calc')
STATUSES = ('agrees', 'disagrees', 'unverified', 'nooracle', 'none')

def _record(value, source='', page=None, reader='regex'):
    empty = value in (None, '', 0) or value == [] or value == {}
    return {'value': value, 'source': (source or '')[:200], 'page': page, 'reader': reader, 'verified_by': None,
            'status': 'none' if empty else 'nooracle', 'detail': ''}

def _fields_init(ex, text):
    """A record per field from what extract() read (statuses 'none' / 'nooracle' until verify runs)."""
    e = ex.get('epma') or {}; o = ex.get('optics') or {}; m = ex.get('method') or {}; bv = ex.get('bv') or {}
    fs = _formulas(text, ex.get('name') or '')
    cells = _paper_cells(text)
    f = {'name': _record(ex.get('name'), 'the paper\'s head'),
         'formula': _record(fs[0][0] if fs else None, fs[0][5] if fs else ''),
         'basis': _record(ex.get('basis'), ex.get('basis_sentence') or ''),
         'method': _record(m.get('h2o') or m.get('charge') or ((m.get('sentences') or [None])[0]), (m.get('sentences') or [''])[0]),
         'epma': _record(len(e['rows']) if e and e.get('rows') else None, e.get('header') or '', e.get('page')),
         'optics.n': _record(o.get('n'), o.get('n_from') or ''),
         'optics.D_meas': _record(o.get('D_meas'), next((s_ for s_ in o.get('sentences') or [] if 'meas' in s_.lower() or 'float' in s_.lower() or 'pycno' in s_.lower()), '')),
         'optics.D_calc': _record(o.get('D_calc'), next((s_ for s_ in o.get('sentences') or [] if 'calc' in s_.lower()), '')),
         'cell': _record(_cell_str(cells[0][1]) if cells else None, cells[0][0] if cells else ''),
         'bv.params': _record(bv.get('params'), (bv.get('sentences') or [''])[0]),
         'pxrd.obs': _record(ex['pxrd']['obs'] or None), 'pxrd.calc': _record(ex['pxrd']['calc'] or None)}
    if ex.get('_table') and ex['_table'][2] and not str(ex.get('_path', '')).lower().endswith('.docx'):
        f['pxrd.obs']['page'] = f['pxrd.calc']['page'] = ex['_table'][2][0]
    return f

def _set(fields, name, status, by=None, detail='', value=None):
    r = fields.get(name)
    if r is None:
        return
    if r['status'] == 'none' and value is None and status != 'none':
        return                                                         # nothing was read: an oracle cannot vouch for an absence
    r['status'] = status; r['verified_by'] = by; r['detail'] = detail[:200]
    if value is not None:
        r['value'] = value

def _same_basis(a, b):
    if not a or not b or a[0] != b[0]:
        return False
    return all(abs(float(x) - float(y)) < 1e-6 if isinstance(x, (int, float)) else x == y for x, y in zip(a[1:], b[1:]))

# --- Gladstone–Dale: the optics against the composition and the density
_GD_CATS = ('superior', 'excellent', 'good', 'fair', 'poor')

def gd_statement(text):
    """The paper's own compatibility statement: {'ci': 1 − K_P/K_C as stated or None, 'category':
    Mandarino's word or None, 'sentence'}. The category word is taken only in the sentence that
    speaks of the compatibility index (every paper has a 'good' somewhere)."""
    t = text.replace('−', '-').replace('–', '-')
    out = {'ci': None, 'category': None, 'sentence': ''}
    for m in re.finditer(r'(?:compatibility index|1\s*-\s*\(?\s*K\s*_?[Pp]\s*/\s*K\s*_?[Cc]\s*\)?)((?:[^.]|\.(?=\d)){0,120})', t):   # the sentence, a decimal point allowed
        seg = m.group(1)
        num = re.search(r'(-?\s?0?\.\d{2,4})', seg)
        if num and out['ci'] is None:
            out['ci'] = float(num.group(1).replace(' ', '')); out['sentence'] = t[max(0, m.start() - 20):m.end()].strip()
        cat = re.search(r'\b(superior|excellent|good|fair|poor)\b', seg, re.I)
        if cat and out['category'] is None:
            out['category'] = cat.group(1).lower(); out['sentence'] = out['sentence'] or t[max(0, m.start() - 20):m.end()].strip()
        if out['ci'] is not None and out['category']:
            break
    if out['category'] is None:
        m = re.search(r'\b(superior|excellent|good|fair|poor)\b[^.]{0,60}(?:compatib|Mandarino|Gladstone)', t, re.I)
        if m:
            out['category'] = m.group(1).lower(); out['sentence'] = out['sentence'] or t[max(0, m.start() - 20):m.end()].strip()
    return out

def gd_check(ex, comp, stmt=None):
    """The mean refractive index and the densities against the composition: K_C from the wt% the
    composition oracle used, K_P = (n − 1)/D, 1 − K_P/K_C for each density the paper gives. Against
    the paper's own stated index or category when it gives one, else against Mandarino's 'poor'
    boundary — a misread n or D moves the index by a tenth. -> {'status': {field: status},
    'ci': {'meas': x, 'calc': y}, 'KC', 'lines', 'red': bool}."""
    from pxrd_review import gd as GD
    o = ex.get('optics') or {}; n = o.get('n'); stmt = stmt or {}
    st = {'optics.n': 'nooracle', 'optics.D_meas': 'nooracle', 'optics.D_calc': 'nooracle'}
    out = {'status': st, 'ci': {}, 'KC': None, 'lines': [], 'red': False}
    wt = (comp or {}).get('wt') or {}
    if not n or not wt:
        out['lines'] = [] if not n else ['Gladstone–Dale: n %.4f read but no wt%% table to check it against' % n]
        return out
    try:
        KC, rows = GD.kc({c: v for c, v in wt.items() if v})
    except Exception as ex_:
        out['lines'] = ['Gladstone–Dale: could not compute K_C (%s)' % ex_]
        return out
    if not KC:
        return out
    out['KC'] = KC
    verified = bool((comp or {}).get('verified'))
    parts = []; best = None
    for key in ('meas', 'calc'):
        D = o.get('D_' + key)
        if not D:
            continue
        ci = 1 - ((n - 1) / D) / KC; out['ci'][key] = ci
        parts.append('D_%s %.3f → %+.3f (%s)' % (key, D, ci, GD.category(ci)))
        if best is None or abs(ci) < abs(best[1]):
            best = (key, ci)
    if best is None:
        out['lines'] = ['Gladstone–Dale: n %.4f and K_C %.4f from the wt%% table, but no density to check them against' % (n, KC)]
        return out
    head = 'Gladstone–Dale: n %.4f (%s), K_C %.4f from the wt%% table — 1 − K_P/K_C with %s' % (n, o.get('n_from') or 'n', KC, '; '.join(parts))
    lines = [head]
    def status_all(s, detail):
        st['optics.n'] = s
        for key in ('meas', 'calc'):
            if o.get('D_' + key):
                st['optics.D_' + key] = s
        out['detail'] = detail
    if stmt.get('ci') is not None:
        theirs = stmt['ci']; key, ci = min(out['ci'].items(), key=lambda kv: abs(kv[1] - theirs))
        gap = abs(ci - theirs)
        # never red: K_C depends on which constituents the table lists and which constants the file
        # has, and the paper's own index often used another density or wt% column — on the corpus a
        # strict comparison flagged 50 papers, most of them with K_C, not the paper, at fault
        if gap <= 0.03:
            status_all('agrees', 'the paper states %+.3f, this gives %+.3f' % (theirs, ci))
            lines.append('the paper states %+.3f (%s): reproduced with D_%s' % (theirs, GD.category(theirs), key))
        else:
            status_all('unverified', 'the paper states %+.3f, this gives %+.3f' % (theirs, ci))
            lines.append('the paper states %+.3f, this gives %+.3f with D_%s — the wt%% column, a constant, or the n or D read may differ from the paper\'s [unverified]' % (theirs, ci, key))
    elif stmt.get('category'):
        theirs = stmt['category']; key, ci = best; ours = GD.category(ci)
        dist = abs(_GD_CATS.index(ours) - _GD_CATS.index(theirs))
        if dist <= 1:
            status_all('agrees', 'the paper says %s, this gives %s' % (theirs, ours))
            lines.append('the paper calls it %s: this gives %s with D_%s' % (theirs, ours, key))
        else:
            status_all('unverified', 'the paper says %s, this gives %s' % (theirs, ours))
            lines.append('the paper calls it %s, this gives %s with D_%s — n, D or the wt%% column may be misread [unverified]' % (theirs, ours, key))
    else:
        key, ci = best
        if abs(ci) < 0.08:
            status_all('agrees', 'consistent: %+.3f (%s)' % (ci, GD.category(ci)))
        else:
            status_all('unverified', 'poor compatibility: %+.3f' % ci)
            lines.append('the optics, density and wt%% read are not compatible (%+.3f, %s) — a misread value, or a genuinely poor index [unverified]' % (ci, GD.category(ci)))
    out['lines'] = lines
    return out

# --- the cell against its own volume, the density, and the .cif
def _volume(cell):
    ca, cb, cg = (math.cos(math.radians(cell[k])) for k in ('α', 'β', 'γ'))
    fac = 1 - ca * ca - cb * cb - cg * cg + 2 * ca * cb * cg
    return cell['a'] * cell['b'] * cell['c'] * math.sqrt(fac) if fac > 0 else None

def _formula_mass_counts(counts):
    M = 0.0
    for el, n in (counts or {}).items():
        w = EP.ATOMIC_WEIGHTS.get(el)
        if w is None:
            return None
        M += w * float(n)
    return M or None

def cell_consistency(text, counts=None, cif=None, D_calc=None, verified_formula=False, masses=None):
    """The cells the paper states, each against itself (V from a b c α β γ vs the V it prints), against the formula (D = 1.66054·Z·M/V
    vs the D_calc it prints — M from the empirical counts, and from every other formula in `masses`
    [(label, M)]: papers compute D_calc from the ideal formula as often as from the empirical one)
    and against the .cif (sorted axes within 3 %, the powder-vs-single-crystal tolerance).
    -> {'status', 'cell_status', 'cell', 'lines', 'detail', 'red', 'D'}. A D_calc off from every
    formula's D is a doubt, never a finding ('red' stays False): on the corpus every red line drawn
    for it dissolved on inspection — an axis the reader had misread, a Z borrowed from another
    statement, a hydrate count the text layer had lost, a garbage ideal-formula mass, and last an
    empirical formula deficient in cations while the paper's D_calc rests on a fuller one (owner,
    2026-09-06). Beyond 15 % the cell
    and the formula are not each other's (another phase's cell, a formula per two units) and it is a
    doubt; so is a volume that does not follow (a misread axis or a misprint). 'cell_status' is the
    cell's own standing (volume, .cif) apart from the density; 'D' what the density check found."""
    from pxrd_review import cell_lambda_check as CL
    out = {'status': 'nooracle', 'cell': None, 'lines': [], 'detail': '', 'red': False, 'D': None, 'cell_status': 'nooracle'}
    cands = CL.find_cells(text)
    if not cands:
        return dict(out, status='none')
    cif_cell = None
    if cif:
        from pxrd_review import extra_checks as X
        cif_cell = _cell_floats((X.parse_cif(cif) or {}).get('cell') or {})
    Ms = [('the formula', _formula_mass_counts(counts))] + [(lab, m) for lab, m in (masses or []) if m]
    Ms = [(lab, m) for lab, m in Ms if m]
    # Z belongs to the structure, not to one cell statement: a powder cell quoted without it takes the
    # single Z the paper states elsewhere — beside its single-crystal cell, or in the crystal-data table
    zs = sorted({int(CL.num_val(c_.Z)) for c_ in cands if c_.Z and CL.num_val(c_.Z)})
    if not zs:
        zs = sorted({int(z_) for z_ in re.findall(r'(?<![A-Za-z])Z\s*=\s*(\d{1,2})(?![\d.])', text)})
    if not zs and cif:
        from pxrd_review import extra_checks as X
        z_cif = (X.parse_cif(cif) or {}).get('Z')                      # the .cif's own Z, when the paper prints none the reader finds
        zs = [int(z_cif)] if z_cif and str(z_cif).isdigit() else []
    Z_paper = zs[0] if len(zs) == 1 and 1 <= zs[0] <= 64 else None
    hydrous = bool((counts or {}).get('H'))
    best = None
    for cc in sorted(cands, key=lambda c: 0 if c.context == 'powder' else 1):
        variants = [{'a': cc.a, 'b': cc.b, 'c': cc.c, 'α': cc.al, 'β': cc.be, 'γ': cc.ga}]
        if cc.b is None and cc.c is not None and cc.ga is None:
            variants = [dict(variants[0], b=cc.a, γ='90'), dict(variants[0], b=cc.a, γ='120')]
        V_stated = CL.num_val(cc.V) if cc.V else None; Z = (CL.num_val(cc.Z) if cc.Z else None) or Z_paper
        z_borrowed = not cc.Z                                          # a borrowed Z supports a verdict of agreement or a doubt, never a red line
        for v in variants:
            cell = _cell_floats(v)
            if not cell:
                continue
            V = _volume(cell)
            if not V:
                continue
            checks = []; status = 'nooracle'; detail = []
            if V_stated:
                dv = abs(V - V_stated) / V_stated
                checks.append(('V', dv <= 0.003)); detail.append('V from the axes %.1f vs %.1f printed' % (V, V_stated))
            D = None; Dlab = ''
            if Z and Ms and D_calc:
                Dlab, D = min(((lab, 1.66054 * Z * m / V) for lab, m in Ms), key=lambda x: abs(x[1] - D_calc))   # the formula the paper used is the one that reproduces its D
            elif Z and Ms:
                Dlab, D = Ms[0][0], 1.66054 * Z * Ms[0][1] / V
            if D and D_calc:
                dd = abs(D - D_calc) / D_calc
                checks.append(('D', dd <= 0.015, dd)); detail.append('D from Z=%g and %s %.3f vs %.3f printed' % (Z, Dlab, D, D_calc))
            cif_ok = None
            if cif_cell:
                ax = sorted([cell['a'], cell['b'], cell['c']]); cx = sorted([cif_cell['a'], cif_cell['b'], cif_cell['c']])
                cif_ok = all(abs(x - y) / y <= 0.03 for x, y in zip(ax, cx)); detail.append('.cif axes %s' % ('agree' if cif_ok else 'differ'))
            v_ok = next((c[1] for c in checks if c[0] == 'V'), None)
            anchored = bool(v_ok) or bool(cif_ok)                          # the cell as read reproduces its printed volume, or the .cif's axes: the axes are not misread
            if not checks:
                status = 'nooracle' if cif_ok is None else ('agrees' if cif_ok else 'unverified')
            elif all(c[1] for c in checks):
                status = 'agrees'
            else:
                status = 'unverified'                                  # a D_calc off from every formula read is a doubt, never a red line: on the corpus each red
                                                                       # dissolved on inspection (a misread axis, a garbage ideal-formula mass, an empirical formula
                                                                       # deficient in cations while the paper's D_calc used a fuller one)
            # the cell's own standing: its printed volume and the .cif anchor it, and a density that follows
            # from it vouches for it too; a density that does not is the formula's or Z's doubt, not the cell's
            d_ok = next((c[1] for c in checks if c[0] == 'D'), None)
            cell_status = 'unverified' if v_ok is False or cif_ok is False else 'agrees' if (v_ok or cif_ok or d_ok) else 'nooracle'
            # the cell statement that stands: one that agrees, else an anchored one with doubts, else a red one
            # (anchored by construction), else a variant whose printed volume the axes do not give (the wrong γ)
            score = (0 if status == 'agrees' else 1 if status == 'unverified' and anchored else 2 if status == 'disagrees' else 3, 0 if cc.context == 'powder' else 1)
            if best is None or score < best[0]:
                best = (score, cc, cell, V, V_stated, Z, D, checks, cif_ok, status, '; '.join(detail), cell_status, anchored)
    if best is None:
        return dict(out, status='none')
    _, cc, cell, V, V_stated, Z, D, checks, cif_ok, status, detail, cell_status, anchored = best
    z_borrowed_best = not cc.Z
    src = {'powder': 'powder', 'single': 'single-crystal'}.get(cc.context, 'stated')
    out.update(status=status, cell=cell, detail=detail, cell_status=cell_status)
    dchk = next((c for c in checks if c[0] == 'D'), None)
    if dchk:
        out['D'] = {'computed': D, 'printed': D_calc, 'Z': Z, 'ok': dchk[1], 'dev': dchk[2]}   # what the density oracle found, for the D_calc record
    lines = ['cell: %s (%s) — %s' % (_cell_str(cell), src, detail or 'nothing printed to check it against')]
    for c in checks:
        if c[0] == 'V' and not c[1]:
            lines.append('the volume printed (%.1f) does not follow from the axes and angles read (%.1f) — a misread axis, or a misprint [unverified]' % (V_stated, V))
        if c[0] == 'D' and not c[1]:
            if c[2] <= 0.15 and verified_formula and len(Ms) >= 2 and anchored and not z_borrowed_best and not hydrous:
                lines.append('D_calc %.3f vs %.3f from the cell, Z = %g and the closest formula read (%+.1f %%) — none of the formulas read gives it; the paper may have used a fuller (ideal or structural) formula [unverified]'
                             % (D_calc, D, Z, (D - D_calc) / D_calc * 100))
            elif z_borrowed_best or hydrous:
                lines.append('D_calc %.3f vs %.3f from the cell, Z = %g and the formula (%+.1f %%)%s [unverified]' % (D_calc, D, Z, (D - D_calc) / D_calc * 100,
                             ' — Z taken from another statement of the paper' if z_borrowed_best else ' — the hydrate count is read from the text layer'))
            elif c[2] > 0.15:
                lines.append('D_calc %.3f vs %.3f from the cell, Z = %g and the formula (%+.0f %%) — not each other\'s: another phase\'s cell, or a formula per two units [unverified]' % (D_calc, D, Z, (D - D_calc) / D_calc * 100))
            else:
                lines.append('D_calc %.3f vs %.3f from the cell, Z = %g and the formula (%+.1f %%) [unverified]' % (D_calc, D, Z, (D - D_calc) / D_calc * 100))
    if cif_ok is False:
        lines.append('the .cif\'s cell (%s) differs from the paper\'s by more than 3 %% on an axis — another determination or setting [unverified]' % _cell_str(cif_cell))
    out['lines'] = lines
    return out

# --- the mineral name against Mindat
def species_check(ex, comp):
    """The species Mindat knows under the name read: its essential elements must all be in the
    paper's table or formula, else the name read is another mineral's (a running head, a related
    phase). -> {'status', 'detail'}; nooracle for a name Mindat does not have (a new mineral)."""
    name = ex.get('name')
    if not name:
        return {'status': 'none', 'detail': ''}
    rec = species_record(name)
    if not rec or not rec.get('elements'):
        return {'status': 'nooracle', 'detail': 'not in the Mindat cache'}
    have = set((comp or {}).get('counts') or {})
    for r in ((ex.get('epma') or {}).get('rows') or []):
        c = r.get('constituent') or ''
        if _parses(c):
            have.add(EP.parse_constituent(c).element)
    if not have:
        return {'status': 'nooracle', 'detail': 'no table or formula to compare the species with'}
    ess = {el for el in rec['elements'] if el not in ('O', 'H')}
    missing = sorted(ess - have)
    if not missing:
        return {'status': 'agrees', 'detail': 'Mindat\'s %s: %s' % (name, rec['formula'])}
    return {'status': 'unverified', 'detail': 'Mindat\'s %s has %s, absent from the paper\'s table and formula — another mineral\'s name?' % (name, ', '.join(missing))}

def _section(lines):
    """An oracle's lines as check_paper prints a section: the head as it is, its findings indented
    under it — the shape the composition, bond-valence and powder sections already have, and what a
    reader of those lines (the manuscript review) takes a head apart from a finding by."""
    return lines[:1] + ['  ' + ln for ln in lines[1:]]

def verify(ex, text, cif=None, comp=None, bv=None, powder=None):
    """Fill the records of ex['fields'] from the oracles (those already run are passed in; the
    Gladstone–Dale, cell and species ones run here). -> {'lines': the new sections' lines, 'gd', 'cell', 'species'}."""
    f = ex.get('fields') or {}
    lines = []
    # the composition oracle: table, formula, basis, method
    if comp:
        s = 'agrees' if comp.get('verified') and comp.get('ok') else 'disagrees' if comp.get('verified') else 'unverified'
        det = '; '.join(comp.get('doubts') or []) or ('the formula follows from the table' if s == 'agrees' else 'the formula deviates from the table')
        _set(f, 'epma', s, 'composition', det); _set(f, 'formula', s, 'composition', det, value=comp.get('formula') or None)
        b = comp.get('basis')
        if comp.get('verified') and b:
            if not ex.get('basis'):
                # the paper states no basis: the one the reduction found reproduces every coefficient, so
                # the formula vouches for it (owner: verified); a basis that reproduces with residuals is a doubt
                _set(f, 'basis', 'agrees' if comp.get('ok') else 'unverified', 'composition', 'inferred — the paper states no basis; %s reproduces the formula' % basis_string(b), value=b)
            elif _same_basis(b, ex.get('basis')):
                _set(f, 'basis', 'agrees' if comp.get('verified') else 'unverified', 'composition', 'the reduction on this basis reproduces the formula' if comp.get('ok') else 'the reduction on this basis is the one the formula rests on; a coefficient is off, not the basis')
            elif comp.get('basis_flag'):
                _set(f, 'basis', 'disagrees', 'composition', 'the paper states %s, but every coefficient follows from %s' % (basis_string(ex['basis']), basis_string(b)), value=b)
            else:
                _set(f, 'basis', 'unverified', 'composition', 'the stated basis does not reproduce the formula; %s does' % basis_string(b), value=b)
        else:
            _set(f, 'basis', 'unverified' if b or ex.get('basis') else 'none', 'composition', det)
        _set(f, 'method', s if comp.get('verified') else 'unverified', 'composition', det)
    elif ex.get('epma'):
        _set(f, 'epma', 'nooracle', None, 'no formula sentence to check the table against')
    # the powder table and the cell
    if powder:
        ps = powder.get('status')
        if ps == 'checked':
            good = powder['loose'] >= 0.8 * powder['n']
            _set(f, 'pxrd.calc', 'agrees' if good else 'unverified', 'cell', powder.get('head', '')[:200])
            nobs = ex['pxrd']['obs']
            um = len(powder.get('unmatched_obs') or []); off = len(powder.get('off_cell') or [])
            # the observed column is vouched for when its lines sit on the cell — paired with a calculated
            # line of the table, or on a reflection the cell allows (a line the table leaves unindexed);
            # a stray or two in ten is a typo or another phase (each is noted), more is a doubt about the column read
            _set(f, 'pxrd.obs', ('agrees' if off <= max(2, 0.1 * nobs) else 'unverified') if (good and nobs) else 'unverified' if nobs else 'none', 'cell',
                 '%d observed lines: %d paired with a calculated line, %d on reflections of the cell, %d on none' % (nobs, nobs - um, um - off, off))
            if good and powder.get('source') not in (None, '.cif'):
                _set(f, 'cell', 'agrees', 'cell', 'the powder table follows it (%d of %d lines within 1.5 %%)' % (powder['loose'], powder['n']), value=_cell_str(powder['cell']))
        elif ps in ('unindexed', 'nocell'):
            _set(f, 'pxrd.calc', 'nooracle' if ps == 'nocell' else 'none', None, powder.get('head', '')[:200])
            _set(f, 'pxrd.obs', 'nooracle', None, 'no cell, or no indices, to check the lines against')
    # the bond-valence parameter set
    if bv:
        bs = bv.get('status')
        if bs == 'checked':
            cited = bv.get('cited') or ('gh', 'burns')
            same = (bv.get('params'), bv.get('u6')) == tuple(cited)
            if bv.get('from_paper'):
                # checked against the structure the paper prints, not a .cif — a doubt, never a verdict
                _set(f, 'bv.params', 'unverified', 'bv', 'checked against the structure the paper itself prints, not a .cif')
            else:
                _set(f, 'bv.params', 'agrees' if same else 'disagrees', 'bv', 'the table agrees best with %s' % bv.get('params'))
        elif bs in ('unmatched', 'foreign'):
            _set(f, 'bv.params', 'unverified', 'bv', bv.get('message') or bv.get('head', '')[:120])
        else:
            _set(f, 'bv.params', 'nooracle', None, (bv.get('message') or '')[:120])
    elif f.get('bv.params', {}).get('status') == 'nooracle':
        f['bv.params']['detail'] = 'no .cif to check the table against'
    # Gladstone–Dale — after a sanity range: no mineral has n outside 1.3–3.0 or D outside 1–25 g/cm³
    o = ex.get('optics') or {}
    gd = gd_check(ex, comp, gd_statement(text))
    for k, s in gd['status'].items():
        v = o.get(k.split('.')[1])
        if v and not ((1.3 <= v <= 3.0) if k == 'optics.n' else (1.0 <= v <= 25.0)):
            _set(f, k, 'unverified', 'range', '%g is outside the range of minerals — misread' % v); continue
        _set(f, k, s, 'gd' if s != 'nooracle' else None, gd.get('detail') or ('no n to check it against' if k != 'optics.n' and not o.get('n') else 'no density, or no wt% table, to check it against' if s == 'nooracle' else ''))
    lines += _section(gd['lines'])
    # the cell against itself and the density
    masses = []
    for ftxt, counts_, _i, _o, kind, _c in _formulas(text, ex.get('name') or ''):
        if counts_ and counts_ != (comp or {}).get('counts'):
            masses.append((kind or 'another formula', _formula_mass_counts(counts_)))
    rec = species_record(ex.get('name')) if ex.get('name') else None
    if rec and rec.get('formula'):
        masses.append(("Mindat's ideal formula", _species_mass(rec)))
    cc = cell_consistency(text, (comp or {}).get('counts'), cif, (ex.get('optics') or {}).get('D_calc'), bool((comp or {}).get('verified') and (comp or {}).get('ok')), masses=masses)
    if f.get('cell') and f['cell']['status'] != 'agrees':                # the powder table's word, when it spoke, outranks the arithmetic
        cs = cc.get('cell_status', cc['status'])                          # the cell's own standing (volume, .cif); the density check speaks through the D_calc record
        _set(f, 'cell', cs, 'cell_consistency' if cs != 'nooracle' else None, cc['detail'], value=_cell_str(cc['cell']) if cc.get('cell') else None)
    lines += _section(cc['lines'])
    # the density from the cell, Z and the formula vouches for the printed D_calc (when the cell check
    # agreed, that is what it agreed on); and the two printed densities vouch for each other: a measured
    # density within 4 % of the calculated one is the usual case, and a misread value (another mineral's
    # density, a digit lost) is far outside it. Only for a field no other oracle adjudicated.
    if f.get('optics.D_calc', {}).get('status') == 'nooracle' and cc.get('D'):
        dd = cc['D']
        _set(f, 'optics.D_calc', 'agrees' if dd['ok'] else 'unverified', 'cell_consistency',
             'D from Z = %g, the cell and the formula %.3f vs %.3f printed (%+.1f %%)' % (dd['Z'], dd['computed'], dd['printed'], (dd['computed'] - dd['printed']) / dd['printed'] * 100))
    Dm, Dc = o.get('D_meas'), o.get('D_calc')
    if Dm and Dc and 1.0 <= Dm <= 25.0 and 1.0 <= Dc <= 25.0:
        gap = abs(Dm - Dc) / Dc
        for key in ('optics.D_meas', 'optics.D_calc'):
            if f.get(key, {}).get('status') == 'nooracle':
                _set(f, key, 'agrees' if gap <= 0.04 else 'unverified', 'density',
                     'D_meas %.3f vs D_calc %.3f (%+.1f %%)%s' % (Dm, Dc, (Dm - Dc) / Dc * 100, '' if gap <= 0.04 else ' — beyond the usual 4 %'))
        if gap > 0.04 and gap <= 0.5:
            lines.append('density: D_meas %.3f vs D_calc %.3f (%+.1f %%) — beyond the usual agreement; a misread value, or a porous or impure sample [unverified]' % (Dm, Dc, (Dm - Dc) / Dc * 100))
    # Mindat's cell for a known species, when nothing in the paper checked its cell: a co-equal proxy — agreement
    # vouches for the reading, a difference is a doubt (another determination, a polytype), never a finding
    if f.get('cell') and f['cell']['status'] == 'nooracle' and cc.get('cell') and ex.get('name'):
        try:
            from pxrd_review import extra_checks as X
            rec_ = X.mindat_struct(ex['name'], exact=True)
            mind = X._cell_lengths((rec_ or {}).get('cell') or {}) if rec_ else None
            mine = X._cell_lengths({k: cc['cell'].get(k) for k in ('a', 'b', 'c')})
        except Exception:
            mind = mine = None
        if mind and mine:
            diff_ = X._len_maxdiff(mine, mind)
            if diff_ < 0.03:
                _set(f, 'cell', 'agrees', 'mindat', "Mindat's cell for %s agrees (axes within %.1f %%)" % (ex['name'], diff_ * 100))
            else:
                _set(f, 'cell', 'unverified', 'mindat', "Mindat's cell for %s differs (%.0f %% on an axis) — another determination or setting" % (ex['name'], diff_ * 100))
                lines.append("cell: Mindat's cell for %s (%s) differs from the paper's by %.0f %% on an axis — another determination, setting or polytype [unverified]" % (ex['name'], ', '.join('%.3f' % x for x in mind), diff_ * 100))
    # the species
    sp = species_check(ex, comp)
    _set(f, 'name', sp['status'], 'species' if sp['status'] != 'nooracle' else None, sp['detail'])
    if sp['status'] == 'unverified':
        lines.append('name: %s' % sp['detail'] + ' [unverified]')
    return {'lines': lines, 'gd': gd, 'cell': cc, 'species': sp}

_READER_LABELS = (('epma', 'table'), ('formula', 'formula'), ('basis', 'basis'), ('optics.n', 'n'), ('optics.D_meas', 'D_meas'), ('optics.D_calc', 'D_calc'),
                  ('cell', 'cell'), ('bv.params', 'bond-valence set'), ('pxrd.calc', 'powder'), ('name', 'name'))
_MARK = {'agrees': '✓', 'disagrees': '✗', 'unverified': '?', 'nooracle': '·', 'none': '—'}

def readers_line(fields):
    """One line on what was read and what vouched for it: 'readers: table ✓ (p6) · formula ✓ · …'."""
    parts = []
    for key, label in _READER_LABELS:
        r = fields.get(key)
        if not r or r['status'] == 'none':
            continue
        page = ' (p%d)' % r['page'] if r.get('page') else ''
        why = '' if r['status'] in ('agrees', 'disagrees') else (' (%s)' % r['detail'][:60] if r['status'] == 'nooracle' and r['detail'] else '')
        parts.append('%s %s%s%s' % (label, _MARK[r['status']], page, why))
    return 'readers: ' + (' · '.join(parts) if parts else 'nothing read')

def _paper_cif(pdf, text, out):
    """A .cif written from the structure the paper itself prints, for a paper that comes without
    one. Note-grade: on the corpus 91 % of its cation sites reproduced a real .cif's bond-valence
    sums, so whatever it supports is a doubt to look at, never a verdict. -> a path or None; the
    caller records it in out['paper_structure'] so it can be cleaned up."""
    if pdf.lower().endswith('.docx'):
        return None
    try:
        from pxrd_review import paper_structure as PS
        st, info = PS.build(pdf, text)
    except Exception:
        return None
    if st is None or not info.get('path'):
        return None
    out['paper_structure'] = info
    return info['path']


def check_paper(pdf, cif=None, out_dir=None):
    """The paper against itself and its .cif: {'extract', 'composition', 'bv', 'bv_status', 'powder',
    'powder_status', 'fields', 'lines'} — the lines are what a manuscript review prints: a 'readers:'
    line on what was read and what vouched for it, then 'composition', 'bond valence', 'powder table',
    'Gladstone–Dale', 'cell' and 'name' sections; ex['fields'] holds the record per field (see verify)."""
    text = text_of(pdf)
    ex = extract(pdf, out_dir, None, write=bool(out_dir))
    out = {'extract': ex, 'composition': check_composition(ex, text), 'bv': None, 'bv_status': None, 'powder': None, 'powder_status': None,
           'fields': ex.get('fields') or {}, 'lines': []}
    if out['composition']:
        out['lines'] += out['composition']['lines']
    elif ex.get('epma'):
        out['lines'].append('composition: an analytical table was read but no empirical formula sentence was found to check it against')
    bc = None
    # no .cif — but roughly half the papers print a structure of their own. It stands in for the
    # BOND-VALENCE check only: the cell check and verify must not be handed the paper's own cell
    # back as if it were independent evidence. Note-grade throughout (91 % of sites on the corpus).
    bv_cif = cif or _paper_cif(pdf, text, out)
    if bv_cif:
        bc = bv_check_paper(pdf, bv_cif, ex)
    if out.get('paper_structure'):                       # bv_check_paper has read it; the file goes now
        from pxrd_review import paper_structure as _PS
        ps_ = out['paper_structure']
        _PS.discard(ps_)
        out['paper_structure'] = {k: v for k, v in ps_.items() if k not in ('path', 'dir')}
        if bc is not None:
            bc['from_paper'] = True                      # verify() must not read this as a verdict
            where = 'the structure the paper itself prints (%s, %s; instability index %.2f vu)' % (
                ps_.get('sym', '?'), _cell_str(dict(zip(('a', 'b', 'c', 'α', 'β', 'γ'), ps_['cell']))), ps_['gii'])
            def _note(t, full):
                t = t.replace('the .cif', where if full else "the paper's own structure")
                t = t.replace('.cif', where if full else "the paper's own structure")
                return t if '[unverified]' in t else t + ' [unverified]'
            for k_ in ('head', 'message'):
                if bc.get(k_):
                    bc[k_] = _note(bc[k_], True)         # the head says which structure, once
            if bc.get('lines'):
                bc['lines'] = [_note(ln, False) for ln in bc['lines']]
    if bc is not None:
        out['bv_status'] = bc['status']
        if bc.get('lines') is None:
            out['lines'].append('bond valence: ' + bc['message'])
        else:
            out['bv'] = {k: bc[k] for k in ('tables', 'lines', 'params', 'u6', 'cited', 'compared', 'disagree')}
            out['lines'] += [bc['head']] + ['  ' + ln for ln in bc['lines']]
    try:
        cc = cell_check(pdf, cif, text, table=ex.get('_table'))
    except Exception as ex_:
        cc = {'status': 'error', 'head': 'powder table: could not check (%s)' % ex_, 'lines': []}
    out['powder_status'] = cc['status']
    if cc.get('head'):
        out['powder'] = cc
        out['lines'] += [cc['head']] + ['  ' + ln for ln in cc['lines']]
    try:
        v = verify(ex, text, cif, comp=out['composition'], bv=bc, powder=cc if cc.get('status') != 'error' else None)
        out['lines'] += v['lines']
    except Exception as ex_:
        out['lines'].append('readers: could not verify the readings (%s)' % ex_)
    out['lines'].insert(0, readers_line(out['fields']))
    return out

_AN_RE = re.compile(r'^(O|OH|OW|Ow|W|Wat|F|Cl|OD|Oh|Hw|H2O)\d*[A-Za-z]?\d*$')

def _bv_norm(t):
    from pxrd_review import bv_check as B
    return B._norm_label(t.replace('−', '-').replace('–', '-'))

def _cation_labels(st):
    """The structure's cation site labels as bv_check normalises them (hydrogen left out)."""
    cats = {_bv_norm(x) for r in st.cations for x in r.label.split('/')} | {_bv_norm(r.label) for r in st.cations}
    return {c for c in cats if not c.startswith('H')}

def _names_cations(rows, st):
    """Whether a Word table belongs to this structure — the rule bv_tables applies to a pdf's
    lines: one of its first three rows names ≥2 distinct cation sites of the .cif (every one of
    them, for a structure with a single cation site). check_bvs_table alone accepts any table with
    one matching header label and an anion row, which scores a stranger's table against the wrong
    .cif when the folder holds only one."""
    cats = _cation_labels(st)
    n_sites = len({r.label for r in st.cations if not _bv_norm(r.label).startswith('H')})   # 'Mg/Mn' is one site, three labels
    for row in rows[:3]:
        hits = {_bv_norm(x) for x in row if _bv_norm(x) in cats}
        if len(hits) >= 2 or (hits and n_sites <= 1):
            return True
    return False

def _bv_like(rows):
    """A table that reads like a bond-valence table (anion-labelled rows, valences below 1 vu)."""
    an = sum(1 for r in rows if r and _AN_RE.match(r[0].strip()))
    vu = sum(1 for r in rows for c in r[1:] if re.match(r'^\s*0\.\d\d', c))
    return an >= 2 and vu >= 4

def bv_check_paper(path, cif, ex):
    """The paper's bond-valence table (a pdf's, read from the page; a manuscript .docx's, from
    its Word table) against the .cif, under every parameter set — the one that agrees best wins,
    the cited one on a tie. -> {'status', …}:
      'checked'   {'head', 'lines', 'tables', 'params', 'u6', 'cited', 'compared', 'disagree'}
      'unmatched' the same keys, compared 0: a table was read but none of its cells matched a
                  bond of the .cif (the labels differ — Ow/OH vs O); its row and column sums are
                  still in the lines, so a wrong Σ is not lost with them
      'foreign'   {'message'}: a bond-valence-like table that names none of the .cif's cations
      'none'      {'message'}: no table;  'error' {'message'}: the .cif will not compute."""
    try:
        from pxrd_review import bv_check as B
        if path.lower().endswith('.docx'):
            all_tabs = [t for t in B.read_tables(path) if len(t) >= 2]
            if not any(sum(1 for r in t for c in r if re.match(r'^\s*0\.\d\d', c)) >= 3 or any(_BVS_HEAD.match((c or '').strip()) for r in t[:3] for c in r) for t in all_tabs):
                return {'status': 'none', 'message': 'no bond-valence table found in the manuscript (nothing to check against the .cif)'}   # no valences and no BVS column: the .cif is not needed, so a broken one is no finding
            st = B.Structure(cif)
            all_tabs = [t for t in (_maybe_transpose(t, st) for t in all_tabs) if len(t) >= 3]   # a grid the other way round has its rows as columns before the transpose
            has_bvs_col = lambda t: any(_BVS_HEAD.match((c or '').strip()) for r in t[:3] for c in r)   # a BVS column of a coordinates / bond table, not a grid
            tabs = [{'page': None, 'rows': t, 'kind': 'grid'} for t in all_tabs if _names_cations(t, st) and not has_bvs_col(t)]
            if not tabs:
                tabs = [{'page': None, 'kind': 'sites', 'head': 'BVS', 'rows': rs} for rs in (_site_rows_from_grid(t, st) for t in all_tabs if has_bvs_col(t)) if len(rs) >= 2]
            if not tabs and any(_bv_like(t) for t in all_tabs):
                return {'status': 'foreign', 'message': 'a bond-valence-like table was read but it names none of the .cif\'s cation sites (%s) — is that the right structure?'
                        % ', '.join(sorted(r.label for r in st.cations if not r.label.upper().startswith('H')))[:80]}
        else:
            st = B.Structure(cif)
            tabs = bv_tables(path, st) or bvs_site_tables(path, st)
        if not tabs:
            return {'status': 'none', 'message': 'no bond-valence table found in the paper (nothing to check against the .cif)'}
        sites = tabs[0].get('kind') == 'sites'
        where = ('the paper\'s %s (p%d)' % ('BVS column' if sites else 'table', tabs[0]['page'])) if tabs[0].get('page') else ('the manuscript\'s %s' % ('BVS column' if sites else 'table'))
        cited = (ex['bv'].get('params') or 'gh', ex['bv'].get('u6') or 'burns')
        best = None; as_cited = None
        # a paper may print several BVS columns (two parameter sets, with and without H): each is judged on
        # its own and the column that agrees best stands for the paper — the others are not findings
        groups = [[t] for t in tabs] if sites else [tabs]
        for key in ('gh', 'bo', 'ba'):
            for u6 in ('burns', 'params'):
                P = B.Params(prefer=key, u6=u6); notes = list(st.notes)
                rk = B.compute(st, P, None, 'oo'); st.notes[:] = notes
                for grp in groups:
                    if sites:
                        lines = B.check_bvs_sites(st, rk[0], rk[1], grp, B.PARAM_NAMES[key], compare_anions=not rk[3])
                    else:
                        lines = B.check_bvs_table(st, rk[0], rk[2], rk[1], [t['rows'] for t in grp], B.PARAM_NAMES[key])
                    hits = [re.search(r'(\d+) cells compared, (\d+) disagree', ln) for ln in lines]
                    n = sum(int(m.group(1)) for m in hits if m); bad = sum(int(m.group(2)) for m in hits if m)
                    if (key, u6) == cited or as_cited is None:
                        as_cited = (key, u6, lines)
                    if n and (best is None or (bad / n, 0 if (key, u6) == cited else 1, -n) < (best[0] / best[6], best[1], -best[6])):
                        best = (bad, 0 if (key, u6) == cited else 1, key, u6, P, lines, n)
        if best is None:
            key, u6, lines = as_cited
            lines = [ln for ln in lines if not ln.startswith('no bond-valence table found')]   # the head below says why
            head = ('bond valence: %s was read but none of its cells matched a bond of the .cif — the site labels may differ '
                    '(Ow/OH in the table vs O in the .cif); the row and column sums below were still checked' % where)
            return {'status': 'unmatched', 'head': head, 'lines': lines, 'tables': len(tabs), 'params': key, 'u6': u6, 'cited': cited, 'compared': 0, 'disagree': 0}
        bad, _, key, u6, P, lines, n = best
        if n >= 4 and 2 * bad >= n:                                   # half or more of the cells differ: not this convention, or not this reading
            head = ('bond valence: %s vs the .cif — %d of %d cells differ under every parameter set: another convention (occupancy-weighted sums, '
                    'other parameters, another site labelling) or a misread table; not compared cell by cell [unverified]' % (where, bad, n))
            lines = [ln for ln in lines if 'cells compared' in ln]
            return {'status': 'unmatched', 'head': head, 'lines': lines, 'tables': len(tabs), 'params': key, 'u6': u6, 'cited': cited, 'compared': n, 'disagree': bad}
        head = 'bond valence: %s vs the .cif — agrees best with %s%s' % (
            where, P.note(), '' if (key, u6) == cited else ' (the paper cites %s%s)' % (B.PARAM_NAMES.get(cited[0], cited[0]), ', U6+ from Burns' if cited[1] == 'burns' else ''))
        return {'status': 'checked', 'head': head, 'lines': lines, 'tables': len(tabs), 'params': key, 'u6': u6, 'cited': cited, 'compared': n, 'disagree': bad}
    except Exception as ex_:
        return {'status': 'error', 'message': 'could not check (%s)' % ex_}

def bv_tables(pdf, st):
    """The paper's bond-valence grids (anion rows × cation columns) as row lists, read from the pdf
    by word positions. The header line names ≥2 of the structure's cations (or ≥2 of its anions,
    for a grid printed the other way round — cations down the rows — which is transposed before it
    is handed on); the table's width is that of the labelled header tokens, so a powder table or the
    other page column printed beside it stays out; a line of bare valences ('0.67×2↓' printed above
    or below its row: a cell with two distances) joins the row nearest to it. Columns are the data
    tokens' x-clusters labelled by the nearest header token. -> [{'page', 'rows', 'kind': 'grid'}]"""
    norm = _bv_norm; cats = _cation_labels(st)
    anions = {norm(x) for a in st.anions for x in a.label.split('/')} | {norm(a.label) for a in st.anions}
    an_re = _AN_RE
    is_an = lambda t: norm(t) in anions or bool(an_re.match(t))
    is_cat = lambda t: norm(t) in cats
    val_re = re.compile(r'^\d\.\d+(?:\s*[×x]\d+)?[↓→]*$|^[×x]?\d*[↓→]+$|^[×x]\d+[↓→]*$|^\(\d+\)$|^[↓→]+$')
    out = []
    for pno, lines in enumerate(_pages(pdf)):
        i = 0
        while i < len(lines):
            toks = [w[4] for w in lines[i]['w']]
            hits = [k for k, t in enumerate(toks) if is_cat(t)]
            an_hits = [k for k, t in enumerate(toks) if is_an(t) and not is_cat(t)]
            transposed = False
            if not (len(hits) >= 2 and len({norm(toks[k]) for k in hits}) >= 2):
                if len(an_hits) >= 2 and len({norm(toks[k]) for k in an_hits}) >= 2 and len(hits) < 2 \
                        and any(is_cat((lines[j]['w'] or [(0, 0, 0, 0, '')])[0][4]) for j in range(i + 1, min(len(lines), i + 4))):
                    transposed = True; hits = an_hits                    # anions across, cations down: 'Site F(1) F(2) … Σcations'
                else:
                    i += 1; continue
            lab_ws = [lines[i]['w'][k] for k in hits]
            k0 = min(hits)
            while k0 > 0 and re.fullmatch(r"[A-Z][a-z]{0,2}\d{0,2}[a-z]?(?:\([^)]{1,6}\))?'*\*?", toks[k0 - 1]):
                k0 -= 1                                                    # 'Pb M1 M2' before the first matched label: the paper's own site names, part of the header
            x_lo = min(lines[i]['w'][k0][0], min(w[0] for w in lab_ws)) - 130; x_hi = max(w[2] for w in lab_ws) + 130   # the row labels to the left, a Σ column to the right
            head = [w for w in lines[i]['w'] if w[0] >= x_lo and w[2] <= x_hi]; j = i + 1
            row_ok = is_cat if transposed else is_an
            while j < len(lines) and j <= i + 2:
                lw = [w for w in lines[j]['w'] if w[0] >= x_lo and w[2] <= x_hi]
                if not lw or row_ok(lw[0][4]) or re.search(r'\d\.\d\d', ' '.join(w[4] for w in lw)):
                    break
                head += lw; j += 1
            raw = []; ys = []; blank = 0; pending = []
            while j < len(lines) and j < i + 90:
                lw = [w for w in lines[j]['w'] if w[0] >= x_lo and w[2] <= x_hi]
                if not lw:
                    blank += 1; j += 1
                    if blank > 3: break
                    continue
                first = lw[0][4]; y = lines[j]['y']
                if row_ok(first) or re.match(r'^(Σ|Sum|Total|Σcat|Σanion|Σan)', first):
                    blank = 0
                    if pending and raw and abs(pending[0][0] - ys[-1]) < abs(y - pending[0][0]):
                        raw[-1] = (raw[-1][0], raw[-1][1] + [w for _, ws in pending for w in ws]); pending = []
                    raw.append((first, lw[1:] + [w for _, ws in pending for w in ws])); ys.append(y); pending = []
                    if re.match(r'^(Σ|Sum|Total)', first) and not re.match(r'^Σan', first):
                        j += 1; break
                elif all(val_re.match(w[4]) for w in lw):
                    pending.append((y, lw))                                # bare valences: the nearest row's cell with two distances
                elif raw:
                    break
                j += 1
            if pending and raw:
                raw[-1] = (raw[-1][0], raw[-1][1] + [w for _, ws in pending for w in ws])
            if len(raw) < 2:
                i += 1; continue
            xs = sorted((w[0] + w[2]) / 2 for _, ws in raw for w in ws)
            cols = []
            for x in xs:
                if cols and x - cols[-1][-1] <= 9.0:
                    cols[-1].append(x)
                else:
                    cols.append([x])
            cols = [sum(c) / len(c) for c in cols]
            if not cols:
                i += 1; continue
            hw = [((w[0] + w[2]) / 2, w[4]) for w in head]
            labels = [next((t for x, t in sorted(hw, key=lambda h: abs(h[0] - cx)) if abs(x - cx) < 28), '?') for cx in cols]
            rows = []
            for first, ws in raw:
                cells = [[] for _ in cols]
                for w in sorted(ws, key=lambda w: w[0]):
                    k = min(range(len(cols)), key=lambda c: abs(cols[c] - (w[0] + w[2]) / 2))
                    cells[k].append(w[4])
                rows.append([first] + [' '.join(c) for c in cells])
            grid = [['Atom'] + labels] + rows
            if transposed:
                grid = [list(col) for col in zip(*[r + [''] * (len(grid[0]) - len(r)) for r in grid])]   # anion rows × cation columns, as the checker reads
                grid[0][0] = 'Atom'
            out.append({'page': pno + 1, 'rows': grid, 'kind': 'grid'})
            i = j
    return out

_BVS_HEAD = re.compile(r'^(BVS|BVSs|BVS\*|BVS[a-c]|ΣBVS?|Σv|BV|BVsum|BVsums|Σs|Σ\(s\)|BVS\d)\*{0,2}$')   # a bare 'Σ' is a grid's sum column, not a BVS column

def bvs_site_tables(pdf, st):
    """Bond-valence SUMS printed as a column of another table — the coordinates table ('Atom x y z
    Ueq BVS'), the bond-distance table ('Atom Distance … BVS'), a per-site table ('Site (a) (b)'):
    one value per site. Each BVS-headed column is one table: [{'page', 'kind': 'sites', 'head',
    'rows': [(site label, value), …]}]. Rows are lines whose first token is a site of the .cif
    (cation or anion) and whose value sits under the header token."""
    norm = _bv_norm; cats = _cation_labels(st)
    anions = {norm(x) for a in st.anions for x in a.label.split('/')} | {norm(a.label) for a in st.anions}
    out = []
    for pno, lines in enumerate(_pages(pdf)):
        for i, ln in enumerate(lines):
            ws = ln['w']; toks = [w[4] for w in ws]
            bcols = [k for k, t in enumerate(toks) if _BVS_HEAD.match(t)]
            if not bcols or not any(re.match(r'^(Atom|Site|Sites?|Atoms?|Label|Cation|Anion|Position)$', t, re.I) for t in toks):
                continue
            for k in bcols:
                x_b = (ws[k][0] + ws[k][2]) / 2
                rows = []; miss = 0; j = i + 1
                while j < len(lines) and j < i + 80:
                    lw = lines[j]['w']
                    if not lw:
                        j += 1; miss += 1
                        if miss > 4: break
                        continue
                    first = lw[0][4]
                    if re.match(r'^(Table|TABLE|Tab\.)$', first) and len(lw) > 1 and re.match(r'^\d', lw[1][4]):
                        break
                    lab_w = None
                    if norm(first) in cats or norm(first) in anions:
                        lab_w = lw[0]; vals = lw[1:]; near = lambda w: abs((w[0] + w[2]) / 2 - x_b) < 26
                    else:                                                     # 'Cr1–O7 1.92 … O1 1.95': the site under the header, its sum beside it (a bond table's BVS column)
                        k_ = next((q for q, w in enumerate(lw) if (norm(w[4]) in cats or norm(w[4]) in anions) and abs((w[0] + w[2]) / 2 - x_b) < 30), None)
                        if k_ is not None:
                            lab_w = lw[k_]; vals = lw[k_ + 1:k_ + 3]; near = lambda w: 0 < w[0] - lab_w[2] < 70
                    if lab_w is not None:
                        cand = [(abs((w[0] + w[2]) / 2 - x_b), w[4]) for w in vals if re.fullmatch(r'\d\.\d+(?:\(\d+\))?', w[4]) and near(w)]
                        if cand:
                            v = float(re.match(r'\d\.\d+', min(cand)[1]).group(0))
                            if 0.1 <= v <= 8.5:
                                rows.append((lab_w[4], v)); miss = 0; j += 1; continue
                    miss += 1
                    if miss > 4: break
                    j += 1
                if len(rows) >= 2 and len({r[0] for r in rows}) >= 2:
                    out.append({'page': pno + 1, 'kind': 'sites', 'head': toks[k], 'rows': rows})
    return out

def _maybe_transpose(rows, st):
    """A Word bond-valence table with the anions across the header and the cations down the rows
    ('Site F(1) F(2) … Σcations' / 'Pb(1a) 0.48 …') -> anion rows × cation columns, as the checker
    reads; any other table unchanged."""
    if not rows or len(rows) < 2:
        return rows
    norm = _bv_norm; cats = _cation_labels(st)
    anions = {norm(x) for a in st.anions for x in a.label.split('/')} | {norm(a.label) for a in st.anions}
    hdr = rows[0]
    an_hits = sum(1 for c in hdr[1:] if norm((c or '').strip()) in anions or _AN_RE.match((c or '').strip()))
    cat_hits = sum(1 for c in hdr[1:] if norm((c or '').strip()) in cats)
    cat_rows = sum(1 for r in rows[1:] if r and norm((r[0] or '').strip()) in cats)
    n_sites = len({r.label for r in st.cations if not norm(r.label).startswith('H')})
    if an_hits >= 2 and cat_hits < 2 and cat_rows >= min(2, max(n_sites, 1)):
        w = max(len(r) for r in rows)
        grid = [list(col) for col in zip(*[list(r) + [''] * (w - len(r)) for r in rows])]
        grid[0][0] = 'Atom'
        return grid
    return rows

def _site_rows_from_grid(rows, st):
    """A Word table with a BVS column: -> [(site label, value), …] or []."""
    norm = _bv_norm; cats = _cation_labels(st)
    anions = {norm(x) for a in st.anions for x in a.label.split('/')} | {norm(a.label) for a in st.anions}
    for ri in range(min(3, len(rows))):
        ks = [ci for ci, c in enumerate(rows[ri]) if _BVS_HEAD.match((c or '').strip())]
        if ks and any(re.match(r'^(Atom|Site|Sites?|Atoms?|Label|Cation|Anion)$', (c or '').strip(), re.I) for c in rows[ri]):
            k = ks[0]; out = []
            for r in rows[ri + 1:]:
                if r and k < len(r) and (norm(r[0].strip()) in cats or norm(r[0].strip()) in anions):
                    m = re.match(r'\s*(\d\.\d+)', r[k] or '')
                    if m:
                        out.append((r[0].strip(), float(m.group(1))))
            return out
    return []

# ----------------------------------------------------------------------------- name, the whole

_NOT_MINERALS = {'despite', 'composite', 'satellite', 'definite', 'infinite', 'opposite', 'favourite', 'favorite', 'suite', 'granite',
                 'quite', 'write', 'white', 'unite', 'ignite', 'recite', 'polite', 'finite', 'excite', 'invite', 'requisite', 'exquisite',
                 'appetite', 'website', 'graphite', 'pegmatite', 'syenite', 'granodiorite', 'diorite', 'rhyolite', 'andesite', 'dolerite',
                 'kimberlite', 'carbonatite', 'peridotite', 'eclogite', 'amphibolite', 'phonolite', 'trachyte', 'aplite', 'skarnite'}

def mineral_name(text):
    """The headline mineral: among the capitalised -ite names of the paper's head, the one the
    whole paper uses most ('Mendigite, a new mineral …' beats a 'Bustamite' in the running head)."""
    head = text[:3000]; low = text.lower()
    cands = []
    for m in re.finditer(r'(?<![^\W\d_])([^\W\d_]{4,}ite(?:-\([A-Za-z]+\))?)(?![^\W\d_])', head):
        w = m.group(1)
        if w[0].isupper() and w.lower().split('-')[0] not in _NOT_MINERALS and w.lower() not in cands:
            cands.append(w.lower())
    if cands:
        return max(cands, key=lambda w: (low.count(w.split('-')[0]), -cands.index(w)))
    names = re.findall(r'\b([a-z]{4,}ite(?:-\([a-z]+\))?)\b', text.lower())
    if names:
        uniq = list(dict.fromkeys(names))                              # first mention first: a set's order is the run's hash seed, and two names tie often (cassiterite and reedmergnerite, 6 each)
        return max(uniq, key=lambda w: (names.count(w), -uniq.index(w)))
    return ''

def extract(pdf, out_dir=None, stem=None, write=True):
    """Everything the tabs can take from the paper: {'name', 'epma': {...}, 'basis', 'basis_sentence',
    'method', 'optics', 'bv', 'pxrd': {'obs': n, 'calc': n}, 'files': {...}, 'notes': [...]}.
    With write=True the data files the EPMA and PXRD tabs read are written to out_dir:
    <stem>_paper_epma.csv (constituents as columns, one row = the means), <stem>_paper_obs.txt and
    <stem>_paper_calc.txt — <stem>_docx_* from a manuscript .docx, so a revised manuscript beside
    the published paper of the same name never overwrites its files (or the other way round)."""
    text = text_of(pdf)
    stem = (stem or os.path.splitext(os.path.basename(pdf))[0]) + ('_docx_' if pdf.lower().endswith('.docx') else '_paper_')
    name = mineral_name(text)
    out = {'name': name, 'epma': epma_table(pdf, name), 'method': method_statements(text), 'optics': optics(text),
           'bv': bv_statement(text), 'files': {}, 'notes': []}
    out['basis'], out['basis_sentence'] = basis_statement(text)
    o, c, hit = pxrd_table(pdf, with_pages=True)
    out['pxrd'] = {'obs': len(o), 'calc': len(c)}; out['_table'] = (o, c, hit)
    if write and out_dir:
        os.makedirs(out_dir, exist_ok=True)
        if out['epma'] and out['epma']['rows']:
            p = os.path.join(out_dir, stem + 'epma.csv')
            rows = out['epma']['rows']
            with open(p, 'w', encoding='utf-8') as f:
                f.write(','.join(r['constituent'] for r in rows) + '\n')
                f.write(','.join('%g' % r['mean'] for r in rows) + '\n')
            out['files']['epma'] = os.path.basename(p)
        if o:
            p = os.path.join(out_dir, stem + 'obs.txt')
            with open(p, 'w', encoding='utf-8') as f:
                f.write('d I\n' + ''.join('%.4f %g\n' % (d, I if I is not None else 0) for d, I in o))
            out['files']['obs'] = os.path.basename(p)
        if c:
            p = os.path.join(out_dir, stem + 'calc.txt')
            with open(p, 'w', encoding='utf-8') as f:
                f.write('d I h k l\n' + ''.join('%.4f %g %d %d %d\n' % (d, I if I is not None else 0, h[0], h[1], h[2]) for d, I, h in c))
            out['files']['calc'] = os.path.basename(p)
    if not out['epma']:
        out['notes'].append('no analytical table found in the paper')
    if not out['basis']:
        out['notes'].append('no normalisation basis stated — the EPMA tab keeps its own')
    out['_path'] = pdf                                                 # the records' page guard: a manuscript's "pages" are its tables, not pages
    out['fields'] = _fields_init(out, text)
    return out

def basis_string(b):
    if not b:
        return ''
    if b[0] == 'O':
        return 'O=%g' % b[1]
    if b[0] == 'cations':
        return 'cations=%g' % b[1]
    return '%s=%g' % (b[1], b[2])

def main(argv=None):
    ap = argparse.ArgumentParser(prog='pxrd paper', description=__doc__.split('\n\n')[1], formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('pdf'); ap.add_argument('--out', help='where the data files go (default <pdf dir>/review_out)')
    ap.add_argument('--check', action='store_true', help="check the paper against itself: its formula from its own table and basis, its bond-valence table against the .cif")
    ap.add_argument('--cif', help='the structure .cif, for the bond-valence check')
    a = ap.parse_args(argv)
    if a.check:
        r = check_paper(a.pdf, a.cif, None)
        print('\n'.join(r['lines']) or 'nothing to check: no analytical table or formula sentence found')
        return 0
    out = extract(a.pdf, a.out or os.path.join(os.path.dirname(os.path.abspath(a.pdf)), 'review_out'))
    e = out['epma']
    print('mineral: %s' % (out['name'] or '?'))
    if e:
        print('analytical table (%s, %d constituents; header: %s):' % (('page %d' % e['page']) if e.get('page') else 'a table of the manuscript', len(e['rows']), e['header'][:80]))
        for r in e['rows']:
            print('  %-8s %8.2f  %s  %s  %s' % (r['constituent'], r['mean'], ('%g–%g' % r['range']) if r['range'] else '', ('sd %g' % r['sd']) if r['sd'] is not None else '', r['standard'] or ''))
        if e['total'] is not None:
            print('  total    %8.2f' % e['total'])
    print('basis: %s   %s' % (basis_string(out['basis']) or '?', ('«%s»' % out['basis_sentence']) if out['basis_sentence'] else ''))
    for s in out['method']['sentences'][:6]:
        print('method: %s' % s[:220])
    print('optics: n %s (%s); D meas %s; D calc %s' % (out['optics']['n'], out['optics']['n_from'], out['optics']['D_meas'], out['optics']['D_calc']))
    print('bond valence: params %s, U6+ %s, H bonds %s' % (out['bv']['params'], out['bv']['u6'], out['bv']['hb']))
    print('powder table: %d observed, %d calculated lines' % (out['pxrd']['obs'], out['pxrd']['calc']))
    for k, v in out['files'].items():
        print('  wrote %s' % v)
    for n in out['notes']:
        print('note: %s' % n)
    return 0

if __name__ == '__main__':
    sys.exit(main())
