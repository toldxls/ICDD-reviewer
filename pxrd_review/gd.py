#!/usr/bin/env python3
"""
gd — the Gladstone–Dale compatibility of a mineral description (Mandarino 1981):

    K_C = Σ k_i · wt%_i / 100         from the composition (ideal or empirical)
    K_P = (n_mean − 1) / D             from the mean refractive index and the density
    compatibility = 1 − K_P / K_C     superior |x| < 0.02, excellent < 0.04, good < 0.06, fair < 0.08, poor beyond

    python3 -m pxrd_review.gd --wt "UO3=63.88,PbO=13.41,SeO2=14.26,H2O=7.0" --n 2.062 --density 6.02
    python3 -m pxrd_review.gd --formula "Pb=1,U=3,Se=2,O=15,H2O=5" --n 2.062 --cif mineral.cif    # D_calc from the cell
    pxrd gd --wt … --n … [--density D | --cif X.cif [--z Z]] [--xlsx] [--out DIR]

Constituents may be given as wt% oxides (--wt) or as an ideal formula in atoms per formula unit
(--formula, converted to wt% of the oxides of the usual valences); the density is measured
(--density) or calculated from the cell (--cif, Z from the .cif or --z) — both are reported when
both are available. Constants come from data/gd_constants.json (each with its source; check the
ones marked 'check'); --k X=0.123 overrides one.
"""
import os, re, sys, json, argparse
from collections import OrderedDict

from pxrd_review import epma as EP

CONST_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data', 'gd_constants.json')

def constants():
    with open(CONST_FILE, encoding='utf-8') as f:
        d = json.load(f)
    return {k: v for k, v in d.items() if not k.startswith('_')}

def category(x):
    a = abs(x)
    return 'superior' if a < 0.02 else 'excellent' if a < 0.04 else 'good' if a < 0.06 else 'fair' if a < 0.08 else 'poor'

USUAL_OXIDE = {'H': 'H2O', 'Li': 'Li2O', 'Na': 'Na2O', 'K': 'K2O', 'Rb': 'Rb2O', 'Cs': 'Cs2O', 'Tl': 'Tl2O', 'Ag': 'Ag2O',
               'Be': 'BeO', 'Mg': 'MgO', 'Ca': 'CaO', 'Sr': 'SrO', 'Ba': 'BaO', 'Mn': 'MnO', 'Fe': 'FeO', 'Co': 'CoO',
               'Ni': 'NiO', 'Cu': 'CuO', 'Zn': 'ZnO', 'Cd': 'CdO', 'Hg': 'HgO', 'Pb': 'PbO', 'B': 'B2O3', 'Al': 'Al2O3',
               'Cr': 'Cr2O3', 'Ga': 'Ga2O3', 'Bi': 'Bi2O3', 'Y': 'Y2O3', 'La': 'La2O3', 'Ce': 'Ce2O3', 'Nd': 'Nd2O3',
               'Sb': 'Sb2O3', 'Si': 'SiO2', 'Ti': 'TiO2', 'Zr': 'ZrO2', 'Sn': 'SnO2', 'Ge': 'GeO2', 'Te': 'TeO2', 'Se': 'SeO2',
               'C': 'CO2', 'P': 'P2O5', 'As': 'As2O5', 'V': 'V2O5', 'Nb': 'Nb2O5', 'Ta': 'Ta2O5', 'S': 'SO3', 'Cr6': 'CrO3',
               'Mo': 'MoO3', 'W': 'WO3', 'U': 'UO3', 'Th': 'ThO2', 'F': 'F', 'Cl': 'Cl', 'Br': 'Br'}

def formula_to_wt(apfu, oxides=None, apfu_by_key=False):
    """{'Pb':1,'U':3,'Se':2,'H2O':5} (O implied) -> {constituent: wt%} and the formula weight.
    Elements map to their usual oxide unless `oxides` says otherwise (e.g. {'Fe': 'Fe2O3', 'S': 'S'})."""
    oxides = oxides or {}
    mass = OrderedDict(); fw = 0.0; by_key = OrderedDict()
    for el, n in apfu.items():
        if el == 'O':
            continue
        # an element maps to its usual oxide (S -> SO3; pass --oxide S=S for a sulfide); a constituent
        # written as such (H2O, F, Cl, CO2, Fe2O3) is kept
        key = oxides.get(el) or (el if el in ('H2O', 'F', 'Cl', 'Br', 'CO2', 'NH3') or re.search(r'\d|O', el) else USUAL_OXIDE.get(el, el))
        c = EP.parse_constituent(key)
        if el == 'H' and c.kind == 'water':
            n = n / 2.0                                          # H atoms, two to a water: 'Mg=1,H=2' is brucite (58.32), not Mg(OH)2·H2O (76.33)
        m = n * c.mw / (c.n_cat if c.kind != 'water' else 1)     # mass of this constituent per formula unit
        mass[key] = mass.get(key, 0.0) + m
        by_key[key] = by_key.get(key, 0.0) + n
    if apfu_by_key:
        return by_key                                          # the apfu each constituent stands for (Fe and Fe2O3 given together add up)
    # oxygen is implied by the oxides; a stated O count only matters for the formula weight
    fw = sum(mass.values())
    hal = sum(mass[k] for k in mass if k in ('F', 'Cl', 'Br'))
    # O=F correction to the formula weight: the oxides double count the oxygen a halogen replaces
    corr = sum(apfu.get(h, 0) * EP.ATOMIC_WEIGHTS['O'] / 2 for h in ('F', 'Cl', 'Br') if h in apfu)
    fw -= corr
    wt = OrderedDict((k, m / fw * 100) for k, m in mass.items())
    if corr:
        wt['O=F,Cl'] = -corr / fw * 100
    return wt, fw

def kc(wt, k_override=None):
    K = constants(); k_override = k_override or {}
    rows = []; total = 0.0
    for key, w in wt.items():
        if key.startswith('O='):
            continue
        k = k_override.get(key, K.get(key, {}).get('k'))
        src = 'override' if key in k_override else K.get(key, {}).get('source', 'MISSING')
        contrib = (k or 0.0) * w / 100
        rows.append((key, w, k, contrib, src)); total += contrib
    return total, rows

def density_from_cell(cif, z=None, fw=None):
    from pxrd_review import bv_check as B
    st = B.Structure(cif, include_h=False)
    zz = z or B._num(st.block['items'].get('_cell_formula_units_z') or '') or None
    if not zz:
        raise ValueError('Z is not in the .cif — give Z')
    return zz * fw / (st.volume * 0.602214), zz, st.volume

def evaluate(wt, n_mean, density=None, cif=None, z=None, k_override=None, fw=None):
    KC, rows = kc(wt, k_override)
    out = {'KC': KC, 'rows': rows, 'n': n_mean, 'fw': fw}
    if cif and fw:
        dc, zz, V = density_from_cell(cif, z, fw)
        out.update(D_calc=dc, Z=zz, V=V)
    if density:
        out['D_meas'] = density
    for key, D in (('meas', out.get('D_meas')), ('calc', out.get('D_calc'))):
        if D:
            KP = (n_mean - 1) / D
            out['KP_' + key] = KP; out['CI_' + key] = 1 - KP / KC if KC else float('nan')
    return out

def report_text(res, name=''):
    L = ['Gladstone–Dale compatibility%s' % (' — ' + name if name else ''), '']
    L.append('  %-10s %8s %8s %10s   %s' % ('constituent', 'wt%', 'k', 'k·wt%/100', 'source'))
    for key, w, k, c, src in res['rows']:
        L.append('  %-10s %8.2f %8s %10.4f   %s' % (key, w, ('%.3f' % k) if k is not None else '—', c, src))
    L.append('  %-10s %8.2f %8s %10.4f   K_C' % ('Σ', sum(r[1] for r in res['rows']), '', res['KC']))
    if res.get('fw'):
        L.append('  formula weight %.2f' % res['fw'])
    if res.get('D_calc'):
        L.append('  D_calc = %.3f g/cm³  (Z = %g, V = %.2f Å³)' % (res['D_calc'], res['Z'], res['V']))
    for key, lab in (('meas', 'measured density'), ('calc', 'calculated density')):
        if 'KP_' + key in res:
            D = res['D_' + key]
            L.append('  n_mean %.3f, %s %.3f → K_P = %.4f, K_C = %.4f, compatibility 1 − K_P/K_C = %+.3f (%s)'
                     % (res['n'], lab, D, res['KP_' + key], res['KC'], res['CI_' + key], category(res['CI_' + key])))
    missing = [r[0] for r in res['rows'] if r[2] is None]
    if missing:
        L.append('  no constant for: %s — add it to data/gd_constants.json or pass --k' % ', '.join(missing))
    return '\n'.join(L)

def summary(res):
    """One line per density: (label, D, K_P, compatibility, category)."""
    out = []
    for key, lab in (('meas', 'measured'), ('calc', 'calculated')):
        if 'KP_' + key in res:
            out.append((lab, res['D_' + key], res['KP_' + key], res['CI_' + key], category(res['CI_' + key])))
    return out

def table(res, name='', journal_key=None):
    """The constants table as a pxrd_review.tables table (for render_html / write_word); the
    K_C / K_P / compatibility lines go in the note."""
    from pxrd_review import tables as T
    J = T.journal(journal_key)
    head = [T.C('Constituent'), T.C('wt%'), T.C(T.R('k', 'i')), T.C(T.R('k', 'i'), T.R('·wt%/100')), T.C('Source')]
    rows = [[T.C(key), T.C('%.2f' % w), T.C(('%.3f' % k) if k is not None else '—'), T.C('%.4f' % c), T.C(src)] for key, w, k, c, src in res['rows']]
    rows.append([T.C(T.R('K', 'i'), T.R('C', 'sub')), T.C('%.2f' % sum(r[1] for r in res['rows'])), T.C(''), T.C('%.4f' % res['KC']), T.C('')])
    note = []
    if res.get('D_calc'):
        note.append('Dcalc = %.3f g/cm³ (Z = %g, V = %.2f Å³, formula weight %.2f).' % (res['D_calc'], res['Z'], res['V'], res['fw']))
    for lab, D, KP, CI, cat in summary(res):
        note.append('n = %.3f, %s density %.3f: KP = %.4f, KC = %.4f, 1 − KP/KC = %+.3f (%s).' % (res['n'], lab, D, KP, res['KC'], CI, cat))
    missing = [r[0] for r in res['rows'] if r[2] is None]
    if missing:
        note.append('No constant for: %s.' % ', '.join(missing))
    cap = 'Gladstone–Dale compatibility%s' % (' of ' + name if name else '')
    return [('gd', {'n': 1, 'label': J['caption'].format(n=1), 'caption': T._title(J, cap), 'journal': J, 'head': head, 'rows': rows,
                    'note': ' '.join(note)})]

def prepare(formula=None, wt=None, oxide=None, n=None, density=None, cif=None, z=None, k=None):
    """CLI-style strings in -> the evaluation dict; ValueError on bad input."""
    if n is None:
        raise ValueError('the mean refractive index n is needed')
    fw = None
    if formula and formula.strip():
        ox_ = dict(kv.split('=', 1) for kv in (oxide or '').split(',') if '=' in kv)
        wt_, fw = formula_to_wt(_parse_pairs(formula), ox_)
        by_key = formula_to_wt(_parse_pairs(formula), ox_, apfu_by_key=True)
    elif wt and wt.strip():
        wt_ = _parse_pairs(wt)
    else:
        raise ValueError('give a formula (apfu) or wt% oxides')
    if not wt_:
        raise ValueError('nothing parsed from the composition — use X=value,X=value')
    res = evaluate(wt_, float(n), float(density) if density else None, cif, float(z) if z else None, _parse_pairs(k) if k else None, fw)
    res['o_corr'] = sum(v for k_, v in wt_.items() if k_.startswith('O='))     # the O ≡ F,Cl deduction: no constant, but part of the total
    if formula and formula.strip():
        res['apfu'] = True; res['apfu_by_key'] = by_key; res['wt_keys'] = [k_ for k_ in wt_ if not k_.startswith('O=')]   # what the workbook derives the wt% from
    return res

def write_xlsx(res, path, name='', paper=None):
    """The whole calculation as live formulas: from a formula, apfu -> mass per formula unit (apfu x MW over
    the cations of the constituent) -> formula weight (less the oxygen F and Cl replace) -> wt%; then
    k·wt%/100 -> K_C; D_calc = Z·FW / (V·0.602214); K_P = (n − 1)/D; 1 − K_P/K_C and its category.
    Change a k, an apfu, n, Z or a density and everything after it follows.
    paper: what the paper itself states — {'ci', 'category', 'D_calc', 'sentence', 'source'} — set beside the tool's on a
    'check' sheet, with what would explain a difference (see _write_gd_check)."""
    import openpyxl
    from openpyxl.styles import Font
    bold = Font(bold=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'GD'
    ws.append(['Gladstone–Dale' + (' — ' + name if name else '')]); ws['A1'].font = bold
    ws.append(['constituent', 'wt%', 'k', 'k·wt%/100', 'source of k']); ws.row_dimensions[2].font = bold
    first = 3; n_rows = len(res['rows']); last = first + n_rows - 1
    r_kc = last + 1; r_n = last + 3
    apfu = res.get('apfu'); fw_cell = None
    # the block a formula's wt% come from sits under the results; the wt% cells above point into it
    n_res = 1 + (3 if res.get('D_meas') else 0) + (7 if res.get('D_calc') else 0)
    d0 = r_n + n_res + 1                                                 # header row of the derivation
    wt_ref = {}
    if apfu:
        keys = [k for k in res['wt_keys']]
        for j, key in enumerate(keys):
            wt_ref[key] = 'E%d' % (d0 + 1 + j)
        r_corr = d0 + 1 + len(keys); r_fw = r_corr + 1
        fw_cell = 'D%d' % r_fw
    for i, (key, w, k, c, src) in enumerate(res['rows']):
        r = first + i
        ws.append([key, ('=' + wt_ref[key]) if key in wt_ref else w, k, '=B%d*C%d/100' % (r, r), src if k is not None else 'NO CONSTANT: this constituent adds nothing to K_C'])
    # Σ wt% is the analysis's total, so it carries the O ≡ F,Cl deduction although that row has no constant: without it
    # fluorite totalled 120.49 and fluorapatite 101.59, and the check sheet 'explained' a paper's index by a normalisation
    corr = ('+E%d' % r_corr) if apfu and any(k_ in ('F', 'Cl', 'Br') for k_ in res['wt_keys']) else ('%+r' % res['o_corr']) if res.get('o_corr') else ''
    ws.append(['K_C', '=SUM(B%d:B%d)%s' % (first, last, corr), '', '=SUM(D%d:D%d)' % (first, last),
               'Σ k·wt%/100 — over the WHOLE analysis: a short list gives a short K_C' + ('; Σ wt% less O ≡ F,Cl' if corr else '')])
    ws.append([]); ws.append(['n (mean index)', res['n']])
    row = r_n + 1; A = {'kc': 'D%d' % r_kc, 'sum': 'B%d' % r_kc, 'n': 'B%d' % r_n, 'first': first, 'last': last}
    cat = lambda cell: '=IF(ISNUMBER(%s),IF(ABS(%s)<0.02,"superior",IF(ABS(%s)<0.04,"excellent",IF(ABS(%s)<0.06,"good",IF(ABS(%s)<0.08,"fair","poor")))),"")' % ((cell,) * 5)   # no index, no category
    if res.get('D_meas'):
        ws.append(['D measured', res['D_meas']]); ws.append(['K_P (measured D)', '=(B%d-1)/B%d' % (r_n, row), '(n − 1) / D'])
        ws.append(['compatibility', '=IF(D%d=0,"no K_C: no constituent of this analysis has a constant",1-B%d/D%d)' % (r_kc, row + 1, r_kc), '1 − K_P/K_C', cat('B%d' % (row + 2))])
        A['meas'] = {'D': 'B%d' % row, 'KP': 'B%d' % (row + 1), 'CI': 'B%d' % (row + 2)}; row += 3
    if res.get('D_calc'):
        ws.append(['formula weight', ('=' + fw_cell) if fw_cell else res['fw']]); ws.append(['Z', res['Z']]); ws.append(['V (Å³)', res['V'], 'from the .cif cell'])
        ws.append(['Avogadro × 10⁻²⁴', 0.602214])
        ws.append(['D calculated', '=B%d*B%d/(B%d*B%d)' % (row + 1, row, row + 2, row + 3), 'Z · FW / (V · 0.602214)'])
        ws.append(['K_P (calculated D)', '=(B%d-1)/B%d' % (r_n, row + 4), '(n − 1) / D'])
        ws.append(['compatibility', '=IF(D%d=0,"no K_C: no constituent of this analysis has a constant",1-B%d/D%d)' % (r_kc, row + 5, r_kc), '1 − K_P/K_C', cat('B%d' % (row + 6))])
        A['calc'] = {'D': 'B%d' % (row + 4), 'KP': 'B%d' % (row + 5), 'CI': 'B%d' % (row + 6)}; row += 7
    if apfu:
        from pxrd_review import epma as EP
        assert row + 1 == d0, (row, d0)
        ws.append([]); ws.append(['from the formula', 'apfu', 'MW of the constituent', 'mass per formula unit', 'wt%', 'cations per formula of the constituent'])
        ws.row_dimensions[d0].font = bold
        for j, key in enumerate(keys):
            r = d0 + 1 + j; c = EP.parse_constituent(key)
            ws.append([key, res['apfu_by_key'][key], EP._mw_formula(c), '=B%d*C%d/F%d' % (r, r, r), '=100*D%d/$D$%d' % (r, r_fw), c.n_cat if c.kind != 'water' else 1])
        hal = [d0 + 1 + j for j, key in enumerate(keys) if key in ('F', 'Cl', 'Br')]
        ws.append(['O ≡ F,Cl', '', '', ('=-(%s)*%s/2' % ('+'.join('B%d' % h for h in hal), EP.ATOMIC_WEIGHTS['O'])) if hal else 0,
                   ('=100*D%d/$D$%d' % (r_corr, r_fw)) if hal else 0, 'half an O per F, Cl: the oxides count an oxygen the halogen replaces'])
        ws.append(['formula weight', '', '', '=SUM(D%d:D%d)' % (d0 + 1, r_corr)])
    for col, w in zip('ABCDEF', (22, 12, 22, 22, 12, 60)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions['E'].width = 44 if not apfu else 14
    if paper and (paper.get('ci') is not None or paper.get('category') or paper.get('D_calc')):
        _write_gd_check(wb, res, paper, A, cat)
    wb.save(path); return path

def paper_statement(path):
    """What a paper states about its own compatibility, for the check sheet — read locally, as every paper is."""
    from pxrd_review import paper_extract as PE
    text = PE.text_of(path)
    g = PE.gd_statement(text); o = PE.optics(text)
    return {'ci': g.get('ci'), 'category': g.get('category'), 'sentence': g.get('sentence') or '', 'D_calc': o.get('D_calc'), 'source': os.path.basename(path)}

def _write_gd_check(wb, res, paper, A, cat):
    """The 'check' sheet: the index the paper states beside the tool's, for each density, and — live — what would explain a
    difference. A Gladstone–Dale difference is NEVER red by itself: K_C is a mean over the whole analysis and rests on
    constants a paper may have taken from another class (Mandarino's variants), so a gap is amber and the line that
    reproduces the paper's number is the finding. Red is for the paper's own arithmetic alone: a category word that is not
    the category of the paper's own number. (A calculated density that is not Z·FW/(V·0.602214) for the values typed here,
    a constituent with no constant: amber — the input's doing as often as the paper's.)"""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    bold = Font(bold=True)
    red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid'); amber = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    green = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    wc = wb.create_sheet('check')
    if not res.get('KC'):
        # an element table, an organic mineral: nothing here has a Gladstone–Dale constant, so there is no index to set beside the paper's
        wc.cell(1, 1, 'PROBLEM — no constituent of this analysis has a Gladstone–Dale constant (an element table? oxides are needed): K_C is 0 and no index follows').fill = red_fill
        wc.column_dimensions['A'].width = 120
        return
    g = lambda a: 'GD!$%s$%s' % (a[0], a[1:])
    wc.cell(1, 1, "The compatibility the paper states against the one its numbers give%s — green: reproduces the paper's; amber: differs; red: arithmetic" % (' (%s)' % paper['source'] if paper.get('source') else '')).font = bold
    for c_, h in enumerate(['', 'value', "paper's", 'difference', 'reading'], 1):
        wc.cell(3, c_, h).font = bold
    row = 3; ci_p = paper.get('ci'); r_ci = {}; c_ci = None                  # c_ci: THE cell that holds the paper's index — edit it and everything follows
    dens = [(k, lab) for k, lab in (('meas', 'measured density'), ('calc', 'calculated density')) if k in A]
    for k, lab in dens:
        row += 1; r_ci[k] = row
        wc.cell(row, 1, '1 − K_P/K_C, %s' % lab); wc.cell(row, 2, '=' + g(A[k]['CI']))
        if ci_p is not None:
            if c_ci is None:
                wc.cell(row, 3, ci_p); c_ci = '$C$%d' % row
            else:
                wc.cell(row, 3, '=' + c_ci)
            wc.cell(row, 4, '=B%d-C%d' % (row, row))
            wc.cell(row, 5, '=IF(ABS(D%d)<=0.005,"ok — reproduces the paper\'s index",IF(ABS(D%d)<=0.03,"ok — within what the constants allow (0.03)","note — differs by "&TEXT(D%d,"+0.000;-0.000")&": see what would explain it, below"))' % (row, row, row))
        for c_ in (2, 3, 4):
            wc.cell(row, c_).number_format = '+0.000;-0.000'
    if ci_p is not None and len(dens) == 2:
        row += 1
        wc.cell(row, 1, 'the density the paper used'); wc.cell(row, 5, '=IF(ABS(D%d)<=ABS(D%d),"its index is nearer the MEASURED density\'s","its index is nearer the CALCULATED density\'s")' % (r_ci['meas'], r_ci['calc']))
    if ci_p is not None and not dens:
        row += 1; r_ci['paper'] = row                                   # no density given: the paper's number still stands beside its word
        wc.cell(row, 1, "1 − K_P/K_C as the paper states it (no density given here: none computed)"); wc.cell(row, 3, ci_p); wc.cell(row, 3).number_format = '+0.000;-0.000'; c_ci = '$C$%d' % row
    if paper.get('category'):
        row += 1
        wc.cell(row, 1, 'category'); wc.cell(row, 3, paper['category'])
        best = r_ci[dens[0][0]] if dens else None
        if best:
            wc.cell(row, 2, cat('B%d' % best))
        if ci_p is not None:
            # the paper's word against the paper's own number: no constant and no density enters into that
            own = cat('C%d' % (best or r_ci['paper']))[1:]
            wc.cell(row, 5, '=IF(%s="%s","ok — the paper\'s word is the category of its own number","PROBLEM — the paper calls "&TEXT(%s,"+0.000;-0.000")&" a ‘%s’ compatibility; Mandarino\'s category for that number is "&%s)' % (
                own, paper['category'], c_ci, paper['category'], own))
    if paper.get('D_calc') and 'calc' in A:
        row += 1
        wc.cell(row, 1, 'D calculated = Z·FW/(V·0.602214)'); wc.cell(row, 2, '=' + g(A['calc']['D'])); wc.cell(row, 3, paper['D_calc']); wc.cell(row, 4, '=B%d-C%d' % (row, row))
        # amber: the paper's density is from its EMPIRICAL formula and the one here is whatever was typed; every corpus red of
        # the paper check's D_calc line dissolved on inspection and it was retired to a note (owner, 2026-09-06)
        wc.cell(row, 5, '=IF(ABS(D%d)<=0.01*C%d,"ok — the paper\'s calculated density follows from Z, the formula weight and V","note — Z·FW/(V·0.602214) for the values HERE differs from the paper\'s by "&TEXT(100*D%d/C%d,"+0.0;-0.0")&" %%: the formula typed (ideal against the paper\'s empirical one, its water above all), Z or the cell")' % (row, row, row, row))
        wc.cell(row, 2).number_format = wc.cell(row, 4).number_format = '0.000'
    row += 1
    wc.cell(row, 1, 'Σ wt% of the analysis'); wc.cell(row, 2, '=' + g(A['sum']))
    # said as a note only where the index is NOT reproduced: on 62 corpus papers whose stated index the check reproduces the
    # total was outside 98.5–101.5 all the same (water by difference, a trace left out) and the amber line said nothing
    ok_rows = [r_ci[k] for k, _l in dens] if ci_p is not None else []
    reproduced = ('OR(%s)' % ','.join('LEFT(E%d,2)="ok"' % r_ for r_ in ok_rows)) if ok_rows else 'FALSE'
    wc.cell(row, 5, '=IF(AND(B%d>=98.5,B%d<=101.5),"ok",IF(%s,"ok — a total of "&TEXT(B%d,"0.00")&", and the paper\'s index is reproduced all the same","note — K_C is a mean over the WHOLE analysis: a total of "&TEXT(B%d,"0.00")&" moves K_C, and the index, by about that proportion"))' % (row, row, reproduced, row, row))
    r_sum = row
    missing = [r[0] for r in res['rows'] if r[2] is None]
    for m in missing:
        row += 1; wc.cell(row, 1, 'no constant for %s' % m); wc.cell(row, 5, 'note — %s adds nothing to K_C here: the index is short by its share (give one with --k)' % m)
    if ci_p is not None and dens:
        row += 2
        wc.cell(row, 1, 'WHAT WOULD EXPLAIN A DIFFERENCE').font = bold
        for c_, h in enumerate(['index it gives', "paper's", 'difference', 'reading'], 2):
            wc.cell(row, c_, h).font = bold
        expl = '=IF(ABS(D%d)<=0.005,IF(ABS($D$%d)<=0.005,"reproduces it too","EXPLAINS IT — %s"),"does not reproduce it")'
        for k, lab in dens:
            base = r_ci[k]; kp = g(A[k]['KP'])
            row += 1
            wc.cell(row, 1, 'the analysis normalised to 100 %% (%s)' % lab)
            wc.cell(row, 2, '=1-%s/(%s*100/%s)' % (kp, g(A['kc']), g(A['sum']))); wc.cell(row, 3, '=' + c_ci); wc.cell(row, 4, '=B%d-C%d' % (row, row))
            wc.cell(row, 5, expl % (row, base, 'the paper normalised its analysis to 100 % before taking K_C'))
            K = constants()
            for i, (key, w, kk, c, src) in enumerate(res['rows']):
                for var in (K.get(key, {}).get('variants') or []):
                    if kk is None or var.get('k') is None or abs(var['k'] - kk) < 1e-9:
                        continue
                    row += 1
                    wc.cell(row, 1, '%s with k = %g (%s) — %s' % (key, var['k'], ILLEGAL_CHARACTERS_RE.sub(' ', str(var.get('for', '')))[:60], lab))
                    wc.cell(row, 2, '=1-%s/(%s+(%r-GD!$C$%d)*GD!$B$%d/100)' % (kp, g(A['kc']), var['k'], A['first'] + i, A['first'] + i))
                    wc.cell(row, 3, '=' + c_ci); wc.cell(row, 4, '=B%d-C%d' % (row, row))
                    wc.cell(row, 5, expl % (row, base, "the paper took Mandarino's other constant for %s" % key))
        row += 2
        wc.cell(row, 1, "WHAT WOULD GIVE THE PAPER'S INDEX").font = bold
        for c_, h in enumerate(['needed', 'used', 'difference'], 2):
            wc.cell(row, c_, h).font = bold
        for k, lab in dens:
            D = g(A[k]['D'])
            # (the paper's index by reference, never pasted: edited in its cell, the 'needed' values follow — and an index of
            # exactly 1 no longer divides by zero silently: IF guards it)
            for what, need, used in (('K_C (%s)' % lab, '=IF(%s=1,"",%s/(1-%s))' % (c_ci, g(A[k]['KP']), c_ci), g(A['kc'])),
                                     ('n (%s)' % lab, '=1+%s*%s*(1-%s)' % (D, g(A['kc']), c_ci), g(A['n'])),
                                     ('D (%s)' % lab, '=IF(%s=1,"",(%s-1)/(%s*(1-%s)))' % (c_ci, g(A['n']), g(A['kc']), c_ci), D)):
                row += 1
                wc.cell(row, 1, what); wc.cell(row, 2, need); wc.cell(row, 3, '=' + used); wc.cell(row, 4, '=B%d-C%d' % (row, row))
                for c_ in (2, 3, 4):
                    wc.cell(row, c_).number_format = '0.0000' if what.startswith('K_C') else '0.000'
    wc.conditional_formatting.add('A4:E%d' % row, FormulaRule(formula=['LEFT($E4,7)="PROBLEM"'], fill=red_fill))
    wc.conditional_formatting.add('A4:E%d' % row, FormulaRule(formula=['LEFT($E4,4)="note"'], fill=amber))
    wc.conditional_formatting.add('A4:E%d' % row, FormulaRule(formula=['LEFT($E4,8)="EXPLAINS"'], fill=green))
    if paper.get('sentence'):
        row += 2; wc.cell(row, 1, 'the paper: «%s»' % ILLEGAL_CHARACTERS_RE.sub(' ', paper['sentence'])[:400])
    wc.column_dimensions['A'].width = 58
    for c_, w in zip('BCDE', (16, 12, 12, 110)):
        wc.column_dimensions[c_].width = w

def _parse_pairs(s):
    out = OrderedDict()
    for tok in (s or '').split(','):
        if '=' in tok:
            k, v = tok.rsplit('=', 1) if tok.strip().startswith('O=') else tok.split('=', 1)   # 'O=F=-1.47': the deduction a table prints
            out[k.strip()] = float(v)
    return out

def main(argv=None):
    ap = argparse.ArgumentParser(prog='pxrd gd', description=__doc__.split('\n\n')[1], formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--wt', help='composition as wt%% oxides: "UO3=63.88,PbO=13.41,H2O=7.0"')
    ap.add_argument('--formula', help='ideal formula as atoms per formula unit: "Pb=1,U=3,Se=2,H2O=5" (O implied)')
    ap.add_argument('--oxide', action='append', default=[], help='element=constituent for the formula route, e.g. Fe=Fe2O3, S=S (repeatable)')
    ap.add_argument('--n', type=float, required=True, help='mean refractive index')
    ap.add_argument('--density', type=float, help='measured density (g/cm³)')
    ap.add_argument('--cif', help='.cif for the calculated density'); ap.add_argument('--z', type=float)
    ap.add_argument('--k', action='append', default=[], help='override a constant: UO3=0.118 (repeatable)')
    ap.add_argument('--name', default=''); ap.add_argument('--xlsx', action='store_true'); ap.add_argument('--out')
    ap.add_argument('--paper', help="the paper (.pdf / .docx): its stated compatibility index, category and calculated density go on the workbook's check sheet")
    ap.add_argument('--paper-ci', type=float, help='the compatibility index the paper states, when there is no file to read it from')
    a = ap.parse_args(argv)
    try:
        res = prepare(a.formula, a.wt, ','.join(a.oxide), a.n, a.density, a.cif, a.z, ','.join(a.k))
    except ValueError as e:
        raise SystemExit('gd: %s' % e)
    print(report_text(res, a.name))
    if a.xlsx:
        out = a.out or os.getcwd(); os.makedirs(out, exist_ok=True)
        try:
            paper = paper_statement(a.paper) if a.paper else {}
        except Exception as ex:
            raise SystemExit('gd: --paper: %s could not be read (%s)' % (a.paper, str(ex)[:120]))
        if a.paper_ci is not None:
            paper['ci'] = a.paper_ci
        p = write_xlsx(res, os.path.join(out, (a.name or 'gd') + '_gd.xlsx'), a.name, paper or None); print('  xlsx → %s' % p)
    return 0

if __name__ == '__main__':
    sys.exit(main())
